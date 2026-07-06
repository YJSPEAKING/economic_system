from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_paper_outline_docx import (
    DSCR_FIG,
    MAIN_FIG,
    SUMMARY_CSV,
    OUT_DIR,
    add_cn_paragraph,
    add_heading,
    read_summary_rows,
    set_cell_text,
    set_paragraph_format,
    set_run_font,
    set_table_borders,
    setup_styles,
    shade_cell,
)


OUT_PATH = OUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版.docx"


def add_plain_heading(doc: Document, text: str, level: int) -> None:
    add_heading(doc, text, level)


def add_intro(doc: Document) -> None:
    add_plain_heading(doc, "1  引言", 1)
    add_cn_paragraph(
        doc,
        "理论经济学研究常需要对宏观经济现象建立数学模型，并以形式化方式描述经济运行中具有普遍性的规律。"
        "这类研究通常依赖大量经济运行数据，但现实经济活动数据存在收集难度大、来源环境不可控、数据质量难以稳定保证等问题。"
        "同时，部分理论经济变量具有较强抽象性，难以在真实经济环境中直接采集。上述限制使得单纯依靠真实数据支撑理论经济建模存在现实困难。",
    )
    add_cn_paragraph(
        doc,
        "为弥补现实数据的不足，研究者可以利用仿真方法构建虚拟经济运行环境，并在可控条件下生成和提取经济运行数据。"
        "多主体仿真方法能够将经济主体表示为具有一定自主性的 agent，并通过主体间交互推动系统状态变化。"
        "在理论经济学研究场景下，该思路可以形成复杂经济系统多主体仿真，即 MAS-E。"
    )
    add_cn_paragraph(
        doc,
        "MAS-E 的关键问题之一是主体决策生成机制。传统固定规则方法依赖人工设定，灵活性和自适应性有限。"
        "强化学习方法能够根据环境反馈训练决策策略，但其训练目标通常围绕奖励最大化展开，容易表现出偏向完全理性的行为模式。"
        "真实市场主体的决策则受到风险偏好、经验判断、有限信息和行为习惯等因素影响，不能简单等同于单一奖励目标的最优化结果。"
    )
    add_cn_paragraph(
        doc,
        "基于上述背景，模仿学习可以作为强化学习在经济主体决策生成中的补充方法。"
        "模仿学习以专家行为数据为参照，目标是使生成行为接近专家数据中体现的决策模式。"
        "本文在课题组既有多主体经济仿真系统基础上，围绕主体决策模型进行改进：一方面，将生成对抗模仿学习与 TD3 结合，利用判别器提供与专家行为相联系的模仿信号；另一方面，在扩展模型中引入 Transformer 编码器，以增强生产企业策略网络对历史状态信息的表征能力。"
    )
    add_cn_paragraph(
        doc,
        "本文的研究重点不是重新定义仿真环境，而是在既有环境与主体设置下比较不同主体决策模型的训练表现。"
        "实验部分以 TD3 为基线方法，比较 GAIL+TD3 与 Transformer+GAIL+TD3 在系统存活天数、企业累计收益、银行累计收益和生产企业偿债风险等指标上的表现。"
    )


def add_environment_and_problem(doc: Document) -> None:
    add_plain_heading(doc, "2  仿真环境与主体决策建模问题", 1)
    add_plain_heading(doc, "2.1  仿真环境", 2)
    add_cn_paragraph(
        doc,
        "本文沿用既有复杂经济系统仿真环境的建模方案[1]。该仿真环境由企业主体、银行主体和市场机制共同构成。"
        "企业主体通过生产、交易和借贷活动参与市场运行，银行主体通过贷款发放与本息回收同企业形成信用关系。"
        "在仿真过程中，虚拟经济主体被置于经济环境中持续交互，主体决策会改变市场状态，市场状态又会影响后续主体决策。"
    )
    add_cn_paragraph(
        doc,
        "一次完整的仿真实验过程被定义为一个“回合”。在每个回合中，系统按照包含 P1 至 P6 六个阶段的小周期循环运作，本文将这样一个小周期称为“一天”。"
        "当系统达到终止条件时，该回合结束，环境状态和主体属性随后被重置，用于开始下一回合。"
        "由于本文重点是主体决策模型改进，具体生产公式、交易规则和银行约束均沿用既有研究[1]，本文不再展开重复推导。"
    )
    add_plain_heading(doc, "2.2  主体决策建模", 2)
    add_cn_paragraph(
        doc,
        "在上述仿真环境中，企业和银行均被建模为具有自主决策能力的主体。企业主体需要根据自身现金、存货、债务、市场价格和历史交易信息生成生产、采购、借贷和定价相关动作；银行主体需要根据企业运行状态和自身资金状态生成贷款相关动作。"
        "因此，主体决策建模可以被表述为连续状态空间和连续动作空间下的多主体决策问题。"
    )
    add_cn_paragraph(
        doc,
        "本文的核心问题是：在同一仿真环境和相同主体设置下，模仿学习信号是否能够改善强化学习主体的训练过程，以及时序表征模块是否能够进一步改善模型对动态经济环境的适应能力。"
    )
    add_plain_heading(doc, "2.3  评价指标设计", 2)
    add_cn_paragraph(
        doc,
        "本文使用系统存活天数评价仿真系统持续运行能力。存活天数越高，表示系统在该回合中维持生产、交易和清算过程的时间越长。"
        "企业累计收益用于评价企业经营表现，其展示指标按 2×总收入−总支出−总利息 计算。银行累计收益按企业累计支付利息之和计算。"
        "此外，本文引入生产企业偿债能力指标 DSCR，用于在稳定运行阶段分析企业现金对当期应还本息的覆盖程度。"
    )


def add_human_data_section(doc: Document) -> None:
    add_plain_heading(doc, "3  专家数据采集与有限理性行为表示", 1)
    add_plain_heading(doc, "3.1  专家数据的作用", 2)
    add_cn_paragraph(
        doc,
        "模仿学习需要以专家数据作为行为参照。对于本文的主体决策任务，专家数据可以表示为状态和动作之间的对应关系，也可以进一步扩展为一段连续决策轨迹。"
        "判别器利用专家数据和模型生成数据之间的差异，为策略学习提供区别于环境奖励的行为约束信号。"
    )
    add_plain_heading(doc, "3.2  人类行为采集系统", 2)
    add_cn_paragraph(
        doc,
        "人类行为采集系统对应 Method2_人类行为采集分支。该系统的目标是让人类操作者在仿真环境中控制经济主体，并记录操作过程中的状态、动作和环境反馈，从而形成可用于模仿学习的专家数据。"
        "正式论文后续需要补充该分支中的交互界面、数据字段和数据清洗流程。"
    )
    add_plain_heading(doc, "3.3  当前实验使用边界", 2)
    add_cn_paragraph(
        doc,
        "根据当前实验材料，真实企业行为数据尚未纳入本轮训练结果分析。"
        "因此，本文在现阶段应将真实企业行为数据写作后续可扩展的数据来源，而不能将其表述为已经完成的实验数据来源。"
    )


def add_algorithm_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    title = p.add_run("Algorithm 1  Online GAIL+TD3 for economic-agent decision learning")
    set_run_font(title, east_asia="宋体", ascii_font="Times New Roman")
    title.font.size = Pt(9)
    title.bold = True

    lines = [
        "Initialize actor μθ, twin critics Qω1,Qω2, target networks μθ-, Qω1-,Qω2-, replay buffer D",
        "Initialize discriminator Dψ and expert data sampler E",
        "for each episode do",
        "    Reset the MAS-E environment and obtain initial state s0",
        "    while the episode is not terminated do",
        "        Select action at = μθ(st) + exploration noise",
        "        Execute at in the environment and observe reward renv,t, next state st+1 and done flag",
        "        Store transition (st, at, renv,t, st+1, done) in D",
        "        if replay buffer size reaches the learning threshold then",
        "            Sample mini-batch B = (s,a,r,s',done) from D",
        "            Sample expert mini-batch BE = (sE,aE) from E",
        "            Update Dψ with expert pairs as soft real labels and replay pairs as soft fake labels",
        "            Compute imitation reward rint = sigmoid(Dψ(s,a))",
        "            Fuse reward rfused = renv + wgail · rint, where wgail increases during warm-up",
        "            Compute TD3 target using target actors, target critics and rfused",
        "            Update twin critics by minimizing TD error",
        "            if policy-delay condition is satisfied then",
        "                Update actor by maximizing Qω1(s, μθ(s))",
        "                Soft-update target actor and target critics",
        "            end if",
        "        end if",
        "    end while",
        "end for",
    ]
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_run_font(r, east_asia="宋体", ascii_font="Times New Roman")
        r.font.size = Pt(8.5)


def add_gail_td3_section(doc: Document) -> None:
    add_plain_heading(doc, "4  基于生成对抗模仿学习与TD3的主体决策模型", 1)
    add_plain_heading(doc, "4.1  模型训练方案", 2)
    add_cn_paragraph(
        doc,
        "生成对抗模仿学习通过生成器与判别器之间的对抗训练学习专家行为模式。"
        "在本文模型中，动作生成网络同时承担强化学习中的 Actor 角色和模仿学习中的生成器角色。"
        "判别器接收状态-动作对，并区分该样本来自专家数据还是来自当前策略生成数据。"
    )
    add_cn_paragraph(
        doc,
        "与单纯离线模仿不同，本文采用在线 GAIL+TD3 训练框架。"
        "主体在仿真环境中持续交互并产生经验，经验池中的样本用于 TD3 的 Critic 与 Actor 更新；同一批生成样本也作为判别器的假样本，与专家样本共同用于判别器训练。"
        "因此，模仿信号通过判别器输出进入价值估计过程，并间接影响 Actor 的更新方向。"
    )
    add_plain_heading(doc, "4.2  状态、动作与经验样本", 2)
    add_cn_paragraph(
        doc,
        "企业主体的动作包括与借贷意愿、采购意愿和定价相关的连续动作；银行主体的动作包括面向企业的贷款发放决策。"
        "在每一次环境交互后，系统将状态、动作、环境奖励、下一状态和终止标记写入经验池。"
        "代码中的学习过程从经验池采样一个 mini-batch，并在同一批状态-动作样本上计算判别器打分和 TD3 目标值。"
    )
    add_plain_heading(doc, "4.3  判别器更新与模仿奖励构造", 2)
    add_cn_paragraph(
        doc,
        "当前代码中，判别器采用专家样本作为真实样本，采用经验池采样得到的策略样本作为生成样本。"
        "判别器训练使用二元交叉熵形式的目标函数，并采用软标签设置：专家样本标签为 0.9，生成样本标签为 0.1。"
        "在计算模仿奖励时，代码将判别器输出 logits 经过 sigmoid 映射到 0 到 1 的区间，并作为内部奖励。"
    )
    add_cn_paragraph(
        doc,
        "融合奖励由环境奖励与加权模仿奖励构成。代码中 GAIL 奖励权重随训练步数经过 warm-up 逐步增大，随后保持在设定权重。"
        "Critic 使用该融合奖励构造 TD 目标值，Actor 则在延迟更新条件满足时通过最大化 Critic 估计值进行更新。"
    )
    add_plain_heading(doc, "4.4  在线训练流程", 2)
    add_cn_paragraph(
        doc,
        "算法 1 给出当前 GAIL+TD3 实现对应的在线训练流程。该伪代码强调训练逻辑，不展开网络层数和具体超参数。"
    )
    add_algorithm_table(doc)


def add_transformer_section(doc: Document) -> None:
    add_plain_heading(doc, "5  基于Transformer表征增强的GAIL+TD3主体决策模型", 1)
    add_plain_heading(doc, "5.1  表征增强动机", 2)
    add_cn_paragraph(
        doc,
        "经济主体的当前决策受到历史交易、历史贷款、历史价格和库存变化等因素影响。"
        "如果策略网络只处理单一时刻状态，模型可能难以充分利用历史窗口内的动态信息。"
        "因此，扩展模型在 GAIL+TD3 框架基础上引入 Transformer 编码器，以学习历史状态序列中的时序表征。"
    )
    add_plain_heading(doc, "5.2  Transformer表征增强网络结构", 2)
    add_cn_paragraph(
        doc,
        "Transformer 编码器用于将历史窗口内的状态序列映射为表征向量，随后由 Actor 网络根据该表征生成连续动作。"
        "根据当前项目描述，Transformer 主要用于生产企业侧的策略网络。该表述在正式定稿前还需要以 run_transformer 分支代码为准。"
    )
    add_plain_heading(doc, "5.3  模型集成与训练方式", 2)
    add_cn_paragraph(
        doc,
        "Transformer+GAIL+TD3 保留生成对抗模仿学习和 TD3 价值更新框架。"
        "其主要差异在于策略网络前端加入时序表征模块，使主体动作不只依赖当前状态，也依赖历史状态窗口中包含的动态信息。"
        "本文当前实验未设置 Transformer 历史序列长度消融实验，因此实验分析部分不单独讨论历史窗口长度对性能的影响。"
    )


def add_metric_summary_table(doc: Document) -> None:
    rows = read_summary_rows()
    if not rows:
        add_cn_paragraph(doc, "当前未读取到四指标汇总 CSV。")
        return
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    headers = ["算法", "指标", "样本数", "第60步均值", "第60步SEM"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, center=True)
        shade_cell(table.rows[0].cells[idx], "EDEDED")
    metric_names = {
        "survival": "存活天数",
        "production": "生产企业累计收益",
        "consumption": "消费企业累计收益",
        "bank": "银行累计收益",
    }
    for row in rows:
        cells = table.add_row().cells
        values = [
            row["algorithm"],
            metric_names.get(row["metric"], row["metric"]),
            row["n"],
            f"{float(row['last_mean']):.4f}",
            f"{float(row['last_sem']):.4f}",
        ]
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, center=idx >= 2)


def add_experiment_section(doc: Document) -> None:
    add_plain_heading(doc, "6  实验设置与结果分析", 1)
    add_plain_heading(doc, "6.1  实验设置", 2)
    add_plain_heading(doc, "6.1.1  对比方法", 3)
    add_cn_paragraph(
        doc,
        "本文设置三类对比方法：TD3、GAIL+TD3 和 Transformer+GAIL+TD3。"
        "TD3 作为基线方法，GAIL+TD3 用于检验模仿学习信号对主体决策学习的影响，Transformer+GAIL+TD3 用于检验时序表征模块对模型表现的影响。"
    )
    add_plain_heading(doc, "6.1.2  训练配置", 3)
    add_cn_paragraph(
        doc,
        "每类方法均采用 8 个随机种子进行实验。本文将每 100 个回合记录一次统计指标，并使用跨 seed 的均值作为曲线值，使用标准误 SEM 表示 seed 间波动。"
        "具体 run id 记录见附录 A。"
    )
    add_plain_heading(doc, "6.1.3  评价指标", 3)
    add_cn_paragraph(
        doc,
        "主要评价指标包括每百回合存活天数、生产企业累计收益、消费企业累计收益、银行累计收益和生产企业偿债能力风险。"
        "其中，偿债风险图选取稳定运行阶段样本，用 DSCR<1 的企业日观测占比表示偿债压力。"
    )
    add_plain_heading(doc, "6.2  三类方法的主指标对比", 2)
    add_cn_paragraph(
        doc,
        "图 1 给出三类方法在四个主要指标上的训练曲线。曲线表示 8 个 seed 的均值，阴影表示 SEM。"
        "该图用于比较三类方法在学习速度、稳定阶段水平和随机种子波动方面的差异。"
    )
    if MAIN_FIG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(MAIN_FIG), width=Inches(5.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(cap, first_line=False, after=6, line_spacing=1.2)
        r = cap.add_run("图 1  TD3、GAIL+TD3 与 Transformer+GAIL+TD3 的四指标对比曲线")
        set_run_font(r)
        r.font.size = Pt(9)
    add_metric_summary_table(doc)
    add_plain_heading(doc, "6.3  稳定阶段偿债风险分析", 2)
    add_cn_paragraph(
        doc,
        "图 2 给出稳定阶段生产企业 DSCR<1 的风险占比。"
        "该指标用于补充说明主体在获得较高存活天数后是否仍存在短期偿债压力。"
        "正式论文需要在图注或方法部分明确 DSCR 的计算公式、稳定阶段样本选择规则和离群值处理方式。"
    )
    if DSCR_FIG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(DSCR_FIG), width=Inches(4.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(cap, first_line=False, after=6, line_spacing=1.2)
        r = cap.add_run("图 2  稳定阶段生产企业 DSCR<1 风险占比")
        set_run_font(r)
        r.font.size = Pt(9)


def add_conclusion(doc: Document) -> None:
    add_plain_heading(doc, "7  总结与展望", 1)
    add_plain_heading(doc, "7.1  总结", 2)
    add_cn_paragraph(
        doc,
        "本文围绕复杂经济系统仿真中的主体决策生成问题，构建了基于 GAIL+TD3 的在线决策模型，并进一步设计了 Transformer 表征增强版本。"
        "实验部分比较 TD3、GAIL+TD3 与 Transformer+GAIL+TD3 在存活天数、累计收益和偿债风险方面的表现。"
        "当前大纲稿已经给出论文结构、方法描述、伪代码、指标定义和实验结果位置，后续需要继续补充完整正文和统计分析。"
    )
    add_plain_heading(doc, "7.2  展望", 2)
    add_cn_paragraph(
        doc,
        "后续研究可以从三个方面展开：第一，补充更完整的人类行为采集实验，并明确专家数据来源；第二，统一 DSCR 后处理规则并增加统计检验；第三，在 Transformer 扩展模型中加入历史窗口长度、网络结构和模仿奖励权重的消融实验。"
    )


def add_references(doc: Document) -> None:
    add_plain_heading(doc, "参考文献", 1)
    refs = [
        'Chen, Bo, Shiqi Ding, Yukun Lin, and Guohong Chen. "Multi-agent Simulation of Complex Economic Systems Driven by Deep Reinforcement Learning." 2025 IEEE 17th International Conference on Computer Research and Development (ICCRD), IEEE, 2025, pp. 47-54. doi:10.1109/ICCRD64588.2025.10963160.',
        'Ho, Jonathan, and Stefano Ermon. "Generative Adversarial Imitation Learning." Advances in Neural Information Processing Systems, 2016. 待补全页码信息。',
        'Fujimoto, Scott, Herke van Hoof, and David Meger. "Addressing Function Approximation Error in Actor-Critic Methods." Proceedings of the 35th International Conference on Machine Learning, 2018. 待补全页码信息。',
        'Vaswani, Ashish, et al. "Attention Is All You Need." Advances in Neural Information Processing Systems, 2017. 待补全页码信息。',
    ]
    for ref in refs:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=False, after=3, line_spacing=1.25)
        r = p.add_run(ref)
        set_run_font(r)
        r.font.size = Pt(9)


def add_appendix(doc: Document) -> None:
    doc.add_page_break()
    add_plain_heading(doc, "附录A  实验 run id 记录表", 1)
    add_cn_paragraph(
        doc,
        "下表根据当前用户确认后的实验记录整理，用于对应三类方法、随机种子和 swanlog 运行目录。",
    )
    groups = {
        "TD3": [
            ("184", "run-20260621_202812-6ig7ywjq5ma7q4r3n4d2k"),
            ("291", "run-20260621_221653-3oxu02lub6ic88gg2tlyv"),
            ("83", "run-20260621_235606-cxmnv919c5vgxd18ag5m2"),
            ("739", "run-20260622_025851-vc8eettjrkns0ask0j149"),
            ("512", "run-20260622_012518-o3wa23ld5k0anno4nhwb9"),
            ("117", "run-20260622_043232-p1wsv6nx8qn3fwe50d8dd"),
            ("894", "run-20260622_063205-sgs3p05et5bj1by7pfa86"),
            ("652", "run-20260622_080822-vx1bs5edgj8z9ihb6xe1x"),
        ],
        "GAIL+TD3": [
            ("184", "run-20260621_193255-5pg7iyrvmcx34frn6ale5"),
            ("291", "run-20260621_223131-dm2s10i96kbbd7fky6owy"),
            ("83", "run-20260622_004812-zslinp7qvacgwbfuqujtw"),
            ("739", "run-20260622_031937-daocyl38dnxhsxwoxtkov"),
            ("894", "run-20260622_053309-3gob8caas2nn2pu2ozthf"),
            ("652", "run-20260622_075227-p7wj1wfq83b3mhkyo053i"),
            ("187", "run-20260622_110022-x8utkar2er1eerentcryu"),
            ("191", "run-20260622_134950-g5g4dxxf5ozv6uufkbfmk"),
        ],
        "Transformer+GAIL+TD3": [
            ("184", "run-20260622_223056-o8iu1b7pexa77lh7sd8cx"),
            ("291", "run-20260623_033112-glurxrgn07nhly7dvjkjp"),
            ("83", "run-20260623_080809-pdrdexgwuux1bd54os63u"),
            ("739", "run-20260623_131038-y6xubdud4wjny573weehp"),
            ("894", "run-20260623_172840-t8vokks22cdly7camzcud"),
            ("652", "run-20260623_212641-cj8cxudsz6o3ilixh1qc8"),
            ("187", "run-20260624_024311-i0350u1ei1u9bcogckxi4"),
            ("192", "run-20260624_073818-an8ie0ctj7ii46dc18f2a"),
        ],
    }
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for idx, header in enumerate(["算法", "seed", "run id"]):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, center=True)
        shade_cell(table.rows[0].cells[idx], "EDEDED")
    for algorithm, items in groups.items():
        for seed, run_id in items:
            row = table.add_row().cells
            set_cell_text(row[0], algorithm, center=True)
            set_cell_text(row[1], seed, center=True)
            set_cell_text(row[2], run_id)


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    add_intro(doc)
    add_environment_and_problem(doc)
    add_human_data_section(doc)
    add_gail_td3_section(doc)
    add_transformer_section(doc)
    add_experiment_section(doc)
    add_conclusion(doc)
    add_references(doc)
    add_appendix(doc)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
