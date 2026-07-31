from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import importlib.util
from pathlib import Path
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

BUNDLED_SITE_PACKAGES = Path(
    r"C:\Users\legion\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\python\Lib\site-packages"
)
if str(BUNDLED_SITE_PACKAGES) not in sys.path:
    sys.path.append(str(BUNDLED_SITE_PACKAGES))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "reference"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_v1.docx"
)
OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT = OUTPUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v28.docx"
EXPECTED_SOURCE_SHA256 = (
    "753659DBA9394BA37EC057DF1C61A53450F779258635D8E2E30C234253B92421"
)
MML2OMML_XSL = Path(
    r"C:\Program Files\Microsoft Office\Root\Office16\MML2OMML.XSL"
)
FIGURE_OUTPUT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "paper_gail_component_validation"
)
NOISY_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "expert_vs_noisy_expert_discriminator_eval_noise_std_0p30"
)
ACTOR_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "actor_snapshot_evolution"
)
DISCRIMINATOR_FIGURE = FIGURE_OUTPUT_DIR / "expert_vs_noisy_expert_seed184_paper.png"
ACTOR_FIGURE = FIGURE_OUTPUT_DIR / "actor_expert_js_evolution_mean_sem_paper.png"
ENGLISH_FIGURE_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "paper_v28_english_figures"
)
TD3_GAIL_FIGURE = ENGLISH_FIGURE_DIR / "td3_vs_gail_td3_three_metrics_sem.png"
GAIL_TRANSFORMER_FIGURE = (
    ENGLISH_FIGURE_DIR / "gail_td3_vs_transformer_three_metrics_sem.png"
)
DSCR_RISK_FIGURE = ENGLISH_FIGURE_DIR / "three_algorithm_dscr_risk.png"
LEARNING_CURVE_SCRIPT = (
    ROOT / "real_System_remake" / "plot_five_metrics_two_comparisons_aulc_sem.py"
)
LEARNING_CURVE_RAW_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "five_metrics_two_comparisons_aulc_sem"
    / "raw_metrics"
)
DSCR_RISK_CSV = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "post_training_dscr"
    / "three_algorithm_recent_100_survival_gt_90"
    / "three_algorithm_dscr_below_1_share.csv"
)


NOISY_ACTION_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mover><mi>a</mi><mo>~</mo></mover><mi>i</mi><mi>E</mi></msubsup>
  <mo>=</mo><mi mathvariant="normal">clip</mi><mfenced>
  <msubsup><mover><mi>a</mi><mo>^</mo></mover><mi>i</mi><mi>E</mi></msubsup>
  <mo>+</mo><msub><mi>ε</mi><mi>i</mi></msub><mo>,</mo><mo>−</mo><mi>b</mi>
  <mo>,</mo><mi>b</mi></mfenced><mo>,</mo>
  <msub><mi>ε</mi><mi>i</mi></msub><mo>∼</mo>
  <mi mathvariant="normal">N</mi><mfenced><mn>0</mn><mo>,</mo>
  <msubsup><mi>σ</mi><mtext>noise</mtext><mn>2</mn></msubsup></mfenced>
</math>
"""

DISCRIMINATOR_SCORE_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mi>d</mi><mi>i</mi><mi>E</mi></msubsup><mo>=</mo>
  <msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mover><mi>a</mi><mo>^</mo></mover><mi>i</mi><mi>E</mi></msubsup></mfenced>
  <mo>,</mo><mspace width="1em"/>
  <msubsup><mi>d</mi><mi>i</mi><mi>N</mi></msubsup><mo>=</mo>
  <msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mover><mi>a</mi><mo>~</mo></mover><mi>i</mi><mi>E</mi></msubsup></mfenced>
</math>
"""

JS_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi>M</mi><mo>=</mo><mfrac><mrow><mi>P</mi><mo>+</mo><mi>Q</mi></mrow><mn>2</mn></mfrac><mo>,</mo><mspace width="1em"/>
  <mi mathvariant="normal">JS</mi><mfenced open="(" close=")">
  <mi>P</mi><mo>∥</mo><mi>Q</mi></mfenced>
  <mo>=</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">KL</mi><mfenced><mi>P</mi><mo>∥</mo><mi>M</mi></mfenced>
  <mo>+</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">KL</mi><mfenced><mi>Q</mi><mo>∥</mo><mi>M</mi></mfenced>
</math>
"""

OVERLAP_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi mathvariant="normal">OVL</mi><mfenced><mi>P</mi><mo>,</mo><mi>Q</mi></mfenced><mo>=</mo>
  <munderover><mo>∑</mo><mrow><mi>h</mi><mo>=</mo><mn>1</mn></mrow><mi>H</mi></munderover>
  <mi mathvariant="normal">min</mi><mfenced>
  <msub><mi>P</mi><mi>h</mi></msub><mo>,</mo>
  <msub><mi>Q</mi><mi>h</mi></msub></mfenced>
</math>
"""

AUC_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi mathvariant="normal">AUC</mi><mo>=</mo>
  <mi mathvariant="normal">Pr</mi><mfenced>
  <msubsup><mi>x</mi><mi>i</mi><mo>+</mo></msubsup><mo>&gt;</mo>
  <msubsup><mi>x</mi><mi>j</mi><mo>−</mo></msubsup></mfenced>
  <mo>+</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">Pr</mi><mfenced>
  <msubsup><mi>x</mi><mi>i</mi><mo>+</mo></msubsup><mo>=</mo>
  <msubsup><mi>x</mi><mi>j</mi><mo>−</mo></msubsup></mfenced>
</math>
"""

ACTOR_SCORE_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mi>a</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup>
  <mo>=</mo><msub><mi>μ</mi><msub><mi>θ</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub></mfenced><mo>,</mo><mspace width="1em"/>
  <msubsup><mi>d</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup>
  <mo>=</mo><msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mi>a</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup></mfenced>
</math>
"""

ACTOR_JS_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mi>J</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub><mo>=</mo>
  <mi mathvariant="normal">JS</mi><mfenced>
  <msubsup><mi>P</mi><mi>r</mi><mi>E</mi></msubsup><mo>∥</mo>
  <msubsup><mi>P</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow><mi>G</mi></msubsup>
  </mfenced>
</math>
"""

ACTOR_MEAN_SEM_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mover><mi>J</mi><mo>¯</mo></mover><mi>k</mi></msub><mo>=</mo>
  <mfrac><mn>1</mn><mi>R</mi></mfrac>
  <munderover><mo>∑</mo><mrow><mi>r</mi><mo>=</mo><mn>1</mn></mrow><mi>R</mi></munderover>
  <msub><mi>J</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub><mo>,</mo><mspace width="1em"/>
  <msub><mi>SEM</mi><mi>k</mi></msub><mo>=</mo><mfrac><msub><mi>s</mi><mi>k</mi></msub><msqrt><mi>R</mi></msqrt></mfrac>
</math>
"""

EQUATIONS = (
    NOISY_ACTION_MATHML,
    DISCRIMINATOR_SCORE_MATHML,
    JS_MATHML,
    OVERLAP_MATHML,
    AUC_MATHML,
    ACTOR_SCORE_MATHML,
    ACTOR_JS_MATHML,
    ACTOR_MEAN_SEM_MATHML,
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def first_run_properties(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    source_rpr = first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def insert_paragraph_after(
    reference: Paragraph,
    text: str = "",
    *,
    format_source: Paragraph,
) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    if format_source._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(format_source._p.pPr))
    if text:
        run = paragraph.add_run(text)
        source_rpr = first_run_properties(format_source)
        if source_rpr is not None:
            run._r.insert(0, source_rpr)
    return paragraph


def mathml_to_omml(mathml: str):
    transform = etree.XSLT(etree.parse(str(MML2OMML_XSL)))
    result = transform(etree.fromstring(mathml.encode("utf-8")))
    return etree.fromstring(etree.tostring(result.getroot()))


def insert_equation_after(
    reference: Paragraph,
    mathml: str,
    *,
    equation_template: Paragraph,
) -> Paragraph:
    paragraph = insert_paragraph_after(
        reference, format_source=equation_template
    )
    paragraph._p.append(mathml_to_omml(mathml))
    return paragraph


def insert_picture_after(
    reference: Paragraph,
    image_path: Path,
    *,
    body_template: Paragraph,
    width_inches: float,
) -> Paragraph:
    paragraph = insert_paragraph_after(reference, format_source=body_template)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    return paragraph


def insert_results_table_after(
    document: Document,
    reference: Paragraph,
    rows: tuple[tuple[str, str], ...],
) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.style = document.tables[2].style
    table.autofit = False
    widths = (4.25, 2.25)
    header_rpr = first_run_properties(document.tables[2].cell(0, 0).paragraphs[0])
    body_rpr = first_run_properties(document.tables[2].cell(1, 0).paragraphs[0])

    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            source_rpr = header_rpr if row_index == 0 else body_rpr
            if source_rpr is not None:
                run._r.insert(0, deepcopy(source_rpr))

    reference._p.addnext(table._tbl)


def configure_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "font.serif": ["Times New Roman", "SimSun"],
            "axes.unicode_minus": False,
            "font.size": 9.5,
            "axes.labelsize": 10,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 7.7,
            "axes.linewidth": 1.0,
            "lines.linewidth": 1.7,
            "savefig.dpi": 300,
        }
    )


def style_axis(axis) -> None:
    axis.grid(True, linestyle="-.", linewidth=0.65, color="#a8a8a8", alpha=0.65)
    for spine in axis.spines.values():
        spine.set_color("black")
        spine.set_linewidth(1.0)


def load_learning_curve_module():
    spec = importlib.util.spec_from_file_location(
        "paper_v28_learning_curves", LEARNING_CURVE_SCRIPT
    )
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load plotting module: {LEARNING_CURVE_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def create_english_learning_curve_figures() -> None:
    plotting = load_learning_curve_module()
    plotting.METRIC_TITLES.update(
        {
            "survival": "(a) Survival Days",
            "production": "(b) Production Agent Return",
            "naulc": "(c) Cumulative nAULC",
        }
    )
    plotting.Y_LABELS["production"] = "Return"
    plotting.CHINESE_TITLE_FONT = plotting.FontProperties(
        family="Times New Roman", size=10.5
    )
    plotting.configure_matplotlib()
    raw_group_dirs = {
        "TD3": "TD3",
        "GAIL+TD3": "GAIL_plus_TD3",
        "Transformer+GAIL+TD3": "Transformer_plus_GAIL_plus_TD3",
    }
    runs_by_group = {}
    for group_name, directory_name in raw_group_dirs.items():
        csv_paths = sorted(
            (LEARNING_CURVE_RAW_DIR / directory_name).glob("*_metrics.csv")
        )
        if len(csv_paths) != 8:
            raise RuntimeError(
                f"Expected 8 cached runs for {group_name}, found {len(csv_paths)}."
            )
        runs_by_group[group_name] = [
            plotting.load_run(csv_path) for csv_path in csv_paths
        ]

    all_survival = [
        run["values"]["survival"]
        for runs in runs_by_group.values()
        for run in runs
    ]
    common_steps = sorted(
        set.intersection(*(set(values.keys()) for values in all_survival))
    )
    for runs in runs_by_group.values():
        for run in runs:
            plotting.add_cumulative_normalized_aulc(run, common_steps)
    data = {
        group_name: {
            metric_id: plotting.aggregate_sem(runs, metric_id)
            for metric_id in ("survival", "production", "naulc")
        }
        for group_name, runs in runs_by_group.items()
    }

    comparisons = (
        (("TD3", "GAIL+TD3"), TD3_GAIL_FIGURE),
        (("GAIL+TD3", "Transformer+GAIL+TD3"), GAIL_TRANSFORMER_FIGURE),
    )
    metric_order = ("survival", "production", "naulc")
    for groups, output_path in comparisons:
        figure, axes = plt.subplots(1, 3, figsize=(12.75, 4.0))
        for axis, metric_id in zip(axes, metric_order):
            plotting.plot_metric(axis, metric_id, groups, data)
        figure.subplots_adjust(
            left=0.055,
            right=0.992,
            bottom=0.19,
            top=0.90,
            wspace=0.34,
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        figure.savefig(output_path, facecolor="white", dpi=300)
        plt.close(figure)


def create_english_dscr_risk_figure() -> None:
    from matplotlib.ticker import MultipleLocator

    rows = pd.read_csv(DSCR_RISK_CSV, encoding="utf-8-sig")
    values_by_algorithm = {
        str(row["algorithm"]): float(row["dscr_below_1_share_percent_all"])
        for _, row in rows.iterrows()
    }
    order = ("TD3", "GAIL+TD3", "Transformer+GAIL+TD3")
    missing = [name for name in order if name not in values_by_algorithm]
    if missing:
        raise RuntimeError(f"Missing DSCR risk rows: {missing}")

    colors = ("#4C86E8", "#F08A5D", "#63B995")
    values = [values_by_algorithm[name] for name in order]
    figure, axis = plt.subplots(figsize=(6.6, 5.6))
    bars = axis.bar(
        range(len(order)),
        values,
        width=0.56,
        color=colors,
        edgecolor="#555555",
        linewidth=0.8,
        zorder=3,
    )
    upper = max(1.0, float(np.ceil(max(values) * 1.22)))
    axis.set_ylim(0, upper)
    axis.yaxis.set_major_locator(MultipleLocator(1))
    axis.set_xticks(range(len(order)))
    axis.set_xticklabels(order)
    axis.set_ylabel("Share (%)", labelpad=10)
    axis.grid(
        axis="y",
        linestyle="-.",
        linewidth=0.65,
        color="#A8A8A8",
        alpha=0.65,
        zorder=0,
    )
    for spine in axis.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.0)
    axis.tick_params(direction="out", width=1.0, length=4, colors="black")
    label_offset = upper * 0.025
    for bar, value in zip(bars, values):
        axis.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + label_offset,
            f"{value:.2f}%",
            ha="center",
            va="bottom",
            fontsize=9,
            color="#222222",
        )
    figure.subplots_adjust(left=0.13, right=0.98, bottom=0.14, top=0.98)
    DSCR_RISK_FIGURE.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(DSCR_RISK_FIGURE, facecolor="white", dpi=300)
    plt.close(figure)


def replace_picture_before_caption(
    document: Document, caption_prefix: str, image_path: Path
) -> None:
    caption = find_paragraph(document, caption_prefix)
    element = caption._p.getprevious()
    while element is not None:
        blips = element.xpath(".//a:blip")
        if blips:
            relationship_id = blips[0].get(qn("r:embed"))
            image_part = document.part.related_parts[relationship_id]
            image_part._blob = image_path.read_bytes()
            return
        element = element.getprevious()
    raise RuntimeError(f"No picture found before caption: {caption_prefix}")


def create_paper_figures() -> None:
    FIGURE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    configure_matplotlib()

    noisy_samples = pd.read_csv(
        NOISY_RESULT_DIR / "seed_184" / "evaluation_samples_and_scores.csv"
    )
    noisy_summary = pd.read_csv(
        NOISY_RESULT_DIR / "three_seed_expert_noise_summary.csv"
    )
    noisy_row = noisy_summary.loc[noisy_summary["seed"].astype(int) == 184].iloc[0]
    figure, axis = plt.subplots(figsize=(6.4, 4.05))
    axis.hist(
        noisy_samples["expert_score"],
        bins=50,
        range=(0.0, 1.0),
        alpha=0.52,
        color="#2ca02c",
        label="Expert",
    )
    axis.hist(
        noisy_samples["noisy_expert_score"],
        bins=50,
        range=(0.0, 1.0),
        alpha=0.42,
        color="#9467bd",
        label="Noisy expert",
    )
    axis.set_xlabel("Discriminator confidence score")
    axis.set_ylabel("Sample count")
    axis.set_xlim(0.0, 1.0)
    style_axis(axis)
    axis.legend(loc="upper left", frameon=True, framealpha=0.95)
    figure.tight_layout()
    figure.savefig(DISCRIMINATOR_FIGURE, bbox_inches="tight")
    plt.close(figure)

    actor_summary = pd.read_csv(ACTOR_RESULT_DIR / "actor_expert_js_mean_sem.csv")
    x = actor_summary["evaluation_step"].to_numpy(dtype=float)
    mean = actor_summary["mean"].to_numpy(dtype=float)
    lower = actor_summary["mean_minus_sem"].to_numpy(dtype=float)
    upper = actor_summary["mean_plus_sem"].to_numpy(dtype=float)
    figure, axis = plt.subplots(figsize=(6.4, 4.05))
    axis.fill_between(x, lower, upper, color="#4c78a8", alpha=0.22, linewidth=0)
    axis.plot(x, mean, color="#1f4e79", linewidth=1.8)
    axis.set_xlabel("Evaluation step (100 episodes)")
    axis.set_ylabel("Actor-expert JS divergence")
    axis.set_xlim(1, 60)
    axis.set_ylim(0.0, max(0.7, float(upper.max()) * 1.03))
    style_axis(axis)
    figure.tight_layout()
    figure.savefig(ACTOR_FIGURE, bbox_inches="tight")
    plt.close(figure)


def renumber_existing_section(document: Document) -> None:
    replacements = (
        ("6.3 GAIL+TD3与Transformer+GAIL+TD3的结果对比", "6.4 GAIL+TD3与Transformer+GAIL+TD3的结果对比"),
        ("基于上一节的结果，本文进一步比较", "在完成生成对抗模仿学习组件有效性检验后，本文进一步比较"),
        ("图6-2给出了两种模型", "图6-4给出了两种模型"),
        ("图6-2(a)中的水平虚线", "图6-4(a)中的水平虚线"),
        ("图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比", "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比"),
        ("从图6-2(a)可以看出", "从图6-4(a)可以看出"),
        ("图6-2(b)显示", "图6-4(b)显示"),
        ("图6-2(c)显示", "图6-4(c)显示"),
        ("图6-3给出了三种方法", "图6-5给出了三种方法"),
        ("图6-3 三种主体训练方案下", "图6-5 三种主体训练方案下"),
        ("图6-3显示", "图6-5显示"),
        ("综合图6-2、图6-3和表6-2", "综合图6-4、图6-5和表6-2"),
    )
    for paragraph in document.paragraphs:
        revised = paragraph.text
        for old, new in replacements:
            revised = revised.replace(old, new)
        if revised != paragraph.text:
            replace_paragraph_text(paragraph, revised)


def insert_metric_definition(document: Document) -> None:
    body_template = find_paragraph(document, "（5）生产企业偿债能力")
    dscr_explanation = find_paragraph(document, "式中，分子表示生产企业第t天")
    equation_template = Paragraph(
        dscr_explanation._p.getprevious(), dscr_explanation._parent
    )
    anchor = dscr_explanation

    paragraphs_and_equations = (
        (
            "text",
            "（6）生成对抗模仿学习组件有效性。本文采用Jensen–Shannon散度、重叠系数和受试者工作特征曲线下面积评价生成对抗模仿学习组件。其中，Jensen–Shannon散度（JS divergence）用于度量两组判别器得分概率分布之间的差异。设P和Q表示两组归一化判别器得分分布，其定义如下：",
        ),
        ("equation", JS_MATHML),
        (
            "text",
            "式中，M表示P和Q的等权混合分布，KL表示Kullback–Leibler散度。JS散度越小表示两组判别器得分分布越接近。重叠系数（OVL: Overlap Coefficient）用于度量两组归一化直方图的共同区域，其定义如下：",
        ),
        ("equation", OVERLAP_MATHML),
        (
            "text",
            "式中，h表示得分区间编号，H表示区间总数，P_h和Q_h分别表示两组样本落入第h个区间的概率。OVL取值范围为[0,1]，数值越大表示两组得分分布的重叠程度越高。受试者工作特征曲线下面积（AUC: Area Under the Receiver Operating Characteristic Curve）用于评价判别器对两类样本的排序能力，其定义如下：",
        ),
        ("equation", AUC_MATHML),
        (
            "text",
            "式中，Pr表示事件发生的概率，x_i^+和x_j^-分别表示从正类与负类样本中抽取的判别器得分。AUC=0.5对应随机排序，AUC大于0.5表示判别器更倾向于给予正类样本较高评分。",
        ),
    )
    for kind, content in paragraphs_and_equations:
        if kind == "text":
            anchor = insert_paragraph_after(
                anchor, content, format_source=body_template
            )
        else:
            anchor = insert_equation_after(
                anchor, content, equation_template=equation_template
            )


def insert_component_validation_section(document: Document) -> None:
    old_section_heading = find_paragraph(
        document, "6.4 GAIL+TD3与Transformer+GAIL+TD3的结果对比"
    )
    heading_template = old_section_heading
    body_template = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    caption_template = find_paragraph(document, "图6-1 TD3与GAIL+TD3训练结果对比")
    table_caption_template = find_paragraph(document, "表6-2 三种主体训练方案的最终性能统计")
    dscr_explanation = find_paragraph(document, "式中，分子表示生产企业第t天")
    equation_template = Paragraph(
        dscr_explanation._p.getprevious(), dscr_explanation._parent
    )

    anchor = Paragraph(old_section_heading._p.getprevious(), old_section_heading._parent)
    heading = insert_paragraph_after(
        anchor,
        "6.3 生成对抗模仿学习组件有效性分析",
        format_source=heading_template,
    )
    anchor = insert_paragraph_after(
        heading,
        "第6.2节的训练结果表明，引入生成对抗模仿学习后，生产企业主体能够以较少的训练回合达到较高的平均存活水平。然而，平均存活天数属于系统运行结果层面的综合指标，同时受到商品交易、信贷关系及其他主体行为等多种因素影响，难以单独刻画生成对抗模仿学习组件的运行状态；若直接逐一计算Actor生成动作与专家动作之间的误差，所得结果又容易受到多维连续动作各分量尺度及样本差异的影响，难以反映两类行为在整体分布上的接近程度。基于此，本文利用固定判别器将不同来源的状态—动作样本映射至统一的得分空间：首先通过JS散度、重叠系数和AUC评价最终判别器对动作扰动的识别能力，随后根据历史Actor与专家样本得分分布之间的JS散度变化，分析Actor生成行为在训练过程中的模仿效果。",
        format_source=body_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "第一步检验最终判别器对动作扰动的识别能力。本文选取一组独立运行对应的最终判别器作为评价尺度，在留出测试集的16952条专家状态—动作样本上叠加标准差为0.30的独立高斯噪声，并将动作裁剪至[-0.5,0.5]。留出测试集包含173个完整专家回合。加噪专家动作定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, NOISY_ACTION_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，i表示留出测试集中的样本序号；hat{a}_i^E表示第i个专家动作；tilde{a}_i^E表示加噪并裁剪后的第i个专家动作；ε_i表示第i个样本对应的零均值独立高斯噪声；σ_noise表示原始动作单位下的噪声标准差；b表示动作边界，本文取b=0.5。固定最终判别器对原始专家样本和加噪专家样本的评分分别定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, DISCRIMINATOR_SCORE_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，D表示判别器；hat{s}_i表示第i个专家状态；带星号的φ_r表示第r次独立运行对应的最终判别器参数；d_i^E和d_i^N分别表示第i个专家样本与加噪专家样本的判别器得分。图6-2给出了两组样本的判别器得分分布。",
        format_source=body_template,
    )
    anchor = insert_picture_after(
        anchor,
        DISCRIMINATOR_FIGURE,
        body_template=body_template,
        width_inches=6.35,
    )
    figure_caption = insert_paragraph_after(
        anchor,
        "图6-2 专家样本与加噪专家样本的判别器得分分布",
        format_source=caption_template,
    )
    analysis = insert_paragraph_after(
        figure_caption,
        "图6-2与表6-3显示，最终判别器对专家样本与加噪专家样本的平均评分分别为0.8988和0.6111，平均评分差为0.2877。两组得分分布的JS散度为0.4100；由于JS散度为0时两组分布完全一致，因此该结果反映出专家样本与加噪专家样本的判别器得分分布存在差异。OVL为18.85%，表示两组归一化得分直方图的共同区域占18.85%，两组分布的重叠程度较低。以专家样本为正类的AUC为0.8186，表示随机抽取一个专家样本和一个加噪专家样本时，判别器给予专家样本更高评分的概率约为81.86%，高于随机排序对应的0.5。上述结果表明，在标准差为0.30的动作扰动条件下，最终判别器能够识别加噪动作相对于专家动作的偏离，并倾向于给予专家样本更高的评分。",
        format_source=body_template,
    )
    table_caption = insert_paragraph_after(
        figure_caption,
        "表6-3 判别器对专家样本与加噪专家样本的评价结果",
        format_source=table_caption_template,
    )
    insert_results_table_after(
        document,
        table_caption,
        (
            ("统计量", "估计值"),
            ("专家样本平均评分", "0.8988"),
            ("加噪专家样本平均评分", "0.6111"),
            ("平均评分差", "0.2877"),
            ("JS散度", "0.4100"),
            ("OVL", "18.85%"),
            ("AUC", "0.8186"),
        ),
    )
    anchor = analysis
    anchor = insert_paragraph_after(
        anchor,
        "第二步检验历史Actor生成行为在训练过程中的变化。本文选取3次独立运行，分别以各次运行对应的最终判别器作为固定评价尺度，并加载每100个训练回合保存的历史Actor快照。所有历史Actor均在同一批留出专家状态上以确定性方式生成动作，不加入探索噪声。对于第r次独立运行的第k个评估时刻，历史Actor生成的动作及其判别器得分定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, ACTOR_SCORE_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，μ表示生产企业Actor；θ_(r,k)表示第r次独立运行在第k个评估时刻保存的Actor参数；上标G表示Actor生成动作或相应判别器得分。根据专家得分分布P_r^E和历史Actor得分分布P_(r,k)^G，Actor—专家JS散度定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, ACTOR_JS_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，J_(r,k)表示第r次独立运行在第k个评估时刻的Actor—专家JS散度。较小的J_(r,k)表示历史Actor生成行为在固定判别器得分表示下更接近专家样本。本文在每个评估时刻计算3次独立运行的JS散度均值及其均值标准误：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, ACTOR_MEAN_SEM_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，带横线的J表示第k个评估时刻3次独立运行的JS散度均值；s_k表示相应JS散度的样本标准差；R表示独立运行次数，本文取R=3。图6-3给出了60个评估时刻的Actor—专家JS散度均值，阴影表示均值上下1个SEM。",
        format_source=body_template,
    )
    anchor = insert_picture_after(
        anchor,
        ACTOR_FIGURE,
        body_template=body_template,
        width_inches=6.35,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-3 Actor与专家样本判别器得分分布的JS散度变化",
        format_source=caption_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-3显示，Actor—专家JS散度均值在第1个评估时刻为0.580±0.041，训练过程中总体下降，并在第19个评估时刻达到0.172±0.033。第60个评估时刻的均值为0.258±0.017，较第1个评估时刻降低约55.5%，且训练后期整体维持在低于训练初期的水平。固定最终判别器下的得分分布结果表明，历史Actor生成行为在训练过程中总体向专家样本接近。",
        format_source=body_template,
    )
    insert_paragraph_after(
        anchor,
        "两步评价分别从判别器对扰动行为的识别能力和Actor生成行为的训练演化两个方面检验了生成对抗模仿学习组件。最终判别器能够识别具有明确动作扰动的样本，历史Actor相对于训练初期则形成了更接近专家得分分布的生成行为。上述结果为第6.2节中模仿学习信号改善生产企业主体前期训练表现提供了组件层面的实验证据。",
        format_source=body_template,
    )


def update_transitions_and_metric_count(document: Document) -> None:
    replace_paragraph_text(
        find_paragraph(document, "在上述训练过程中，本文跟踪"),
        "在上述训练过程中，本文跟踪平均存活天数、生产企业累计收益、累计归一化学习曲线下面积、学习速度、生产企业偿债能力和生成对抗模仿学习组件有效性六项评价指标。各项指标的定义如下。",
    )
    replace_paragraph_text(
        find_paragraph(document, "不过，GAIL+TD3的最终稳定平均存活天数为80.71天"),
        "不过，GAIL+TD3的最终稳定平均存活天数为80.71天，与专家轨迹平均值仍相差17.28天。该结果说明，模仿学习信号缩短了模型达到80天所需的训练过程，但稳定阶段的存活水平仍有进一步接近专家轨迹的空间。下一节首先从判别器与Actor生成器两个方面检验生成对抗模仿学习组件的有效性，随后在第6.4节引入历史状态序列表征，以考察模型利用连续交互信息后能否进一步缩小这一差距。",
    )


def verify(document: Document, source_equation_count: int) -> None:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    text = "\n".join(paragraphs)
    required = (
        "（6）生成对抗模仿学习组件有效性",
        "6.3 生成对抗模仿学习组件有效性分析",
        "图6-2 专家样本与加噪专家样本的判别器得分分布",
        "图6-3 Actor与专家样本判别器得分分布的JS散度变化",
        "6.4 GAIL+TD3与Transformer+GAIL+TD3的结果对比",
        "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        "图6-5 三种主体训练方案下DSCR 低于 1 的生产企业日观测占比",
        "表6-3 判别器对专家样本与加噪专家样本的评价结果",
        "第60个评估时刻的均值为0.258±0.017",
        "OVL为18.85%",
    )
    for value in required:
        if value not in text:
            raise RuntimeError(f"Required content is missing: {value}")
    forbidden = (
        "6.3 GAIL+TD3与Transformer+GAIL+TD3的结果对比",
        "图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        "图6-3 三种主体训练方案下DSCR 低于 1",
        "并非单调下降",
        "随机种子184",
        "随机种子652",
        "随机种子187",
        "p=0.0002",
        "95%区间[0.4012,0.4191]",
        "95%区间[0.8073,0.8293]",
    )
    for value in forbidden:
        if value in text:
            raise RuntimeError(f"Former numbering remains: {value}")
    if len(document.tables) != 5:
        raise RuntimeError(f"Expected 5 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 11:
        raise RuntimeError(f"Expected 11 inline figures, found {len(document.inline_shapes)}.")
    equation_count = len(document.element.body.xpath(".//m:oMath"))
    if equation_count != source_equation_count + len(EQUATIONS):
        raise RuntimeError(
            f"Expected {source_equation_count + len(EQUATIONS)} equations, "
            f"found {equation_count}."
        )


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The source DOCX changed after this builder was prepared.")
    if not MML2OMML_XSL.exists():
        raise FileNotFoundError(MML2OMML_XSL)
    create_paper_figures()
    create_english_learning_curve_figures()
    create_english_dscr_risk_figure()
    if not DISCRIMINATOR_FIGURE.exists() or not ACTOR_FIGURE.exists():
        raise RuntimeError("Publication figures were not generated.")
    for figure_path in (
        TD3_GAIL_FIGURE,
        GAIL_TRANSFORMER_FIGURE,
        DSCR_RISK_FIGURE,
    ):
        if not figure_path.exists():
            raise RuntimeError(f"English publication figure was not generated: {figure_path}")

    document = Document(SOURCE)
    source_equation_count = len(document.element.body.xpath(".//m:oMath"))
    renumber_existing_section(document)
    update_transitions_and_metric_count(document)
    insert_metric_definition(document)
    insert_component_validation_section(document)
    replace_picture_before_caption(
        document,
        "图6-1 TD3与GAIL+TD3训练结果对比",
        TD3_GAIL_FIGURE,
    )
    replace_picture_before_caption(
        document,
        "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        GAIL_TRANSFORMER_FIGURE,
    )
    replace_picture_before_caption(
        document,
        "图6-5 三种主体训练方案下DSCR 低于 1 的生产企业日观测占比",
        DSCR_RISK_FIGURE,
    )
    verify(document, source_equation_count)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The source DOCX was modified unexpectedly.")
    verify(Document(OUTPUT), source_equation_count)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
