from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT / "paper_drafts" / "rendered_v6"
OUT_DIR = ROOT / "paper_drafts" / "rendered_v7"
REF_DIR = ROOT / "paper_drafts" / "reference"


def find_docx():
    files = list(SRC_DIR.glob("*.docx"))
    if not files:
        raise FileNotFoundError(f"No docx found in {SRC_DIR}")
    return files[0]


def find_image(name):
    path = REF_DIR / name
    if not path.exists():
        raise FileNotFoundError(f"Missing image: {path}")
    return path


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_paragraph_text(paragraph, text, indent=None):
    clear_paragraph(paragraph)
    paragraph.add_run(text)
    if indent is not None:
        paragraph.paragraph_format.first_line_indent = indent
    paragraph.alignment = None


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def clone_paragraph_format(ppr, dst):
    if dst._p.pPr is not None:
        dst._p.remove(dst._p.pPr)
    if ppr is not None:
        dst._p.insert(0, deepcopy(ppr))


def add_body_after(anchor, text, style_name, ppr, indent):
    p = insert_paragraph_after(anchor, text, style_name)
    clone_paragraph_format(ppr, p)
    p.paragraph_format.first_line_indent = indent
    p.alignment = None
    return p


def add_picture_after(anchor, image_path, width_inches, caption, caption_style_name, caption_ppr):
    pic_p = insert_paragraph_after(anchor, style=caption_style_name)
    clone_paragraph_format(caption_ppr, pic_p)
    pic_p.paragraph_format.first_line_indent = None
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(image_path), width=Inches(width_inches))

    cap_p = insert_paragraph_after(pic_p, caption, caption_style_name)
    clone_paragraph_format(caption_ppr, cap_p)
    cap_p.paragraph_format.first_line_indent = None
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return cap_p


def main():
    src = find_docx()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / src.name.replace("v6", "v7")
    shutil.copy2(src, out)

    img1 = find_image("人类数据采集1.png")
    img2 = find_image("人类数据采集2.png")
    img3 = find_image("人类数据采集3.png")

    doc = Document(out)

    paragraphs = doc.paragraphs
    idx_52 = next(i for i, p in enumerate(paragraphs) if "".join(r.text for r in p.runs).strip().startswith("5.2 "))
    idx_53 = next(i for i, p in enumerate(paragraphs) if "".join(r.text for r in p.runs).strip().startswith("5.3 "))

    body_template = paragraphs[idx_52 + 1]
    caption_template = paragraphs[idx_52 + 4]
    body_style_name = body_template.style.name
    caption_style_name = caption_template.style.name
    body_ppr = deepcopy(body_template._p.pPr)
    caption_ppr = deepcopy(caption_template._p.pPr)
    indent = body_template.paragraph_format.first_line_indent

    # Remove the old 5.2 body and placeholder figure paragraphs.
    for p in list(paragraphs[idx_52 + 1:idx_53]):
        delete_paragraph(p)

    heading_52 = doc.paragraphs[idx_52]

    p1 = add_body_after(
        heading_52,
        "网页实验系统采用“任务说明—每日决策—规则明细”的页面组织方式。任务描述页如图5-1所示。页面首先说明仿真系统由甲公司、乙公司、银行和第三方市场构成，并明确参与者扮演甲公司经理。页面将参与者的操作概括为查看今天状态、做出经营决策、系统自动运行并进入下一天三个步骤。该设计把参与者需要完成的决策任务与系统自动完成的交易、生产和清算过程区分开。页面还记录参与者编号、年龄段、受教育程度、性别、经济或管理相关背景、经营或策略类游戏经验，并通过访问口令控制实验入口。",
        body_style_name,
        body_ppr,
        indent,
    )
    cap1 = add_picture_after(p1, img1, 6.2, "图5-1 网页实验系统任务描述页", caption_style_name, caption_ppr)

    p2 = add_body_after(
        cap1,
        "每日经营决策页如图5-2所示。页面左侧提供当天可见信息，包括贷款信息、上一天两家公司买到的原料数量、定价与销售信息、经营小结和净利润折线图。页面右侧集中放置四个动作输入项，分别为原料A采购需求、原料B采购需求、申请贷款金额和产品A销售价格。系统在输入框下方提供最小值、默认值、最大值以及小幅增减按钮，以降低连续多天填写时的操作负担。页面下方展示当前进度、自动跳转条件和重新决策说明。",
        body_style_name,
        body_ppr,
        indent,
    )
    cap2 = add_picture_after(p2, img2, 6.3, "图5-2 网页实验系统每日经营决策页", caption_style_name, caption_ppr)

    p3 = add_body_after(
        cap2,
        "“更多规则和成交明细”的展开内容如图5-3所示。该部分补充展示产品制作与市场流转示意图、市场购买规则、仿真环境六个阶段和上一天成交明细。示意图将甲公司和乙公司的生产循环并列呈现：甲公司从普通市场和第三方市场购买原料A与原料B，并使用二者生产产品A；产品A在下一天进入普通市场，成为原料A。乙公司以同样方式生产产品B，产品B在下一天进入普通市场，成为原料B。图中还给出产品A产量计算关系，即“产品A产量 = 2.5 × min(买到的原料A, 买到的原料B)”，并说明原料当天购买、当天投入生产、不跨天保存。第三方市场在该系统中承担补充供给功能，可理解为政府宏观调控部门。",
        body_style_name,
        body_ppr,
        indent,
    )
    cap3 = add_picture_after(p3, img3, 5.3, "图5-3 每日经营决策页中的规则和成交明细", caption_style_name, caption_ppr)

    add_body_after(
        cap3,
        "网页端对四个动作输入值设置上下界，避免转换后的连续动作超过模型训练取值范围。系统将业务含义的输入值转换为连续动作向量，并保存为专家样本。页面还提供阶段跳转功能。参与者连续完成 8 天人工决策后，系统最多自动推进 15 天，并在现金、债务、库存、采购参考量或销售价格出现明显变化时重新交回人工决策。",
        body_style_name,
        body_ppr,
        indent,
    )

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
