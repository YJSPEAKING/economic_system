from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

BUNDLED_SITE_PACKAGES = Path(
    r"C:\Users\legion\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\python\Lib\site-packages"
)
if str(BUNDLED_SITE_PACKAGES) not in sys.path:
    sys.path.append(str(BUNDLED_SITE_PACKAGES))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "reference"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_v1.docx"
)
OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT = OUTPUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v26.docx"
EXPECTED_SOURCE_SHA256 = (
    "753659DBA9394BA37EC057DF1C61A53450F779258635D8E2E30C234253B92421"
)
MML2OMML_XSL = Path(
    r"C:\Program Files\Microsoft Office\Root\Office16\MML2OMML.XSL"
)
FIGURE_OUTPUT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "paper_gail_component_validation"
)
NOISY_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "expert_vs_noisy_expert_discriminator_eval_noise_std_0p30"
)
ACTOR_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "actor_snapshot_evolution"
)
DISCRIMINATOR_FIGURE = FIGURE_OUTPUT_DIR / "expert_vs_noisy_expert_seed184_paper.png"
ACTOR_FIGURE = FIGURE_OUTPUT_DIR / "actor_expert_js_evolution_mean_sem_paper.png"


NOISY_ACTION_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mover><mi>a</mi><mo>~</mo></mover><mi>i</mi><mi>E</mi></msubsup>
  <mo>=</mo><mi mathvariant="normal">clip</mi><mfenced>
  <msubsup><mover><mi>a</mi><mo>^</mo></mover><mi>i</mi><mi>E</mi></msubsup>
  <mo>+</mo><msub><mi>ε</mi><mi>i</mi></msub><mo>,</mo><mo>−</mo><mi>b</mi>
  <mo>,</mo><mi>b</mi></mfenced><mo>,</mo>
  <msub><mi>ε</mi><mi>i</mi></msub><mo>∼</mo>
  <mi mathvariant="normal">N</mi><mfenced><mn>0</mn><mo>,</mo>
  <msubsup><mi>σ</mi><mtext>noise</mtext><mn>2</mn></msubsup></mfenced>
</math>
"""

DISCRIMINATOR_SCORE_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mi>d</mi><mi>i</mi><mi>E</mi></msubsup><mo>=</mo>
  <msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mover><mi>a</mi><mo>^</mo></mover><mi>i</mi><mi>E</mi></msubsup></mfenced>
  <mo>,</mo><mspace width="1em"/>
  <msubsup><mi>d</mi><mi>i</mi><mi>N</mi></msubsup><mo>=</mo>
  <msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mover><mi>a</mi><mo>~</mo></mover><mi>i</mi><mi>E</mi></msubsup></mfenced>
</math>
"""

JS_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi>M</mi><mo>=</mo><mfrac><mrow><msup><mi>P</mi><mi>E</mi></msup><mo>+</mo>
  <msup><mi>P</mi><mi>N</mi></msup></mrow><mn>2</mn></mfrac><mo>,</mo><mspace width="1em"/>
  <mi mathvariant="normal">JS</mi><mfenced open="(" close=")">
  <msup><mi>P</mi><mi>E</mi></msup><mo>∥</mo><msup><mi>P</mi><mi>N</mi></msup></mfenced>
  <mo>=</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">KL</mi><mfenced><msup><mi>P</mi><mi>E</mi></msup><mo>∥</mo><mi>M</mi></mfenced>
  <mo>+</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">KL</mi><mfenced><msup><mi>P</mi><mi>N</mi></msup><mo>∥</mo><mi>M</mi></mfenced>
</math>
"""

OVERLAP_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi mathvariant="normal">OVL</mi><mfenced><msup><mi>P</mi><mi>E</mi></msup><mo>,</mo>
  <msup><mi>P</mi><mi>N</mi></msup></mfenced><mo>=</mo>
  <munderover><mo>∑</mo><mrow><mi>h</mi><mo>=</mo><mn>1</mn></mrow><mi>H</mi></munderover>
  <mi mathvariant="normal">min</mi><mfenced>
  <msubsup><mi>P</mi><mi>h</mi><mi>E</mi></msubsup><mo>,</mo>
  <msubsup><mi>P</mi><mi>h</mi><mi>N</mi></msubsup></mfenced>
</math>
"""

AUC_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi mathvariant="normal">AUC</mi><mo>=</mo>
  <mi mathvariant="normal">Pr</mi><mfenced>
  <msubsup><mi>d</mi><mi>i</mi><mi>E</mi></msubsup><mo>&gt;</mo>
  <msubsup><mi>d</mi><mi>j</mi><mi>N</mi></msubsup></mfenced>
  <mo>+</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">Pr</mi><mfenced>
  <msubsup><mi>d</mi><mi>i</mi><mi>E</mi></msubsup><mo>=</mo>
  <msubsup><mi>d</mi><mi>j</mi><mi>N</mi></msubsup></mfenced>
</math>
"""

ACTOR_SCORE_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mi>a</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup>
  <mo>=</mo><msub><mi>μ</mi><msub><mi>θ</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub></mfenced><mo>,</mo><mspace width="1em"/>
  <msubsup><mi>d</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup>
  <mo>=</mo><msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mi>a</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup></mfenced>
</math>
"""

ACTOR_JS_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mi>J</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub><mo>=</mo>
  <mi mathvariant="normal">JS</mi><mfenced>
  <msubsup><mi>P</mi><mi>r</mi><mi>E</mi></msubsup><mo>∥</mo>
  <msubsup><mi>P</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow><mi>G</mi></msubsup>
  </mfenced>
</math>
"""

ACTOR_MEAN_SEM_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mover><mi>J</mi><mo>¯</mo></mover><mi>k</mi></msub><mo>=</mo>
  <mfrac><mn>1</mn><mi>R</mi></mfrac>
  <munderover><mo>∑</mo><mrow><mi>r</mi><mo>=</mo><mn>1</mn></mrow><mi>R</mi></munderover>
  <msub><mi>J</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub><mo>,</mo><mspace width="1em"/>
  <msub><mi>SEM</mi><mi>k</mi></msub><mo>=</mo><mfrac><msub><mi>s</mi><mi>k</mi></msub><msqrt><mi>R</mi></msqrt></mfrac>
</math>
"""

EQUATIONS = (
    NOISY_ACTION_MATHML,
    DISCRIMINATOR_SCORE_MATHML,
    JS_MATHML,
    OVERLAP_MATHML,
    AUC_MATHML,
    ACTOR_SCORE_MATHML,
    ACTOR_JS_MATHML,
    ACTOR_MEAN_SEM_MATHML,
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def first_run_properties(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    source_rpr = first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def insert_paragraph_after(
    reference: Paragraph,
    text: str = "",
    *,
    format_source: Paragraph,
) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    if format_source._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(format_source._p.pPr))
    if text:
        run = paragraph.add_run(text)
        source_rpr = first_run_properties(format_source)
        if source_rpr is not None:
            run._r.insert(0, source_rpr)
    return paragraph


def mathml_to_omml(mathml: str):
    transform = etree.XSLT(etree.parse(str(MML2OMML_XSL)))
    result = transform(etree.fromstring(mathml.encode("utf-8")))
    return etree.fromstring(etree.tostring(result.getroot()))


def insert_equation_after(
    reference: Paragraph,
    mathml: str,
    *,
    equation_template: Paragraph,
) -> Paragraph:
    paragraph = insert_paragraph_after(
        reference, format_source=equation_template
    )
    paragraph._p.append(mathml_to_omml(mathml))
    return paragraph


def insert_picture_after(
    reference: Paragraph,
    image_path: Path,
    *,
    body_template: Paragraph,
    width_inches: float,
) -> Paragraph:
    paragraph = insert_paragraph_after(reference, format_source=body_template)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    return paragraph


def configure_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "font.serif": ["Times New Roman", "SimSun"],
            "axes.unicode_minus": False,
            "font.size": 9.5,
            "axes.labelsize": 10,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 7.7,
            "axes.linewidth": 1.0,
            "lines.linewidth": 1.7,
            "savefig.dpi": 300,
        }
    )


def style_axis(axis) -> None:
    axis.grid(True, linestyle="-.", linewidth=0.65, color="#a8a8a8", alpha=0.65)
    for spine in axis.spines.values():
        spine.set_color("black")
        spine.set_linewidth(1.0)


def create_paper_figures() -> None:
    FIGURE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    configure_matplotlib()

    noisy_samples = pd.read_csv(
        NOISY_RESULT_DIR / "seed_184" / "evaluation_samples_and_scores.csv"
    )
    noisy_summary = pd.read_csv(
        NOISY_RESULT_DIR / "three_seed_expert_noise_summary.csv"
    )
    noisy_row = noisy_summary.loc[noisy_summary["seed"].astype(int) == 184].iloc[0]
    figure, axis = plt.subplots(figsize=(6.4, 4.05))
    axis.hist(
        noisy_samples["expert_score"],
        bins=50,
        range=(0.0, 1.0),
        alpha=0.52,
        color="#2ca02c",
        label=f"Expert (Mean={noisy_row['expert_mean']:.3f})",
    )
    axis.hist(
        noisy_samples["noisy_expert_score"],
        bins=50,
        range=(0.0, 1.0),
        alpha=0.42,
        color="#9467bd",
        label=(
            f"Noisy expert (Mean={noisy_row['noisy_expert_mean']:.3f}, "
            f"JS={noisy_row['js_divergence']:.3f}, "
            f"Overlap={noisy_row['histogram_overlap_coefficient']:.3f}, "
            f"AUC={noisy_row['roc_auc']:.3f})"
        ),
    )
    axis.set_xlabel("Discriminator confidence score")
    axis.set_ylabel("Sample count")
    axis.set_xlim(0.0, 1.0)
    style_axis(axis)
    axis.legend(loc="upper left", frameon=True, framealpha=0.95)
    figure.tight_layout()
    figure.savefig(DISCRIMINATOR_FIGURE, bbox_inches="tight")
    plt.close(figure)

    actor_summary = pd.read_csv(ACTOR_RESULT_DIR / "actor_expert_js_mean_sem.csv")
    x = actor_summary["evaluation_step"].to_numpy(dtype=float)
    mean = actor_summary["mean"].to_numpy(dtype=float)
    lower = actor_summary["mean_minus_sem"].to_numpy(dtype=float)
    upper = actor_summary["mean_plus_sem"].to_numpy(dtype=float)
    figure, axis = plt.subplots(figsize=(6.4, 4.05))
    axis.fill_between(x, lower, upper, color="#4c78a8", alpha=0.22, linewidth=0)
    axis.plot(x, mean, color="#1f4e79", linewidth=1.8)
    axis.set_xlabel("Evaluation step (100 episodes)")
    axis.set_ylabel("Actor-expert JS divergence")
    axis.set_xlim(1, 60)
    axis.set_ylim(0.0, max(0.7, float(upper.max()) * 1.03))
    style_axis(axis)
    figure.tight_layout()
    figure.savefig(ACTOR_FIGURE, bbox_inches="tight")
    plt.close(figure)


def renumber_existing_section(document: Document) -> None:
    replacements = (
        ("6.3 GAIL+TD3与Transformer+GAIL+TD3的结果对比", "6.4 GAIL+TD3与Transformer+GAIL+TD3的结果对比"),
        ("基于上一节的结果，本文进一步比较", "在完成生成对抗模仿学习组件有效性检验后，本文进一步比较"),
        ("图6-2给出了两种模型", "图6-4给出了两种模型"),
        ("图6-2(a)中的水平虚线", "图6-4(a)中的水平虚线"),
        ("图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比", "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比"),
        ("从图6-2(a)可以看出", "从图6-4(a)可以看出"),
        ("图6-2(b)显示", "图6-4(b)显示"),
        ("图6-2(c)显示", "图6-4(c)显示"),
        ("图6-3给出了三种方法", "图6-5给出了三种方法"),
        ("图6-3 三种主体训练方案下", "图6-5 三种主体训练方案下"),
        ("图6-3显示", "图6-5显示"),
        ("综合图6-2、图6-3和表6-2", "综合图6-4、图6-5和表6-2"),
    )
    for paragraph in document.paragraphs:
        revised = paragraph.text
        for old, new in replacements:
            revised = revised.replace(old, new)
        if revised != paragraph.text:
            replace_paragraph_text(paragraph, revised)


def insert_metric_definition(document: Document) -> None:
    body_template = find_paragraph(document, "（5）生产企业偿债能力")
    dscr_explanation = find_paragraph(document, "式中，分子表示生产企业第t天")
    equation_template = Paragraph(
        dscr_explanation._p.getprevious(), dscr_explanation._parent
    )
    anchor = dscr_explanation

    paragraphs_and_equations = (
        (
            "text",
            "（6）生成对抗模仿学习组件有效性。本文分别从最终判别器对扰动专家动作的识别能力和历史Actor生成行为在训练过程中的变化两个方面评价生成对抗模仿学习组件。评价样本来自按完整回合划分的留出测试集，评价期间固定判别器与Actor参数，不进行网络更新。",
        ),
        (
            "text",
            "对于第r次独立运行，本文在留出专家状态—动作样本的动作分量上叠加独立高斯噪声，并将扰动后的动作限制在合法动作区间内。加噪专家动作定义如下：",
        ),
        ("equation", NOISY_ACTION_MATHML),
        (
            "text",
            "式中，i表示留出测试集中的样本序号；带尖号的s和带尖号的a分别表示记录的专家状态与专家动作；带波浪号的a表示加噪并裁剪后的专家动作；ε表示零均值独立高斯噪声；σ_noise表示原始动作单位下的噪声标准差；b表示动作边界，本文取b=0.5。固定最终判别器对原始专家样本和加噪专家样本的评分分别定义如下：",
        ),
        ("equation", DISCRIMINATOR_SCORE_MATHML),
        (
            "text",
            "式中，D表示判别器，带星号的φ表示第r次独立运行中固定的最终判别器参数；d的上标E和N分别表示专家样本得分和加噪专家样本得分。本文将得分区间[0,1]划分为H=100个等宽区间，并将两组得分直方图归一化为概率分布P^E和P^N。两组分布之间的Jensen–Shannon散度（JS divergence）定义如下：",
        ),
        ("equation", JS_MATHML),
        (
            "text",
            "式中，M表示两组概率分布的等权混合分布，KL表示Kullback–Leibler散度。JS散度越大表示两组判别器得分分布之间的差异越大。本文同时计算归一化直方图的重叠系数：",
        ),
        ("equation", OVERLAP_MATHML),
        (
            "text",
            "式中，h表示得分区间编号；H表示区间总数；P_h^E和P_h^N分别表示两组样本落入第h个区间的概率。重叠系数取值范围为[0,1]，较小的数值表示两组得分分布的重叠程度较低。判别器的排序能力采用以专家样本为正类的受试者工作特征曲线下面积（AUC: Area Under the Receiver Operating Characteristic Curve）衡量：",
        ),
        ("equation", AUC_MATHML),
        (
            "text",
            "式中，Pr表示事件发生的概率；i和j分别表示从专家得分分布与加噪专家得分分布中抽取的样本序号。AUC=0.5对应随机排序，AUC大于0.5表示判别器更倾向于给予专家样本较高评分。本文以完整专家回合为统计单位，采用单侧配对符号翻转置换检验评价专家回合平均得分是否高于加噪专家回合平均得分，并采用回合聚类Bootstrap计算JS散度和AUC的区间估计。",
        ),
        (
            "text",
            "在Actor训练过程评价中，本文每完成100个训练回合保存一次历史Actor快照。对于第r次独立运行的第k个评估时刻，历史Actor在同一批留出专家状态上以确定性方式生成动作，固定最终判别器随后计算相应得分：",
        ),
        ("equation", ACTOR_SCORE_MATHML),
        (
            "text",
            "式中，μ表示生产企业Actor；θ_(r,k)表示第r次独立运行在第k个评估时刻保存的Actor参数；上标G表示Actor生成动作或相应判别器得分。根据专家得分分布P_r^E和历史Actor得分分布P_(r,k)^G，定义Actor与专家的判别器得分分布JS散度：",
        ),
        ("equation", ACTOR_JS_MATHML),
        (
            "text",
            "式中，J_(r,k)表示第r次独立运行在第k个评估时刻的Actor—专家JS散度。较小的J_(r,k)表示历史Actor生成行为在固定判别器得分表示下更接近专家样本。本文选取R=3次独立运行，在每个评估时刻计算JS散度均值及其均值标准误：",
        ),
        ("equation", ACTOR_MEAN_SEM_MATHML),
        (
            "text",
            "式中，带横线的J表示第k个评估时刻三个独立运行的JS散度均值；s_k表示三个JS散度的样本标准差；R表示独立运行次数，本文取R=3。对应曲线中的实线表示均值，阴影表示均值上下1个SEM。",
        ),
    )
    for kind, content in paragraphs_and_equations:
        if kind == "text":
            anchor = insert_paragraph_after(
                anchor, content, format_source=body_template
            )
        else:
            anchor = insert_equation_after(
                anchor, content, equation_template=equation_template
            )


def insert_component_validation_section(document: Document) -> None:
    old_section_heading = find_paragraph(
        document, "6.4 GAIL+TD3与Transformer+GAIL+TD3的结果对比"
    )
    heading_template = old_section_heading
    body_template = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    caption_template = find_paragraph(document, "图6-1 TD3与GAIL+TD3训练结果对比")

    anchor = Paragraph(old_section_heading._p.getprevious(), old_section_heading._parent)
    heading = insert_paragraph_after(
        anchor,
        "6.3 生成对抗模仿学习组件有效性分析",
        format_source=heading_template,
    )
    anchor = insert_paragraph_after(
        heading,
        "第6.2节的训练结果表明，引入生成对抗模仿学习后，生产企业主体能够以较少的训练回合达到较高的平均存活水平。为进一步检验该训练方案中判别器与Actor生成器的工作状态，本文按照“最终判别器识别扰动行为—历史Actor逐步接近专家行为”的顺序开展两步评价。",
        format_source=body_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "第一步检验最终判别器对非专家动作扰动的识别能力。本文选取随机种子184对应的最终判别器作为展示样例，在留出测试集的16952条专家状态—动作样本上叠加标准差为0.30的独立高斯噪声，并将动作裁剪至[-0.5,0.5]。留出测试集包含173个完整专家回合。图6-2给出了原始专家样本与加噪专家样本的判别器得分分布。",
        format_source=body_template,
    )
    anchor = insert_picture_after(
        anchor,
        DISCRIMINATOR_FIGURE,
        body_template=body_template,
        width_inches=6.35,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-2 专家样本与加噪专家样本的判别器得分分布",
        format_source=caption_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-2显示，最终判别器对专家样本与加噪专家样本的平均评分分别为0.8988和0.6111，两者相差0.2877。两组得分分布的JS散度为0.4100，其回合聚类Bootstrap 95%区间为[0.4012,0.4191]；归一化直方图重叠系数为18.85%；以专家样本为正类的AUC为0.8186，其95%区间为[0.8073,0.8293]。单侧配对符号翻转置换检验的p值为0.0002。上述结果表明，在标准差为0.30的动作扰动条件下，最终判别器对专家样本赋予了更高的总体评分，并能够识别加噪动作相对于专家动作的偏离。",
        format_source=body_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "第二步检验历史Actor生成行为在训练过程中的变化。对于随机种子184、652和187，本文分别以对应的最终判别器作为固定评价尺度，并加载每100个训练回合保存的历史Actor快照。所有历史Actor均在同一批留出专家状态上以确定性方式生成动作，不加入探索噪声。图6-3给出了三个独立运行在60个评估时刻的Actor—专家JS散度均值，阴影表示均值上下1个SEM。",
        format_source=body_template,
    )
    anchor = insert_picture_after(
        anchor,
        ACTOR_FIGURE,
        body_template=body_template,
        width_inches=6.35,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-3 Actor与专家样本判别器得分分布的JS散度变化",
        format_source=caption_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-3显示，Actor—专家JS散度均值在第1个评估时刻为0.580±0.041，随后总体下降，并在第19个评估时刻达到最低值0.172±0.033。此后曲线出现一定回升，并在训练后期稳定于较低区间；第60个评估时刻的均值为0.258±0.017，较第1个评估时刻降低约55.5%。因此，固定最终判别器下的得分分布结果表明，历史Actor生成行为在训练过程中总体向专家样本接近，但这一变化具有阶段性波动，并非单调下降。",
        format_source=body_template,
    )
    insert_paragraph_after(
        anchor,
        "两步评价分别从判别器对扰动行为的识别能力和Actor生成行为的训练演化两个方面检验了生成对抗模仿学习组件。最终判别器能够识别具有明确动作扰动的样本，历史Actor相对于训练初期则形成了更接近专家得分分布的生成行为。上述结果为第6.2节中模仿学习信号改善生产企业主体前期训练表现提供了组件层面的实验证据。",
        format_source=body_template,
    )


def update_transitions_and_metric_count(document: Document) -> None:
    replace_paragraph_text(
        find_paragraph(document, "在上述训练过程中，本文跟踪"),
        "在上述训练过程中，本文跟踪平均存活天数、生产企业累计收益、累计归一化学习曲线下面积、学习速度、生产企业偿债能力和生成对抗模仿学习组件有效性六项评价指标。各项指标的定义如下。",
    )
    replace_paragraph_text(
        find_paragraph(document, "不过，GAIL+TD3的最终稳定平均存活天数为80.71天"),
        "不过，GAIL+TD3的最终稳定平均存活天数为80.71天，与专家轨迹平均值仍相差17.28天。该结果说明，模仿学习信号缩短了模型达到80天所需的训练过程，但稳定阶段的存活水平仍有进一步接近专家轨迹的空间。下一节首先从判别器与Actor生成器两个方面检验生成对抗模仿学习组件的有效性，随后在第6.4节引入历史状态序列表征，以考察模型利用连续交互信息后能否进一步缩小这一差距。",
    )


def verify(document: Document, source_equation_count: int) -> None:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    text = "\n".join(paragraphs)
    required = (
        "（6）生成对抗模仿学习组件有效性",
        "6.3 生成对抗模仿学习组件有效性分析",
        "图6-2 专家样本与加噪专家样本的判别器得分分布",
        "图6-3 Actor与专家样本判别器得分分布的JS散度变化",
        "6.4 GAIL+TD3与Transformer+GAIL+TD3的结果对比",
        "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        "图6-5 三种主体训练方案下DSCR 低于 1 的生产企业日观测占比",
        "第60个评估时刻的均值为0.258±0.017",
        "归一化直方图重叠系数为18.85%",
    )
    for value in required:
        if value not in text:
            raise RuntimeError(f"Required content is missing: {value}")
    forbidden = (
        "6.3 GAIL+TD3与Transformer+GAIL+TD3的结果对比",
        "图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        "图6-3 三种主体训练方案下DSCR 低于 1",
    )
    for value in forbidden:
        if value in text:
            raise RuntimeError(f"Former numbering remains: {value}")
    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 11:
        raise RuntimeError(f"Expected 11 inline figures, found {len(document.inline_shapes)}.")
    equation_count = len(document.element.body.xpath(".//m:oMath"))
    if equation_count != source_equation_count + len(EQUATIONS):
        raise RuntimeError(
            f"Expected {source_equation_count + len(EQUATIONS)} equations, "
            f"found {equation_count}."
        )


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The source DOCX changed after this builder was prepared.")
    if not MML2OMML_XSL.exists():
        raise FileNotFoundError(MML2OMML_XSL)
    create_paper_figures()
    if not DISCRIMINATOR_FIGURE.exists() or not ACTOR_FIGURE.exists():
        raise RuntimeError("Publication figures were not generated.")

    document = Document(SOURCE)
    source_equation_count = len(document.element.body.xpath(".//m:oMath"))
    renumber_existing_section(document)
    update_transitions_and_metric_count(document)
    insert_metric_definition(document)
    insert_component_validation_section(document)
    verify(document, source_equation_count)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The source DOCX was modified unexpectedly.")
    verify(Document(OUTPUT), source_equation_count)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
