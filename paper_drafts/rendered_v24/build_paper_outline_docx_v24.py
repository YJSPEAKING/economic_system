from __future__ import annotations

import csv
import math
import shutil
import statistics
import sys
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from typing import Dict, Iterable, List, Sequence

import matplotlib

BUNDLED_SITE_PACKAGES = Path(
    r"C:\Users\legion\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\python\Lib\site-packages"
)
if str(BUNDLED_SITE_PACKAGES) not in sys.path:
    sys.path.append(str(BUNDLED_SITE_PACKAGES))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from matplotlib.font_manager import FontProperties


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "reference"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_v1.docx"
)
OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT = OUTPUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v24.docx"
TEMP_FIGURE_DIR = OUTPUT_DIR / "_generated_figures"
RAW_METRIC_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "five_metrics_two_comparisons_aulc_sem"
    / "raw_metrics"
)
MML2OMML_XSL = Path(
    r"C:\Program Files\Microsoft Office\Root\Office16\MML2OMML.XSL"
)

GROUP_DIR_NAMES = {
    "TD3": "TD3",
    "GAIL+TD3": "GAIL_plus_TD3",
    "Transformer+GAIL+TD3": "Transformer_plus_GAIL_plus_TD3",
}
COLORS = {
    "TD3": "#1f77b4",
    "GAIL+TD3": "#ff7f0e",
    "Transformer+GAIL+TD3": "#2ca02c",
}
LEGEND_LABELS = {
    "TD3": "TD3",
    "GAIL+TD3": "GAIL+TD3",
    "Transformer+GAIL+TD3": "Trans.+GAIL+TD3",
}
COMPARISONS = (
    (
        ("TD3", "GAIL+TD3"),
        "td3_vs_gail_td3_three_metrics_aulc_sem.png",
    ),
    (
        ("GAIL+TD3", "Transformer+GAIL+TD3"),
        "gail_td3_vs_transformer_three_metrics_aulc_sem.png",
    ),
)
EXPERT_SAMPLE_ROWS = 168_548
EXPERT_SUCCESSFUL_EPISODES = 1_720
EXPERT_REFERENCE_LEVEL = EXPERT_SAMPLE_ROWS / EXPERT_SUCCESSFUL_EPISODES
MAX_EPISODE_DAYS = 100.0
EXPECTED_RUNS_PER_GROUP = 8
FINAL_STABLE_STEPS = 20
SPEED_THRESHOLDS = (80.0, 90.0)
SPEED_CONSECUTIVE_STEPS = 3

MEAN_CURVE_MATHML = (
    '<math xmlns="http://www.w3.org/1998/Math/MathML">'
    '<msub><mi>μ</mi><mi>k</mi></msub><mo>=</mo>'
    '<mfrac><mn>1</mn><mi>S</mi></mfrac>'
    '<munderover><mo>∑</mo><mrow><mi>r</mi><mo>=</mo><mn>1</mn></mrow>'
    '<mi>S</mi></munderover>'
    '<mover><msub><mi>L</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow>'
    '</msub><mo>¯</mo></mover>'
    '</math>'
)
SPEED_MATHML = (
    '<math xmlns="http://www.w3.org/1998/Math/MathML">'
    '<msub><mi>T</mi><mi>q</mi></msub><mo>=</mo><mi>min</mi>'
    '<mfenced open="{" close="}"><mrow><mi>k</mi><mo>|</mo>'
    '<msub><mi>μ</mi><mi>k</mi></msub><mo>≥</mo><mi>q</mi><mo>,</mo>'
    '<msub><mi>μ</mi><mrow><mi>k</mi><mo>+</mo><mn>1</mn></mrow></msub>'
    '<mo>≥</mo><mi>q</mi><mo>,</mo>'
    '<msub><mi>μ</mi><mrow><mi>k</mi><mo>+</mo><mn>2</mn></mrow></msub>'
    '<mo>≥</mo><mi>q</mi></mrow></mfenced><mo>,</mo>'
    '<mi>q</mi><mo>∈</mo><mfenced open="{" close="}"><mrow>'
    '<mn>80</mn><mo>,</mo><mn>90</mn></mrow></mfenced>'
    '</math>'
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def sample_sem(values: Sequence[float]) -> float:
    if len(values) < 2:
        return 0.0
    return statistics.stdev(values) / math.sqrt(len(values))


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def first_run_properties(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    run_properties = first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def insert_paragraph_after(
    reference: Paragraph,
    text: str = "",
    copy_format_from: Paragraph | None = None,
) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    format_source = copy_format_from or reference
    if format_source._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(format_source._p.pPr))
    if text:
        run_properties = first_run_properties(format_source)
        run = paragraph.add_run(text)
        if run_properties is not None:
            run._r.insert(0, run_properties)
    return paragraph


def mathml_to_omml(mathml: str):
    transform = etree.XSLT(etree.parse(str(MML2OMML_XSL)))
    result = transform(etree.fromstring(mathml.encode("utf-8")))
    return etree.fromstring(etree.tostring(result.getroot()))


def insert_equation_after(
    reference: Paragraph,
    mathml: str,
    template: Paragraph,
) -> Paragraph:
    equation = insert_paragraph_after(reference, copy_format_from=template)
    equation._p.append(mathml_to_omml(mathml))
    return equation


def replace_cell_text(cell, text: str) -> None:
    replace_paragraph_text(cell.paragraphs[0], text)
    for extra in list(cell.paragraphs[1:]):
        remove_paragraph(extra)


def find_table_by_header(document: Document, headers: Sequence[str]) -> Table:
    matches = [
        table
        for table in document.tables
        if table.rows
        and [cell.text.strip() for cell in table.rows[0].cells] == list(headers)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table with header {headers}, found {len(matches)}.")
    return matches[0]


def rebuild_table_rows(table: Table, rows: Sequence[Sequence[str]]) -> None:
    if len(table.rows) < 2:
        raise RuntimeError("The table has no body-row formatting template.")
    row_template = deepcopy(table.rows[1]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for values in rows:
        new_row = deepcopy(row_template)
        table._tbl.append(new_row)
        row = table.rows[-1]
        if len(row.cells) != len(values):
            raise RuntimeError("Table row and supplied values have different widths.")
        for cell, value in zip(row.cells, values):
            replace_cell_text(cell, value)


def replace_figure_before_caption(
    document: Document,
    caption_prefix: str,
    image_path: Path,
) -> None:
    caption = find_paragraph(document, caption_prefix)
    image_element = caption._p.getprevious()
    if image_element is None or image_element.tag != qn("w:p"):
        raise RuntimeError(f"No image paragraph found before {caption_prefix!r}.")
    image_paragraph = Paragraph(image_element, caption._parent)
    if not image_paragraph._p.xpath(".//w:drawing"):
        raise RuntimeError(f"Previous paragraph is not a figure for {caption_prefix!r}.")
    for child in list(image_paragraph._p):
        if child.tag != qn("w:pPr"):
            image_paragraph._p.remove(child)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.35))


def load_runs() -> Dict[str, List[dict]]:
    groups: Dict[str, List[dict]] = {}
    for group_name, directory_name in GROUP_DIR_NAMES.items():
        paths = sorted((RAW_METRIC_DIR / directory_name).glob("*_metrics.csv"))
        if len(paths) != EXPECTED_RUNS_PER_GROUP:
            raise RuntimeError(
                f"{group_name} should have {EXPECTED_RUNS_PER_GROUP} cached runs; "
                f"found {len(paths)}."
            )
        runs: List[dict] = []
        for path in paths:
            run = {"run": "", "seed": "", "values": {"survival": {}, "production": {}}}
            with path.open("r", encoding="utf-8-sig", newline="") as file:
                for row in csv.DictReader(file):
                    metric_id = row["metric_id"]
                    if metric_id not in run["values"]:
                        continue
                    run["run"] = row["run"]
                    run["seed"] = row["seed"]
                    run["values"][metric_id][int(row["step"])] = float(row["value"])
            if not run["values"]["survival"] or not run["values"]["production"]:
                raise RuntimeError(f"Incomplete metrics in {path}.")
            runs.append(run)
        groups[group_name] = runs

    all_survival = [
        run["values"]["survival"]
        for runs in groups.values()
        for run in runs
    ]
    common_steps = sorted(set.intersection(*(set(values) for values in all_survival)))
    if common_steps != list(range(1, 61)):
        raise RuntimeError(f"Expected common evaluation steps 1..60, found {common_steps}.")
    for runs in groups.values():
        for run in runs:
            add_cumulative_naulc(run, common_steps)
    return groups


def add_cumulative_naulc(run: dict, steps: Sequence[int]) -> None:
    survival = run["values"]["survival"]
    horizon = steps[-1] - steps[0]
    area = 0.0
    values = {steps[0]: 0.0}
    for previous_step, step in zip(steps[:-1], steps[1:]):
        area += (step - previous_step) * (
            survival[previous_step] + survival[step]
        ) / 2.0
        values[step] = area / (horizon * MAX_EPISODE_DAYS)
    run["values"]["naulc"] = values


def aggregate_runs(runs: Sequence[dict], metric_id: str) -> dict:
    steps = sorted(
        set.intersection(*(set(run["values"][metric_id]) for run in runs))
    )
    means = []
    sems = []
    for step in steps:
        values = [run["values"][metric_id][step] for run in runs]
        means.append(statistics.mean(values))
        sems.append(sample_sem(values))
    return {"steps": steps, "means": means, "sems": sems}


def group_mean_curve(runs: Sequence[dict]) -> Dict[int, float]:
    steps = sorted(
        set.intersection(*(set(run["values"]["survival"]) for run in runs))
    )
    return {
        step: statistics.mean(run["values"]["survival"][step] for run in runs)
        for step in steps
    }


def threshold_step(mean_curve: Dict[int, float], threshold: float) -> int | None:
    steps = sorted(mean_curve)
    for index, step in enumerate(steps):
        window = steps[index : index + SPEED_CONSECUTIVE_STEPS]
        if len(window) != SPEED_CONSECUTIVE_STEPS:
            break
        if window != list(range(step, step + SPEED_CONSECUTIVE_STEPS)):
            continue
        if all(mean_curve[item] >= threshold for item in window):
            return step
    return None


def compute_statistics(groups: Dict[str, List[dict]]) -> dict:
    statistics_by_group = {}
    for group_name, runs in groups.items():
        group_statistics = {}
        for metric_id in ("survival", "production"):
            common_steps = sorted(
                set.intersection(*(set(run["values"][metric_id]) for run in runs))
            )
            stable_steps = common_steps[-FINAL_STABLE_STEPS:]
            seed_values = [
                statistics.mean(run["values"][metric_id][step] for step in stable_steps)
                for run in runs
            ]
            mean_value = statistics.mean(seed_values)
            sem_value = sample_sem(seed_values)
            group_statistics[metric_id] = {
                "mean": mean_value,
                "sem": sem_value,
                "relative_sem": sem_value / abs(mean_value) * 100.0,
            }

        final_step = min(max(run["values"]["naulc"]) for run in runs)
        naulc_values = [run["values"]["naulc"][final_step] for run in runs]
        naulc_mean = statistics.mean(naulc_values)
        naulc_sem = sample_sem(naulc_values)
        group_statistics["naulc"] = {
            "mean": naulc_mean,
            "sem": naulc_sem,
            "relative_sem": naulc_sem / abs(naulc_mean) * 100.0,
        }

        mean_curve = group_mean_curve(runs)
        group_statistics["speed"] = {
            threshold: threshold_step(mean_curve, threshold)
            for threshold in SPEED_THRESHOLDS
        }
        statistics_by_group[group_name] = group_statistics
    return statistics_by_group


def configure_matplotlib() -> None:
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "font.serif": ["Times New Roman", "SimSun"],
            "font.sans-serif": ["Microsoft YaHei", "SimHei", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "font.size": 9.5,
            "axes.titlesize": 10.5,
            "axes.labelsize": 10,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 7.8,
            "axes.linewidth": 1.0,
            "lines.linewidth": 1.7,
            "savefig.dpi": 300,
        }
    )


def style_axis(ax) -> None:
    ax.grid(True, linestyle="-.", linewidth=0.65, color="#a8a8a8", alpha=0.65)
    for spine in ax.spines.values():
        spine.set_color("black")
        spine.set_linewidth(1.0)


def plot_metric(ax, metric_id: str, group_names: Sequence[str], groups: dict) -> None:
    from matplotlib.ticker import FuncFormatter, MultipleLocator

    titles = {
        "survival": "(a) 每百回合/存活天数",
        "production": "(b) 每百回合/累计收益/生产企业",
        "naulc": "(c) 累计归一化学习曲线下面积",
    }
    y_labels = {"survival": "Days", "production": "Income", "naulc": "nAULC"}
    title_font = FontProperties(family="Microsoft YaHei", size=10.5)
    upper_values: List[float] = []

    for group_name in group_names:
        aggregate = aggregate_runs(groups[group_name], metric_id)
        lower = [
            max(0.0, mean - sem)
            for mean, sem in zip(aggregate["means"], aggregate["sems"])
        ]
        upper = [
            mean + sem
            for mean, sem in zip(aggregate["means"], aggregate["sems"])
        ]
        upper_values.extend(upper)
        ax.fill_between(
            aggregate["steps"],
            lower,
            upper,
            color=COLORS[group_name],
            alpha=0.12,
            linewidth=0,
            zorder=1,
        )
        ax.plot(
            aggregate["steps"],
            aggregate["means"],
            color=COLORS[group_name],
            label=LEGEND_LABELS[group_name],
            zorder=2,
        )

    if metric_id == "survival":
        ax.axhline(
            EXPERT_REFERENCE_LEVEL,
            color="#4d4d4d",
            linestyle=(0, (5, 3)),
            linewidth=1.15,
            label=f"Expert trajectories ({EXPERT_REFERENCE_LEVEL:.2f} days)",
            zorder=3,
        )

    ax.set_title(titles[metric_id], pad=7, fontproperties=title_font)
    ax.set_xlabel("Evaluation Step (100 Episodes)")
    ax.set_ylabel(y_labels[metric_id], labelpad=10)
    ax.set_xlim(0, 60)
    ax.set_xticks(range(0, 61, 10))
    if metric_id == "survival":
        ax.set_ylim(0, 100)
        ax.yaxis.set_major_locator(MultipleLocator(20))
    elif metric_id == "production":
        upper = max(max(upper_values, default=1.0), 1.0)
        upper = math.ceil(upper / 50_000.0) * 50_000.0
        ax.set_ylim(0, upper)
        ax.yaxis.set_major_locator(MultipleLocator(50_000))
        ax.yaxis.set_major_formatter(
            FuncFormatter(lambda value, _position: f"{int(round(value))}")
        )
    else:
        ax.set_ylim(0, 1.0)
        ax.yaxis.set_major_locator(MultipleLocator(0.2))
    style_axis(ax)
    ax.legend(
        loc="lower right" if metric_id in ("survival", "naulc") else "upper left",
        frameon=True,
        fancybox=False,
        framealpha=0.88,
        edgecolor="#333333",
        borderpad=0.4,
        labelspacing=0.25,
        handlelength=2.2,
    )


def render_figures(groups: Dict[str, List[dict]]) -> Dict[str, Path]:
    import matplotlib.pyplot as plt

    configure_matplotlib()
    TEMP_FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    paths = {}
    for group_names, filename in COMPARISONS:
        figure, axes = plt.subplots(1, 3, figsize=(13.2, 4.15))
        for axis, metric_id in zip(axes, ("survival", "production", "naulc")):
            plot_metric(axis, metric_id, group_names, groups)
        figure.subplots_adjust(
            left=0.065,
            right=0.99,
            bottom=0.17,
            top=0.92,
            wspace=0.37,
        )
        path = TEMP_FIGURE_DIR / filename
        figure.savefig(path, facecolor="white", dpi=300)
        plt.close(figure)
        paths[filename] = path
    return paths


def format_stat(
    value: float,
    relative_sem: float,
    scale: float = 1.0,
    suffix: str = "",
    digits: int = 2,
) -> str:
    return f"{value / scale:.{digits}f}{suffix}（±{relative_sem:.2f}%）"


def update_chapter_five(document: Document) -> None:
    data_sources = find_paragraph(document, "本文使用两类数据来源")
    replace_paragraph_text(
        data_sources,
        "本文使用两类数据来源。第一类为网页人工采集数据。网页人工采集实验共有63名参与者。参与者依据网页提供的生产企业经营状态完成连续决策，并可独立进行多个仿真回合。经存活天数筛选后，人工采集部分共保留285个有效回合和27987条状态—动作样本。第二类为Codex交互式辅助采集数据。Codex 5.5通过Control Chrome读取同一网页中的可见信息和动作边界，依据第5.3节所述分层提示逐日生成并提交四项业务动作。为控制环境随机性和非目标主体策略变化对采集结果的影响，两类采集均采用统一的实验随机种子，并固定消费企业和银行的预训练Actor参数。采集阶段不对Codex、消费企业或银行进行在线训练，仅执行模型推理和环境交互，从而避免实时训练带来的额外计算和等待开销。所有轨迹均按照相同的环境规则执行，并依据存活天数进行筛选。",
    )

    table = find_table_by_header(
        document,
        ("数据来源", "成功回合数", "样本行数", "平均存活天数"),
    )
    replace_cell_text(table.rows[2].cells[0], "Codex交互式辅助采集数据")


def update_chapter_six(
    document: Document,
    figures: Dict[str, Path],
    stats: dict,
) -> None:
    timing = find_paragraph(document, "本文以回合作为训练进度单位")
    replace_paragraph_text(
        timing,
        "本文以回合作为训练进度单位，并将每个评估时刻所包含的连续训练回合数记为E。模型每完成E个连续训练回合，本文汇总一次评价指标，并将该汇总位置记为一个评估时刻（evaluation step）。第k个评估时刻对应第k组连续训练回合的统计结果。本文比较前60个评估时刻，共计6000个训练回合。表6-1列出了纳入统计的训练回合数、评估时刻数、随机种子组数、每个评估时刻对应回合数、单回合最大天数和统计方法。",
    )

    table_6_1 = find_table_by_header(document, ("项目", "设定"))
    rebuild_table_rows(
        table_6_1,
        (
            ("纳入统计的训练回合数", "6000个回合"),
            ("评估时刻数", "60个"),
            ("随机种子组数", "每种模型8组"),
            ("每个评估时刻对应回合数", "E=100个连续训练回合"),
            ("单回合最大天数", "100天"),
            ("统计方法", "8次独立运行的均值及均值上下1个SEM"),
        ),
    )

    metric_intro = find_paragraph(document, "在上述训练过程中")
    replace_paragraph_text(
        metric_intro,
        "在上述训练过程中，本文跟踪平均存活天数、生产企业累计收益、累计归一化学习曲线下面积、学习速度和生产企业偿债能力五项评价指标。各项指标的定义如下。",
    )

    income_intro = find_paragraph(document, "（2）累计收益")
    replace_paragraph_text(
        income_intro,
        "（2）生产企业累计收益。本文首先在单个回合内汇总生产企业的收入、支出与利息，再对第k个评估时刻内E个回合的生产企业累计收益取算术平均值，其定义如下：",
    )
    income_explanation = find_paragraph(document, "式中，J表示企业在对应评估时刻")
    replace_paragraph_text(
        income_explanation,
        "式中，J表示生产企业在对应评估时刻的平均累计收益；R、C和I分别表示单个回合内的累计收入、累计支出和累计利息；k和n分别表示评估时刻编号和该评估时刻内的回合序号。",
    )

    bank_intro = find_paragraph(document, "生产企业和消费企业分别依据上述定义")
    bank_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is bank_intro._p
    )
    bank_equation = document.paragraphs[bank_index + 1]
    bank_explanation = document.paragraphs[bank_index + 2]
    if not bank_equation._p.xpath(".//m:oMath"):
        raise RuntimeError("The bank-income equation was not found.")
    remove_paragraph(bank_explanation)
    remove_paragraph(bank_equation)
    remove_paragraph(bank_intro)

    naulc_intro = find_paragraph(document, "（3）累计归一化学习曲线下面积")
    replace_paragraph_text(
        naulc_intro,
        "（3）累计归一化学习曲线下面积。归一化学习曲线下面积（nAULC: normalized area under the learning curve）用于描述平均存活天数在完整训练区间内的累积表现。本文对每次独立运行的平均存活天数曲线采用梯形法积分，并以完整评估区间和单回合最大天数进行归一化。截至第k个评估时刻的累计归一化面积定义如下：",
    )
    naulc_explanation = find_paragraph(document, "式中，r表示独立运行编号")
    replace_paragraph_text(
        naulc_explanation,
        "式中，r表示独立运行编号；k表示当前评估时刻编号；K表示图中纳入比较的评估时刻总数；j表示积分区间编号；带上横线的L表示相应独立运行在对应评估时刻的平均存活天数；L的最大值表示表6-1所列的单回合最大天数。本文令第1个评估时刻的累计归一化面积为0。在相同训练区间内，较大的累计归一化面积表示模型获得了较高的累积平均存活表现。",
    )

    naulc_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is naulc_explanation._p
    )
    equation_template = document.paragraphs[naulc_index - 1]
    speed_intro = insert_paragraph_after(
        naulc_explanation,
        "（4）学习速度。本文采用平均存活天数均值曲线达到目标存活水平所需的评估时刻衡量学习速度。设第k个评估时刻的跨随机种子均值为μ，计算方式如下：",
        copy_format_from=naulc_intro,
    )
    mean_equation = insert_equation_after(speed_intro, MEAN_CURVE_MATHML, equation_template)
    speed_definition = insert_paragraph_after(
        mean_equation,
        "为减少均值曲线短时波动对阈值到达时刻的影响，本文将连续3个评估时刻均不低于目标值时的第1个评估时刻定义为达到该目标值的时刻，其定义如下：",
        copy_format_from=naulc_explanation,
    )
    speed_equation = insert_equation_after(speed_definition, SPEED_MATHML, equation_template)
    insert_paragraph_after(
        speed_equation,
        "式中，S表示独立运行次数，本文取S=8；r表示独立运行编号；μ表示第k个评估时刻的平均存活天数均值；q表示目标存活天数，本文分别取80天和90天；T表示均值曲线达到目标值所需的评估时刻。较小的T表示模型以较少的训练回合达到目标存活水平。若均值曲线截至第60个评估时刻仍未满足连续3个评估时刻均不低于目标值的条件，则记为“未达到”。",
        copy_format_from=naulc_explanation,
    )

    dscr_intro = find_paragraph(document, "（4）生产企业偿债能力")
    replace_paragraph_text(
        dscr_intro,
        dscr_intro.text.replace("（4）", "（5）", 1),
    )

    section_6_2_intro = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    replace_paragraph_text(
        section_6_2_intro,
        "本文首先比较TD3与GAIL+TD3的训练过程，以分析模仿学习信号对生产企业主体决策训练的影响。图6-1给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标上的均值曲线及SEM阴影。其中，图6-1(a)中的水平虚线表示表5-1所列专家轨迹的合计平均存活天数，其数值为97.99天。为定量比较三种模型的训练表现，本文对每次独立运行在最后20个评估时刻的平均存活天数和生产企业累计收益分别取均值，再在8次独立运行之间计算总体均值和SEM；累计归一化学习曲线下面积采用第60个评估时刻的数值；学习速度采用平均存活天数均值曲线达到80天和90天所需的评估时刻。表6-2汇总了三种模型的相应统计结果。",
    )
    replace_figure_before_caption(
        document,
        "图6-1 TD3与GAIL+TD3",
        figures["td3_vs_gail_td3_three_metrics_aulc_sem.png"],
    )

    section_6_2_figure = find_paragraph(document, "从图6-1(a)可以看出")
    replace_paragraph_text(
        section_6_2_figure,
        "从图6-1(a)可以看出，引入模仿学习信号后，GAIL+TD3的平均存活天数在训练前期快速提高，并在第17个评估时刻达到80天；TD3在第44个评估时刻达到相同水平。两种模型的均值曲线在前60个评估时刻内均未达到连续3个评估时刻不低于90天的条件。图6-1(b)显示，GAIL+TD3的生产企业累计收益较早进入快速提升阶段并转入相对稳定状态，TD3的生产企业累计收益在训练后期达到较高水平。图6-1(c)显示，GAIL+TD3的累计归一化学习曲线下面积整体高于TD3，且两条曲线的差距主要在训练前期形成并延续至训练末期。",
    )
    section_6_2_table = find_paragraph(document, "表6-2中的具体数值显示")
    replace_paragraph_text(
        section_6_2_table,
        "表6-2中的具体数值显示，GAIL+TD3达到80天所需的评估时刻由TD3的44缩短至17，提前27个评估时刻，即2700个训练回合。GAIL+TD3在第60个评估时刻的累计归一化学习曲线下面积由TD3的0.435提高至0.683，增幅约为56.8%，表明该模型在相同训练区间内取得了更高的累积平均存活表现。生产企业累计收益曲线在训练前期的上升过程也与平均存活天数的变化保持一致，说明模仿奖励能够较早引导生产企业主体形成有效策略。",
    )
    section_6_2_transition = find_paragraph(document, "不过，GAIL+TD3的最终稳定")
    replace_paragraph_text(
        section_6_2_transition,
        "不过，GAIL+TD3的最终稳定平均存活天数为80.71天，与专家轨迹平均值仍相差17.28天，且其均值曲线在前60个评估时刻内未达到90天。该结果表明，模仿学习信号缩短了模型达到80天所需的训练过程，但稳定阶段的存活水平仍有进一步接近专家轨迹的空间。因此，下一节引入历史状态序列表征，以考察模型利用连续交互信息后能否缩小这一差距。",
    )

    table_6_2 = find_table_by_header(
        document,
        ("评价指标", "TD3", "GAIL+TD3", "Transformer+GAIL+TD3"),
    )
    rows = []
    for metric_label, metric_id, scale, suffix, digits in (
        ("平均存活天数/天", "survival", 1.0, "", 2),
        ("生产企业累计收益", "production", 1000.0, "K", 2),
        ("第60个评估时刻的nAULC", "naulc", 1.0, "", 3),
    ):
        row = [metric_label]
        for group_name in GROUP_DIR_NAMES:
            values = stats[group_name][metric_id]
            row.append(
                format_stat(
                    values["mean"],
                    values["relative_sem"],
                    scale=scale,
                    suffix=suffix,
                    digits=digits,
                )
            )
        rows.append(tuple(row))
    rows.extend(
        (
            (
                "达到80天所需评估时刻",
                str(stats["TD3"]["speed"][80.0]),
                str(stats["GAIL+TD3"]["speed"][80.0]),
                str(stats["Transformer+GAIL+TD3"]["speed"][80.0]),
            ),
            (
                "达到90天所需评估时刻",
                "未达到" if stats["TD3"]["speed"][90.0] is None else str(stats["TD3"]["speed"][90.0]),
                "未达到" if stats["GAIL+TD3"]["speed"][90.0] is None else str(stats["GAIL+TD3"]["speed"][90.0]),
                "未达到" if stats["Transformer+GAIL+TD3"]["speed"][90.0] is None else str(stats["Transformer+GAIL+TD3"]["speed"][90.0]),
            ),
        )
    )
    rebuild_table_rows(table_6_2, rows)

    caption_6_2 = find_paragraph(document, "表6-2 三种主体训练方案")
    replace_paragraph_text(caption_6_2, "表6-2 三种主体训练方案的训练表现统计")
    note_6_2 = find_paragraph(document, "注：平均存活天数及三项累计收益")
    replace_paragraph_text(
        note_6_2,
        "注：平均存活天数和生产企业累计收益为每次独立运行最后20个评估时刻的均值，再对8次运行进行统计；nAULC采用第60个评估时刻的数值。括号内为相对SEM，K表示10³。学习速度依据8次独立运行的平均存活天数均值曲线计算；“未达到”表示均值曲线截至第60个评估时刻仍未连续3个评估时刻达到相应目标值。",
    )

    section_6_3_intro = find_paragraph(document, "基于上一节的结果")
    replace_paragraph_text(
        section_6_3_intro,
        "基于上一节的结果，本文进一步比较GAIL+TD3与Transformer+GAIL+TD3，以分析历史状态序列表征对生产企业主体训练过程的影响。图6-2给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标上的均值曲线及SEM阴影，表6-2列出了对应的训练表现统计。图6-2(a)中的水平虚线同样表示专家轨迹的合计平均存活天数，用于比较两种模型与专家轨迹存活水平之间的差距。",
    )
    replace_figure_before_caption(
        document,
        "图6-2 GAIL+TD3与Transformer+GAIL+TD3",
        figures["gail_td3_vs_transformer_three_metrics_aulc_sem.png"],
    )
    section_6_3_figure = find_paragraph(document, "从图6-2(a)可以看出")
    replace_paragraph_text(
        section_6_3_figure,
        "从图6-2(a)可以看出，Transformer+GAIL+TD3在第14个评估时刻达到80天，比GAIL+TD3提前3个评估时刻；该模型随后在第20个评估时刻达到90天，并在训练后期保持较接近专家参考线的平均水平。图6-2(b)显示，Transformer+GAIL+TD3的生产企业累计收益在训练前期较快提高，稳定阶段的均值高于GAIL+TD3。图6-2(c)显示，Transformer+GAIL+TD3的累计归一化学习曲线下面积在多数评估时刻高于GAIL+TD3。上述曲线表明，历史状态序列为生产企业主体提供了单步状态之外的连续交互信息，使模型能够更充分地利用模仿奖励所包含的专家行为信息。",
    )
    section_6_3_table = find_paragraph(document, "表6-2进一步显示")
    replace_paragraph_text(
        section_6_3_table,
        "表6-2进一步显示，引入Transformer编码器后，最终稳定平均存活天数由80.71天提高至90.24天，与专家轨迹平均值的差距由17.28天缩小至7.75天，缩小幅度约为55.2%。第60个评估时刻的累计归一化学习曲线下面积由0.683提高至0.773，增幅约为13.2%；生产企业最终稳定累计收益由126.35K提高至180.43K。学习速度指标显示，Transformer+GAIL+TD3不仅较早达到80天，而且在三种模型中率先达到并连续保持90天。上述结果说明，历史状态表征提高了生产企业主体对跨时段信息的利用能力，并使模型在训练全过程的累积平均存活表现、稳定阶段存活水平和生产企业经营结果方面进一步接近专家行为所形成的运行状态。",
    )
    section_6_3_conclusion = find_paragraph(document, "综合图6-2、图6-3和表6-2的结果")
    replace_paragraph_text(
        section_6_3_conclusion,
        "综合图6-2、图6-3和表6-2的结果，Transformer+GAIL+TD3达到80天和90天所需的训练回合较少，最终稳定平均存活天数与专家轨迹平均值的差距由17.28天缩小至7.75天，并将DSCR低于1的生产企业日观测占比降至0.00%。历史状态序列表征使生产企业主体能够利用跨时段交互信息，从而在学习速度、训练全过程的累积平均存活表现、稳定阶段存活水平和债务偿付状态等方面进一步利用专家行为信息。",
    )


def update_chapter_seven(document: Document) -> None:
    overview = find_paragraph(document, "本文围绕复杂经济系统多主体仿真环境")
    replace_paragraph_text(
        overview,
        overview.text.replace("主体累计收益", "生产企业累计收益"),
    )
    first = find_paragraph(document, "首先，生成对抗模仿学习信号")
    replace_paragraph_text(
        first,
        "首先，生成对抗模仿学习信号缩短了生产企业主体达到较高存活水平所需的训练过程。与TD3相比，GAIL+TD3达到80天所需的评估时刻由44缩短至17，提前27个评估时刻；第60个评估时刻的累计归一化学习曲线下面积由0.435提高至0.683。专家状态—动作样本经由判别器形成的模仿奖励为生产企业主体提供了额外的学习信息，使模型以较少的训练回合达到80天，并在完整训练区间内取得较高的累积平均存活表现。",
    )
    second = find_paragraph(document, "其次，引入Transformer编码器后")
    replace_paragraph_text(
        second,
        "其次，引入Transformer编码器后，模型的学习速度和稳定阶段存活水平得到进一步改善。Transformer+GAIL+TD3在第14个评估时刻达到80天，并在第20个评估时刻达到90天；其最终稳定平均存活天数由80.71天提高至90.24天，与专家轨迹平均值的差距由17.28天缩小至7.75天，第60个评估时刻的累计归一化学习曲线下面积由0.683提高至0.773。上述结果说明，历史状态序列增强了生产企业主体对跨时段交互信息的表征，并使模型在平均存活天数方面进一步接近专家轨迹。",
    )


def verify(document: Document, stats: dict) -> None:
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required = (
        "网页人工采集实验共有63名参与者",
        "纳入统计的训练回合数",
        "（4）学习速度",
        "用于描述平均存活天数在完整训练区间内的累积表现",
        "图6-1给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标",
        "图6-2给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标",
        "达到80天所需的评估时刻由TD3的44缩短至17",
        "在第20个评估时刻达到90天",
    )
    for item in required:
        if item not in text:
            raise RuntimeError(f"Required revision is missing: {item}")

    forbidden = (
        "用于综合衡量平均存活天数的提升速度与持续水平",
        "该结果从整个训练区间对图6-1所反映的学习速度差异进行了量化",
        "图6-1(e)",
        "图6-2(e)",
        "三项累计收益",
        "消费企业最终稳定累计收益",
        "银行最终稳定累计收益",
        "总训练天数",
    )
    for item in forbidden:
        if item in text:
            raise RuntimeError(f"Outdated wording remains: {item}")

    table_6_1 = find_table_by_header(document, ("项目", "设定"))
    expected_6_1 = [
        "纳入统计的训练回合数",
        "评估时刻数",
        "随机种子组数",
        "每个评估时刻对应回合数",
        "单回合最大天数",
        "统计方法",
    ]
    if [row.cells[0].text for row in table_6_1.rows[1:]] != expected_6_1:
        raise RuntimeError("Table 6-1 rows are incorrect.")

    table_6_2 = find_table_by_header(
        document,
        ("评价指标", "TD3", "GAIL+TD3", "Transformer+GAIL+TD3"),
    )
    expected_6_2 = [
        "平均存活天数/天",
        "生产企业累计收益",
        "第60个评估时刻的nAULC",
        "达到80天所需评估时刻",
        "达到90天所需评估时刻",
    ]
    if [row.cells[0].text for row in table_6_2.rows[1:]] != expected_6_2:
        raise RuntimeError("Table 6-2 rows are incorrect.")

    expected_speed = {
        "TD3": {80.0: 44, 90.0: None},
        "GAIL+TD3": {80.0: 17, 90.0: None},
        "Transformer+GAIL+TD3": {80.0: 14, 90.0: 20},
    }
    for group_name, values in expected_speed.items():
        if stats[group_name]["speed"] != values:
            raise RuntimeError(
                f"Unexpected learning-speed result for {group_name}: "
                f"{stats[group_name]['speed']}"
            )

    if len(document.inline_shapes) != 9:
        raise RuntimeError(f"Expected 9 inline figures, found {len(document.inline_shapes)}.")
    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.element.body.xpath(".//m:oMath")) != 117:
        raise RuntimeError("The Word equation-object count is incorrect.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not MML2OMML_XSL.exists():
        raise FileNotFoundError(MML2OMML_XSL)
    source_hash = file_sha256(SOURCE)

    groups = load_runs()
    stats = compute_statistics(groups)
    figures = render_figures(groups)

    try:
        document = Document(SOURCE)
        update_chapter_five(document)
        update_chapter_six(document, figures, stats)
        update_chapter_seven(document)
        verify(document, stats)
        document.save(OUTPUT)

        if file_sha256(SOURCE) != source_hash:
            raise RuntimeError("The source DOCX changed during generation.")
        check = Document(OUTPUT)
        verify(check, stats)
    finally:
        shutil.rmtree(TEMP_FIGURE_DIR, ignore_errors=True)

    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
