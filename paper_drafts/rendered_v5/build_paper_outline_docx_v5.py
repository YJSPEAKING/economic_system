from __future__ import annotations

import math
import zipfile
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
ROOT = OUT_DIR.parent.parent
FIG_DIR = OUT_DIR / "generated_figures"
OUT_PATH = OUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v5.docx"
SIM_RULE_FIG = FIG_DIR / "sim_rule_figure.png"
TRAIN_FLOW_FIG = FIG_DIR / "gail_td3_training_flow.png"
GAIL_TD3_REF_FIG = ROOT / "paper_drafts" / "reference" / "GAIL+TD3v0.png"


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman") -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)


def set_paragraph_format(
    paragraph,
    *,
    first_line: bool = False,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.5,
    align=None,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    if align is not None:
        paragraph.alignment = align


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if level == 1:
        set_paragraph_format(paragraph, before=10, after=6, line_spacing=1.25)
        size = 12
    elif level == 2:
        set_paragraph_format(paragraph, before=6, after=4, line_spacing=1.25)
        size = 10.5
    else:
        set_paragraph_format(paragraph, before=4, after=3, line_spacing=1.25)
        size = 10.5
    run = paragraph.add_run(text)
    set_run_font(run, east_asia="黑体")
    run.font.size = Pt(size)
    run.bold = True


def add_body(doc: Document, text: str, *, first_line: bool = True):
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        first_line=first_line,
        after=3,
        line_spacing=1.5,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    run = paragraph.add_run(text)
    set_run_font(run)
    run.font.size = Pt(10.5)
    return paragraph


def add_body_runs(doc: Document, parts: list[tuple[str, bool]], *, first_line: bool = True):
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        first_line=first_line,
        after=3,
        line_spacing=1.5,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    for text, is_subscript in parts:
        run = paragraph.add_run(text)
        set_run_font(run)
        run.font.size = Pt(10.5)
        run.font.subscript = is_subscript
    return paragraph


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        first_line=False,
        after=6,
        line_spacing=1.2,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    run = paragraph.add_run(text)
    set_run_font(run)
    run.font.size = Pt(9)


def add_formula(doc: Document, latex: str) -> None:
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        first_line=False,
        before=3,
        after=3,
        line_spacing=1.2,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    run = paragraph.add_run(latex)
    set_run_font(run, east_asia="Times New Roman", ascii_font="Times New Roman")
    run.font.size = Pt(10.5)
    run.italic = True


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "808080")


def set_cell_text(cell, text: str, *, bold: bool = False, center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run)
    run.font.size = Pt(9)
    run.bold = bold


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, "F2F2F2")
        set_cell_text(cell, header, bold=True, center=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, center=(idx == 0))
    doc.add_paragraph()


def find_font(size: int, bold: bool = False):
    names = [
        "msyhbd.ttc" if bold else "msyh.ttc",
        "simhei.ttf",
        "simsun.ttc",
        "arial.ttf",
    ]
    for name in names:
        path = Path(r"C:\Windows\Fonts") / name
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font,
    fill: str = "#111111",
) -> None:
    max_width = box[2] - box[0] - 24
    lines: list[str] = []
    for raw_line in text.split("\n"):
        current = ""
        for char in raw_line:
            trial = current + char
            if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = char
        if current:
            lines.append(current)
    line_height = draw.textbbox((0, 0), "国", font=font)[3] + 6
    total_height = line_height * len(lines)
    y = box[1] + (box[3] - box[1] - total_height) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        x = box[0] + (box[2] - box[0] - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += line_height


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str = "#2F5597") -> None:
    draw.line([start, end], fill=fill, width=3)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 12
    points = [
        end,
        (
            int(end[0] - size * math.cos(angle - math.pi / 6)),
            int(end[1] - size * math.sin(angle - math.pi / 6)),
        ),
        (
            int(end[0] - size * math.cos(angle + math.pi / 6)),
            int(end[1] - size * math.sin(angle + math.pi / 6)),
        ),
    ]
    draw.polygon(points, fill=fill)


def make_training_flow_figure() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (2000, 1100), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = find_font(38, bold=True)
    box_font = find_font(25)
    small_font = find_font(22)
    line = "#2F5597"
    fill_main = "#EAF2FF"
    fill_disc = "#FFF2CC"
    fill_value = "#E2F0D9"
    fill_data = "#FCE4D6"

    draw.text((565, 30), "生产企业主体的 GAIL+TD3 在线训练流程", font=title_font, fill="#111111")
    boxes = {
        "env": (90, 155, 390, 280, fill_main, "仿真环境\n状态 s_t"),
        "actor": (535, 155, 835, 280, fill_main, "Actor / Generator\n生成动作 a_t"),
        "step": (980, 155, 1280, 280, fill_main, "环境交互\n获得 R_env, s'"),
        "replay": (1425, 155, 1735, 280, fill_data, "经验池\n存储 (s,a,r,s')"),
        "expert": (90, 430, 390, 555, fill_data, "专家样本\nB_E=(S_E,A_E)"),
        "disc": (535, 430, 835, 555, fill_disc, "判别器 D\n专家标签 1\n生成样本标签 0"),
        "batch": (1425, 430, 1735, 555, fill_data, "同一批经验样本\nB=(S,A)"),
        "rint": (980, 430, 1280, 555, fill_disc, "模仿奖励 R_int\n由 D(S,A) 转换"),
        "fuse": (980, 710, 1280, 835, fill_value, "融合奖励\nR=R_env+w_gail*R_int"),
        "critic": (535, 710, 835, 835, fill_value, "Critic\n计算 TD-Error\n并更新"),
        "actor_update": (90, 710, 390, 835, fill_value, "Actor 更新\n最大化 Critic\n估计价值"),
    }
    for x1, y1, x2, y2, color, text in boxes.values():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=color, outline=line, width=3)
        draw_centered_text(draw, (x1, y1, x2, y2), text, box_font)

    arrow(draw, (390, 218), (535, 218))
    arrow(draw, (835, 218), (980, 218))
    arrow(draw, (1280, 218), (1425, 218))
    arrow(draw, (1580, 280), (1580, 430))
    arrow(draw, (1425, 493), (1280, 493))
    arrow(draw, (390, 493), (535, 493))
    arrow(draw, (835, 493), (980, 493))
    arrow(draw, (1130, 555), (1130, 710))
    arrow(draw, (980, 773), (835, 773))
    arrow(draw, (535, 773), (390, 773))
    draw.text((1160, 640), "R_env 与 R_int 在同一 TD 目标中融合", font=small_font, fill="#666666")
    draw.text((600, 875), "同一个经验 Batch 同时用于判别器打分和 Critic 目标值计算", font=small_font, fill="#666666")
    image.save(TRAIN_FLOW_FIG, quality=95)


def extract_sim_rule_figure_from_previous_docx() -> None:
    if SIM_RULE_FIG.exists():
        return
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    candidates = sorted(OUT_DIR.glob("*v4.docx"))
    if not candidates:
        return
    chosen_bytes = None
    chosen_ext = ".png"
    with zipfile.ZipFile(candidates[-1]) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            data = archive.read(name)
            try:
                with Image.open(BytesIO(data)) as image:
                    width, height = image.size
                    ratio = width / max(height, 1)
                    if width < 1200 and 1.4 <= ratio <= 2.4:
                        chosen_bytes = data
                        chosen_ext = Path(name).suffix
                        break
            except Exception:
                continue
    if chosen_bytes is not None:
        target = SIM_RULE_FIG.with_suffix(chosen_ext)
        target.write_bytes(chosen_bytes)
        if target != SIM_RULE_FIG:
            with Image.open(target) as image:
                image.save(SIM_RULE_FIG)


def add_intro(doc: Document) -> None:
    add_heading(doc, "1  引言", 1)
    add_body(
        doc,
        "理论经济学研究经常需要对经济运行过程进行抽象建模，并以模型推演解释经济系统中的一般性规律。"
        "这类研究通常需要较多经济运行数据作为支撑，但是现实经济活动数据的获取成本较高，数据来源受到真实环境和采集条件的限制。"
        "许多理论经济模型中涉及的变量，例如社会总产品、商品流通速度等，是理论层面的抽象表达，并不能在真实的经济环境中加以采集。"
        "因此，单纯依赖现实数据支撑理论经济学建模会遇到明显的数据约束。"
    )
    add_body(
        doc,
        "为了弥补现实数据的不足，研究者可以利用仿真方法构建虚拟经济运行环境，并根据研究需要生成和提取虚拟经济数据。"
        "经过一段时间的发展，这类仿真方法中的多主体仿真（MAS: Multi-agent Simulation）逐渐成为主流形式。"
        "MAS 是由多个相互作用的智能主体（agent）组成的系统。"
        "将 MAS 方法应用于理论经济学研究领域，可以形成复杂经济系统多主体仿真（MAS-E: Multi-agent Simulation of Complex Economic System）。"
        "在 MAS-E 中，经济主体通过多轮交互改变系统状态，系统状态又会影响主体后续决策。"
    )
    add_body(
        doc,
        "MAS-E 的一个关键问题在于智能主体的决策驱动机制。"
        "传统的固定规则方法依赖人工设定，模型灵活性和自适应能力有限。"
        "强化学习（RL: Reinforcement Learning）方法能够根据环境反馈训练决策策略，但是其训练目标通常围绕奖励最大化展开，容易使主体表现出偏向完全理性的行为模式。"
        "真实市场主体的决策受到风险偏好、经验判断、有限信息和行为习惯等因素影响，因此真实决策不能简单等同于单一奖励目标下的最优结果。"
    )
    add_body(
        doc,
        "基于上述背景，模仿学习（IL: Imitation Learning）可以作为强化学习在经济主体决策生成中的补充方法。"
        "模仿学习以专家行为数据为参照，目标是使生成行为接近专家数据中体现的决策模式。"
        "本文在原有多主体经济仿真系统基础上，围绕生产企业主体的决策模型进行改进。"
        "一方面，本文将生成对抗模仿学习（GAIL: Generative Adversarial Imitation Learning）与双延迟深度确定性策略梯度算法（TD3: Twin Delayed Deep Deterministic Policy Gradient）结合，利用判别器提供与专家行为相关的模仿信号。"
        "另一方面，本文在扩展模型中将 Transformer 编码器（Transformer Encoder）加入生产企业策略网络的状态表征环节，使 Actor 在输出动作前能够结合当前观测与历史状态序列。"
        "该扩展仍然保持 GAIL 与 Actor-Critic 结构结合的训练方式，即判别器结果不直接作为 Actor 的监督目标，而是通过 Critic 对价值估计的影响间接作用于策略更新。"
    )
    add_body(
        doc,
        "通过上述改进，本文旨在为理论经济学研究突破“数据困境”提供支持。"
        "本文也希望通过这一具体实践，获得在有限理性决策生成问题上运用模仿学习方法的指导性结论，并为同类领域研究提供参考。"
    )


def add_environment_and_problem(doc: Document) -> None:
    add_heading(doc, "2  仿真环境与主体决策建模问题", 1)
    add_heading(doc, "2.1  仿真环境", 2)
    add_body(
        doc,
        "本文沿用既有复杂经济系统仿真环境的建模方案[1]。"
        "该环境包含企业主体和银行主体两类抽象微观经济主体。"
        "企业主体具有生产企业和消费企业两个实例，生产企业产出生产资料，消费企业产出消费品（劳动力的等效表示）。"
        "本文将企业所生产的生产资料及消费品统称为“产品”，生产资料及消费品被投放到市场中进行流通时也被称为“商品”。"
        "银行主体具有一个实例，银行通过贷款发放和本息回收与企业形成信用关系。"
        "企业根据自身持有的生产资料和消费品进行生产，并将产品投放至市场。"
        "企业再通过市场交易购买商品，以支持下一轮生产。"
    )
    if SIM_RULE_FIG.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(SIM_RULE_FIG), width=Inches(5.4))
        add_caption(doc, "图2-1  初步设计的仿真运行规则示意图")
    add_body(
        doc,
        "如图2-1所示，在仿真过程中，虚拟经济主体被置于新的环境中进行交互，并推动环境状态不断变化。"
        "当系统触发终止条件时，一次完整的仿真运行结束。"
        "本文将这样一个大周期称为一个“回合”。"
        "终止条件包括以下三种情况：某个企业破产，市场上已经无法购买到支撑后续生产的某种商品，或者仿真达到预制的天数上限 T。"
        "一个回合结束后，整体环境与各经济主体的个体属性会被重置，并开始新的回合。"
    )
    add_body(
        doc,
        "在每个回合中，系统按照含有 6 个阶段（P1-P6）的小周期循环运行。"
        "本文将一个小周期称为“一天”。"
        "一天中的 6 个阶段依次为 P1 企业决策阶段、P2 银行决策阶段、P3 贷款发放阶段、P4 商品交易阶段、P5 生产阶段和 P6 结算阶段。"
        "由于本文重点是主体决策模型的训练方案，具体生产公式、交易规则和银行约束沿用既有研究[1]，本文不展开重复推导。"
    )

    add_heading(doc, "2.2  主体决策建模", 2)
    add_body(
        doc,
        "决策模型是智能主体建立“从观察到行动”映射的计算结构。"
        "在本文所沿用的仿真环境中，Actor 是主体执行决策的策略网络。"
        "Actor 以主体在当前时刻能够观测到的状态向量为输入，并输出连续动作向量。"
        "动作向量随后进入仿真环境，并通过贷款、交易、生产和结算过程影响下一时刻状态。"
    )
    add_body(
        doc,
        "根据既有环境定义与当前代码结构，企业主体的状态输入包含企业经营状态、生产要素需求信息、市场交易记录、市场价格信息和时间状态变量等信息，维度为 33。"
        "企业主体的动作输出为 4 维，分别对应贷款意愿（WNDF）、生产资料购买意愿（K）、消费品购买意愿（L）和次日定价（P）。"
        "银行主体的状态输入同时包含银行经营状态和企业主体信息，维度为 51。"
    )
    add_body_runs(
        doc,
        [
            ("银行主体的动作输出为 2 维，对应向两类企业发放贷款的决策量。其中，WNDB", False),
            ("A", True),
            (" 表示银行向生产企业发放贷款的意愿，WNDB", False),
            ("B", True),
            (" 表示银行向消费企业发放贷款的意愿。", False),
        ],
    )
    add_caption(doc, "表2-1  企业与银行的状态与动作空间")
    add_table(
        doc,
        ["主体", "Actor 输入", "Actor 输出"],
        [
            [
                "企业主体",
                "33 维状态向量，包含经营状态、需求信息、交易记录、价格信息和时间状态变量",
                "4 维连续动作：贷款意愿、生产资料购买意愿、消费品购买意愿、次日定价",
            ],
            [
                "银行主体",
                "51 维状态向量，包含银行经营状态与企业主体信息",
                "2 维连续动作：向生产企业和消费企业发放贷款的意愿",
            ],
        ],
    )
    add_body(
        doc,
        "在进行上述决策时，主体能够利用的信息主要来自自身经营状态、上一阶段决策结果以及可观测的市场交互信息。"
        "如果模型需要进一步利用历史信息，则可以在 Actor 的状态表征环节引入时序编码模块。"
        "该模块不会改变动作空间本身，而是改变 Actor 接收状态信息的表达方式，使策略网络能够在生成动作前利用更长时间窗口内的交互信息。"
    )


def add_gail_td3_section_legacy(doc: Document) -> None:
    add_heading(doc, "3  专家数据采集与有限理性行为表示", 1)

    add_heading(doc, "4  基于生成对抗模仿学习与TD3的主体训练方案", 1)
    add_heading(doc, "4.1  模型训练方案", 2)
    add_body(
        doc,
        "本文的 GAIL+TD3 训练方案以生产企业主体为主要改进对象。"
        "在日内交互过程中，生产企业 Actor 从仿真环境获得状态，并生成连续动作。"
        "该 Actor 在 TD3 中是策略网络，在 GAIL 结构中也承担生成器（Generator）的作用。"
        "Actor 生成的动作首先作用于仿真环境，并得到环境奖励、下一状态和终止标记。"
        "同一动作也会与状态共同组成状态-动作样本，进入判别器与 Critic 的后续更新过程。"
    )
    add_body(
        doc,
        "当经验池达到学习起点后，算法在每一次网络参数更新阶段从经验池采样一个状态-动作批量 B=(S,A)。"
        "这个批量来自生产企业主体与仿真环境在线交互产生的经验。"
        "判别器将 B 作为生成样本，并将专家数据采样得到的 B_E=(S_E,A_E) 作为专家样本。"
        "判别器的任务是区分状态-动作对来自专家数据还是来自当前 Actor 生成的交互经验。"
        "判别器更新完成后，同一个 B 用于计算判别器打分，并将该打分转换为模仿奖励 R_int。"
        "因此，进入判别器作为生成样本的批量和进入 Critic 计算目标值的批量保持为同一个经验批量。"
    )
    add_body(
        doc,
        "生产企业 Critic 使用融合奖励计算 TD3 的 Critic 目标值。"
        "融合奖励由环境奖励 R_env 和加权模仿奖励 R_int 组成。"
        "模仿奖励不会直接作为 Actor 的监督标签，也不会把专家动作直接送入 Actor 损失函数。"
        "它通过 TD3 的 Critic 目标值影响价值估计，再由 Actor 最大化 Critic 估计价值完成策略更新。"
        "这种结构保留了 TD3 的在线强化学习能力，同时让判别器对专家行为模式的判断通过价值网络间接影响生产企业策略。"
    )
    if TRAIN_FLOW_FIG.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(TRAIN_FLOW_FIG), width=Inches(6.2))
        add_caption(doc, "图4-1  基于 GAIL+TD3 的生产企业主体训练流程示意图")

    add_heading(doc, "4.2  状态、动作与经验样本", 2)
    add_body(
        doc,
        "生产企业主体的状态向量维度为 33，动作向量维度为 4。"
        "专家数据文件中的每一条记录由状态字段和动作字段组成，其中前 33 个字段对应专家状态，后 4 个字段对应专家动作。"
        "在线交互过程中，生产企业 Actor 输出的动作与当前状态共同构成生成样本。"
        "这些生成样本一方面进入经验池，用于 TD3 的 Critic 和 Actor 更新；另一方面在网络参数更新阶段被取出，用作判别器的生成样本。"
    )
    add_body(
        doc,
        "为了使专家样本和生成样本在判别器中具有一致的数值尺度，本文根据专家数据计算状态均值、状态方差、动作均值和动作方差。"
        "专家状态、专家动作以及经验池采样得到的生成动作会按照这些统计量进行标准化。"
        "生产企业与环境交互时，输入 Actor 的状态也进行标准化处理。"
        "这种处理的目标是降低状态量纲差异对判别器和 Actor 输入分布的影响。"
    )
    add_body(
        doc,
        "经验池样本的使用方式需要在论文中明确区分。"
        "环境交互产生的状态、动作、环境奖励、下一状态和终止信息用于 TD3 更新。"
        "判别器只使用状态-动作对判断样本来源。"
        "因此，判别器的输入不是完整经验样本，而是从同一批经验样本中取出的状态和动作部分。"
        "这种写法可以避免“同一批生成样本”这一表述造成歧义。"
    )

    add_heading(doc, "4.3  判别器更新与模仿奖励构造", 2)
    add_body(
        doc,
        "判别器 D 的输入为状态-动作对 (s,a)，输出为一个未经过 Sigmoid 的 logits 值。"
        "判别器训练中，专家样本标签设为 1，生成样本标签设为 0。"
        "判别器通过二元交叉熵损失学习区分专家样本与当前生产企业 Actor 生成的样本。"
        "该过程可以表示为：L_D = -E_(s,a)~Expert[log D(s,a)] - E_(s,a)~Replay[log(1-D(s,a))]。"
    )
    add_body(
        doc,
        "判别器更新后，同一个从经验池采样得到的生产企业状态-动作批量 B=(S,A) 会再次输入判别器。"
        "代码使用 Sigmoid 函数将判别器 logits 转换到 0 到 1 的区间，并将该值作为模仿奖励 R_int。"
        "随后，算法按照 R = R_env + w_gail R_int 的形式构造融合奖励，其中 w_gail 是模仿奖励权重。"
        "在当前代码中，w_gail 还包含 warm-up 缩放项，用于避免训练初期判别器信号过强。"
    )
    add_body(
        doc,
        "TD3 的 Critic 目标值使用融合奖励和双 Critic 目标网络计算。"
        "其基本形式可以写为 y = R + γ(1-d) min(Q_1'(s', μ'(s')), Q_2'(s', μ'(s'))) ，其中 d 表示终止标记。"
        "Critic 根据目标值 y 与当前估计值之间的 TD-Error 更新参数。"
        "Actor 的更新目标是最大化 Critic 对当前策略动作的价值估计，等价于最小化 -E[Q_1(s, μ(s))]。"
        "因此，在本文实现中，专家信息通过判别器奖励和 Critic 目标值进入策略学习过程，而不是以专家动作监督 Actor 输出。"
    )

    add_heading(doc, "5  基于Transformer表征增强的GAIL+TD3主体训练方案", 1)
    add_heading(doc, "6  实验设置与结果分析", 1)


def add_gail_td3_section(doc: Document) -> None:
    add_heading(doc, "3  专家数据采集与有限理性行为表示", 1)

    add_heading(doc, "4  基于生成对抗模仿学习与TD3的主体训练方案", 1)
    add_heading(doc, "4.1  模型训练方案", 2)
    add_body(
        doc,
        "本文将生成对抗模仿学习（GAIL: Generative Adversarial Imitation Learning）引入生产企业主体的 TD3 训练过程。"
        "在该方案中，生产企业 Actor 同时承担 TD3 策略网络和 GAIL 生成器的作用。"
        "Actor 根据当前状态生成连续动作，动作一方面作用于仿真环境并产生环境奖励、下一状态和终止标记，另一方面与状态共同形成状态-动作样本。"
        "这些样本随后进入经验池，并在学习阶段同时服务于判别器训练和 Critic 目标值计算。"
    )
    add_body(
        doc,
        "图4-1给出了基于生成对抗模仿学习与 TD3 的主体训练方案。"
        "在该方案中，专家行为样本为判别器提供专家分布参照，生产企业在线交互样本构成生成样本。"
        "判别器根据两类状态-动作样本之间的差异产生模仿奖励。"
        "本文实现将该模仿奖励与环境奖励融合，并将融合奖励用于 TD3 的 Critic 目标值计算，使 GAIL 信号通过价值估计影响 Actor 更新。"
        "模仿学习信号进入策略学习过程的路径可以概括为“判别器—融合奖励—Critic—Actor”。"
    )
    if GAIL_TD3_REF_FIG.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(GAIL_TD3_REF_FIG), width=Inches(6.2))
        add_caption(doc, "图4-1  基于生成对抗模仿学习与TD3的主体训练方案")

    add_heading(doc, "4.2  状态、动作与经验样本", 2)
    add_body(
        doc,
        "生产企业主体的状态向量维度为 33，动作向量维度为 4。"
        "专家数据文件中的每一条记录由状态字段和动作字段组成，其中前 33 个字段对应专家状态，后 4 个字段对应专家动作。"
        "在线交互过程中，生产企业 Actor 输出的动作与当前状态共同构成生成样本。"
        "这些生成样本一方面进入经验池，用于 TD3 的 Critic 和 Actor 更新；另一方面在网络参数更新阶段被取出，用作判别器的生成样本。"
        "令一次网络参数更新阶段采样得到的经验池批量为："
    )
    add_formula(doc, r"\mathcal{B}=\{(s_i,a_i,r_i^{env},s'_i,d_i)\}_{i=1}^{N}")
    add_body(
        doc,
        "其中，s_i 表示第 i 个样本的生产企业状态，a_i 表示对应动作，r_i^{env} 表示环境奖励，s'_i 表示下一状态，d_i 表示终止标记，N 表示批量大小。"
        "判别器不使用完整经验样本，而只使用同一批量中的状态-动作部分作为生成样本："
    )
    add_formula(doc, r"\mathcal{B}_{G}=\{(s_i,a_i)\}_{i=1}^{N}")
    add_body(
        doc,
        "专家样本从上述专家数据文件中独立采样得到。"
        "专家批量定义为："
    )
    add_formula(doc, r"\mathcal{B}_{E}=\{(s_i^{E},a_i^{E})\}_{i=1}^{N}")
    add_body(
        doc,
        "其中，s_i^{E} 和 a_i^{E} 分别表示第 i 个专家状态和专家动作。"
        "本文还根据专家数据计算状态均值、状态方差、动作均值和动作方差，并用这些统计量对专家样本与生成样本进行标准化。"
        "这一处理用于降低量纲差异对判别器输入分布的影响。"
    )
    add_body(
        doc,
        "本文实现中的一个重要细节是批量对齐。"
        "每一次网络参数更新阶段只从经验池采样一次生产企业经验批量。"
        "该批量的状态-动作部分用于判别器中的生成样本，该批量的完整经验样本又用于 TD3 的 Critic 目标值计算。"
        "因此，判别器给出的模仿奖励与 Critic 使用的环境奖励和下一状态在样本索引上保持一致。"
    )

    add_heading(doc, "4.3  判别器目标与模仿奖励构造", 2)
    add_body(
        doc,
        "判别器 D_φ 的输入为经标准化处理的状态-动作对，输出为一个 logits 值。"
        "Sigmoid 函数将 logits 转换为区间在 0 到 1 之间的判别概率。"
        "判别器概率定义为："
    )
    add_formula(doc, r"D_{\phi}(s,a)=\sigma(f_{\phi}(s,a))")
    add_body(
        doc,
        "其中，f_φ(s,a) 表示判别器网络在参数 φ 下输出的 logits，σ(·) 表示 Sigmoid 函数，D_φ(s,a) 表示样本来自专家分布的判别概率。"
        "判别器训练中，专家样本标签设为 1，生成样本标签设为 0。"
        "判别器训练采用二元交叉熵目标，使专家状态-动作对对应较高的判别概率，使生产企业生成的状态-动作对对应较低的判别概率。"
        "判别器损失定义为："
    )
    add_formula(doc, r"L_D(\phi)=-\frac{1}{N}\sum_{i=1}^{N}\left[\log D_{\phi}(s_i^{E},a_i^{E})+\log\left(1-D_{\phi}(s_i,a_i)\right)\right]")
    add_body(
        doc,
        "该式中，L_D(φ) 表示判别器损失。"
        "第一项提高专家状态-动作对被判定为专家样本的概率，第二项降低生成状态-动作对被判定为专家样本的概率。"
        "生产企业生成样本对应的模仿奖励由判别器概率直接给出："
    )
    add_formula(doc, r"r_i^{int}=D_{\phi}(s_i,a_i)=\sigma(f_{\phi}(s_i,a_i))")
    add_body(
        doc,
        "其中，r_i^{int} 表示第 i 个生成样本对应的模仿奖励。"
        "该奖励表示判别器认为生成状态-动作对接近专家分布的概率。"
        "在本文训练流程中，生产企业主体与仿真环境交互产生的经验首先进入经验池。"
        "每次网络参数更新阶段从经验池采样一个经验批量，并将其中的状态-动作对作为生成样本。"
        "同一批生成样本与专家样本共同用于判别器参数更新，随后判别器对该批生成样本输出模仿奖励。"
        "模仿奖励与环境奖励形成融合奖励，并进入 TD3 的 Critic 目标值计算；Actor 的参数更新由 Critic 对当前策略动作的价值估计驱动。"
        "本文采用在线判别器训练方式，判别器参数随生产企业主体交互样本和策略变化持续更新。"
    )

    add_heading(doc, "4.4  融合奖励与TD3参数更新", 2)
    add_body(
        doc,
        "生产企业 Critic 使用环境奖励与模仿奖励构造融合奖励。"
        "本文采用随训练步数递增的 warm-up 模仿奖励权重。"
        "在训练早期，判别器尚未形成稳定的区分边界；如果模仿奖励权重过大，判别器的不可靠打分会放大 Critic 目标值偏移，并通过价值估计影响 Actor 更新。"
        "因此，训练初期使用较小的模仿奖励权重，随后逐步提升至基础权重。"
        "融合奖励定义为："
    )
    add_formula(doc, r"\tilde{r}_i=r_i^{env}+w_t^{GAIL}r_i^{int}")
    add_formula(doc, r"w_t^{GAIL}=\lambda_{GAIL}\cdot \min\left(1,\frac{n_t}{n_{warm}}\right)")
    add_body(
        doc,
        "其中，\\tilde{r}_i 表示融合奖励，r_i^{env} 表示仿真环境给出的奖励，r_i^{int} 表示判别器产生的模仿奖励，w_t^{GAIL} 表示第 t 次网络参数更新阶段采用的模仿奖励权重。"
        "\\lambda_{GAIL} 表示模仿奖励基础权重，n_t 表示第 t 次更新时超过学习起点后的训练步数，n_{warm} 表示 warm-up 长度。"
        "在融合奖励确定后，目标 Actor 根据下一状态生成动作，并加入裁剪后的目标策略平滑噪声；该动作随后被限制在动作边界内。"
        "目标 Critic 根据下一状态和目标动作输出两个 Q 值估计，TD3 取两个估计值中的较小值，并与融合奖励共同构造当前 Critic 更新所需的目标值。"
        "目标动作与目标值定义如下："
    )
    add_formula(doc, r"\tilde{a}'_i=\operatorname{clip}\left(\mu_{\bar{\psi}}(s'_i)+\operatorname{clip}(\epsilon_i,-c,c),-a_{\max},a_{\max}\right)")
    add_formula(doc, r"y_i=\tilde{r}_i+\gamma(1-d_i)\min_{k=1,2}Q_{\bar{\theta}_k}(s'_i,\tilde{a}'_i)")
    add_body(
        doc,
        "其中，\\mu_{\\bar{\\psi}} 表示目标 Actor，Q_{\\bar{\\theta}_k} 表示目标 Critic 模块中的第 k 个 Q 值估计分支，k=1,2。"
        "该目标值形式沿用 TD3 中取双 Q 估计较小值的目标值构造方式。"
        "\\epsilon_i 表示目标策略平滑噪声，c 表示噪声裁剪边界，a_{\\max} 表示动作边界，\\gamma 表示折扣因子。"
        "d_i 为终止标记；当样本为终止状态时，后续状态价值不再进入目标值。"
        "Critic 根据当前 Q 值与目标值之间的误差更新参数："
    )
    add_formula(doc, r"L_Q(\theta_1,\theta_2)=\frac{1}{N}\sum_{i=1}^{N}\sum_{k=1}^{2}\ell\left(Q_{\theta_k}(s_i,a_i),y_i\right)")
    add_body(
        doc,
        "其中，L_Q 表示 Critic 损失，Q_{\\theta_k} 表示在线 Critic 模块中的第 k 个 Q 值估计分支，k=1,2。"
        "\\ell(\\cdot) 表示本文采用的 Critic 误差函数。"
        "在 Critic 更新之后，生产企业 Actor 通过最大化当前 Critic 对其策略动作的价值估计完成参数更新。"
        "Actor 损失函数定义为："
    )
    add_formula(doc, r"L_{\mu}(\psi)=-\frac{1}{N}\sum_{i=1}^{N}Q_{\theta_1}(s_i,\mu_{\psi}(s_i))")
    add_body(
        doc,
        "其中，L_\\mu 表示 Actor 损失，\\mu_{\\psi} 表示当前训练的生产企业 Actor，即生产企业主体的策略网络。"
        "在该训练方案中，生产企业 Actor 的梯度信号来自 Critic 对当前策略动作的价值估计。"
        "专家动作在该训练方案中用于构造判别器的专家样本，并通过判别器奖励间接进入策略学习过程。"
        "判别器信息通过融合奖励进入 TD3 的 Critic 目标值计算，并由价值估计影响 Actor 的策略更新。"
        "目标网络采用软更新形式："
    )
    add_formula(doc, r"\bar{\theta}_k\leftarrow \tau\theta_k+(1-\tau)\bar{\theta}_k,\quad \bar{\psi}\leftarrow \tau\psi+(1-\tau)\bar{\psi}")
    add_body(
        doc,
        "其中，\\tau 表示软更新系数。"
    )

    add_heading(doc, "5  基于Transformer表征增强的GAIL+TD3主体训练方案", 1)
    add_heading(doc, "6  实验设置与结果分析", 1)


def add_references(doc: Document) -> None:
    add_heading(doc, "参考文献", 1)
    refs = [
        '[1] B. Chen, S. Ding, Y. Lin and G. Chen, "Multi-agent Simulation of Complex Economic Systems Driven by Deep Reinforcement Learning," 2025 IEEE 17th International Conference on Computer Research and Development (ICCRD), Shangrao, China, 2025, pp. 47-54, doi: 10.1109/ICCRD64588.2025.10963160.'
    ]
    for ref in refs:
        add_body(doc, ref, first_line=False)


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    extract_sim_rule_figure_from_previous_docx()

    doc = Document()
    setup_document(doc)
    add_intro(doc)
    add_environment_and_problem(doc)
    add_gail_td3_section(doc)
    add_references(doc)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
