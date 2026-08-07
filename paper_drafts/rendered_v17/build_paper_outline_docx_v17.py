from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v16"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v16.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v17.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "E7D0C71EA040D45FF9A79EB92FD88781E8F59D2798FC4540FAA3088A09FBE964"
)


FIGURE_6_1_ANALYSIS = (
    "从图6-1可以看出，GAIL+TD3的各项训练曲线整体早于TD3上升，并较早进入相对稳定阶段。"
    "在平均存活天数方面，GAIL+TD3在训练前期快速提高，并在第20个评估时刻之前进入相对稳定区间；"
    "TD3的提升过程相对缓慢，直至训练后期才接近GAIL+TD3的平均水平。"
    "在主体收益方面，GAIL+TD3较早提高生产企业、消费企业和银行的累计收益，说明生产企业主体所获得的模仿学习信号会通过产品交易与信贷关系影响其他经济主体。"
    "TD3在训练后期的生产企业和消费企业累计收益均值高于GAIL+TD3，但其对应的SEM阴影也更宽，反映出两种训练方案在前期学习速度、后期收益水平和跨随机种子波动之间存在差异。"
    "银行累计收益则呈现不同结果，GAIL+TD3在训练前期完成提升后保持较高均值，而TD3的提升时间较晚，训练后期均值仍低于GAIL+TD3。"
    "图6-1(e)进一步显示，第60个评估时刻GAIL+TD3与TD3的累计归一化学习曲线下面积均值分别为0.683和0.435，前者在相同训练区间内形成了更大的存活天数累积面积。"
    "综合来看，GAIL+TD3能够加快系统存活天数和主体收益的前期提升，并使训练曲线较早进入相对稳定阶段；"
    "TD3在部分企业收益指标上具有较高的后期均值，说明学习速度、系统存活水平和主体收益水平之间需要结合具体评价目标进行分析。"
)


FIGURE_6_2_ANALYSIS = (
    "从图6-2可以看出，Transformer+GAIL+TD3在多项指标上的训练曲线整体早于GAIL+TD3上升。"
    "在平均存活天数方面，Transformer+GAIL+TD3更早进入快速提升阶段，并在训练后期保持较高的平均水平，说明历史状态序列表征能够为生产企业主体提供当前状态之外的近期交互信息。"
    "在企业收益方面，Transformer+GAIL+TD3的生产企业累计收益提升更早，稳定阶段均值也高于GAIL+TD3；"
    "其消费企业累计收益在多数评估时刻同样较高，表明生产企业策略的变化会经由商品交易和资金流转影响消费企业的经营结果。"
    "与此同时，Transformer+GAIL+TD3在部分评估时刻的企业收益SEM阴影较宽，说明不同随机种子下的收益结果仍存在一定波动。"
    "银行累计收益的变化方向与企业收益并不完全一致，Transformer+GAIL+TD3在训练前期出现较大波动，进入相对稳定阶段后的均值低于GAIL+TD3，反映出企业经营策略变化对信贷关系产生了间接影响。"
    "图6-2(e)显示，第60个评估时刻Transformer+GAIL+TD3与GAIL+TD3的累计归一化学习曲线下面积均值分别为0.773和0.683，前者在相同训练区间内形成了更大的存活天数累积面积。"
    "综合来看，Transformer编码器所提供的历史状态表征改善了系统存活水平和企业经营收益，并进一步提高了模型的前期学习速度；"
    "银行收益及跨随机种子波动的结果也表明，不同主体目标之间仍然存在表现差异。"
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph, text: str) -> None:
    source_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            source_rpr = deepcopy(run._r.rPr)
            break

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def merge_figure_analysis(
    document: Document,
    first_prefix: str,
    remaining_prefixes: tuple[str, ...],
    replacement: str,
) -> None:
    first = find_paragraph(document, first_prefix)
    replace_paragraph_text(first, replacement)
    for prefix in remaining_prefixes:
        remove_paragraph(find_paragraph(document, prefix))


def verify(document: Document) -> None:
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    expected = [
        FIGURE_6_1_ANALYSIS,
        FIGURE_6_2_ANALYSIS,
        "图6-3显示，TD3、GAIL+TD3和Transformer+GAIL+TD3中DSCR小于1",
    ]
    for fragment in expected:
        if fragment not in all_text:
            raise RuntimeError(f"Missing expected text: {fragment[:40]!r}")

    forbidden_prefixes = (
        "图6-1(b)显示",
        "图6-1(c)显示",
        "图6-1(d)显示",
        "图6-1(e)显示",
        "图6-2(b)显示",
        "图6-2(c)显示",
        "图6-2(d)显示",
        "图6-2(e)显示",
    )
    for prefix in forbidden_prefixes:
        if any(paragraph.text.startswith(prefix) for paragraph in document.paragraphs):
            raise RuntimeError(f"Separate subfigure paragraph remains: {prefix!r}")

    figure_6_1_result = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("从图6-1可以看出")
    ]
    figure_6_2_result = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("从图6-2可以看出")
    ]
    if len(figure_6_1_result) != 1 or len(figure_6_2_result) != 1:
        raise RuntimeError("Each learning-curve figure must have one result paragraph.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v16 source DOCX changed after the v17 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)
    merge_figure_analysis(
        document,
        "图6-1(a)显示",
        (
            "图6-1(b)显示",
            "图6-1(c)显示",
            "图6-1(d)显示",
            "图6-1(e)显示",
        ),
        FIGURE_6_1_ANALYSIS,
    )
    merge_figure_analysis(
        document,
        "图6-2(a)显示",
        (
            "图6-2(b)显示",
            "图6-2(c)显示",
            "图6-2(d)显示",
            "图6-2(e)显示",
        ),
        FIGURE_6_2_ANALYSIS,
    )
    verify(document)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v16 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
