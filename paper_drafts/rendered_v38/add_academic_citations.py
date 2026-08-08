from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
V4_PATH = next((ROOT / "paper_drafts" / "reference").glob("*_v4.docx"))
V38_PATH = next((ROOT / "paper_drafts" / "rendered_v38").glob("*v38.docx"))


REFERENCES = (
    (
        "[1] Drechsler M. Improving models of coordination incentives for "
        "biodiversity conservation by fitting a multi-agent simulation model "
        "to a lab experiment[J]. Journal of Behavioral and Experimental "
        "Economics, 2023, 102: 101967."
    ),
    (
        "[2] Arulkumaran K, Deisenroth M P, Brundage M, et al. Deep "
        "Reinforcement Learning: A Brief Survey[J]. IEEE Signal Processing "
        "Magazine, 2017, 34(6): 26-38."
    ),
    (
        "[3] Moerland T M, Broekens J, Plaat A, et al. Model-based "
        "Reinforcement Learning: A Survey[J]. Foundations and Trends in "
        "Machine Learning, 2023, 16(1): 1-118."
    ),
    (
        "[4] 刘全, 翟建伟, 章宗长, 等. 深度强化学习综述[J]. 计算机学报, "
        "2018, 41(1): 1-27."
    ),
    (
        "[5] Zheng B, Verma S, Zhou J, et al. Imitation Learning: Progress, "
        "Taxonomies and Challenges[J]. IEEE Transactions on Neural Networks "
        "and Learning Systems, 2024, 35(5): 6322-6337."
    ),
    (
        "[6] Ho J, Ermon S. Generative Adversarial Imitation Learning[C]//"
        "Advances in Neural Information Processing Systems. 2016, 29."
    ),
    (
        "[7] Fujimoto S, van Hoof H, Meger D. Addressing Function Approximation "
        "Error in Actor-Critic Methods[C]//Proceedings of the 35th International "
        "Conference on Machine Learning. PMLR, 2018, 80: 1587-1596."
    ),
    (
        "[8] Vaswani A, Shazeer N, Parmar N, et al. Attention Is All You Need[C]//"
        "Advances in Neural Information Processing Systems. 2017, 30."
    ),
    (
        "[9] B. Chen, S. Ding, Y. Lin and G. Chen, \"Multi-agent Simulation of "
        "Complex Economic Systems Driven by Deep Reinforcement Learning,\" 2025 "
        "IEEE 17th International Conference on Computer Research and Development "
        "(ICCRD), Shangrao, China, 2025, pp. 47-54, doi: "
        "10.1109/ICCRD64588.2025.10963160."
    ),
    (
        "[10] 中华人民共和国科学技术部等. “关于印发《科技伦理审查办法（试行）》的通知.” "
        "2023年10月8日, https://www.most.gov.cn/xxgk/xinxifenlei/fdzdgknr/"
        "fgzc/gfxwj/gfxwj2023/202310/t20231008_188309.html."
    ),
    (
        "[11] Lin J. Divergence Measures Based on the Shannon Entropy[J]. IEEE "
        "Transactions on Information Theory, 1991, 37(1): 145-151."
    ),
    (
        "[12] Inman H F, Bradley E L Jr. The Overlapping Coefficient as a Measure "
        "of Agreement Between Probability Distributions and Point Estimation "
        "of the Overlap of Two Normal Densities[J]. Communications in "
        "Statistics-Theory and Methods, 1989, 18(10): 3851-3874."
    ),
    (
        "[13] Fawcett T. An Introduction to ROC Analysis[J]. Pattern Recognition "
        "Letters, 2006, 27(8): 861-874."
    ),
)


CITATION_INSERTIONS = (
    ("为了弥补现实数据的不足", ((None, "[1]"),)),
    (
        "MAS-E 的一个关键问题在于智能主体的决策驱动机制",
        (("强化学习虽可依据环境反馈形成策略", "[2-4]"),),
    ),
    (
        "基于上述背景，模仿学习",
        (
            ("模仿学习（IL: Imitation Learning）", "[5]"),
            (
                "生成对抗模仿学习（GAIL: Generative Adversarial Imitation Learning）",
                "[6]",
            ),
            (
                "双延迟深度确定性策略梯度算法（TD3: Twin Delayed Deep Deterministic Policy Gradient）",
                "[7]",
            ),
            ("Transformer 编码器（Transformer Encoder）", "[8]"),
        ),
    ),
    ("本文沿用既有复杂经济系统仿真环境", (("建模方案", "[9]"),)),
    (
        "具体生产公式、交易规则和银行约束沿用既有研究",
        (("沿用既有研究", "[9]"),),
    ),
    (
        "多主体经济系统仿真的关键问题之一",
        (("解决现实或仿真环境中的决策问题。", "[2-4]"),),
    ),
    (
        "基于上述局限，本文进一步引入有限理性仿真目标",
        ((None, "[5]"),),
    ),
    ("判别器用于评价标准化后的专家状态", ((None, "[6]"),)),
    ("本节在第3.3节所定义模仿奖励的基础上", ((None, "[7]"),)),
    ("本文将线性序列映射、可学习位置嵌入", ((None, "[8]"),)),
    ("依据知情同意和数据保护原则", ((None, "[10]"),)),
    ("为定量评价判别器对专家行为与加噪专家行为", ((None, "[11]"),)),
    ("式中，M表示P和Q的等权混合分布", ((None, "[12]"),)),
    ("式中，h表示得分区间编号", ((None, "[13]"),)),
)


def set_run_text(run_element, text: str) -> None:
    for child in list(run_element):
        if child.tag != qn("w:rPr"):
            run_element.remove(child)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run_element.append(text_element)


def apply_superscript(run_element, *, highlight: bool) -> None:
    properties = run_element.find(qn("w:rPr"))
    if properties is None:
        properties = OxmlElement("w:rPr")
        run_element.insert(0, properties)

    for old in properties.findall(qn("w:vertAlign")):
        properties.remove(old)
    vertical_alignment = OxmlElement("w:vertAlign")
    vertical_alignment.set(qn("w:val"), "superscript")
    properties.append(vertical_alignment)

    if highlight:
        for old in properties.findall(qn("w:highlight")):
            properties.remove(old)
        highlight_element = OxmlElement("w:highlight")
        highlight_element.set(qn("w:val"), "yellow")
        properties.append(highlight_element)


def highlight_paragraph(paragraph: Paragraph) -> None:
    for run in paragraph._p.findall(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            run.insert(0, properties)
        for old in properties.findall(qn("w:highlight")):
            properties.remove(old)
        element = OxmlElement("w:highlight")
        element.set(qn("w:val"), "yellow")
        properties.append(element)


def paragraph_has_superscript(paragraph: Paragraph, marker: str) -> bool:
    for run in paragraph._p.findall(qn("w:r")):
        text = "".join(node.text or "" for node in run.findall(qn("w:t")))
        if marker not in text:
            continue
        properties = run.find(qn("w:rPr"))
        if properties is None:
            return False
        vertical_alignment = properties.find(qn("w:vertAlign"))
        return (
            vertical_alignment is not None
            and vertical_alignment.get(qn("w:val")) == "superscript"
        )
    return False


def replace_marker_with_superscript(
    paragraph: Paragraph, marker: str, *, highlight: bool
) -> None:
    if paragraph_has_superscript(paragraph, marker):
        return

    for run in paragraph._p.findall(qn("w:r")):
        text = "".join(node.text or "" for node in run.findall(qn("w:t")))
        if marker not in text:
            continue
        before, after = text.split(marker, 1)
        set_run_text(run, before)

        citation_run = deepcopy(run)
        set_run_text(citation_run, marker)
        apply_superscript(citation_run, highlight=highlight)
        run.addnext(citation_run)

        if after:
            trailing_run = deepcopy(run)
            set_run_text(trailing_run, after)
            citation_run.addnext(trailing_run)
        return
    raise RuntimeError(f"Citation marker {marker!r} was not found in paragraph")


def append_superscript_citation(
    paragraph: Paragraph, marker: str, *, highlight: bool
) -> None:
    if marker in paragraph.text:
        replace_marker_with_superscript(paragraph, marker, highlight=highlight)
        return

    direct_runs = paragraph._p.findall(qn("w:r"))
    if direct_runs:
        citation_run = deepcopy(direct_runs[-1])
    else:
        citation_run = OxmlElement("w:r")
    set_run_text(citation_run, marker)
    apply_superscript(citation_run, highlight=highlight)
    paragraph._p.append(citation_run)


def insert_superscript_citation_after_text(
    paragraph: Paragraph, anchor: str, marker: str, *, highlight: bool
) -> None:
    runs = paragraph._p.findall(qn("w:r"))
    run_texts = [
        "".join(node.text or "" for node in run.findall(qn("w:t")))
        for run in runs
    ]
    full_text = "".join(run_texts)
    if anchor not in full_text:
        raise RuntimeError(
            f"Citation insertion anchor {anchor!r} was not found in paragraph"
        )

    insertion_offset = full_text.index(anchor) + len(anchor)
    traversed = 0
    for run, text in zip(runs, run_texts):
        run_end = traversed + len(text)
        if insertion_offset > run_end:
            traversed = run_end
            continue

        local_offset = insertion_offset - traversed
        citation_run = deepcopy(run)
        set_run_text(citation_run, marker)
        apply_superscript(citation_run, highlight=highlight)

        if local_offset == 0:
            run.addprevious(citation_run)
        elif local_offset == len(text):
            run.addnext(citation_run)
        else:
            before = text[:local_offset]
            after = text[local_offset:]
            set_run_text(run, before)
            run.addnext(citation_run)
            trailing_run = deepcopy(run)
            set_run_text(trailing_run, after)
            citation_run.addnext(trailing_run)
        return

    raise RuntimeError("Unable to resolve the citation insertion position")


def find_unique_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}"
        )
    return matches[0]


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    source_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        source_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = OxmlElement("w:r")
    if source_properties is not None:
        run.append(source_properties)
    set_run_text(run, text)
    paragraph._p.append(run)


def insert_reference_after(reference: Paragraph, text: str) -> Paragraph:
    element = deepcopy(reference._p)
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    replace_paragraph_text(paragraph, text)
    return paragraph


def ensure_reference_list(document: Document, *, highlight: bool) -> None:
    heading = find_unique_paragraph(document, "参考文献")
    existing = {
        p.text.strip().split(" ", 1)[0]: p
        for p in document.paragraphs
        if p._p.getparent() is heading._p.getparent()
        and p.text.strip().startswith("[")
    }
    anchor = existing.get("[1]")
    if anchor is None:
        raise RuntimeError("Reference [1] was not found")

    for index, text in enumerate(REFERENCES, start=1):
        label = text.split(" ", 1)[0]
        paragraph = existing.get(label)
        if paragraph is None:
            paragraph = insert_reference_after(anchor, text)
        else:
            replace_paragraph_text(paragraph, text)
        if highlight:
            highlight_paragraph(paragraph)
        if index == 1:
            anchor = paragraph
        else:
            anchor = paragraph


def ensure_body_citations(document: Document, *, highlight: bool) -> None:
    reference_heading = find_unique_paragraph(document, "参考文献")
    body_paragraphs = []
    for paragraph in document.paragraphs:
        if paragraph._p is reference_heading._p:
            break
        body_paragraphs.append(paragraph)

    known_markers = {
        "[1]", "[2]", "[3]", "[4]", "[5]", "[6]", "[7]", "[8]", "[9]",
        "[10]", "[11]", "[12]", "[13]", "[1-2]", "[2-4]", "[3-4]",
        "[4-7]", "[5-6]", "[6-9]",
    }

    for phrase, insertions in CITATION_INSERTIONS:
        matches = [
            paragraph
            for paragraph in body_paragraphs
            if phrase in paragraph.text
        ]
        if len(matches) != 1:
            raise RuntimeError(
                f"Expected one citation target containing {phrase!r}, "
                f"found {len(matches)}"
            )
        paragraph = matches[0]
        citation_runs = []
        for run in paragraph._p.findall(qn("w:r")):
            run_text = "".join(
                node.text or "" for node in run.findall(qn("w:t"))
            )
            if run_text in known_markers:
                citation_runs.append(run)
        for citation_run in citation_runs:
            paragraph._p.remove(citation_run)

        for anchor, marker in insertions:
            if anchor is None:
                append_superscript_citation(
                    paragraph, marker, highlight=highlight
                )
            else:
                insert_superscript_citation_after_text(
                    paragraph, anchor, marker, highlight=highlight
                )


def revise_document(path: Path, *, highlight: bool) -> None:
    document = Document(path)
    ensure_body_citations(document, highlight=highlight)
    ensure_reference_list(document, highlight=highlight)
    temporary = path.with_suffix(".citations.tmp.docx")
    document.save(temporary)
    os.replace(temporary, path)


def main() -> None:
    revise_document(V4_PATH, highlight=True)
    revise_document(V38_PATH, highlight=False)
    print(V4_PATH)
    print(V38_PATH)


if __name__ == "__main__":
    main()
