from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import importlib.util
from pathlib import Path
import sys

BUNDLED_SITE_PACKAGES = Path(
    r"C:\Users\legion\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\python\Lib\site-packages"
)
if str(BUNDLED_SITE_PACKAGES) not in sys.path:
    sys.path.append(str(BUNDLED_SITE_PACKAGES))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from docx import Document
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

from apply_teacher_feedback_layout_and_statistics import revise_document_contents
from add_academic_citations import ensure_body_citations, ensure_reference_list


ROOT = Path(__file__).resolve().parents[2]
SOURCE_DIR = ROOT / "paper_drafts" / "rendered_v37"
SOURCE = next(
    path
    for path in SOURCE_DIR.glob("*.docx")
    if not path.name.startswith("~$")
)
EXPECTED_SOURCE_SHA256 = (
    "F11C827E5EAEB26454946BB4499352643002F8C6FF23F5BFB4CF9BE8790C80EB"
)
OUTPUT = Path(__file__).resolve().parent / (
    "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v38.docx"
)

V37_BUILDER = SOURCE_DIR / "build_paper_outline_docx_v37.py"
FIGURE_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "paper_gail_component_validation_v38"
)
ACTOR_JS_FIGURE = FIGURE_DIR / "actor_expert_js_with_noisy_reference.png"
ACTOR_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "actor_snapshot_evolution"
)
NOISY_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "expert_vs_noisy_expert_discriminator_eval_noise_std_0p30"
)


GENERIC_MULTI_SEED_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mover><mi>M</mi><mo>¯</mo></mover><mi>k</mi></msub><mo>=</mo>
  <mfrac><mn>1</mn><mi>S</mi></mfrac>
  <munderover><mo>∑</mo><mrow><mi>r</mi><mo>=</mo><mn>1</mn></mrow><mi>S</mi></munderover>
  <msub><mi>M</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub>
  <mo>,</mo><mspace width="1em"/>
  <msub><mi>SEM</mi><mi>k</mi></msub><mo>=</mo>
  <mfrac><msub><mi>s</mi><mi>k</mi></msub><msqrt><mi>S</mi></msqrt></mfrac>
</math>
"""

RELATIVE_JS_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mi>ρ</mi><mi>k</mi></msub><mo>=</mo>
  <mfrac>
    <msubsup><mover><mi>J</mi><mo>¯</mo></mover><mi>k</mi><mi>G</mi></msubsup>
    <msup><mover><mi>J</mi><mo>¯</mo></mover><mi>N</mi></msup>
  </mfrac>
</math>
"""


def load_v37_helpers():
    spec = importlib.util.spec_from_file_location("paper_v37", V37_BUILDER)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load {V37_BUILDER}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def body_index(document: Document, paragraph: Paragraph) -> int:
    return list(document.element.body).index(paragraph._p)


def move_range_before(
    document: Document,
    start: Paragraph,
    end_exclusive: Paragraph,
    target: Paragraph,
) -> None:
    body = document.element.body
    children = list(body)
    start_index = children.index(start._p)
    end_index = children.index(end_exclusive._p)
    nodes = children[start_index:end_index]
    for node in nodes:
        target._p.addprevious(node)


def next_paragraph(paragraph: Paragraph) -> Paragraph:
    element = paragraph._p.getnext()
    while element is not None and element.tag != qn("w:p"):
        element = element.getnext()
    if element is None:
        raise RuntimeError("No following paragraph found.")
    return Paragraph(element, paragraph._parent)


def replace_text_nodes(paragraph: Paragraph, old: str, new: str) -> int:
    count = 0
    for node in paragraph._p.xpath(".//w:t"):
        if node.text and old in node.text:
            count += node.text.count(old)
            node.text = node.text.replace(old, new)
    return count


def normalize_state_action_dashes(document: Document) -> None:
    for p_element in document.element.body.xpath(".//w:p"):
        nodes = p_element.xpath(".//w:t")
        for index, node in enumerate(nodes):
            value = node.text or ""
            if "状态-动作" in value:
                node.text = value.replace("状态-动作", "状态—动作")
                continue
            if value != "-":
                continue
            before = "".join((item.text or "") for item in nodes[:index])
            after = "".join((item.text or "") for item in nodes[index + 1 :])
            if before.endswith("状态") and after.startswith("动作"):
                node.text = "—"


def math_run(text: str, *, normal: bool = False):
    run = OxmlElement("m:r")
    if normal:
        math_rpr = OxmlElement("m:rPr")
        math_rpr.append(OxmlElement("m:nor"))
        run.append(math_rpr)
    word_rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Cambria Math")
    fonts.set(qn("w:hAnsi"), "Cambria Math")
    word_rpr.append(fonts)
    run.append(word_rpr)
    token = OxmlElement("m:t")
    token.text = text
    run.append(token)
    return run


def math_n_eval():
    subscript = OxmlElement("m:sSub")
    base = OxmlElement("m:e")
    base.append(math_run("N"))
    sub = OxmlElement("m:sub")
    sub.append(math_run("eval", normal=True))
    subscript.append(base)
    subscript.append(sub)
    return subscript


def replace_naulc_total_count_symbol(document: Document) -> None:
    indicator = find_paragraph(document, "（3）累计归一化学习曲线下面积")
    equation = next_paragraph(indicator)
    matches = [node for node in equation._p.xpath(".//m:t") if node.text == "K-1"]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one K-1 token, found {len(matches)}")
    old_run = matches[0].getparent()
    parent = old_run.getparent()
    position = parent.index(old_run)
    parent.remove(old_run)
    parent.insert(position, math_n_eval())
    parent.insert(position + 1, math_run("-1"))

    explanation = next_paragraph(equation)
    direct_runs = explanation._p.findall(qn("w:r"))
    target = None
    for index, run in enumerate(direct_runs[:-1]):
        texts = run.findall(qn("w:t"))
        next_texts = direct_runs[index + 1].findall(qn("w:t"))
        if (
            len(texts) == 1
            and texts[0].text == "K"
            and next_texts
            and (next_texts[0].text or "").startswith("表示图中纳入比较的评估时刻总数")
        ):
            target = run
            break
    if target is None:
        raise RuntimeError("Could not locate explanatory K in the nAULC paragraph")
    target.find(qn("w:t")).text = "N"
    eval_run = deepcopy(target)
    eval_run.find(qn("w:t")).text = "eval"
    rpr = eval_run.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        eval_run.insert(0, rpr)
    vert_align = OxmlElement("w:vertAlign")
    vert_align.set(qn("w:val"), "subscript")
    rpr.append(vert_align)
    target.addnext(eval_run)


def apply_consistency_updates(document: Document, helpers) -> None:
    helpers.replace_paragraph_text(
        find_paragraph(document, "MAS-E 的一个关键问题在于智能主体的决策驱动机制"),
        "MAS-E 的一个关键问题在于智能主体的决策驱动机制。固定规则方法的灵活性和自适应能力有限；强化学习虽可依据环境反馈形成策略，但以奖励最大化为核心的训练目标与真实市场主体受有限信息和经验影响的决策方式存在差异。如何生成更贴近有限理性特征的主体行为，因而成为经济系统仿真需要解决的问题。",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "为简化记号，本节后续公式以状态s和动作a表示"),
        "为简化记号，本节后续公式以状态s和动作a表示已经完成标准化处理的判别器输入。记判别器在参数φ下对状态—动作向量输出的logits为f_φ(s,a)，判别概率定义为：",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "其中，r_i^int表示第i个生成样本对应的模仿奖励"),
        "其中，r_i^int表示第i个生成样本对应的模仿奖励，ε表示防止对数输入为零的正数下界，r_clip^int表示模仿奖励裁剪上限。当判别器对生成状态—动作对给出更高的专家样本判别概率时，该样本获得更高的模仿奖励。判别器参数更新时，算法分别从专家数据训练集和近期生产企业交互样本中抽取状态—动作批量，判别器参数随生产企业交互样本和策略变化持续更新。",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "本文仅将存活天数大于90天的完整回合"),
        "人工采集轨迹与Codex辅助采集轨迹采用统一的数据格式保存，后续筛选条件及整理结果见第5.4节。",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "第一步验证评价判别器对专家行为与受控非专家行为"),
        "第一步验证训练结束时判别器对专家行为与受控非专家行为的区分能力。本文将第r次独立运行在训练结束时依据验证集保存并固定参数的判别器称为评价判别器，本次评价共选取3个评价判别器。每次评价均使用相同的留出专家测试集、得分区间和直方图划分。本文在16952条专家状态—动作样本的动作分量上叠加标准差为0.30的独立高斯噪声，并将动作裁剪至[-0.5,0.5]，由此构造加噪专家样本，作为受控非专家行为参照。留出测试集包含173个完整专家回合。加噪专家动作定义如下：",
    )
    scoring = find_paragraph(document, "式中，i表示留出测试集中的样本序号")
    if replace_text_nodes(scoring, "固定评价判别器", "评价判别器") != 1:
        raise RuntimeError("Expected one fixed-evaluation-discriminator phrase")
    helpers.replace_paragraph_text(
        find_paragraph(document, "综合图6-4、图6-5和表6-2的结果"),
        "综合图6-4、图6-5和表6-2的结果，Transformer+GAIL+TD3以较少的训练回合达到目标存活水平，并在稳定阶段平均存活水平、训练全过程的累积平均存活表现和生产企业债务偿付状态等方面取得了更好的统计结果。历史状态序列表征使生产企业主体能够利用跨时段交互信息，并进一步发挥专家行为信息对策略形成的引导作用。",
    )
    replace_naulc_total_count_symbol(document)
    normalize_state_action_dashes(document)


def find_next_table(paragraph: Paragraph) -> Table:
    element = paragraph._p.getnext()
    while element is not None:
        if element.tag == qn("w:tbl"):
            return Table(element, paragraph._parent)
        element = element.getnext()
    raise RuntimeError(f"No table found after {paragraph.text!r}")


def remove_paragraph_and_following_equation(
    helpers, paragraph: Paragraph
) -> None:
    equation = next_paragraph(paragraph)
    if not equation._p.xpath(".//m:oMath"):
        raise RuntimeError(f"Expected equation after {paragraph.text!r}")
    helpers.remove_paragraph(equation)


def configure_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "font.size": 11,
            "axes.labelsize": 12,
            "xtick.labelsize": 10.5,
            "ytick.labelsize": 10.5,
            "axes.linewidth": 1.1,
            "savefig.dpi": 300,
        }
    )


def create_actor_js_figure() -> tuple[float, float, float, float]:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    actor = pd.read_csv(ACTOR_RESULT_DIR / "actor_expert_js_mean_sem.csv")
    noisy = pd.read_csv(NOISY_RESULT_DIR / "three_seed_expert_noise_summary.csv")
    noisy_values = noisy["js_divergence"].to_numpy(dtype=float)
    noisy_mean = float(noisy_values.mean())
    noisy_sem = float(noisy_values.std(ddof=1) / np.sqrt(noisy_values.size))

    stable = actor.tail(20)
    actor_stable_mean = float(stable["mean"].mean())
    actor_stable_sem = float(stable["sem"].mean())

    x = actor["evaluation_step"].to_numpy(dtype=float)
    mean = actor["mean"].to_numpy(dtype=float)
    lower = actor["mean_minus_sem"].to_numpy(dtype=float)
    upper = actor["mean_plus_sem"].to_numpy(dtype=float)

    configure_matplotlib()
    figure, axis = plt.subplots(figsize=(6.4, 4.05))
    axis.fill_between(x, lower, upper, color="#4c78a8", alpha=0.22, linewidth=0)
    axis.plot(x, mean, color="#1f4e79", linewidth=1.9)
    axis.axhspan(
        noisy_mean - noisy_sem,
        noisy_mean + noisy_sem,
        color="#777777",
        alpha=0.10,
        linewidth=0,
    )
    axis.axhline(
        noisy_mean,
        color="#555555",
        linewidth=1.4,
        linestyle=(0, (5, 3)),
    )
    axis.text(
        59.2,
        noisy_mean + noisy_sem + 0.012,
        "Noisy-expert reference",
        ha="right",
        va="bottom",
        fontsize=10.2,
        color="#444444",
    )
    axis.set_xlabel("Evaluation step (100 episodes)")
    axis.set_ylabel("Actor-expert JS divergence")
    axis.set_xlim(1, 60)
    axis.set_ylim(0.0, max(0.7, float(upper.max()) * 1.03))
    axis.grid(True, color="#b8b8b8", linestyle="--", linewidth=0.65, alpha=0.70)
    for spine in axis.spines.values():
        spine.set_color("black")
        spine.set_linewidth(1.1)
    figure.tight_layout()
    figure.savefig(ACTOR_JS_FIGURE, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return noisy_mean, noisy_sem, actor_stable_mean, actor_stable_sem


def update_table_cell(helpers, cell, text: str) -> None:
    helpers.replace_paragraph_text(cell.paragraphs[0], text)


def revise_chapter_six(document: Document, helpers) -> None:
    body_template = find_paragraph(
        document, "本文首先比较TD3与GAIL+TD3的训练过程"
    )
    equation_template = next_paragraph(
        find_paragraph(document, "（1）平均存活天数")
    )

    helpers.replace_paragraph_text(
        find_paragraph(document, "6.1 实验设置与评价指标"),
        "6.1 实验设置与统计方法",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "本节实验从模仿学习信号与历史状态表征两个层面"),
        "本节实验从模仿学习信号与历史状态表征两个层面，通过比较不同主体决策模型在复杂经济系统仿真中的训练表现，分析两类机制对生产企业主体决策形成过程及系统运行结果的影响。本文设置两组对比实验：第一组比较TD3与GAIL+TD3，用于考察模仿学习信号对生产企业主体运行效果与行为模仿效果的影响；第二组比较GAIL+TD3与Transformer+GAIL+TD3，用于考察历史状态序列表征对生产企业主体训练过程的影响。三种模型中，TD3作为基线模型；GAIL+TD3在生产企业主体的TD3训练过程中引入生成对抗模仿学习信号；Transformer+GAIL+TD3在GAIL+TD3基础上加入Transformer编码器。本次实验中的模仿学习与历史表征改进均应用于生产企业主体，消费企业主体和银行主体在三种模型中保持既有TD3决策流程。GAIL+TD3和Transformer+GAIL+TD3使用第五章整理得到的同一组专家数据。上述实验分别从系统运行结果和行为分布两个层面，评价生产企业主体能否在有限可观测信息与有限训练经验下形成接近专家行为的决策模式。",
    )

    metric_intro = find_paragraph(document, "在上述训练过程中，本文跟踪")
    helpers.remove_paragraph(metric_intro)

    metric_1 = find_paragraph(document, "（1）平均存活天数")
    metric_5 = find_paragraph(document, "（5）生产企业偿债能力")
    metric_6 = find_paragraph(document, "（6）生成对抗模仿学习有效性")
    seed_stats = find_paragraph(document, "神经网络参数初始化、动作探索")
    comparison_intro = find_paragraph(
        document, "本文首先比较TD3与GAIL+TD3的训练过程"
    )
    first_discriminator_step = find_paragraph(
        document, "第一步检验判别器对专家行为与非专家行为的区分能力"
    )
    dscr_analysis = find_paragraph(
        document, "为进一步考察三种主体训练方案下生产企业的当期偿债风险"
    )

    move_range_before(document, metric_1, metric_5, comparison_intro)
    move_range_before(document, metric_5, metric_6, dscr_analysis)
    move_range_before(document, metric_6, seed_stats, first_discriminator_step)

    helpers.replace_paragraph_text(
        seed_stats,
        "为降低单次运行中随机因素对模型比较结果的影响，本文采用多随机种子（multi-seed）统计方法汇总训练曲线。神经网络参数初始化、动作探索、训练样本抽取和经验池采样等过程均包含随机因素。本文将用于控制上述随机因素的联合配置统称为“随机种子”。TD3与GAIL+TD3分别在8组随机种子下独立运行，Transformer+GAIL+TD3的统计分析纳入7组独立运行。设M_(r,k)表示第r次独立运行在第k个评估时刻的某项指标值，同一模型在该评估时刻的跨随机种子均值及均值标准误定义如下：",
    )
    generic_equation = next_paragraph(seed_stats)
    helpers.replace_equation_paragraph(generic_equation, GENERIC_MULTI_SEED_MATHML)
    helpers.replace_paragraph_text(
        find_paragraph(document, "式中，s表示同一模型各次独立运行结果的样本标准差"),
        "式中，S表示相应模型纳入统计的独立运行次数；r表示独立运行编号；k表示评估时刻编号；带横线的M表示同一模型在第k个评估时刻的跨随机种子均值；s_k表示相应指标在该评估时刻的样本标准差。训练曲线中的实线表示跨随机种子均值，阴影表示均值上下1个SEM。后续各项评价均沿用该多随机种子统计方法。涉及稳定阶段表现时，本文先对每次独立运行最后20个评估时刻的指标取均值，再按上述方法汇总不同随机种子的结果。",
    )

    for table in document.tables:
        if not table.rows:
            continue
        if table.rows[0].cells[0].text.strip() == "项目":
            row_names = [row.cells[0].text.strip() for row in table.rows[1:]]
            if "总训练回合数" in row_names and "统计方法" in row_names:
                statistics_row = next(
                    row
                    for row in table.rows[1:]
                    if row.cells[0].text.strip() == "统计方法"
                )
                table._tbl.append(deepcopy(statistics_row._tr))
                update_table_cell(helpers, statistics_row.cells[0], "单回合最大天数")
                update_table_cell(helpers, statistics_row.cells[1], "100天")
                new_statistics_row = table.rows[-1]
                update_table_cell(helpers, new_statistics_row.cells[0], "统计方法")
                update_table_cell(
                    helpers,
                    new_statistics_row.cells[1],
                    "多随机种子均值及均值上下1个SEM",
                )

    metric_1 = find_paragraph(document, "（1）平均存活天数")
    metric_context = helpers.insert_paragraph_before(
        metric_1,
        "本小节从系统运行效果层面比较TD3与GAIL+TD3。平均存活天数和生产企业累计收益分别描述系统运行时长与生产企业经营结果，累计归一化学习曲线下面积用于汇总完整训练区间内的累积平均存活表现，达到目标存活水平所需的评估时刻用于衡量学习速度。各项指标定义如下。",
        format_source=body_template,
    )

    helpers.replace_paragraph_text(
        find_paragraph(document, "（2）累计收益"),
        "（2）生产企业累计收益。本文首先在单个回合内汇总生产企业的收入、支出与利息，再对第k个评估时刻所含E个回合的累计收益取算术平均值，其定义如下：",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "式中，J表示企业在对应评估时刻的平均累计收益"),
        "式中，J表示生产企业在对应评估时刻的平均累计收益；R、C和I分别表示单个回合内的累计收入、累计支出和累计利息；i表示生产企业编号；k和n分别表示评估时刻编号和该评估时刻内的回合序号。",
    )
    bank_metric = find_paragraph(document, "生产企业和消费企业分别依据上述定义")
    bank_equation = next_paragraph(bank_metric)
    bank_explanation = next_paragraph(bank_equation)
    helpers.remove_paragraph(bank_metric)
    helpers.remove_paragraph(bank_equation)
    helpers.remove_paragraph(bank_explanation)

    speed_metric = find_paragraph(document, "（4）学习速度")
    helpers.replace_paragraph_text(
        speed_metric,
        "（4）学习速度。本文采用平均存活天数的跨随机种子均值曲线达到目标存活水平所需的评估时刻衡量学习速度。平均存活天数均值按照6.1所述多随机种子统计方法计算。",
    )
    speed_mean_equation = next_paragraph(speed_metric)
    if not speed_mean_equation._p.xpath(".//m:oMath"):
        raise RuntimeError("The learning-speed mean equation was not found.")
    helpers.remove_paragraph(speed_mean_equation)
    helpers.replace_paragraph_text(
        find_paragraph(document, "式中，S表示对应模型纳入统计的独立运行次数"),
        "式中，k表示评估时刻编号；μ_k表示第k个评估时刻平均存活天数的跨随机种子均值；q表示目标存活天数，本文分别取80天和90天；T_q表示均值曲线达到目标q所需的评估时刻。较小的T_q表示模型以较少的训练回合达到目标存活水平。若均值曲线截至第60个评估时刻仍未满足连续3个评估时刻均不低于目标值的条件，则记为“未达到”。",
    )

    helpers.replace_paragraph_text(
        comparison_intro,
        "本文首先比较TD3与GAIL+TD3的训练过程，以分析模仿学习信号对生产企业主体决策训练的影响。图6-1给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标上的跨随机种子均值曲线及SEM阴影。其中，图6-1(a)中的水平虚线表示表5-1所列专家轨迹的合计平均存活天数，其数值为97.99天。为定量比较三种模型的训练表现，本文对每次独立运行最后20个评估时刻的平均存活天数和生产企业累计收益分别取均值，再按照6.1所述方法计算跨随机种子均值和SEM；累计归一化学习曲线下面积采用第60个评估时刻的数值；学习速度采用平均存活天数均值曲线达到80天和90天所需的评估时刻。表6-2汇总了三种模型的相应统计结果。",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "表6-2 三种主体训练方案的最终性能统计"),
        "表6-2 三种模型的稳定阶段运行效果与学习速度统计",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "注：平均存活天数和生产企业累计收益"),
        "注：平均存活天数和生产企业累计收益为最后20个评估时刻的稳定阶段统计结果；括号内为相对SEM，K表示10³；“未达到”表示截至第60个评估时刻，均值曲线尚未连续3个评估时刻达到相应目标值。",
    )

    helpers.replace_paragraph_text(
        find_paragraph(document, "（2）生成对抗模仿学习有效性"),
        "（2）生成对抗模仿学习的模仿效果评价",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "前述运行结果表明，引入生成对抗模仿学习后"),
        "第6.2节第（1）部分的运行结果表明，引入生成对抗模仿学习后，生产企业主体能够以较少的训练回合达到较高的平均存活水平。为判断该运行效果是否对应着对专家行为的有效模仿，本文进一步从行为分布层面对模仿效果进行评价。专家数据来源于训练前采集的历史仿真轨迹，在线训练过程不存在与各评估时刻Actor动作逐一配对的同步专家动作，因而无法直接计算相同决策条件下的动作差异。本文采用参数固定的判别器将不同来源的状态—动作样本映射至统一得分空间，并以得分分布差异评价Actor生成行为与专家行为的接近程度。由于该评价依赖判别器输出，本文首先验证评价判别器对专家行为与受控非专家行为的区分能力，再使用相同判别器评价历史Actor的行为分布变化。",
    )

    metric_js = find_paragraph(document, "（6）生成对抗模仿学习有效性")
    helpers.replace_paragraph_text(
        metric_js,
        "上述两步评价均以判别器得分及其分布为分析对象。为分别刻画两组得分分布的差异程度、重叠程度以及判别器对两类样本的排序能力，本文采用Jensen–Shannon散度、重叠系数和受试者工作特征曲线下面积作为评价指标。Jensen–Shannon散度（JS divergence）用于度量两组判别器得分概率分布之间的差异。设P和Q表示两组归一化判别器得分分布，其定义如下：",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "式中，M表示P和Q的等权混合分布"),
        "式中，M表示P和Q的等权混合分布，KL表示Kullback–Leibler散度。JS散度越小表示两组判别器得分分布越接近。在衡量分布差异的基础上，本文采用重叠系数（OVL: Overlap Coefficient）描述两组归一化直方图的共同区域，其定义如下：",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "式中，h表示得分区间编号"),
        "式中，h表示得分区间编号，H表示区间总数，P_h和Q_h分别表示两组样本落入第h个区间的概率。OVL取值范围为[0,1]，数值越大表示两组得分分布的重叠程度越高。为进一步考察判别器对两类样本的排序能力，本文采用受试者工作特征曲线下面积（AUC: Area Under the Receiver Operating Characteristic Curve）进行评价，其定义如下：",
    )

    helpers.replace_paragraph_text(
        first_discriminator_step,
        "第一步验证评价判别器对专家行为与受控非专家行为的区分能力。本文选取3次独立运行在训练结束时依据验证集保存的判别器，并分别固定其网络参数。每次评价均使用相同的留出专家测试集、得分区间和直方图划分。本文在16952条专家状态—动作样本的动作分量上叠加标准差为0.30的独立高斯噪声，并将动作裁剪至[-0.5,0.5]，由此构造加噪专家样本，作为受控非专家行为参照。留出测试集包含173个完整专家回合。加噪专家动作定义如下：",
    )
    score_explanation = find_paragraph(document, "式中，D表示判别器；hat{s}_i")
    helpers.replace_paragraph_text(
        score_explanation,
        "式中，D表示判别器；hat{s}_i表示第i个专家状态；带星号的φ_r表示第r次独立运行中保存并固定用于评价的判别器参数；d_i^E和d_i^N分别表示第i个专家样本与加噪专家样本的判别器得分。图6-2以其中一次独立运行为例给出两组样本的判别器得分分布。为使加噪专家参照与历史Actor评价保持一致，本文采用相同的3个评价判别器、同一留出专家测试集以及[0,1]区间内100个固定得分区间。按照前述定义，本文以J_r^N表示第r个评价判别器下专家得分分布P_r^E与加噪专家得分分布P_r^N之间的JS散度，并将其作为加噪专家参照。本文按照6.1所述多随机种子统计方法汇总3次独立运行的评价结果。",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "图6-2 专家样本与加噪专家样本的判别器得分分布"),
        "图6-2 一组独立运行中专家样本与加噪专家样本的判别器得分分布",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "表6-3 判别器对专家样本与加噪专家样本的评价结果"),
        "表6-3 三个评价判别器对专家样本与加噪专家样本的评价结果",
    )

    table_6_3 = find_next_table(
        find_paragraph(document, "表6-3 三个评价判别器")
    )
    rows = (
        ("统计量", "三次独立运行均值（±SEM）"),
        ("专家样本平均评分", "0.8989（±0.0005）"),
        ("加噪专家样本平均评分", "0.6736（±0.0402）"),
        ("平均评分差", "0.2253（±0.0398）"),
        ("JS散度", "0.3542（±0.0282）"),
        ("OVL", "27.87%（±5.08%）"),
        ("AUC", "0.7808（±0.0329）"),
    )
    if len(table_6_3.rows) != len(rows):
        raise RuntimeError("Unexpected Table 6-3 row count.")
    for row, values in zip(table_6_3.rows, rows):
        update_table_cell(helpers, row.cells[0], values[0])
        update_table_cell(helpers, row.cells[1], values[1])

    helpers.replace_paragraph_text(
        find_paragraph(document, "图6-2与表6-3显示"),
        "图6-2展示的一组独立运行中，评价判别器给予原始专家样本的得分整体高于加噪专家样本。表6-3进一步汇总了3次独立运行的评价结果：专家样本与加噪专家样本的平均评分分别为0.8989±0.0005和0.6736±0.0402，平均评分差为0.2253±0.0398；两组得分分布的JS散度为0.3542±0.0282，表明两组得分分布并不重合；OVL为27.87%±5.08%，表示两组归一化得分直方图的共同区域约占27.87%；以专家样本为正类的AUC为0.7808±0.0329，表示评价判别器将随机抽取的专家样本排在加噪专家样本之前的概率约为78.08%。上述结果表明，3个评价判别器均能对原始专家样本与加噪专家样本形成有方向的得分区分，因此可以作为后续历史Actor行为分布评价的固定尺度。",
    )

    helpers.replace_paragraph_text(
        find_paragraph(document, "第二步评价历史Actor生成行为在训练过程中的模仿效果"),
        "第二步评价历史Actor生成行为在训练过程中的模仿效果。本文沿用第一步中的3次独立运行及其评价判别器，并加载每100个训练回合保存的历史Actor快照。所有历史Actor均在同一批留出专家状态上以确定性方式生成动作，不加入探索噪声。对于第r次独立运行的第k个评估时刻，历史Actor生成的动作及其判别器得分定义如下：",
    )
    actor_parameter_explanation = find_paragraph(document, "式中，μ表示生产企业Actor")
    helpers.replace_paragraph_text(
        actor_parameter_explanation,
        "式中，μ表示生产企业Actor；θ_(r,k)表示第r次独立运行在第k个评估时刻保存的Actor参数；上标G表示Actor生成动作或相应判别器得分。在同一记号体系下，本文以J_(r,k)^G表示专家得分分布P_r^E与第r次独立运行第k个评估时刻的历史Actor得分分布P_(r,k)^G之间的JS散度。",
    )
    actor_specific_js_equation = next_paragraph(actor_parameter_explanation)
    if not actor_specific_js_equation._p.xpath(".//m:oMath"):
        raise RuntimeError("The repeated Actor-specific JS equation was not found.")
    helpers.remove_paragraph(actor_specific_js_equation)
    actor_js_explanation = find_paragraph(document, "式中，J_(r,k)表示第r次独立运行")
    helpers.replace_paragraph_text(
        actor_js_explanation,
        "较小的J_(r,k)^G表示历史Actor生成行为在固定判别器得分表示下更接近专家样本。本文在每个评估时刻按照6.1所述多随机种子统计方法计算3次独立运行的JS散度均值及其SEM。",
    )
    actor_mean_equation = next_paragraph(actor_js_explanation)
    if not actor_mean_equation._p.xpath(".//m:oMath"):
        raise RuntimeError("The repeated Actor mean/SEM equation was not found.")
    helpers.remove_paragraph(actor_mean_equation)
    actor_curve_explanation = find_paragraph(document, "式中，带横线的J表示第k个评估时刻")
    helpers.replace_paragraph_text(
        actor_curve_explanation,
        "图6-3给出了60个评估时刻的Actor—专家JS散度跨随机种子均值，蓝色阴影表示均值上下1个SEM。图中的灰色虚线及其阴影分别表示同一组评价判别器下加噪专家参照JS散度的跨随机种子均值及均值上下1个SEM。为比较两类JS散度，本文进一步定义相对JS指标：",
    )
    relative_equation = helpers.insert_equation_after(
        actor_curve_explanation,
        RELATIVE_JS_MATHML,
        equation_template=equation_template,
    )
    helpers.insert_paragraph_after(
        relative_equation,
        "式中，带横线的J_k^G表示第k个评估时刻的Actor—专家JS散度均值，带横线的J^N表示加噪专家参照JS散度均值。ρ_k小于1表示在相同评价判别器与得分空间下，Actor生成样本的得分分布比加噪专家样本更接近原始专家样本。",
        format_source=body_template,
    )

    noisy_mean, noisy_sem, actor_stable_mean, actor_stable_sem = (
        create_actor_js_figure()
    )
    helpers.replace_picture_before_caption(
        document,
        "图6-3 Actor与专家样本判别器得分分布的JS散度变化",
        ACTOR_JS_FIGURE,
        width_inches=6.35,
    )
    ratio = actor_stable_mean / noisy_mean
    reduction = (0.5800147962167418 - actor_stable_mean) / 0.5800147962167418 * 100
    reference_reduction = (noisy_mean - actor_stable_mean) / noisy_mean * 100
    helpers.replace_paragraph_text(
        find_paragraph(document, "图6-3显示，Actor—专家JS散度均值"),
        f"图6-3显示，Actor—专家JS散度均值在第1个评估时刻为0.580±0.041，随后总体下降。最后20个评估时刻的稳定阶段均值为{actor_stable_mean:.3f}±{actor_stable_sem:.3f}，较第1个评估时刻降低约{reduction:.1f}%。同一评价口径下，加噪专家参照JS散度为{noisy_mean:.3f}±{noisy_sem:.3f}；稳定阶段的相对JS指标为{ratio:.3f}，即Actor—专家JS散度比加噪专家参照低约{reference_reduction:.1f}%。该结果表明，训练后期Actor生成行为的判别器得分分布比受控加噪专家行为更接近原始专家样本。",
    )
    helpers.replace_paragraph_text(
        find_paragraph(document, "两步评价分别从判别器对专家与非专家行为的区分能力"),
        "两步评价分别考察了评价判别器对专家行为与受控非专家行为的区分能力，以及历史Actor生成行为与专家行为的分布接近程度。评价判别器能够对原始专家样本与加噪专家样本形成稳定的得分区分；历史Actor的行为分布在训练过程中逐步接近专家样本，并在稳定阶段达到低于加噪专家参照的JS散度。上述结果从行为分布层面支持了生成对抗模仿学习的模仿效果，也表明生产企业主体能够在有限可观测状态和有限训练经验下形成更接近专家行为分布的决策模式，从而为有限理性决策行为的生成提供实验依据。",
    )

    helpers.replace_paragraph_text(
        find_paragraph(document, "在完成上述生成对抗模仿学习有效性检验后"),
        "在完成上述生成对抗模仿学习模仿效果评价后，本文进一步比较GAIL+TD3与Transformer+GAIL+TD3，以分析历史状态序列表征对生产企业主体训练过程的影响。图6-4给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标上的跨随机种子均值曲线及SEM阴影，表6-2列出了对应的稳定阶段运行效果与学习速度统计。图6-4(a)中的水平虚线同样表示专家轨迹的合计平均存活天数，用于比较两种模型与专家轨迹存活水平之间的差距。",
    )

    dscr_metric = find_paragraph(document, "（5）生产企业偿债能力")
    helpers.replace_paragraph_text(
        dscr_metric,
        "在完成存活水平、学习速度和生产企业收益比较后，本文进一步从债务偿付角度评价三种主体训练方案下生产企业的经营状态。生产企业偿债能力采用偿债覆盖率（DSCR: Debt Service Coverage Ratio）度量。该指标以生产企业当日可支配现金与当期应还本息之比表示，其定义如下：",
    )
    helpers.replace_paragraph_text(
        dscr_analysis,
        "基于上述定义，本文进一步统计三种主体训练方案在稳定运行回合中的生产企业日观测，并计算DSCR小于1的观测占比，以考察生产企业的当期偿债风险。图6-5给出了三种方法的偿债风险统计结果。",
    )

    in_chapter_six = False
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("6  实验设置与结果分析"):
            in_chapter_six = True
        if paragraph.text.startswith("7  总结与展望"):
            break
        if not in_chapter_six:
            continue
        revised = paragraph.text.replace(
            "最终稳定平均存活天数", "稳定阶段平均存活天数"
        ).replace(
            "生产企业最终稳定累计收益", "生产企业稳定阶段累计收益"
        )
        if revised != paragraph.text:
            helpers.replace_paragraph_text(paragraph, revised)


def verify(document: Document) -> None:
    text = "\n".join(p.text for p in document.paragraphs)
    required = (
        "6.1 实验设置",
        "多随机种子（multi-seed）统计方法",
        "（1）生成对抗模仿学习的主体决策运行效果",
        "（2）生成对抗模仿学习的模仿效果评价",
        "表6-2 三种模型的稳定阶段运行效果与学习速度统计",
        "表6-3 三个评价判别器对专家样本与加噪专家样本的评价结果",
        "加噪专家参照JS散度为0.354±0.028",
        "稳定阶段的相对JS指标为0.728",
        "有限理性决策行为的生成提供实验依据",
        "生产企业偿债能力采用偿债覆盖率",
        "为定量评价判别器对专家行为与加噪专家行为的区分能力",
    )
    forbidden = (
        "6.1 实验设置与评价指标",
        "在上述训练过程中，本文跟踪平均存活天数",
        "（6）生成对抗模仿学习有效性",
        "（2）生成对抗模仿学习有效性",
        "第19个评估时刻达到0.172",
        "表6-2 三种主体训练方案的最终性能统计",
        "在完成上述生成对抗模仿学习有效性检验后",
        "6.1 实验设置与统计方法",
        "受控非专家行为",
        "非专家行为参照",
    )
    for phrase in required:
        if phrase not in text:
            raise RuntimeError(f"Required Chapter 6 content is missing: {phrase}")
    for phrase in forbidden:
        if phrase in text:
            raise RuntimeError(f"Obsolete Chapter 6 content remains: {phrase}")
    if len(document.tables) != 6:
        raise RuntimeError(f"Expected 6 tables, found {len(document.tables)}")
    if len(document.inline_shapes) != 15:
        raise RuntimeError(
            f"Expected 15 inline figures, found {len(document.inline_shapes)}"
        )

    chapter_6 = text[text.index("6  实验设置与结果分析") : text.index("7  总结与展望")]
    if chapter_6.count("Jensen–Shannon散度（JS divergence）用于度量") != 1:
        raise RuntimeError("The generic JS definition must appear exactly once.")
    if chapter_6.count("多随机种子（multi-seed）统计方法") != 1:
        raise RuntimeError("The multi-seed statistical method must be introduced once.")


def build() -> None:
    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The v37 source DOCX changed after this builder was prepared.")
    helpers = load_v37_helpers()
    document = Document(SOURCE)
    revise_chapter_six(document, helpers)
    apply_consistency_updates(document, helpers)
    revise_document_contents(document, highlight=False, transformer_runs=7)
    ensure_body_citations(document, highlight=False)
    ensure_reference_list(document, highlight=False)
    verify(document)
    document.save(OUTPUT)
    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The v37 source DOCX was modified unexpectedly.")
    verify(Document(OUTPUT))
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
