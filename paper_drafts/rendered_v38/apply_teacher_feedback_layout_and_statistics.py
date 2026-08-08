from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import os
import struct

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
V4_PATH = next((ROOT / "paper_drafts" / "reference").glob("*_v4.docx"))
V38_PATH = next((ROOT / "paper_drafts" / "rendered_v38").glob("*v38.docx"))
FIGURE_PATH = (
    ROOT
    / "paper_drafts"
    / "figures"
    / "transformer_three_branch_structure"
    / "transformer_three_branch_structure.png"
)

TITLE_TEXT = "基于生成对抗模仿学习与Transformer表征的复杂经济系统主体有限理性决策生成方法"
ABSTRACT_TEXT = (
    "复杂经济系统多主体仿真能够为理论经济研究提供虚拟数据支持，但强化学习主体通常围绕奖励最大化进行策略优化，"
    "难以充分反映真实市场主体的有限理性决策特征。针对这一问题，本文在既有复杂经济系统仿真模型基础上，将生成对抗"
    "模仿学习（GAIL）与双延迟深度确定性策略梯度（TD3）算法相结合，以专家行为数据为参照构建生产企业主体决策模型。"
    "实验结果表明，引入模仿学习能够提高主体的学习效率，并使生成行为逐步接近专家行为分布。此外，通过引入Transformer"
    "编码器表征历史交互信息，进一步改善了模型的学习速度和稳定阶段存活水平。本文探索了模仿学习在复杂经济系统主体决策"
    "生成中的应用，为有限理性决策生成及复杂经济系统多主体仿真研究提供参考。"
)
KEYWORDS_TEXT = "复杂经济系统；多主体仿真；生成对抗模仿学习；TD3；Transformer；有限理性"


def find_paragraph(document, predicate) -> Paragraph:
    for paragraph in document.paragraphs:
        if predicate(paragraph.text.strip()):
            return paragraph
    raise ValueError("Target paragraph was not found")


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    paragraph._p.append(run)


def insert_paragraph_after(reference: Paragraph, text: str) -> Paragraph:
    element = deepcopy(reference._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    set_paragraph_text(paragraph, text)
    return paragraph


def move_paragraph_block_after(
    document, start: Paragraph, end: Paragraph, target: Paragraph
) -> list[Paragraph]:
    paragraphs = document.paragraphs
    start_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph._p is start._p
    )
    end_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph._p is end._p
    )
    if end_index < start_index:
        raise ValueError("Invalid paragraph block")
    block = paragraphs[start_index : end_index + 1]
    for paragraph in reversed(block):
        target._p.addnext(paragraph._p)
    return block


def replace_text_in_section(document, start_heading: str, end_heading: str) -> list[Paragraph]:
    paragraphs = document.paragraphs
    start_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip().startswith(start_heading)
    )
    end_index = next(
        index for index, paragraph in enumerate(paragraphs[start_index + 1 :], start_index + 1)
        if paragraph.text.strip().startswith(end_heading)
    )
    changed = []
    replacements = (
        ("受控加噪专家行为", "加噪专家行为"),
        ("受控非专家行为", "加噪专家行为"),
        ("非专家行为参照", "加噪专家行为参照"),
        ("非专家行为", "加噪专家行为"),
    )
    for paragraph in paragraphs[start_index:end_index]:
        text = paragraph.text
        revised = text
        for old, new in replacements:
            revised = revised.replace(old, new)
        if revised != text:
            set_paragraph_text(paragraph, revised)
            changed.append(paragraph)
    return changed


def highlight_paragraph(paragraph: Paragraph) -> None:
    for run in paragraph._p.xpath(".//w:r | .//m:r"):
        run_properties = run.find(qn("w:rPr"))
        if run_properties is None:
            run_properties = OxmlElement("w:rPr")
            run.insert(0, run_properties)
        highlight = run_properties.find(qn("w:highlight"))
        if highlight is None:
            highlight = OxmlElement("w:highlight")
            run_properties.append(highlight)
        highlight.set(qn("w:val"), "yellow")


def set_run_font(run, *, east_asia: str, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run_properties = run._r.get_or_add_rPr()
    run_fonts = run_properties.find(qn("w:rFonts"))
    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, run_fonts)
    run_fonts.set(qn("w:ascii"), "Times New Roman")
    run_fonts.set(qn("w:hAnsi"), "Times New Roman")
    run_fonts.set(qn("w:eastAsia"), east_asia)


def add_front_matter_paragraph(document, first: Paragraph) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    first._p.addprevious(paragraph._p)
    return paragraph


def ensure_front_matter(document, *, highlight: bool) -> list[Paragraph]:
    paragraphs = document.paragraphs
    title = next((p for p in paragraphs if p.text.strip() == TITLE_TEXT), None)
    abstract = next((p for p in paragraphs if p.text.strip().startswith("摘要：")), None)
    keywords = next((p for p in paragraphs if p.text.strip().startswith("关键词：")), None)
    changed: list[Paragraph] = []

    if title is None or abstract is None or keywords is None:
        first = document.paragraphs[0]
        if title is None:
            title = add_front_matter_paragraph(document, first)
        if abstract is None:
            abstract = add_front_matter_paragraph(document, first)
        if keywords is None:
            keywords = add_front_matter_paragraph(document, first)

    title.clear()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run(TITLE_TEXT)
    set_run_font(title_run, east_asia="黑体", size=18, bold=True)

    abstract.clear()
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.first_line_indent = Pt(0)
    abstract.paragraph_format.line_spacing = 1.5
    abstract.paragraph_format.space_before = Pt(0)
    abstract.paragraph_format.space_after = Pt(6)
    abstract_label = abstract.add_run("摘要：")
    set_run_font(abstract_label, east_asia="黑体", size=10.5, bold=True)
    abstract_body = abstract.add_run(ABSTRACT_TEXT)
    set_run_font(abstract_body, east_asia="宋体", size=10.5)

    keywords.clear()
    keywords.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    keywords.paragraph_format.first_line_indent = Pt(0)
    keywords.paragraph_format.line_spacing = 1.5
    keywords.paragraph_format.space_before = Pt(0)
    keywords.paragraph_format.space_after = Pt(12)
    keywords_label = keywords.add_run("关键词：")
    set_run_font(keywords_label, east_asia="黑体", size=10.5, bold=True)
    keywords_body = keywords.add_run(KEYWORDS_TEXT)
    set_run_font(keywords_body, east_asia="宋体", size=10.5)

    changed.extend((title, abstract, keywords))
    if highlight:
        for paragraph in changed:
            highlight_paragraph(paragraph)
    return changed


def replace_text_nodes(paragraph: Paragraph, old: str, new: str) -> int:
    replacements = 0
    for text_element in paragraph._p.xpath(".//w:t"):
        if text_element.text and old in text_element.text:
            text_element.text = text_element.text.replace(old, new)
            replacements += 1
    return replacements


def append_text_run(paragraph_element, text: str, run_properties=None) -> None:
    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(deepcopy(run_properties))
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    paragraph_element.append(run)


def append_m_subscript(paragraph_element, base: str, subscript: str) -> None:
    math = OxmlElement("m:oMath")
    subscript_element = OxmlElement("m:sSub")
    base_element = OxmlElement("m:e")
    base_run = OxmlElement("m:r")
    base_text = OxmlElement("m:t")
    base_text.text = base
    base_run.append(base_text)
    base_element.append(base_run)
    sub_element = OxmlElement("m:sub")
    sub_run = OxmlElement("m:r")
    sub_text = OxmlElement("m:t")
    sub_text.text = subscript
    sub_run.append(sub_text)
    sub_element.append(sub_run)
    subscript_element.append(base_element)
    subscript_element.append(sub_element)
    math.append(subscript_element)
    paragraph_element.append(math)


def ensure_seed_statistics_paragraph(
    document, statistics_intro: Paragraph
) -> tuple[Paragraph, bool]:
    for paragraph in document.paragraphs:
        if "同一模型在该评估时刻的跨随机种子均值及均值标准误定义如下" in paragraph.text:
            return paragraph, False

    element = deepcopy(statistics_intro._p)
    run_properties = None
    if statistics_intro.runs and statistics_intro.runs[0]._r.rPr is not None:
        run_properties = statistics_intro.runs[0]._r.rPr
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    append_text_run(
        element,
        "本小节比较TD3与GAIL+TD3。两种模型均在8组随机种子下独立运行。"
        "对于任一评价指标，每次独立运行形成一条跨评估时刻的指标序列，"
        "因此每种模型对应8条指标序列。设",
        run_properties,
    )
    append_m_subscript(element, "M", "r,k")
    append_text_run(
        element,
        "表示第r次独立运行在第k个评估时刻的某项指标值，"
        "同一模型在该评估时刻的跨随机种子均值及均值标准误定义如下：",
        run_properties,
    )
    statistics_intro._p.addnext(element)
    return Paragraph(element, statistics_intro._parent), True


def revise_seed_scope(document, *, transformer_runs: int) -> list[Paragraph]:
    changed: list[Paragraph] = []
    statistics_paragraph = find_paragraph(
        document,
        lambda text: "同一模型在该评估时刻的跨随机种子均值及均值标准误定义如下" in text,
    )
    new_prefix = (
        "本小节比较TD3与GAIL+TD3。两种模型均在8组随机种子下独立运行。"
        "对于任一评价指标，每次独立运行形成一条跨评估时刻的指标序列，"
        "因此每种模型对应8条指标序列。"
    )
    if not statistics_paragraph.text.startswith(new_prefix):
        run_properties = None
        if statistics_paragraph.runs and statistics_paragraph.runs[0]._r.rPr is not None:
            run_properties = statistics_paragraph.runs[0]._r.rPr
        for child in list(statistics_paragraph._p):
            if child.tag != qn("w:pPr"):
                statistics_paragraph._p.remove(child)
        append_text_run(statistics_paragraph._p, new_prefix + "设", run_properties)
        append_m_subscript(statistics_paragraph._p, "M", "r,k")
        append_text_run(
            statistics_paragraph._p,
            "表示第r次独立运行在第k个评估时刻的某项指标值，"
            "同一模型在该评估时刻的跨随机种子均值及均值标准误定义如下：",
            run_properties,
        )
        changed.append(statistics_paragraph)

    comparison_paragraph = find_paragraph(
        document,
        lambda text: text.startswith("在完成上述生成对抗模仿学习模仿效果评价后"),
    )
    if transformer_runs == 8:
        run_sentence = (
            "该组比较中，GAIL+TD3与Transformer+GAIL+TD3均纳入8次独立运行，"
            "各项曲线与稳定阶段指标沿用第6.2节第（1）部分所述统计方法。"
        )
    else:
        run_sentence = (
            "该组比较中，GAIL+TD3纳入8次独立运行，"
            "Transformer+GAIL+TD3纳入7次独立运行，"
            "各项曲线与稳定阶段指标沿用第6.2节第（1）部分所述统计方法。"
        )
    if run_sentence not in comparison_paragraph.text:
        revised = comparison_paragraph.text.replace(
            "图6-4给出了",
            run_sentence + "图6-4给出了",
            1,
        )
        set_paragraph_text(comparison_paragraph, revised)
        changed.append(comparison_paragraph)
    return changed


def png_dimensions(path: Path) -> tuple[int, int]:
    data = path.read_bytes()
    if data[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError("Figure must be a PNG file")
    return struct.unpack(">II", data[16:24])


def replace_figure_4_1(document) -> None:
    image_paragraph = None
    paragraphs = document.paragraphs
    for caption_index, caption in enumerate(paragraphs):
        if not caption.text.strip().startswith("图4-1"):
            continue
        for paragraph in reversed(paragraphs[max(0, caption_index - 8) : caption_index]):
            if paragraph._p.xpath(".//a:blip"):
                image_paragraph = paragraph
                break
        if image_paragraph is not None:
            break
    if image_paragraph is None:
        raise ValueError("Figure 4-1 image paragraph was not found")

    new_blob = FIGURE_PATH.read_bytes()
    width_px, height_px = png_dimensions(FIGURE_PATH)
    for blip in image_paragraph._p.xpath(".//a:blip"):
        relationship_id = blip.get(qn("r:embed"))
        document.part.rels[relationship_id].target_part._blob = new_blob

    extent = image_paragraph._p.xpath(".//wp:extent")
    if extent:
        width_emu = int(extent[0].get("cx"))
        height_emu = round(width_emu * height_px / width_px)
        for element in extent:
            element.set("cy", str(height_emu))
        for element in image_paragraph._p.xpath(".//a:xfrm/a:ext"):
            element.set("cx", str(width_emu))
            element.set("cy", str(height_emu))


def revise_document_contents(
    document, *, highlight: bool, transformer_runs: int
) -> None:
    changed: list[Paragraph] = []
    changed.extend(ensure_front_matter(document, highlight=False))

    heading_61 = find_paragraph(document, lambda text: text.startswith("6.1 "))
    set_paragraph_text(heading_61, "6.1 实验设置")
    changed.append(heading_61)

    statistics_start = find_paragraph(
        document,
        lambda text: "神经网络参数初始化、动作探索、训练样本抽取和经验池采样" in text
        or text.startswith("为降低单次运行中随机因素对模型比较结果的影响"),
    )
    statistics_end = find_paragraph(
        document, lambda text: "涉及稳定阶段表现时" in text
    )
    learning_speed_end = find_paragraph(
        document,
        lambda text: "若均值曲线截至第60个评估时刻" in text,
    )
    moved_statistics = move_paragraph_block_after(
        document, statistics_start, statistics_end, learning_speed_end
    )
    changed.extend(moved_statistics)

    statistics_intro = moved_statistics[0]
    if transformer_runs == 8:
        intro_text = (
            "上述四项指标分别从系统运行时长、生产企业经营结果、累积存活表现和目标存活水平到达时刻描述单次训练过程。"
            "为降低单次运行中随机因素对模型比较结果的影响，本文采用多随机种子（multi-seed）统计方法汇总训练曲线。"
            "神经网络参数初始化、动作探索、训练样本抽取和经验池采样等过程均包含随机因素。"
            "本文将用于控制上述随机因素的联合配置统称为“随机种子”。"
        )
    else:
        intro_text = (
            "上述四项指标分别从系统运行时长、生产企业经营结果、累积存活表现和目标存活水平到达时刻描述单次训练过程。"
            "为降低单次运行中随机因素对模型比较结果的影响，本文采用多随机种子（multi-seed）统计方法汇总训练曲线。"
            "神经网络参数初始化、动作探索、训练样本抽取和经验池采样等过程均包含随机因素。"
            "本文将用于控制上述随机因素的联合配置统称为“随机种子”。"
        )
    set_paragraph_text(statistics_intro, intro_text)
    seed_statistics, was_inserted = ensure_seed_statistics_paragraph(
        document, statistics_intro
    )
    if was_inserted:
        changed.append(seed_statistics)
    changed.extend(revise_seed_scope(document, transformer_runs=transformer_runs))

    learning_speed_definition = find_paragraph(
        document, lambda text: text.startswith("（4）学习速度。")
    )
    set_paragraph_text(
        learning_speed_definition,
        learning_speed_definition.text.replace(
            "平均存活天数均值按照6.1所述多随机种子统计方法计算。",
            "平均存活天数均值按照下文所述多随机种子统计方法计算。",
        ),
    )
    changed.append(learning_speed_definition)

    for paragraph in document.paragraphs:
        if "再按照6.1所述方法计算跨随机种子均值和SEM" in paragraph.text:
            set_paragraph_text(
                paragraph,
                paragraph.text.replace(
                    "再按照6.1所述方法计算跨随机种子均值和SEM",
                    "再按照上述方法计算跨随机种子均值和SEM",
                ),
            )
            changed.append(paragraph)

    changed.extend(replace_text_in_section(document, "（2）生成对抗模仿学习", "6.3 "))

    intro = find_paragraph(
        document,
        lambda text: text.startswith("第6.2节第（1）部分的运行结果表明"),
    )
    revised_intro = intro.text.replace("受控非专家行为", "加噪专家行为")
    set_paragraph_text(intro, revised_intro)
    changed.append(intro)

    metric_start = find_paragraph(
        document, lambda text: text.startswith("上述两步评价均以判别器得分")
    )
    metric_end = find_paragraph(
        document, lambda text: text.startswith("式中，Pr表示事件发生的概率")
    )

    first_step = find_paragraph(
        document,
        lambda text: text.startswith("第一步验证") or text.startswith("第一步检验"),
    )
    if transformer_runs == 8:
        first_step_text = (
            "第一步验证评价判别器对专家行为与加噪专家行为的区分能力。"
            "本文将第r次独立运行在训练结束时得到并固定参数的判别器称为评价判别器。"
            "每次评价均使用相同的留出专家测试集、得分区间和直方图划分。"
            "本文在16952条专家状态—动作样本的动作分量上叠加标准差为0.30的独立高斯噪声，并将动作裁剪至[-0.5,0.5]，由此构造加噪专家样本，其对应行为记为加噪专家行为。"
            "留出测试集包含173个完整专家回合。加噪专家动作定义如下："
        )
    else:
        first_step_text = (
            "第一步验证评价判别器对专家行为与加噪专家行为的区分能力。"
            "本文将第r次独立运行在训练结束时依据验证集保存并固定参数的判别器称为评价判别器，本次评价共选取3个评价判别器。"
            "每次评价均使用相同的留出专家测试集、得分区间和直方图划分。"
            "本文在16952条专家状态—动作样本的动作分量上叠加标准差为0.30的独立高斯噪声，并将动作裁剪至[-0.5,0.5]，由此构造加噪专家样本，其对应行为记为加噪专家行为。"
            "留出测试集包含173个完整专家回合。加噪专家动作定义如下："
        )
    set_paragraph_text(first_step, first_step_text)
    changed.append(first_step)

    score_explanation = find_paragraph(
        document,
        lambda text: text.startswith("式中，D表示判别器") and "图6-2" in text,
    )
    score_prefix = score_explanation.text.split("按照前述定义")[0]
    score_prefix = score_prefix.replace(
        "为使加噪专家参照与历史Actor评价保持一致",
        "为保持第一步与历史Actor评价的得分口径一致",
    )
    set_paragraph_text(score_explanation, score_prefix.rstrip())
    changed.append(score_explanation)

    moved_metrics = move_paragraph_block_after(
        document, metric_start, metric_end, score_explanation
    )
    changed.extend(moved_metrics)
    metric_intro = moved_metrics[0]
    set_paragraph_text(
        metric_intro,
        "为定量评价判别器对专家行为与加噪专家行为的区分能力，"
        "本文从判别器得分分布的差异程度、重叠程度和样本排序能力三个方面引入评价指标。"
        "Jensen–Shannon散度（JS divergence）用于度量两组判别器得分概率分布之间的差异。"
        "设P和Q表示两组归一化判别器得分分布，其定义如下：",
    )

    if transformer_runs == 8:
        reference_text = (
            "在上述指标定义基础上，本文以J_r^N表示第r个评价判别器下专家得分分布P_r^E与加噪专家得分分布P_r^N之间的JS散度，并将其作为加噪专家参照。"
            "本文按照本节前述多随机种子统计方法汇总各次独立运行的评价结果。"
        )
    else:
        reference_text = (
            "在上述指标定义基础上，本文以J_r^N表示第r个评价判别器下专家得分分布P_r^E与加噪专家得分分布P_r^N之间的JS散度，并将其作为加噪专家参照。"
            "本文按照本节前述多随机种子统计方法汇总3次独立运行的评价结果。"
        )
    reference_paragraph = insert_paragraph_after(moved_metrics[-1], reference_text)
    changed.append(reference_paragraph)

    for paragraph in document.paragraphs:
        if "按照6.1所述多随机种子统计方法" in paragraph.text:
            set_paragraph_text(
                paragraph,
                paragraph.text.replace(
                    "按照6.1所述多随机种子统计方法",
                    "按照本节前述多随机种子统计方法",
                ),
            )
            changed.append(paragraph)

    replace_figure_4_1(document)

    if highlight:
        for paragraph in changed:
            highlight_paragraph(paragraph)


def revise_document(path: Path, *, highlight: bool, transformer_runs: int) -> None:
    document = Document(path)
    revise_document_contents(
        document, highlight=highlight, transformer_runs=transformer_runs
    )

    temporary_path = path.with_suffix(".latest.tmp.docx")
    document.save(temporary_path)
    os.replace(temporary_path, path)


def main() -> None:
    revise_document(V4_PATH, highlight=True, transformer_runs=8)
    revise_document(V38_PATH, highlight=False, transformer_runs=7)
    print(V4_PATH)
    print(V38_PATH)


if __name__ == "__main__":
    main()
