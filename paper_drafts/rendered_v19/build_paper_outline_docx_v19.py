from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v18"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v18.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v19.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "45BA93C9C53D129EDBF1FA17F7504AC57FD282B185D53539F2E6A8A64504D35C"
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def find_table_after_caption(caption):
    element = caption._p.getnext()
    while element is not None:
        if element.tag.endswith("}tbl"):
            return element
        if element.tag.endswith("}p"):
            raise RuntimeError("A paragraph appears between the Table 6-2 caption and table.")
        element = element.getnext()
    raise RuntimeError("Table 6-2 was not found after its caption.")


def move_after(reference, element) -> None:
    reference.addnext(element)


def move_before(reference, element) -> None:
    reference.addprevious(element)


def verify(document: Document) -> None:
    definition = find_paragraph(document, "为定量比较三种模型在训练后期的表现")
    caption = find_paragraph(document, "表6-2 三种主体训练方案")
    note = find_paragraph(document, "注：平均存活天数及三项累计收益")
    heading_6_2 = find_paragraph(document, "6.2 TD3与GAIL+TD3")
    heading_6_3 = find_paragraph(
        document, "6.3 GAIL+TD3与Transformer+GAIL+TD3"
    )
    table_element = find_table_after_caption(caption)

    body = document._body._element
    positions = {element: index for index, element in enumerate(list(body))}
    if not (
        positions[heading_6_2._p]
        < positions[definition._p]
        < positions[caption._p]
        < positions[table_element]
        < positions[note._p]
        < positions[heading_6_3._p]
    ):
        raise RuntimeError("The Chapter 6.2 result-table elements are in the wrong order.")

    between_6_1_and_6_2 = list(body)[
        positions[find_paragraph(document, "式中，s表示8次运行结果")._p] + 1 :
        positions[heading_6_2._p]
    ]
    if between_6_1_and_6_2:
        raise RuntimeError("Experimental result content remains at the end of Section 6.1.")

    if heading_6_2._p.getnext() is not definition._p:
        raise RuntimeError("The Table 6-2 statistical definition is not at the start of Section 6.2.")
    if note._p.getnext() is not heading_6_3._p:
        raise RuntimeError("Table 6-2 is not the final content block of Section 6.2.")

    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 9:
        raise RuntimeError(f"Expected 9 inline images, found {len(document.inline_shapes)}.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v18 source DOCX changed after the v19 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)
    definition = find_paragraph(document, "为定量比较三种模型在训练后期的表现")
    caption = find_paragraph(document, "表6-2 三种主体训练方案")
    note = find_paragraph(document, "注：平均存活天数及三项累计收益")
    heading_6_2 = find_paragraph(document, "6.2 TD3与GAIL+TD3")
    heading_6_3 = find_paragraph(
        document, "6.3 GAIL+TD3与Transformer+GAIL+TD3"
    )
    table_element = find_table_after_caption(caption)

    move_after(heading_6_2._p, definition._p)
    for element in (caption._p, table_element, note._p):
        move_before(heading_6_3._p, element)

    verify(document)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v18 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
