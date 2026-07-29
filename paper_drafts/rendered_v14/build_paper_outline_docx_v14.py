from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "reference"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_v1.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v14.docx"
)
FIGURE_6_1 = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "four_metrics_two_comparisons_sem_2x2"
    / "td3_vs_gail_td3_four_metrics_sem_2x2.png"
)
FIGURE_6_2 = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "four_metrics_two_comparisons_sem_2x2"
    / "gail_td3_vs_transformer_four_metrics_sem_2x2.png"
)

EXPECTED_SOURCE_SHA256 = (
    "03C54ECD85F84578066BB56EEBEFA8921DF40C90CCF80BF1D4286F46EF56CE5A"
)

EXPERT_ROWS = [
    ("网页人工采集数据", 285, 27_987),
    ("Codex交互式辅助采集数据", 1_435, 140_561),
    ("合计", 1_720, 168_548),
]


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def replace_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    runs = list(paragraph.runs)
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)

    for extra_paragraph in list(cell.paragraphs[1:]):
        extra_paragraph._element.getparent().remove(extra_paragraph._element)


def append_formatted_sentence(paragraph, sentence: str) -> None:
    if sentence in paragraph.text:
        return

    previous_run = paragraph.runs[-1] if paragraph.runs else None
    run = paragraph.add_run(sentence)
    if previous_run is not None and previous_run._r.rPr is not None:
        run._r.insert(0, deepcopy(previous_run._r.rPr))


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_picture_before_caption(
    document: Document,
    caption_prefix: str,
    image_path: Path,
) -> None:
    paragraphs = document.paragraphs
    caption_indexes = [
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text.startswith(caption_prefix)
    ]
    if len(caption_indexes) != 1:
        raise RuntimeError(
            f"Expected one caption beginning with {caption_prefix!r}, "
            f"found {len(caption_indexes)}."
        )

    image_paragraph = paragraphs[caption_indexes[0] - 1]
    blips = image_paragraph._p.xpath(".//a:blip")
    if len(blips) != 1:
        raise RuntimeError(
            f"Expected one image before {caption_prefix!r}, found {len(blips)}."
        )

    relationship_id = blips[0].get(qn("r:embed"))
    image_part = document.part.related_parts[relationship_id]
    image_part._blob = image_path.read_bytes()


def update_table_5_1(document: Document) -> None:
    matching_tables = [
        table
        for table in document.tables
        if table.rows
        and [cell.text.strip() for cell in table.rows[0].cells[:3]]
        == ["数据来源", "成功回合数", "样本行数"]
    ]
    if len(matching_tables) != 1:
        raise RuntimeError(
            f"Expected one Table 5-1 candidate, found {len(matching_tables)}."
        )

    table = matching_tables[0]
    if len(table.rows) != len(EXPERT_ROWS) + 1 or len(table.columns) != 4:
        raise RuntimeError("Table 5-1 geometry does not match the expected 4x4 layout.")

    replace_cell_text(table.cell(0, 3), "平均存活天数")
    for row_index, (source_name, successful_episodes, sample_rows) in enumerate(
        EXPERT_ROWS,
        start=1,
    ):
        existing = [table.cell(row_index, col).text.strip() for col in range(3)]
        expected = [source_name, str(successful_episodes), str(sample_rows)]
        if existing != expected:
            raise RuntimeError(
                f"Unexpected Table 5-1 values in row {row_index}: {existing!r}."
            )
        average_survival_days = sample_rows / successful_episodes
        replace_cell_text(table.cell(row_index, 3), f"{average_survival_days:.2f}")


def build() -> None:
    for required_path in (SOURCE, FIGURE_6_1, FIGURE_6_2):
        if not required_path.exists():
            raise FileNotFoundError(required_path)

    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The reference DOCX changed after the v14 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)
    update_table_5_1(document)

    append_formatted_sentence(
        find_paragraph(document, "表5-1为筛选后的专家数据收集与整理结果。"),
        "表中平均存活天数按照样本行数除以成功回合数计算，网页人工采集数据、"
        "Codex辅助采集数据和合计数据对应的结果分别为98.20天、97.95天和97.99天。",
    )
    append_formatted_sentence(
        find_paragraph(document, "本文首先比较TD3与GAIL+TD3的训练过程"),
        "图6-1(a)中的水平虚线表示表5-1所列专家轨迹的合计平均存活天数，"
        "其数值为97.99天。",
    )
    append_formatted_sentence(
        find_paragraph(
            document,
            "本文进一步比较GAIL+TD3与Transformer+GAIL+TD3",
        ),
        "图6-2(a)中的水平虚线同样表示专家轨迹的合计平均存活天数，"
        "用于比较两种训练方案与专家轨迹存活水平之间的差距。",
    )

    replace_picture_before_caption(
        document,
        "图6-1 TD3与GAIL+TD3训练结果对比",
        FIGURE_6_1,
    )
    replace_picture_before_caption(
        document,
        "图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        FIGURE_6_2,
    )

    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    table = next(
        table
        for table in check.tables
        if table.rows and table.cell(0, 0).text.strip() == "数据来源"
    )
    expected_last_column = ["平均存活天数", "98.20", "97.95", "97.99"]
    actual_last_column = [row.cells[3].text.strip() for row in table.rows]
    if actual_last_column != expected_last_column:
        raise RuntimeError(
            f"Table 5-1 verification failed: {actual_last_column!r}."
        )

    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
