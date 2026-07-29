from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


SRC = next(Path("paper_drafts/rendered_v5").glob("*.docx"))
OUT_DIR = Path("paper_drafts/rendered_v6")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v6.docx"


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_paragraph_text(paragraph, text):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)


def paragraph_by_text(doc, text):
    target = " ".join(text.split())
    for paragraph in doc.paragraphs:
        if " ".join(paragraph.text.split()) == target:
            return paragraph
    raise ValueError(f"paragraph not found: {text}")


def insert_paragraphs_before(marker, entries):
    for text, style in entries:
        marker.insert_paragraph_before(text, style=style)


def insert_table_before(doc, marker, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or c_idx in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = "宋体"
    tbl = table._element
    tbl.getparent().remove(tbl)
    marker._p.addprevious(tbl)
    return table


doc = Document(str(SRC))

# Minimal compliance edits in existing text.
for paragraph in doc.paragraphs:
    if "Transformer 编码器（Transformer 编码器）" in paragraph.text:
        set_paragraph_text(
            paragraph,
            paragraph.text.replace("Transformer 编码器（Transformer 编码器）", "Transformer 编码器（Transformer Encoder）"),
        )
    if "根据既有环境定义与当前代码结构" in paragraph.text:
        set_paragraph_text(
            paragraph,
            paragraph.text.replace("根据既有环境定义与当前代码结构", "根据既有环境定义与本文实验设置"),
        )

# Add section 2.3 before the old empty chapter-3 heading, then remove that heading.
old_ch3 = paragraph_by_text(doc, "3 专家数据采集与有限理性行为表示")
section_23 = [
    ("2.3 外部评价目标与有限理性专家数据表示", "Heading 2"),
    (
        "仿真主体的外部评价需要先区分完全理性目标与部分理性目标。完全理性评价通常关注主体能否在给定环境规则下取得较高收益、延长存活时间并形成稳定运行状态。这类评价适合检验仿真环境是否可运行，也适合检验强化学习主体是否能够利用奖励信号学习到高收益策略。",
        "Normal",
    ),
    (
        "部分理性评价的目标不同。有限理性主体并不要求在每个状态下输出全局最优动作，而是要求其行为能够体现人类决策中的信息不完全、经验依赖、风险权衡和路径依赖等特征。因此，部分理性主体的外部评价不能只依赖收益最大化指标，还需要考察生成动作与专家行为分布之间的一致性。对本文而言，有限理性评价主要服务于生产企业主体的决策生成问题。",
        "Normal",
    ),
    (
        "专家数据是连接有限理性行为与模型训练目标的中间表示。立项报告将专家数据来源概括为人类参与的仿真实验、大语言模型（LLM: Large Language Model）辅助决策和真实企业行为数据。本文当前整理的数据来源于前两类实验材料，即网页实验系统中的参与式采集数据，以及在同一仿真接口下形成的大模型辅助批量采集数据。本文未将真实企业行为数据写入当前专家样本文件；该来源的数据对齐方法仍需后续补充。",
        "Normal",
    ),
    (
        "在数据结构上，专家样本以状态-动作对表示。每条记录的前 33 个字段为生产企业主体在决策时可观测的状态信息，后 4 个字段为对应动作。4 个动作分别表示贷款意愿 WNDF、原料A采购需求、原料B采购需求和产品A定价。网页实验系统允许参与者输入具有业务含义的原始数值，并在保存专家样本时转换为模型训练所需的连续动作取值。该处理使人工输入能够进入生成对抗模仿学习（GAIL: Generative Adversarial Imitation Learning）与双延迟深度确定性策略梯度（TD3: Twin Delayed Deep Deterministic Policy Gradient）的统一训练数据格式。",
        "Normal",
    ),
]
insert_paragraphs_before(old_ch3, section_23)
delete_paragraph(old_ch3)

# Renumber former chapters 4 and 5 to 3 and 4.
replacements = {
    "4 基于生成对抗模仿学习与TD3的主体训练方案": "3 基于生成对抗模仿学习与TD3的主体训练方案",
    "4.1 模型训练方案": "3.1 模型训练方案",
    "4.2 状态、动作与经验样本": "3.2 状态、动作与经验样本",
    "4.3 判别器目标与模仿奖励构造": "3.3 判别器目标与模仿奖励构造",
    "4.4 融合奖励与TD3参数更新": "3.4 融合奖励与TD3参数更新",
    "5 基于Transformer表征增强的GAIL+TD3主体训练方案": "4 基于Transformer表征增强的GAIL+TD3主体训练方案",
    "5.1 时序状态表征的引入动机": "4.1 时序状态表征的引入动机",
    "5.2 Transformer 编码器结构": "4.2 Transformer 编码器结构",
    "5.3 表征增强模块的训练接入与方法边界": "4.3 表征增强模块的训练接入与方法边界",
}
text_replacements = {
    "图4-1": "图3-1",
    "图5-1": "图4-1",
    "图5-2": "图4-2",
    "第4章": "第3章",
}
for paragraph in doc.paragraphs:
    stripped = " ".join(paragraph.text.split())
    if stripped in replacements:
        set_paragraph_text(paragraph, replacements[stripped])
        continue
    new_text = paragraph.text
    for old, new in text_replacements.items():
        new_text = new_text.replace(old, new)
    if new_text != paragraph.text:
        set_paragraph_text(paragraph, new_text)

# Replace the existing result-assertive sentence in the former chapter 5.
for paragraph in doc.paragraphs:
    if paragraph.text.startswith("后续实验将进一步表明") or paragraph.text.startswith("第6章对比实验进一步表明"):
        set_paragraph_text(
            paragraph,
            "后续实验将进一步检验单步状态表征是否能够充分刻画动态交互过程中的历史依赖。基于这一问题设定，本文在第3章 GAIL+TD3 训练方案基础上引入 Transformer 编码器。该扩展不改变 TD3 的基本更新机制，而是在生产企业主体的状态表征环节加入历史窗口编码，使策略网络和价值网络能够在更新时利用局部历史序列信息。",
        )

# Insert new chapter 5 before the existing chapter 6.
ch6 = paragraph_by_text(doc, "6 实验设置与结果分析")
chapter_5_before_table = [
    ("5 专家数据采集系统与数据整理", "Heading 1"),
    ("5.1 专家数据来源与采集边界", "Heading 2"),
    (
        "本文围绕生产企业主体构建专家数据采集流程。采集对象不是现实企业经营记录，而是在 MAS-E 仿真环境中形成的主体决策轨迹。该设定使状态空间、动作空间和环境反馈与后续模型训练保持一致，也避免了真实企业数据与仿真变量之间的大规模范式转换问题。",
        "Normal",
    ),
    (
        "本文专家数据包括两类来源。第一类为网页实验系统采集的数据，参与者在页面中扮演甲公司经理，并依据当天可见经营状态填写动作。第二类为大模型辅助批量采集数据，辅助程序以已通过筛选的成功轨迹为参考，在同一网页实验系统接口下提交带有随机扰动的候选动作。两类数据均经过仿真环境执行和存活天数筛选。本文当前数据文件未纳入真实企业行为数据。[待补充：正式人类参与者招募人数、实验时长和伦理说明]",
        "Normal",
    ),
    ("5.2 网页实验系统设计", "Heading 2"),
    (
        "网页实验系统首先向参与者说明任务背景。系统中包含甲公司、乙公司、银行和第三方市场。参与者扮演甲公司的经营者，目标是在每天经营状态变化后做出贷款、采购和定价决策，使企业经营更稳定并延长回合存活时间。系统在开始页面记录参与者编号、年龄段、受教育程度、性别、经济或管理相关背景、经营或策略类游戏经验，并通过访问口令控制实验入口。",
        "Normal",
    ),
    (
        "每日决策页面将信息分为贷款信息、上一天两家公司买到的原料数量、定价与销售信息、经营小结和更多规则与成交明细。贷款信息包括当前现金、今天还款、总欠款和昨天贷款。原料数量模块同时展示甲公司和乙公司实际买到的原料A与原料B数量。定价与销售信息展示普通市场和第三方市场价格，以及产品A的上一天售出数量、上一天产出数量和当前库存。经营小结以甲公司和乙公司的净利润变化为核心，并在出现价格过高、原料不匹配、采购不足或还款风险时给出提示。",
        "Normal",
    ),
    (
        "参与者每天需要填写四个动作：申请贷款金额、原料A采购需求、原料B采购需求和产品A销售价格。网页端对输入值设置上下界，避免动作转换后超过模型训练所需的取值范围。系统将具有业务含义的输入值转换为连续动作向量，并保存为专家样本。页面还提供“完成本段，跳到下一段”的功能。参与者连续完成 8 天人工决策后，系统最多自动推进 15 天，并在现金、债务、库存、采购参考量或销售价格出现明显变化时重新交回人工决策。",
        "Normal",
    ),
    ("[待补充：插入网页实验系统任务描述页截图]", "Normal"),
    ("图5-1 网页实验系统任务描述页", "Normal"),
    ("[待补充：插入网页实验系统每日经营决策页截图]", "Normal"),
    ("图5-2 网页实验系统每日经营决策页", "Normal"),
    ("5.3 大模型辅助批量采集方法", "Heading 2"),
    (
        "大模型辅助批量采集用于提高专家样本的生成效率。辅助程序先读取采集目录中已经存在的 meta 文件和 expert 文件，并按参与者编号与回合编号组织轨迹。只有回合长度达到 90 天以上的轨迹会被作为参考轨迹。随后，辅助程序从参考轨迹中抽取同一天位置的动作信息，并结合当前现金、还款压力、原料购买结果、产品售出比例、库存比例和第三方市场价格等页面可见信息生成候选动作。",
        "Normal",
    ),
    (
        "该辅助程序没有直接写入专家文件，而是通过网页实验系统的 /api/start、/api/step 和 /api/next_episode 接口提交动作。每个动作仍需经过网页系统的输入边界检查、动作转换、环境推进、回合终止判断和 summary 记录。该处理保证辅助采集数据与页面采集数据共享同一环境规则和保存格式。辅助策略中包含随机扰动，因此不同回合之间的采购量、贷款金额和产品定价并非固定序列。",
        "Normal",
    ),
    (
        "需要说明的是，大模型辅助批量采集数据不能直接等同于真实人类行为数据。该数据更适合用于扩大专家样本规模、测试训练流程和构造对照数据。后续实验应区分网页人工采集数据与辅助批量采集数据，并从存活天数、行为分布、判别器评分和策略稳定性等指标检验不同数据来源对模型训练的影响。",
        "Normal",
    ),
    ("5.4 数据筛选与专家样本文件整理", "Heading 2"),
    (
        "本文采用回合存活天数作为专家样本的基础筛选条件。只有 summary 文件中 survival_days 大于 90 的回合会进入整理后的专家样本文件。随后，程序从对应 meta 文件中提取生产企业状态字段和模型动作字段，形成与原 expert_data_production1_cleaned.csv 相同的 37 列格式。其中，前 33 列为状态，后 4 列为动作。",
        "Normal",
    ),
    ("表5-1 专家数据收集与整理结果", "Normal"),
]
insert_paragraphs_before(ch6, chapter_5_before_table)

table_rows = [
    ["数据来源", "成功回合数", "样本行数", "说明"],
    ["网页实验系统采集数据", "8", "782", "由页面交互形成，并通过 survival_days > 90 筛选。"],
    ["大模型辅助批量采集数据", "1712", "167766", "由辅助程序经网页系统接口提交动作，并通过同一筛选条件保留。"],
    ["合计", "1720", "168548", "整理为 expert_data_production1_collected.csv，无表头，列数为 37。"],
]
insert_table_before(doc, ch6, table_rows)

chapter_5_after_table = [
    (
        "由表5-1可知，整理后的专家数据文件共包含 1720 个存活超过 90 天的回合，对应 168548 条状态-动作样本。该文件命名为 expert_data_production1_collected.csv。该文件与既有 expert_data_production1_cleaned.csv 保持一致的列组织方式，可直接作为后续 GAIL+TD3 及其 Transformer 表征增强模型的专家数据输入。实验结果部分将进一步检验不同专家数据来源对训练过程和生成行为的影响。",
        "Normal",
    ),
]
insert_paragraphs_before(ch6, chapter_5_after_table)

doc.save(str(OUT))
print(OUT)
