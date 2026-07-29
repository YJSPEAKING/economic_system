from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v22"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v22.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v23.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "3C8E77D115C2DA2A3C21B33184BB5EB171032A0E5B0B0445066E6622C14D90AB"
)


SECTION_6_2_FIGURE_ANALYSIS = (
    "以图6-1(a)中的专家轨迹平均存活天数为参照，GAIL+TD3的平均存活天数在训练前期快速提高，"
    "并在第20个评估时刻之前接近其训练后期水平；TD3则直至训练后期才进入相近区间。"
    "这表明，判别器依据专家状态—动作样本产生的模仿奖励为生产企业主体提供了额外的策略更新信息，"
    "使模型更快地从初始策略向较高存活水平移动。生产企业、消费企业和银行累计收益曲线也呈现相似的前期差异。"
    "图6-1(e)进一步显示，GAIL+TD3的累计归一化学习曲线下面积整体高于TD3，"
    "且两条曲线的差距主要在训练前期形成，并延续至训练末期。"
)

SECTION_6_2_TABLE_ANALYSIS = (
    "表6-2中的具体数值显示，GAIL+TD3在第60个评估时刻的累计归一化学习曲线下面积均值由TD3的0.435提高至0.683，"
    "增幅约为56.8%。该结果从整个训练区间对图6-1所反映的学习速度差异进行了量化，"
    "说明GAIL+TD3在相同训练预算内积累了更多的平均存活天数表现。"
    "银行最终稳定累计收益均值由5.28K提高至10.31K，"
    "表明生产企业策略的提前形成还通过信贷关系影响了银行主体的经营结果。"
)

SECTION_6_2_TRANSITION = (
    "不过，GAIL+TD3的最终稳定平均存活天数为80.71天，与专家轨迹平均值仍相差17.28天。"
    "该结果说明，模仿学习信号主要改善了生产企业主体的前期学习速度，"
    "但尚未使其稳定阶段的存活水平充分接近专家轨迹。"
    "因此，下一节进一步引入历史状态序列表征，"
    "以考察模型利用连续交互信息后能否缩小这一差距。"
)

SECTION_6_3_FIGURE_ANALYSIS = (
    "从图6-2(a)可以看出，Transformer+GAIL+TD3比GAIL+TD3更早进入平均存活天数的快速提升阶段，"
    "并在训练后期保持了更接近专家参考线的平均水平。生产企业和消费企业累计收益曲线在多数评估时刻也整体高于GAIL+TD3。"
    "图6-2(e)显示，Transformer+GAIL+TD3的累计归一化学习曲线下面积整体高于GAIL+TD3。"
    "这些曲线变化表明，历史状态序列为生产企业主体提供了单步状态之外的连续交互信息，"
    "使模型能够更充分地利用模仿奖励所包含的专家行为信息。"
)

SECTION_6_3_TABLE_ANALYSIS = (
    "表6-2进一步显示，引入Transformer编码器后，最终稳定平均存活天数由80.71天提高至90.24天，"
    "与专家轨迹平均值的差距由17.28天缩小至7.75天，缩小幅度约为55.2%。"
    "第60个评估时刻的累计归一化学习曲线下面积由0.683提高至0.773，增幅约为13.2%。"
    "生产企业和消费企业最终稳定累计收益也分别由126.35K提高至180.43K、由150.01K提高至169.48K。"
    "上述结果说明，历史状态表征不仅进一步提高了模型的学习速度，"
    "也使稳定阶段的平均存活天数在结果层面更加接近专家轨迹。"
)

SECTION_6_3_CONCLUSION = (
    "综合图6-2、图6-3和表6-2的结果，Transformer+GAIL+TD3在前期学习效率、"
    "最终稳定平均存活天数和企业累计收益方面均较GAIL+TD3取得了更高的均值表现，"
    "并将DSCR低于1的生产企业日观测占比降至0.00%。"
    "其中，最终稳定平均存活天数与专家轨迹平均值的差距由17.28天缩小至7.75天，"
    "说明历史状态序列表征使生产企业主体能够利用跨时段交互信息，"
    "并在平均存活天数这一结果指标上进一步接近专家轨迹。"
    "在本组实验设置下，Transformer编码器不仅加快了模型进入较高存活水平的过程，"
    "也改善了稳定阶段的存活表现，进一步弥补了GAIL+TD3采用单步状态表征时存在的不足。"
)


CHAPTER_7_PARAGRAPHS = (
    ("h1", "7 总结与展望"),
    ("h2", "7.1 总结"),
    (
        "body",
        "本文围绕复杂经济系统多主体仿真环境中的智能主体决策生成问题，"
        "在既有TD3主体决策模型基础上引入生成对抗模仿学习，"
        "并进一步采用Transformer编码器增强生产企业主体的历史状态表征。"
        "本文通过人工行为采集系统整理专家状态—动作数据，"
        "并从学习速度、专家轨迹存活水平接近程度、主体累计收益和生产企业偿债能力等方面，"
        "比较TD3、GAIL+TD3和Transformer+GAIL+TD3三种主体训练方案。主要研究结论如下。",
    ),
    (
        "body",
        "首先，生成对抗模仿学习信号加快了生产企业主体的前期策略形成过程。"
        "与TD3相比，GAIL+TD3在第20个评估时刻之前已接近其训练后期的平均存活水平，"
        "第60个评估时刻的累计归一化学习曲线下面积由0.435提高至0.683。"
        "这一结果表明，专家状态—动作样本经由判别器形成的模仿奖励为生产企业主体提供了额外的学习信息，"
        "使系统能够更快进入较高的平均存活水平。",
    ),
    (
        "body",
        "其次，引入Transformer编码器后，模型的学习速度和稳定阶段存活水平得到进一步改善。"
        "Transformer+GAIL+TD3的最终稳定平均存活天数由80.71天提高至90.24天，"
        "与专家轨迹平均值的差距由17.28天缩小至7.75天；"
        "第60个评估时刻的累计归一化学习曲线下面积由0.683提高至0.773。"
        "生产企业和消费企业的最终稳定累计收益也分别由126.35K提高至180.43K、"
        "由150.01K提高至169.48K。"
        "上述结果说明，历史状态序列增强了生产企业主体对跨时段交互信息的表征，"
        "并使模型在平均存活天数这一结果指标上进一步接近专家轨迹。",
    ),
    (
        "body",
        "最后，生产企业偿债风险结果表明，在本文选取的稳定运行观测中，"
        "TD3、GAIL+TD3和Transformer+GAIL+TD3的DSCR低于1的生产企业日观测占比分别为6.52%、0.23%和0.00%。"
        "该结果说明，模仿学习信号与历史状态表征的引入对应着更低的当期偿债缺口观测比例，"
        "Transformer+GAIL+TD3在生产企业现金流与债务偿付协调方面取得了更稳定的结果。",
    ),
    ("h2", "7.2 展望"),
    (
        "body",
        "本文仍存在若干需要进一步研究的问题。第一，本文主要通过平均存活天数和累计归一化学习曲线下面积评价模型对专家经验的利用效果。"
        "这类指标能够反映模型结果与专家轨迹存活水平之间的接近程度，"
        "但不能直接量化生产企业动作分布与专家动作分布的一致程度。"
        "后续研究可以引入状态—动作占用分布距离、判别器得分分布和动作偏差等指标，"
        "从行为分布层面对模仿效果进行直接评价。",
    ),
    (
        "body",
        "第二，GAIL+TD3的最终稳定平均存活天数尚未充分接近专家轨迹平均值，"
        "不同主体的累计收益与跨随机种子波动也未呈现完全一致的变化方向。"
        "后续研究可以进一步考察环境奖励与模仿奖励的融合方式，"
        "并结合生产企业、消费企业和银行主体的经营目标改进多主体协同训练过程，"
        "以协调学习速度、系统存活水平和不同主体收益之间的关系。",
    ),
    (
        "body",
        "第三，现有专家数据覆盖的行为主体与经济场景仍然有限，"
        "历史状态表征目前主要应用于生产企业主体，且尚未系统比较不同历史序列长度的影响。"
        "后续研究可以扩展专家行为数据的来源和场景覆盖范围，"
        "将历史状态表征推广至消费企业与银行主体，"
        "并通过序列长度和网络结构的消融实验识别不同模块的具体作用。"
        "研究还可以增加独立运行次数并引入相应的统计检验，"
        "以进一步评估不同主体训练方案的稳定性。",
    ),
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def first_run_properties(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def replace_paragraph_text(paragraph, text: str) -> None:
    source_rpr = first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def insert_cloned_paragraph_before(reference, template, text: str) -> None:
    new_p = deepcopy(template._p)
    for child in list(new_p):
        if child.tag != qn("w:pPr"):
            new_p.remove(child)

    new_r = OxmlElement("w:r")
    source_rpr = first_run_properties(template)
    if source_rpr is not None:
        new_r.append(source_rpr)
    new_t = OxmlElement("w:t")
    new_t.text = text
    new_r.append(new_t)
    new_p.append(new_r)
    reference._p.addprevious(new_p)


def verify(document: Document) -> None:
    paragraphs = [p.text for p in document.paragraphs]
    all_text = "\n".join(paragraphs)

    expected_texts = (
        SECTION_6_2_FIGURE_ANALYSIS,
        SECTION_6_2_TABLE_ANALYSIS,
        SECTION_6_2_TRANSITION,
        SECTION_6_3_FIGURE_ANALYSIS,
        SECTION_6_3_TABLE_ANALYSIS,
        SECTION_6_3_CONCLUSION,
        "7 总结与展望",
        "7.1 总结",
        "7.2 展望",
        CHAPTER_7_PARAGRAPHS[-1][1],
    )
    for expected in expected_texts:
        if expected not in all_text:
            raise RuntimeError(f"Missing revised text: {expected[:60]!r}")

    forbidden_prefixes = (
        "从图6-1可以看出",
        "不过，表6-2和图6-1也显示",
        "从图6-2可以看出",
    )
    for prefix in forbidden_prefixes:
        if any(text.startswith(prefix) for text in paragraphs):
            raise RuntimeError(f"Former Section 6 wording remains: {prefix!r}")

    if paragraphs.count("7 总结与展望") != 1:
        raise RuntimeError("Chapter 7 heading is missing or duplicated.")
    if paragraphs.index("7 总结与展望") >= paragraphs.index("参考文献"):
        raise RuntimeError("Chapter 7 was not inserted before the references.")
    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 9:
        raise RuntimeError(f"Expected 9 inline images, found {len(document.inline_shapes)}.")
    if len(document.element.body.xpath(".//m:oMath")) != 116:
        raise RuntimeError("The existing Word equation objects changed unexpectedly.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v22 source DOCX changed after the v23 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)

    replace_paragraph_text(
        find_paragraph(document, "从图6-1可以看出"),
        SECTION_6_2_FIGURE_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "表6-2中的具体数值显示"),
        SECTION_6_2_TABLE_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "不过，表6-2和图6-1也显示"),
        SECTION_6_2_TRANSITION,
    )
    replace_paragraph_text(
        find_paragraph(document, "从图6-2可以看出"),
        SECTION_6_3_FIGURE_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "表6-2进一步显示"),
        SECTION_6_3_TABLE_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "综合图6-2、图6-3和表6-2的结果"),
        SECTION_6_3_CONCLUSION,
    )

    references = find_paragraph(document, "参考文献")
    heading_1_template = find_paragraph(document, "6  实验设置与结果分析")
    heading_2_template = find_paragraph(document, "6.1 实验设置与评价指标")
    body_template = find_paragraph(document, "综合图6-2、图6-3和表6-2的结果")
    templates = {
        "h1": heading_1_template,
        "h2": heading_2_template,
        "body": body_template,
    }
    for kind, text in CHAPTER_7_PARAGRAPHS:
        insert_cloned_paragraph_before(references, templates[kind], text)

    verify(document)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v22 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
