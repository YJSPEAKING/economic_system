from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v15"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v15.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v16.docx"
)
FIGURE_6_1 = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "five_metrics_two_comparisons_aulc_sem"
    / "td3_vs_gail_td3_five_metrics_aulc_sem.png"
)
FIGURE_6_2 = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "five_metrics_two_comparisons_aulc_sem"
    / "gail_td3_vs_transformer_five_metrics_aulc_sem.png"
)
MML2OMML_XSL = Path(
    r"C:\Program Files\Microsoft Office\Root\Office16\MML2OMML.XSL"
)
EXPECTED_SOURCE_SHA256 = (
    "9CC2305BE445E9F16D853027BE187AD0C5F530D6207BC9F3EF2BF813600A09EF"
)


AULC_MATHML = (
    '<math xmlns="http://www.w3.org/1998/Math/MathML">'
    '<msubsup><mi>A</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow>'
    '<mtext>norm</mtext></msubsup><mo>=</mo>'
    '<mfrac><mn>1</mn><mrow><mfenced><mrow><mi>K</mi><mo>−</mo><mn>1</mn>'
    '</mrow></mfenced><msub><mi>L</mi><mtext>max</mtext></msub></mrow></mfrac>'
    '<munderover><mo>∑</mo><mrow><mi>j</mi><mo>=</mo><mn>1</mn></mrow>'
    '<mrow><mi>k</mi><mo>−</mo><mn>1</mn></mrow></munderover>'
    '<mfrac><mrow>'
    '<mover><msub><mi>L</mi><mrow><mi>r</mi><mo>,</mo><mi>j</mi></mrow>'
    '</msub><mo>¯</mo></mover><mo>+</mo>'
    '<mover><msub><mi>L</mi><mrow><mi>r</mi><mo>,</mo><mi>j</mi><mo>+</mo>'
    '<mn>1</mn></mrow></msub><mo>¯</mo></mover>'
    '</mrow><mn>2</mn></mfrac>'
    '</math>'
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    source_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            source_rpr = deepcopy(run._r.rPr)
            break

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def insert_paragraph_after(
    reference: Paragraph,
    text: str = "",
    copy_format_from: Paragraph | None = None,
) -> Paragraph:
    new_p = OxmlElement("w:p")
    reference._p.addnext(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    format_source = copy_format_from or reference
    if format_source._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(format_source._p.pPr))
    if text:
        source_rpr = None
        for run in format_source.runs:
            if run._r.rPr is not None:
                source_rpr = deepcopy(run._r.rPr)
                break
        run = paragraph.add_run(text)
        if source_rpr is not None:
            run._r.insert(0, source_rpr)
    return paragraph


def mathml_to_omml(mathml: str):
    transform = etree.XSLT(etree.parse(str(MML2OMML_XSL)))
    result = transform(etree.fromstring(mathml.encode("utf-8")))
    return etree.fromstring(etree.tostring(result.getroot()))


def insert_equation_after(reference: Paragraph, mathml: str, template: Paragraph) -> Paragraph:
    equation = insert_paragraph_after(reference, copy_format_from=template)
    equation._p.append(mathml_to_omml(mathml))
    return equation


def find_table_6_1(document: Document):
    matching = [
        table
        for table in document.tables
        if table.rows
        and [cell.text.strip() for cell in table.rows[0].cells] == ["项目", "设定"]
        and any(
            row.cells[0].text.strip() == "对比模型"
            for row in table.rows[1:]
        )
    ]
    if len(matching) != 1:
        raise RuntimeError(f"Expected one Table 6-1, found {len(matching)}.")
    return matching[0]


def replace_cell_text(cell, text: str) -> None:
    replace_paragraph_text(cell.paragraphs[0], text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def update_and_move_table_6_1(document: Document, after_paragraph: Paragraph) -> None:
    caption = find_paragraph(document, "表6-1 实验与统计设置")
    table = find_table_6_1(document)
    for row in table.rows[1:]:
        label = row.cells[0].text.strip()
        if label == "每个评估时刻对应回合数":
            replace_cell_text(row.cells[1], "E=100个连续训练回合")
        elif label == "单回合最大天数":
            replace_cell_text(row.cells[0], "单个回合最大运行时长")
            replace_cell_text(row.cells[1], "100天")

    after_paragraph._p.addnext(caption._p)
    caption._p.addnext(table._tbl)


def replace_figure_before_caption(
    document: Document,
    caption_prefix: str,
    image_path: Path,
) -> None:
    caption = find_paragraph(document, caption_prefix)
    image_element = caption._p.getprevious()
    if image_element is None or image_element.tag != qn("w:p"):
        raise RuntimeError(f"No image paragraph found before {caption_prefix!r}.")
    image_paragraph = Paragraph(image_element, caption._parent)
    if not image_paragraph._p.xpath(".//w:drawing"):
        raise RuntimeError(f"Previous paragraph is not a figure for {caption_prefix!r}.")

    for child in list(image_paragraph._p):
        if child.tag != qn("w:pPr"):
            image_paragraph._p.remove(child)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.35))


def chapter_six_rewrite(document: Document) -> None:
    opening = find_paragraph(document, "本节实验从模仿学习信号与历史状态表征两个层面")
    replace_paragraph_text(
        opening,
        "本节实验从模仿学习信号与历史状态表征两个层面，通过比较不同主体决策模型在复杂经济系统仿真中的训练表现，分析两类机制对模型训练过程及系统运行结果的影响。本文设置两组对比实验：第一组比较TD3与GAIL+TD3，用于考察模仿学习信号对生产企业主体训练过程的影响；第二组比较GAIL+TD3与Transformer+GAIL+TD3，用于考察历史状态序列表征对生产企业主体训练过程的影响。三种模型中，TD3作为基线模型；GAIL+TD3在生产企业主体的TD3训练过程中引入生成对抗模仿学习信号；Transformer+GAIL+TD3在GAIL+TD3基础上加入Transformer编码器。GAIL+TD3和Transformer+GAIL+TD3使用第五章整理得到的同一组专家数据，三种模型采用相同的仿真环境、训练进度尺度和统计方法。",
    )

    timing = find_paragraph(document, "为统一三种模型的训练进度与评价口径")
    replace_paragraph_text(
        timing,
        "本文以回合作为训练进度单位，并将每个评估时刻所包含的连续训练回合数记为E。模型每完成E个连续训练回合，本文汇总一次评价指标，并将该汇总位置记为一个评估时刻（evaluation step）。第k个评估时刻对应第k组连续训练回合的统计结果。表6-1列出了实验采用的模型、评估间隔、单个回合最大运行时长和统计方法。",
    )
    update_and_move_table_6_1(document, timing)

    metric_intro = find_paragraph(document, "本文在各评估时刻跟踪平均存活天数")
    replace_paragraph_text(
        metric_intro,
        "在上述训练过程中，本文跟踪系统运行、主体收益、学习速度和生产企业偿债风险四类评价指标。各项指标的定义如下。",
    )

    survival = find_paragraph(document, "平均存活天数用于衡量系统")
    replace_paragraph_text(
        survival,
        "（1）平均存活天数。该指标用于衡量系统在一个评估时刻内维持运行的平均时长，等于第k个评估时刻所含E个回合存活天数的算术平均值，其定义如下：",
    )

    survival_explanation = find_paragraph(document, "式中，L表示第k个评估时刻内")
    replace_paragraph_text(
        survival_explanation,
        "式中，n表示第k个评估时刻内的回合序号，n=1,2,…,E；L表示对应回合的存活天数。",
    )
    income_intro = insert_paragraph_after(
        survival_explanation,
        "（2）累计收益。累计收益包括生产企业累计收益、消费企业累计收益和银行累计收益。对于企业主体，本文首先在单个回合内汇总收入、支出与利息，再对第k个评估时刻内E个回合的企业累计收益取算术平均值，其定义如下：",
        copy_format_from=survival,
    )

    enterprise_equation = document.paragraphs[
        next(
            index
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph._p is survival_explanation._p
        )
        + 2
    ]
    if not enterprise_equation._p.xpath(".//m:oMath"):
        raise RuntimeError("Enterprise-income equation was not found after insertion.")

    income_explanation = find_paragraph(document, "式中，J表示企业在对应评估时刻")
    replace_paragraph_text(
        income_explanation,
        "式中，J表示企业在对应评估时刻的平均累计收益；R、C和I分别表示单个回合内的累计收入、累计支出和累计利息；i、k和n分别表示企业编号、评估时刻编号和该评估时刻内的回合序号。",
    )

    bank_explanation = find_paragraph(document, "式中，J表示银行在对应评估时刻")
    replace_paragraph_text(
        bank_explanation,
        "式中，J表示银行在对应评估时刻的平均累计收益；m表示系统中的企业主体数量；I表示对应企业在单个回合内支付的累计利息；k和n分别表示评估时刻编号和该评估时刻内的回合序号。",
    )

    aulc_intro = insert_paragraph_after(
        bank_explanation,
        "（3）累计归一化学习曲线下面积。归一化学习曲线下面积（nAULC: normalized area under the learning curve）用于综合衡量平均存活天数的提升速度与持续水平。本文对每次独立运行的平均存活天数曲线采用梯形法积分，并以完整评估区间和单个回合最大运行时长进行归一化。截至第k个评估时刻的累计归一化面积定义如下：",
        copy_format_from=survival,
    )
    equation_template = find_paragraph(document, "式中，J表示银行在对应评估时刻")
    # The paragraph immediately before the bank explanation contains a centered Word equation.
    bank_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is bank_explanation._p
    )
    equation_template = document.paragraphs[bank_index - 1]
    aulc_equation = insert_equation_after(aulc_intro, AULC_MATHML, equation_template)
    insert_paragraph_after(
        aulc_equation,
        "式中，r表示独立运行编号；k表示当前评估时刻编号；K表示图中纳入比较的评估时刻总数；j表示积分区间编号；带上横线的L表示相应独立运行在对应评估时刻的平均存活天数；L的最大值表示表6-1所列的单个回合最大运行时长。本文令第1个评估时刻的累计归一化面积为0。在相同评估时刻下，较大的累计归一化面积表示模型在此前训练过程中较早或较持续地取得了较高的平均存活天数。",
        copy_format_from=survival_explanation,
    )

    dscr = find_paragraph(document, "生产企业偿债能力采用偿债覆盖率")
    replace_paragraph_text(
        dscr,
        "（4）生产企业偿债能力。本文采用偿债覆盖率（DSCR: Debt Service Coverage Ratio）度量生产企业的当期偿债能力。该指标以生产企业当日可支配现金与当期应还本息之比表示，其定义如下：",
    )

    sem_explanation = find_paragraph(document, "式中，s表示8次运行结果的样本标准差")
    replace_paragraph_text(
        sem_explanation,
        "式中，s表示8次运行结果的样本标准差，S表示独立运行次数，本文取S=8。训练曲线中的实线表示8次独立运行的均值，阴影表示均值上下1个SEM。",
    )

    figure_6_1_intro = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    replace_paragraph_text(
        figure_6_1_intro,
        "本文首先比较TD3与GAIL+TD3的训练过程，以分析模仿学习信号对生产企业主体决策训练的影响。图6-1给出了两种方法在平均存活天数、生产企业累计收益、消费企业累计收益、银行累计收益和累计归一化学习曲线下面积五项指标上的均值曲线及SEM阴影。图6-1(a)中的水平虚线表示表5-1所列专家轨迹的合计平均存活天数，其数值为97.99天。",
    )
    replace_figure_before_caption(document, "图6-1 TD3与GAIL+TD3", FIGURE_6_1)

    figure_6_1_d = find_paragraph(document, "图6-1(d)显示")
    replace_paragraph_text(
        figure_6_1_d,
        "图6-1(d)显示，GAIL+TD3的银行累计收益在训练前期迅速提高，之后围绕较高水平波动；TD3的银行累计收益提升时间较晚，训练后期均值仍低于GAIL+TD3。银行收益来自企业支付的利息，因此，该指标记录了企业经营决策对信贷关系的间接影响。",
    )
    insert_paragraph_after(
        figure_6_1_d,
        "图6-1(e)显示，GAIL+TD3的累计归一化学习曲线下面积在训练过程中高于TD3，第60个评估时刻的均值分别为0.683和0.435。该指标累积了此前各评估时刻的平均存活表现，两条曲线之间的差异说明GAIL+TD3在相同训练进度内形成了更大的存活天数累积面积。综合五项指标，GAIL+TD3的主要表现特征是较快提高系统存活天数和各主体收益，并较早形成相对稳定的训练曲线；TD3在部分企业累计收益指标上的后期均值较高，体现了不同决策目标之间的表现差异。",
        copy_format_from=figure_6_1_d,
    )

    figure_6_2_intro = find_paragraph(document, "本文进一步比较GAIL+TD3与Transformer+GAIL+TD3")
    replace_paragraph_text(
        figure_6_2_intro,
        "本文进一步比较GAIL+TD3与Transformer+GAIL+TD3，以分析历史状态序列表征对生产企业主体训练过程的影响。图6-2给出了两种方法在五项指标上的均值曲线及SEM阴影。图6-2(a)中的水平虚线同样表示专家轨迹的合计平均存活天数，用于比较两种训练方案与专家轨迹存活水平之间的差距。",
    )
    replace_figure_before_caption(
        document,
        "图6-2 GAIL+TD3与Transformer+GAIL+TD3",
        FIGURE_6_2,
    )

    figure_6_2_d = find_paragraph(document, "图6-2(d)显示")
    replace_paragraph_text(
        figure_6_2_d,
        "图6-2(d)显示，Transformer+GAIL+TD3的银行累计收益在训练前期出现较大波动，之后进入相对稳定区间；其训练后期均值低于GAIL+TD3。银行累计收益由企业支付利息形成，生产企业贷款需求、现金流状况和存活时间的变化会共同改变银行收益。该结果表明，生产企业历史状态表征对企业经营表现的影响会进一步传导至企业与银行之间的信贷关系。",
    )
    insert_paragraph_after(
        figure_6_2_d,
        "图6-2(e)显示，Transformer+GAIL+TD3的累计归一化学习曲线下面积在训练过程中高于GAIL+TD3，第60个评估时刻的均值分别为0.773和0.683。该结果与图6-2(a)中Transformer+GAIL+TD3较早提高并维持较高平均存活天数的曲线特征一致，说明历史状态序列表征扩大了相同训练区间内的存活天数累积面积。",
        copy_format_from=figure_6_2_d,
    )


def verify(document: Document) -> None:
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    expected = [
        "（1）平均存活天数",
        "（2）累计收益",
        "（3）累计归一化学习曲线下面积",
        "（4）生产企业偿债能力",
        "n表示第k个评估时刻内的回合序号",
        "图6-1给出了两种方法在平均存活天数、生产企业累计收益、消费企业累计收益、银行累计收益和累计归一化学习曲线下面积五项指标",
        "第60个评估时刻的均值分别为0.683和0.435",
        "第60个评估时刻的均值分别为0.773和0.683",
    ]
    for fragment in expected:
        if fragment not in all_text:
            raise RuntimeError(f"Missing expected text: {fragment!r}")
    forbidden = [
        "取E=100",
        "单个回合的最大运行时长为100天",
        "综合四项指标",
        "在四项指标上的均值曲线",
    ]
    for fragment in forbidden:
        if fragment in all_text:
            raise RuntimeError(f"Unexpected text remains: {fragment!r}")

    table = find_table_6_1(document)
    rows = {
        row.cells[0].text.strip(): row.cells[1].text.strip()
        for row in table.rows[1:]
    }
    if rows.get("每个评估时刻对应回合数") != "E=100个连续训练回合":
        raise RuntimeError("Table 6-1 E row is incorrect.")
    if rows.get("单个回合最大运行时长") != "100天":
        raise RuntimeError("Table 6-1 maximum-duration row is incorrect.")

    body_children = list(document._element.body)
    caption = find_paragraph(document, "表6-1 实验与统计设置")
    metric_intro = find_paragraph(document, "在上述训练过程中")
    if body_children.index(caption._p) >= body_children.index(metric_intro._p):
        raise RuntimeError("Table 6-1 was not moved before the evaluation metrics.")

    if len(document.inline_shapes) != 9:
        raise RuntimeError(
            f"Expected 9 inline figures after replacement, found {len(document.inline_shapes)}."
        )
    if len(document._element.body.xpath(".//m:oMath")) < 16:
        raise RuntimeError("The inserted AULC Word equation was not found.")


def build() -> None:
    for required_path in (
        SOURCE,
        FIGURE_6_1,
        FIGURE_6_2,
        MML2OMML_XSL,
    ):
        if not required_path.exists():
            raise FileNotFoundError(required_path)

    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v15 source DOCX changed after the v16 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)
    chapter_six_rewrite(document)
    verify(document)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v15 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
