from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
PROJECT_DIR = Path(__file__).resolve().parent
OUT_DIR = ROOT / "paper_drafts"
OUT_PATH = OUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿.docx"
MAIN_FIG = (
    PROJECT_DIR
    / "analysis_plots"
    / "four_metrics_three_algorithms_sem_2x2"
    / "td3_gail_transformer_four_metrics_sem_2x2.png"
)
DSCR_FIG = (
    PROJECT_DIR
    / "analysis_plots"
    / "post_training_dscr"
    / "three_algorithm_recent_100_survival_gt_90"
    / "three_algorithm_dscr_boxplot_and_risk.png"
)
SUMMARY_CSV = (
    PROJECT_DIR
    / "analysis_plots"
    / "four_metrics_three_algorithms_sem_2x2"
    / "four_metrics_sem_summary.csv"
)
DSCR_RISK_CSV = (
    PROJECT_DIR
    / "analysis_plots"
    / "post_training_dscr"
    / "three_algorithm_recent_100_survival_gt_90"
    / "three_algorithm_dscr_below_1_share.csv"
)
DSCR_BOX_CSV = (
    PROJECT_DIR
    / "analysis_plots"
    / "post_training_dscr"
    / "three_algorithm_recent_100_survival_gt_90"
    / "three_algorithm_boxplot_statistics.csv"
)


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman") -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)


def set_paragraph_format(
    paragraph,
    *,
    first_line: bool = False,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.5,
    align: int | None = None,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    if align is not None:
        paragraph.alignment = align


def set_cell_text(cell, text: str, *, bold: bool = False, center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.clear()
    r = p.add_run(text)
    set_run_font(r)
    r.font.size = Pt(9)
    r.bold = bold


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "808080")


def add_cn_paragraph(doc: Document, text: str = "", *, style: str | None = None, first_line: bool = True):
    paragraph = doc.add_paragraph(style=style)
    set_paragraph_format(paragraph, first_line=first_line, after=3, line_spacing=1.5)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
        run.font.size = Pt(10.5)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if level == 1:
        set_paragraph_format(paragraph, before=10, after=6, line_spacing=1.25)
        size = 12
    elif level == 2:
        set_paragraph_format(paragraph, before=6, after=4, line_spacing=1.25)
        size = 10.5
    else:
        set_paragraph_format(paragraph, before=4, after=3, line_spacing=1.25)
        size = 10.5
    run = paragraph.add_run(text)
    set_run_font(run, east_asia="黑体")
    run.font.size = Pt(size)
    run.bold = True
    return paragraph


def add_note(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    shade_cell(cell, "F2F2F2")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r)
    r.font.size = Pt(9)


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(3)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    footer_run._r.append(fld_begin)
    footer_run._r.append(instr)
    footer_run._r.append(fld_end)


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, after=8, line_spacing=1.2)
    r = p.add_run("基于生成对抗模仿学习的复杂经济系统主体决策模型研究")
    set_run_font(r, east_asia="黑体")
    r.font.size = Pt(16)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, after=4, line_spacing=1.2)
    r = p.add_run("——小论文大纲稿")
    set_run_font(r, east_asia="宋体")
    r.font.size = Pt(12)

    for text in [
        "作者：待补充",
        "单位：福州大学，待补充学院及专业",
        "版本说明：本稿为结构化大纲，不作为最终投稿正文。",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(p, after=2, line_spacing=1.2)
        r = p.add_run(text)
        set_run_font(r)
        r.font.size = Pt(10.5)


def add_abstract(doc: Document) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, before=8, after=4, line_spacing=1.5)
    r = p.add_run("摘  要：")
    set_run_font(r, east_asia="黑体")
    r.font.size = Pt(10.5)
    r.bold = True
    r = p.add_run(
        "复杂经济系统仿真可以在可控环境中生成经济运行数据，但主体决策模型的设计会直接影响仿真结果的真实性、稳定性和解释性。"
        "本研究拟在既有多主体经济仿真系统基础上，引入生成对抗模仿学习与 TD3 相结合的在线决策框架，并进一步讨论 Transformer 编码器在时序状态表征中的作用。"
        "当前大纲稿将论文问题、仿真环境、专家数据来源、GAIL+TD3 模型、Transformer+GAIL+TD3 模型、评价指标和实验记录组织为可继续扩写的论文框架。"
        "正式论文需要进一步补充算法细节、消融实验、统计检验和未实现内容的边界说明。"
    )
    set_run_font(r)
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, after=8, line_spacing=1.5)
    r = p.add_run("关键词：")
    set_run_font(r, east_asia="黑体")
    r.bold = True
    r.font.size = Pt(10.5)
    r = p.add_run("复杂经济系统仿真；多主体系统；生成对抗模仿学习；TD3；Transformer；偿债能力")
    set_run_font(r)
    r.font.size = Pt(10.5)


def add_outline_overview(doc: Document) -> None:
    add_heading(doc, "0  写作边界与待核对事项", 1)
    add_cn_paragraph(
        doc,
        "本文拟讨论的实验主体包括 TD3、GAIL+TD3 和 Transformer+GAIL+TD3。"
        "其中，TD3 作为基线方法，主要依据既有研究和实验日志进行背景描述；GAIL+TD3 对应分支为 v1.20-new-work；"
        "Transformer+GAIL+TD3 对应分支为 run_transformer；人类行为采集系统对应分支为 Method2_人类行为采集。",
    )
    add_cn_paragraph(
        doc,
        "正式论文不宜把未完成的规划写成已完成的实验结果。根据目前材料，真实企业行为数据采集、固有知识与约束嵌入等内容应作为研究背景、设计目标或后续工作处理。",
    )
    add_note(
        doc,
        "实验记录核对项：用户提供的 TD3 记录中，seed_739 与 seed_512 暂写为同一个 run id；当前绘图脚本中 seed_512 对应另一个 run 目录。正式定稿前需要以 swanlog 元数据或 raw_metrics 为准。"
    )


def add_intro(doc: Document) -> None:
    add_heading(doc, "1  引言", 1)
    add_heading(doc, "1.1  研究背景", 2)
    add_cn_paragraph(
        doc,
        "复杂经济系统具有主体异质性、非线性反馈和宏微观相互作用等特征。"
        "真实经济数据通常受到采集成本、环境不可控和变量不可观测等因素限制，因此，多主体经济仿真可以作为研究经济运行机制的一类辅助工具。"
    )
    add_cn_paragraph(
        doc,
        "本课题组前期构建了一个多主体经济仿真系统来实现上述目标。既有研究已经在该系统中使用 TD3 方法驱动企业与银行主体，并验证了深度强化学习方法可以提高系统运行表现[1]。"
    )
    add_heading(doc, "1.2  问题提出", 2)
    add_cn_paragraph(
        doc,
        "仅依靠强化学习训练主体时，策略会围绕环境奖励进行优化。该机制有利于提高累计收益，但它未必能充分刻画经济主体的有限理性行为。"
        "如果研究目标不仅是获得较高收益，还包括让主体行为更接近专家行为模式，并让系统更快达到较高存活水平，则需要在强化学习之外引入模仿学习信息。"
    )
    add_heading(doc, "1.3  本文思路与贡献", 2)
    add_cn_paragraph(
        doc,
        "本文拟采用生成对抗模仿学习与 TD3 相结合的框架。动作生成网络在环境交互中生成动作，判别器根据专家数据与生成动作的差异提供模仿奖励，Critic 网络则结合环境反馈与模仿奖励估计动作价值。"
    )
    add_cn_paragraph(
        doc,
        "本文进一步讨论 Transformer 编码器对时序状态表征的作用。该部分不应简单声称 Transformer 一定优于其他结构，而应依据实验曲线、稳定性和风险指标给出受数据支持的结论。"
    )


def add_environment_section(doc: Document) -> None:
    add_heading(doc, "2  复杂经济系统仿真环境", 1)
    add_heading(doc, "2.1  主体与市场结构", 2)
    add_cn_paragraph(
        doc,
        "本文沿用既有复杂经济系统仿真环境的建模方案[1]。仿真环境包含企业主体和银行主体。"
        "当前实验主要涉及生产企业、消费企业和银行。生产企业生产资本品，消费企业生产消费品，银行负责向企业发放贷款并收取本息。"
    )
    add_heading(doc, "2.2  “天”与“回合”的定义", 2)
    add_cn_paragraph(
        doc,
        "本文将一个回合定义为一次完整的系统运行过程。一个回合由若干个离散的“天”组成。"
        "每一天包括主体决策、贷款发放、商品交易、生产、清算等阶段。当企业破产或达到预设最大运行天数时，该回合结束。"
    )
    add_heading(doc, "2.3  终止条件与评价对象", 2)
    add_cn_paragraph(
        doc,
        "系统存活天数用于度量一个回合内经济系统能够持续运行的时间长度。由于任一关键企业破产都会影响系统运行，存活天数可以作为系统稳定性的直接指标。"
        "企业累计收益和银行累计收益用于衡量主体在仿真中的经营表现。"
    )
    add_note(
        doc,
        "本节建议只概述仿真环境。生产函数、交易规则、银行放贷约束等细节可写为“沿用文献[1]的建模方案”，以避免重复展开既有工作。"
    )


def add_human_collection_section(doc: Document) -> None:
    add_heading(doc, "3  人类行为采集系统与专家数据", 1)
    add_heading(doc, "3.1  专家数据在模仿学习中的作用", 2)
    add_cn_paragraph(
        doc,
        "模仿学习以专家数据作为行为范本。专家数据通常由状态和动作构成，其核心作用是为策略网络提供“在给定状态下如何行动”的行为参照。"
    )
    add_heading(doc, "3.2  当前实现边界", 2)
    add_cn_paragraph(
        doc,
        "根据当前项目描述，人类行为采集代码位于 Method2_人类行为采集分支。本文可以介绍该系统的设计目标、交互流程和专家数据格式。"
        "但是，如果真实企业行为数据尚未采集，正式论文不能写成已经使用真实企业数据完成训练。"
    )
    add_heading(doc, "3.3  专家数据来源的论文写法", 2)
    add_cn_paragraph(
        doc,
        "本文建议将专家数据来源写为三类：一是仿真系统中的人工操作数据；二是可由大语言模型辅助生成的类人决策数据；三是未来可探索的真实企业行为数据。"
        "其中，本文实际使用的数据来源需要在正式稿中依据代码和 CSV 文件名称进一步确认。"
    )


def add_gail_td3_section(doc: Document) -> None:
    add_heading(doc, "4  GAIL+TD3 主体决策模型", 1)
    add_heading(doc, "4.1  模型结构", 2)
    add_cn_paragraph(
        doc,
        "GAIL+TD3 模型将动作生成网络作为在线策略网络。该网络根据环境状态输出动作，动作一方面进入仿真环境并获得环境奖励，另一方面与专家动作共同进入判别器。"
        "判别器输出用于构造模仿奖励，Critic 网络则根据环境奖励和模仿奖励估计动作价值。"
    )
    add_heading(doc, "4.2  奖励与价值更新逻辑", 2)
    add_cn_paragraph(
        doc,
        "当前代码中的展示指标不直接等同于训练奖励。企业展示用累计收益采用 2×总收入−总支出−总利息 的形式。银行展示用累计收益采用企业累计支付利息之和。"
        "正式论文需要区分“训练奖励函数”和“论文评价指标”，避免把展示指标误写成训练目标。"
    )
    add_heading(doc, "4.3  需要在正文中补充的算法细节", 2)
    add_cn_paragraph(
        doc,
        "本节后续需要补充 GAIL 奖励的具体形式、判别器损失函数、TD3 Actor-Critic 更新过程、经验池采样方式、目标网络更新方式和关键超参数。"
        "如果这些设置在不同版本中发生变化，正式论文应选择最终实验版本，并在附录中列出主要代码分支和提交信息。"
    )


def add_transformer_section(doc: Document) -> None:
    add_heading(doc, "5  基于 Transformer 的 GAIL+TD3 模型", 1)
    add_heading(doc, "5.1  引入 Transformer 的动机", 2)
    add_cn_paragraph(
        doc,
        "经济仿真中的主体状态具有时间依赖性。当前状态不仅反映即时现金、债务和库存，也受到历史价格、交易、贷款和生产过程影响。"
        "因此，策略网络前端可以引入 Transformer 编码器，用于处理状态序列并提取时序表征。"
    )
    add_heading(doc, "5.2  当前实验中的实现边界", 2)
    add_cn_paragraph(
        doc,
        "根据此前项目讨论，Transformer 主要加入生产企业侧。正式论文需要通过 run_transformer 分支代码确认该结构是否只作用于生产企业，以及消费企业和银行是否保持原有网络结构。"
    )
    add_heading(doc, "5.3  与 GAIL+TD3 的比较假设", 2)
    add_cn_paragraph(
        doc,
        "本文可以将 Transformer+GAIL+TD3 的研究假设写为：在相同仿真环境与相近训练设置下，时序表征模块有望提高策略早期学习效率或稳定运行阶段的表现。"
        "该假设需要通过多 seed 曲线、SEM 阴影和稳定阶段风险指标验证。"
    )


def read_summary_rows() -> list[dict]:
    if not SUMMARY_CSV.exists():
        return []
    with SUMMARY_CSV.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def add_metrics_table(doc: Document, rows: list[dict]) -> None:
    if not rows:
        add_cn_paragraph(doc, "当前未读取到四指标汇总 CSV。正式稿需要补充实验结果汇总表。")
        return
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    headers = ["算法", "指标", "样本数", "第60步均值", "第60步SEM"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, center=True)
        shade_cell(table.rows[0].cells[idx], "EDEDED")
    metric_names = {
        "survival": "存活天数",
        "production": "生产企业累计收益",
        "consumption": "消费企业累计收益",
        "bank": "银行累计收益",
    }
    for row in rows:
        cells = table.add_row().cells
        values = [
            row["algorithm"],
            metric_names.get(row["metric"], row["metric"]),
            row["n"],
            f"{float(row['last_mean']):.4f}",
            f"{float(row['last_sem']):.4f}",
        ]
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, center=idx >= 2)


def add_experiment_section(doc: Document) -> None:
    add_heading(doc, "6  实验设计与结果分析", 1)
    add_heading(doc, "6.1  实验设置", 2)
    add_cn_paragraph(
        doc,
        "本文比较 TD3、GAIL+TD3 和 Transformer+GAIL+TD3 三类方法。每类方法当前均按 8 个 seed 记录训练曲线。"
        "TD3 作为基线方法；GAIL+TD3 用于检验模仿学习是否能加速学习；Transformer+GAIL+TD3 用于检验时序表征是否进一步改善稳定性或收益表现。"
    )
    add_heading(doc, "6.2  评价指标", 2)
    add_cn_paragraph(
        doc,
        "每百回合存活天数定义为最近 100 个回合的存活天数均值。该定义由 Logger.py 中 swanlab_log 的日志逻辑给出。"
    )
    add_cn_paragraph(
        doc,
        "企业每百回合累计收益定义为最近 100 个回合中 2×总收入−总支出−总利息 的均值。生产企业和消费企业分别记录。"
    )
    add_cn_paragraph(
        doc,
        "银行每百回合累计收益定义为最近 100 个回合内各企业累计支付利息之和的均值。该指标反映银行从放贷活动中获得的利息收入。"
    )
    add_cn_paragraph(
        doc,
        "偿债能力 DSCR 在环境运行中按 money/(should_payback+iDebt) 记录；代码仅在达到债务周期且当日应还本息大于 1e-6 时统计。"
        "当前 DSCR 后训练图使用最近 100 个存活天数大于 90 的回合进行计算。该图中的不同算法处理规则需要在正式论文中进一步统一或明确说明。"
    )
    add_heading(doc, "6.3  主实验结果", 2)
    add_cn_paragraph(
        doc,
        "图 1 展示三类方法在四个主要指标上的均值曲线，阴影表示 SEM。该图用于观察算法在学习速度、稳定阶段表现和 seed 间波动方面的差异。"
    )
    if MAIN_FIG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(MAIN_FIG), width=Inches(5.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(cap, first_line=False, after=6, line_spacing=1.2)
        r = cap.add_run("图 1  TD3、GAIL+TD3 与 Transformer+GAIL+TD3 的四指标对比曲线")
        set_run_font(r)
        r.font.size = Pt(9)
    add_metrics_table(doc, read_summary_rows())
    add_cn_paragraph(
        doc,
        "从当前汇总表可以直接描述第 60 个评价步的均值和 SEM。正式论文还应补充达到 80 天、90 天存活水平所需 step，以及早期 AUC 等学习速度指标。"
    )
    add_heading(doc, "6.4  偿债能力风险分析", 2)
    add_cn_paragraph(
        doc,
        "DSCR 风险图用于补充说明生产企业在稳定运行阶段是否存在偿债压力。当前后处理脚本选取最近 100 个存活天数大于 90 的回合，并统计生产企业日级 DSCR 小于 1 的观测占比。"
    )
    if DSCR_FIG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(DSCR_FIG), width=Inches(4.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(cap, first_line=False, after=6, line_spacing=1.2)
        r = cap.add_run("图 2  稳定阶段生产企业 DSCR<1 风险占比")
        set_run_font(r)
        r.font.size = Pt(9)
    add_note(
        doc,
        "正式论文需要注意：当前 DSCR 图的后处理规则包含算法差异，例如 TD3 曾使用 DSCR 除以 2 的变体，Transformer+GAIL+TD3 的某组图曾去除 DSCR<2 的观测。若该图用于投稿正文，必须统一公式或在图注中完整说明。"
    )


def add_conclusion(doc: Document) -> None:
    add_heading(doc, "7  结论与展望", 1)
    add_cn_paragraph(
        doc,
        "本文拟围绕复杂经济系统仿真中的主体决策生成问题，讨论 GAIL+TD3 和 Transformer+GAIL+TD3 相对于 TD3 基线的表现。"
        "现阶段可以谨慎表述为：模仿学习信息有助于观察主体在早期训练阶段的存活表现变化；Transformer 模块是否带来稳定优势，需要结合最终多 seed 曲线和风险指标判断。"
    )
    add_cn_paragraph(
        doc,
        "后续工作应包括三部分。第一，统一 DSCR 公式和后处理规则。第二，补充严格的统计检验和消融实验。第三，进一步完善专家数据采集来源，避免把仿真专家数据直接等同于真实企业行为数据。"
    )


def add_references(doc: Document) -> None:
    add_heading(doc, "参考文献", 1)
    refs = [
        'Chen, Bo, Shiqi Ding, Yukun Lin, and Guohong Chen. "Multi-agent Simulation of Complex Economic Systems Driven by Deep Reinforcement Learning." 2025 IEEE 17th International Conference on Computer Research and Development (ICCRD), IEEE, 2025, pp. 47-54. doi:10.1109/ICCRD64588.2025.10963160.',
        "Ho, Jonathan, and Stefano Ermon. “Generative Adversarial Imitation Learning.” Advances in Neural Information Processing Systems, 2016. [待核对并补全页码或出版信息]",
        "Fujimoto, Scott, Herke van Hoof, and David Meger. “Addressing Function Approximation Error in Actor-Critic Methods.” Proceedings of the 35th International Conference on Machine Learning, 2018. [待核对并补全页码]",
        "Vaswani, Ashish, et al. “Attention Is All You Need.” Advances in Neural Information Processing Systems, 2017. [待核对并补全页码]",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=False, after=3, line_spacing=1.25)
        r = p.add_run(ref)
        set_run_font(r)
        r.font.size = Pt(9)


def add_appendix(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "附录A  实验 run id 记录表", 1)
    add_cn_paragraph(
        doc,
        "下表根据用户提供的信息整理。正式定稿前，应使用 swanlog 元数据或 raw_metrics 文件再次核对 seed 与 run id 的一一对应关系。",
    )
    groups = {
        "TD3": [
            ("184", "run-20260621_202812-6ig7ywjq5ma7q4r3n4d2k"),
            ("291", "run-20260621_221653-3oxu02lub6ic88gg2tlyv"),
            ("83", "run-20260621_235606-cxmnv919c5vgxd18ag5m2"),
            ("739", "run-20260622_012518-o3wa23ld5k0anno4nhwb9"),
            ("512", "run-20260622_012518-o3wa23ld5k0anno4nhwb9（待核对）"),
            ("117", "run-20260622_043232-p1wsv6nx8qn3fwe50d8dd"),
            ("894", "run-20260622_063205-sgs3p05et5bj1by7pfa86"),
            ("652", "run-20260622_080822-vx1bs5edgj8z9ihb6xe1x"),
        ],
        "GAIL+TD3": [
            ("184", "run-20260621_193255-5pg7iyrvmcx34frn6ale5"),
            ("291", "run-20260621_223131-dm2s10i96kbbd7fky6owy"),
            ("83", "run-20260622_004812-zslinp7qvacgwbfuqujtw"),
            ("739", "run-20260622_031937-daocyl38dnxhsxwoxtkov"),
            ("894", "run-20260622_053309-3gob8caas2nn2pu2ozthf"),
            ("652", "run-20260622_075227-p7wj1wfq83b3mhkyo053i"),
            ("187", "run-20260622_110022-x8utkar2er1eerentcryu"),
            ("191", "run-20260622_134950-g5g4dxxf5ozv6uufkbfmk"),
        ],
        "Transformer+GAIL+TD3": [
            ("184", "run-20260622_223056-o8iu1b7pexa77lh7sd8cx"),
            ("291", "run-20260623_033112-glurxrgn07nhly7dvjkjp"),
            ("83", "run-20260623_080809-pdrdexgwuux1bd54os63u"),
            ("739", "run-20260623_131038-y6xubdud4wjny573weehp"),
            ("894", "run-20260623_172840-t8vokks22cdly7camzcud"),
            ("652", "run-20260623_212641-cj8cxudsz6o3ilixh1qc8"),
            ("187", "run-20260624_024311-i0350u1ei1u9bcogckxi4"),
            ("192", "run-20260624_073818-an8ie0ctj7ii46dc18f2a"),
        ],
    }
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for idx, header in enumerate(["算法", "seed", "run id"]):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, center=True)
        shade_cell(table.rows[0].cells[idx], "EDEDED")
    for algorithm, items in groups.items():
        for seed, run_id in items:
            row = table.add_row().cells
            set_cell_text(row[0], algorithm, center=True)
            set_cell_text(row[1], seed, center=True)
            set_cell_text(row[2], run_id)


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)
    add_abstract(doc)
    add_outline_overview(doc)
    add_intro(doc)
    add_environment_section(doc)
    add_human_collection_section(doc)
    add_gail_td3_section(doc)
    add_transformer_section(doc)
    add_experiment_section(doc)
    add_conclusion(doc)
    add_references(doc)
    add_appendix(doc)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
