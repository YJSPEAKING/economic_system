from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from lxml import etree


SCRIPT_DIR = Path(__file__).resolve().parent
PAPER_DIR = SCRIPT_DIR.parent
PROJECT_ROOT = PAPER_DIR.parent

SOURCE_DOCX = PAPER_DIR / "reference" / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_v1.docx"
OUTPUT_DOCX = SCRIPT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v13.docx"

FIGURE_6_1 = (
    PROJECT_ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "four_metrics_two_comparisons_sem_2x2"
    / "td3_vs_gail_td3_four_metrics_sem_2x2.png"
)
FIGURE_6_2 = (
    PROJECT_ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "four_metrics_two_comparisons_sem_2x2"
    / "gail_td3_vs_transformer_four_metrics_sem_2x2.png"
)
FIGURE_6_3 = (
    PROJECT_ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "post_training_dscr"
    / "three_algorithm_recent_100_survival_gt_90"
    / "three_algorithm_dscr_boxplot_and_risk.png"
)

SOURCE_SHA256 = "4ea22a205bc4eb575e66c1ef440a31b8db0c172afdff451f855c267692d0ba39"
MML2OMML_XSL = Path(r"C:\Program Files\Microsoft Office\Root\Office16\MML2OMML.XSL")


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_inputs() -> None:
    required = [SOURCE_DOCX, FIGURE_6_1, FIGURE_6_2, FIGURE_6_3, MML2OMML_XSL]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required files:\n" + "\n".join(missing))
    actual_hash = file_sha256(SOURCE_DOCX)
    if actual_hash.lower() != SOURCE_SHA256:
        raise RuntimeError(
            "The source document changed after template distillation. "
            f"Expected {SOURCE_SHA256}, got {actual_hash}."
        )


def move_before_reference(block, reference_paragraph) -> None:
    reference_paragraph._p.addprevious(block)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if value and keep_next is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and keep_next is not None:
        p_pr.remove(keep_next)


def set_run_font(run, *, size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")


def add_heading(doc: Document, reference_paragraph, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    move_before_reference(paragraph._p, reference_paragraph)


def add_body(doc: Document, reference_paragraph, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    move_before_reference(paragraph._p, reference_paragraph)


def add_caption(doc: Document, reference_paragraph, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    set_keep_with_next(paragraph, True)
    run = paragraph.add_run(text)
    set_run_font(run, size=9)
    move_before_reference(paragraph._p, reference_paragraph)


def add_figure(
    doc: Document,
    reference_paragraph,
    image_path: Path,
    caption: str,
    width_inches: float,
) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    set_keep_with_next(paragraph, True)
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    move_before_reference(paragraph._p, reference_paragraph)
    add_caption(doc, reference_paragraph, caption)


def mathml_to_omml(mathml: str):
    transform = etree.XSLT(etree.parse(str(MML2OMML_XSL)))
    result = transform(etree.fromstring(mathml.encode("utf-8")))
    return etree.fromstring(etree.tostring(result.getroot()))


def add_equation(doc: Document, reference_paragraph, mathml: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph._p.append(mathml_to_omml(mathml))
    move_before_reference(paragraph._p, reference_paragraph)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, centered: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=bold)


def add_settings_table(doc: Document, reference_paragraph) -> None:
    add_caption(doc, reference_paragraph, "表6-1 实验与统计设置")
    rows = [
        ("对比方法", "TD3、GAIL+TD3、Transformer+GAIL+TD3"),
        ("随机种子组数", "每种方法8组"),
        ("每个评估时刻对应回合数", "100个训练回合"),
        ("单回合最大天数", "100天"),
        ("统计方法", "8次独立运行的均值及均值上下1个SEM"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(5.0), Cm(10.6)]
    for index, text in enumerate(("项目", "设定")):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        shade_cell(cell, "F2F2F2")
        set_cell_text(cell, text, bold=True, centered=True)
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = widths[0]
        cells[1].width = widths[1]
        set_cell_text(cells[0], left, centered=True)
        set_cell_text(cells[1], right)
    move_before_reference(table._tbl, reference_paragraph)


def remove_existing_chapter_six_body(doc: Document, chapter_heading, reference_heading) -> None:
    body = doc._body._element
    chapter_index = body.index(chapter_heading._p)
    reference_index = body.index(reference_heading._p)
    for element in list(body)[chapter_index + 1 : reference_index]:
        body.remove(element)


def find_unique_heading(doc: Document, text: str):
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one heading {text!r}, found {len(matches)}.")
    return matches[0]


def build_document() -> None:
    validate_inputs()
    SCRIPT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_DOCX, OUTPUT_DOCX)
    doc = Document(OUTPUT_DOCX)
    chapter_heading = find_unique_heading(doc, "6  实验设置与结果分析")
    reference_heading = find_unique_heading(doc, "参考文献")
    remove_existing_chapter_six_body(doc, chapter_heading, reference_heading)

    add_heading(doc, reference_heading, "6.1 实验设置与评价指标", 2)
    add_body(
        doc,
        reference_heading,
        "为考察模仿学习信号与历史状态表征对复杂经济系统主体训练过程的影响，本文设置三种对比方法。第一种方法为双延迟深度确定性策略梯度算法TD3，并将其作为基线方法。第二种方法为GAIL+TD3，该方法在生产企业主体的TD3训练过程中引入生成对抗模仿学习信号。第三种方法为Transformer+GAIL+TD3，该方法在GAIL+TD3的基础上加入Transformer编码器，以利用生产企业主体的历史状态序列。三种方法均使用第五章整理得到的专家数据，并采用统一的训练条件和统计口径。",
    )
    add_body(
        doc,
        reference_heading,
        "神经网络参数初始化、动作探索、训练样本抽取和经验池采样等过程均包含随机因素。本文将用于控制上述随机因素的联合配置统称为“随机种子”。三种方法分别在8组随机种子下独立运行。",
    )
    add_body(
        doc,
        reference_heading,
        "本文每完成100个训练回合汇总一次评价指标，并将该汇总位置记为一个评估时刻（evaluation step）。第k个评估时刻对应第k组连续100个训练回合的统计结果。本文在训练过程中跟踪平均存活天数、主体累计收益和生产企业偿债能力。",
    )
    add_body(
        doc,
        reference_heading,
        "平均存活天数反映一个评估时刻内系统维持运行的平均时长。该指标等于第k个评估时刻所包含100个回合的存活天数算术平均值，其定义如下：",
    )
    add_equation(
        doc,
        reference_heading,
        '<math xmlns="http://www.w3.org/1998/Math/MathML"><mover><msub><mi>L</mi><mi>k</mi></msub><mo>¯</mo></mover><mo>=</mo><mfrac><mn>1</mn><mn>100</mn></mfrac><munderover><mo>∑</mo><mrow><mi>n</mi><mo>=</mo><mn>1</mn></mrow><mn>100</mn></munderover><msub><mi>L</mi><mrow><mi>k</mi><mo>,</mo><mi>n</mi></mrow></msub></math>',
    )
    add_body(
        doc,
        reference_heading,
        "累计收益分别包括生产企业累计收益、消费企业累计收益和银行累计收益。对于企业主体，本文首先在单个回合内分别汇总收入、支出和利息，再按照下式计算企业累计收益：",
    )
    add_equation(
        doc,
        reference_heading,
        '<math xmlns="http://www.w3.org/1998/Math/MathML"><msubsup><mi>J</mi><mrow><mi>i</mi><mo>,</mo><mi>n</mi></mrow><mtext>firm</mtext></msubsup><mo>=</mo><mn>2</mn><msub><mi>R</mi><mrow><mi>i</mi><mo>,</mo><mi>n</mi></mrow></msub><mo>−</mo><msub><mi>C</mi><mrow><mi>i</mi><mo>,</mo><mi>n</mi></mrow></msub><mo>−</mo><msub><mi>I</mi><mrow><mi>i</mi><mo>,</mo><mi>n</mi></mrow></msub></math>',
    )
    add_body(
        doc,
        reference_heading,
        "式中，J表示企业累计收益，R表示累计收入，C表示累计支出，I表示累计利息，i和n分别表示企业编号与回合编号。",
    )
    add_body(
        doc,
        reference_heading,
        "生产企业和消费企业分别依据上述定义计算累计收益。银行累计收益按照企业向银行支付的利息之和计算，其定义如下：",
    )
    add_equation(
        doc,
        reference_heading,
        '<math xmlns="http://www.w3.org/1998/Math/MathML"><msubsup><mi>J</mi><mi>n</mi><mtext>bank</mtext></msubsup><mo>=</mo><munderover><mo>∑</mo><mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow><mi>m</mi></munderover><msub><mi>I</mi><mrow><mi>i</mi><mo>,</mo><mi>n</mi></mrow></msub></math>',
    )
    add_body(
        doc,
        reference_heading,
        "式中，J表示银行累计收益，m表示系统中的企业主体数量，I表示对应企业在该回合内支付的累计利息。",
    )
    add_body(
        doc,
        reference_heading,
        "生产企业偿债能力采用偿债覆盖率（DSCR: Debt Service Coverage Ratio）进行度量。该指标以生产企业当日可支配现金与当期应还本息之比表示，其定义如下：",
    )
    add_equation(
        doc,
        reference_heading,
        '<math xmlns="http://www.w3.org/1998/Math/MathML"><msub><mi>DSCR</mi><mi>t</mi></msub><mo>=</mo><mfrac><msub><mi>money</mi><mi>t</mi></msub><mrow><msub><mi>should_payback</mi><mi>t</mi></msub><mo>+</mo><msub><mi>iDebt</mi><mi>t</mi></msub><mo>+</mo><mi>ε</mi></mrow></mfrac></math>',
    )
    add_body(
        doc,
        reference_heading,
        "式中，分子表示生产企业第t天可支配现金；分母中的前两项分别表示当日应归还本金与当日应支付利息；ε表示防止分母为零的正小常数。DSCR小于1表示生产企业当日现金不足以覆盖当期应还本息。",
    )
    add_body(
        doc,
        reference_heading,
        "每种方法在8组随机种子下形成8条指标序列。本文在每个评估时刻计算8次独立运行结果的算术平均值，并以该均值作为曲线值。均值标准误（SEM: Standard Error of the Mean）用于描述8次独立运行均值的不确定性，其定义为：",
    )
    add_equation(
        doc,
        reference_heading,
        '<math xmlns="http://www.w3.org/1998/Math/MathML"><mi>SEM</mi><mo>=</mo><mfrac><mi>s</mi><msqrt><mi>S</mi></msqrt></mfrac></math>',
    )
    add_body(
        doc,
        reference_heading,
        "式中，s表示8次运行结果的样本标准差，S表示独立运行次数，本文取S=8。训练曲线中的实线表示8次独立运行的均值，阴影表示均值上下1个SEM。表6-1汇总了本文采用的实验与统计设置。",
    )
    add_settings_table(doc, reference_heading)

    add_heading(doc, reference_heading, "6.2 TD3与GAIL+TD3的结果对比", 2)
    add_body(
        doc,
        reference_heading,
        "本文首先比较TD3与GAIL+TD3的训练过程，以分析模仿学习信号对生产企业主体决策训练的影响。图6-1给出了两种方法在平均存活天数、生产企业累计收益、消费企业累计收益和银行累计收益四项指标上的均值曲线及SEM阴影。",
    )
    add_figure(
        doc,
        reference_heading,
        FIGURE_6_1,
        "图6-1 TD3与GAIL+TD3训练结果对比",
        6.15,
    )
    add_body(
        doc,
        reference_heading,
        "图6-1(a)显示，GAIL+TD3的平均存活天数在训练前期迅速提高，并在第20个评估时刻之前进入相对稳定区间。TD3的平均存活天数在训练前期上升较缓，随后持续提高，并在训练后期接近GAIL+TD3的水平。该结果说明，判别器提供的模仿奖励能够在训练早期为生产企业策略提供与专家行为分布相关的引导，使系统较早形成可维持较长运行时间的主体决策组合。",
    )
    add_body(
        doc,
        reference_heading,
        "图6-1(b)显示，GAIL+TD3的生产企业累计收益在训练前期较快上升，并较早进入相对稳定阶段。TD3的生产企业累计收益在训练中后期快速提高，其训练后期均值高于GAIL+TD3，同时对应的SEM阴影更宽。该现象反映了两种训练方案在收敛速度、收益水平与跨随机种子波动之间的表现差异。系统存活时间与单一主体累计收益对应不同的决策目标，因此，训练后期较高的企业收益不能单独代表系统整体运行质量。",
    )
    add_body(
        doc,
        reference_heading,
        "图6-1(c)显示，消费企业累计收益虽然未直接接受模仿学习信号，但其曲线随生产企业策略变化而发生同步调整。GAIL+TD3在训练前期较快提高消费企业累计收益并较早进入稳定区间，TD3则在训练后期达到更高的累计收益均值。该结果表明，生产企业主体的策略更新会通过产品交易和主体间交互影响消费企业的经营结果。",
    )
    add_body(
        doc,
        reference_heading,
        "图6-1(d)显示，GAIL+TD3的银行累计收益在训练前期迅速提高，之后围绕较高水平波动；TD3的银行累计收益提升时间较晚，训练后期均值仍低于GAIL+TD3。银行收益来自企业支付的利息，因此，该指标记录了企业经营决策对信贷关系的间接影响。综合四项指标，GAIL+TD3的主要表现特征是较快提高系统存活天数和各主体收益，并较早形成相对稳定的训练曲线；TD3在部分企业累计收益指标上的后期均值较高，体现了不同决策目标之间的表现差异。",
    )

    add_heading(doc, reference_heading, "6.3 GAIL+TD3与Transformer+GAIL+TD3的结果对比", 2)
    add_body(
        doc,
        reference_heading,
        "本文进一步比较GAIL+TD3与Transformer+GAIL+TD3，以分析历史状态序列表征对生产企业主体训练过程的影响。图6-2给出了两种方法在四项指标上的均值曲线及SEM阴影。",
    )
    add_figure(
        doc,
        reference_heading,
        FIGURE_6_2,
        "图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        6.15,
    )
    add_body(
        doc,
        reference_heading,
        "图6-2(a)显示，Transformer+GAIL+TD3的平均存活天数在训练前期更早进入快速上升阶段，并在稳定阶段保持高于GAIL+TD3的平均水平。两种方法在训练前期均能较快提升系统存活时间，但Transformer编码器对历史状态序列的整合使生产企业主体能够同时利用当前状态与近期交互信息，从而较早形成较稳定的决策。",
    )
    add_body(
        doc,
        reference_heading,
        "图6-2(b)显示，Transformer+GAIL+TD3的生产企业累计收益在训练前期上升更快，稳定阶段的均值也高于GAIL+TD3。该方法对应的SEM阴影在部分评估时刻较宽，说明不同随机种子下的收益水平仍存在波动。该结果表明，历史状态表征有助于生产企业主体在产量、采购、贷款和定价等连续决策之间建立跨时刻联系，但收益表现仍受到不同训练轨迹的影响。",
    )
    add_body(
        doc,
        reference_heading,
        "图6-2(c)显示，Transformer+GAIL+TD3的消费企业累计收益在多数评估时刻高于GAIL+TD3，并在训练后期保持较高均值。消费企业未直接加入Transformer编码器，其收益变化主要来自生产企业策略对商品交易、资金流转和系统运行时长的影响。因此，该子图反映了生产企业历史状态表征对其他企业主体经营结果的间接作用。",
    )
    add_body(
        doc,
        reference_heading,
        "图6-2(d)显示，Transformer+GAIL+TD3的银行累计收益在训练前期出现较大波动，之后进入相对稳定区间；其训练后期均值低于GAIL+TD3。银行累计收益由企业支付利息形成，生产企业贷款需求、现金流状况和存活时间的变化会共同改变银行收益。该结果说明，历史状态表征提高企业累计收益和系统存活水平的同时，也会改变企业与银行之间的信贷关系，银行收益不会与企业收益保持完全一致的变化方向。",
    )
    add_body(
        doc,
        reference_heading,
        "为进一步考察三种主体训练方案下生产企业的当期偿债风险，本文统计稳定运行回合中的生产企业日观测，并计算DSCR小于1的观测占比。图6-3给出了三种方法的偿债风险统计结果。",
    )
    add_figure(
        doc,
        reference_heading,
        FIGURE_6_3,
        "图6-3 三种主体训练方案下生产企业偿债风险对比",
        5.2,
    )
    add_body(
        doc,
        reference_heading,
        "图6-3显示，TD3、GAIL+TD3和Transformer+GAIL+TD3中DSCR小于1的生产企业日观测占比分别为6.52%、0.23%和0.00%。DSCR小于1表示生产企业当日可支配现金不足以覆盖当期应还本金与利息，因此，较低的风险占比对应更少的当期偿债缺口。在本组稳定运行观测中，GAIL+TD3的偿债风险占比低于TD3，Transformer+GAIL+TD3未出现DSCR小于1的日观测。该结果与图6-2(a)中的稳定阶段存活水平共同说明，历史状态表征能够帮助生产企业主体在经营扩张、现金持有和债务偿付之间形成更稳定的决策关系。",
    )

    doc.save(OUTPUT_DOCX)
    if file_sha256(SOURCE_DOCX).lower() != SOURCE_SHA256:
        raise RuntimeError("The source document was modified during generation.")
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
