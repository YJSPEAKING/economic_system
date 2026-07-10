from pathlib import Path
import shutil

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT / "paper_drafts" / "rendered_v7"
OUT_DIR = ROOT / "paper_drafts" / "rendered_v8"


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_paragraph_text(paragraph, text):
    clear_runs(paragraph)
    paragraph.add_run(text)


def set_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    clear_runs(paragraph)
    paragraph.add_run(text)
    for extra in cell.paragraphs[1:]:
        clear_runs(extra)


def main():
    src_files = list(SRC_DIR.glob("*.docx"))
    if not src_files:
        raise FileNotFoundError(f"No docx found in {SRC_DIR}")
    src = src_files[0]

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / src.name.replace("v7", "v8")
    shutil.copy2(src, out)

    doc = Document(out)

    for paragraph in doc.paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip()
        if text.startswith("2  "):
            set_paragraph_text(paragraph, "2  仿真环境、评价目标与主体决策建模问题")
        elif text.startswith("仿真主体的外部评价需要先区分完全理性目标"):
            set_paragraph_text(
                paragraph,
                "复杂经济系统仿真一方面服务于理论经济学研究中的数据生成需求，另一方面需要给出可解释的主体决策机制。真实经济运行数据往往存在采集成本高、环境变量难以控制、理论抽象变量难以直接观测等问题，因此，MAS-E 仿真通过可控规则生成虚拟经济运行数据。仿真主体的评价目标需要与研究任务一致。若研究任务关注规则环境下的最优响应，则可以围绕完全理性目标展开，考察主体能否利用奖励信号获得较高收益、延长存活时间并维持系统运行。",
            )
        elif text.startswith("部分理性评价的目标不同。有限理性"):
            set_paragraph_text(
                paragraph,
                "有限理性目标并不要求主体在每个状态下输出全局最优动作。真实市场主体会受到信息不完全、风险态度、经验依赖、偏好差异和路径依赖等因素影响，其行为通常不能简化为单一收益最大化。因而，有限理性目标更关注生成行为是否接近人类或专家样本中体现的决策模式，并进一步考察动作分布、状态覆盖和动态适应性。对本文而言，有限理性目标主要服务于生产企业主体的行为生成，模仿学习因此成为强化学习之外的重要补充。",
            )
        elif text.startswith("专家数据是连接有限理性行为与模型训练目标的中间表示。前期项目研究材料"):
            set_paragraph_text(
                paragraph,
                "专家数据是连接有限理性行为与模型训练目标的中间表示。本文具体使用两类专家数据：第一类来自网页实验系统中的人工参与式采集，参与者扮演甲公司经理，并依据每日可见状态填写贷款、采购和定价动作；第二类来自同一仿真接口下的大模型辅助批量采集，该采集过程仍经过环境执行、动作边界检查和存活天数筛选。两类数据均以生产企业的状态-动作对保存，并用于后续 GAIL 与 TD3 训练。",
            )

    if doc.tables:
        table = doc.tables[0]
        set_cell_text(
            table.rows[1].cells[1],
            "33 维状态向量：包含经营状态、需求信息、交易记录、价格信息和时间状态变量",
        )
        set_cell_text(
            table.rows[2].cells[1],
            "51 维状态向量：包含银行经营状态与企业主体信息",
        )

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
