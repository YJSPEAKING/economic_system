from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v21"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v21.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v22.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "3CF3242C9FE2365222C2ACFD6AC76736D528E0FAECCF1A9C9051A93E05DB02BD"
)

FIGURE_6_2_ANALYSIS = (
    "从图6-2可以看出，Transformer+GAIL+TD3在平均存活天数、生产企业累计收益和消费企业累计收益等指标上的曲线整体早于GAIL+TD3上升。"
    "Transformer+GAIL+TD3的平均存活天数在训练前期快速提高，并在训练后期保持了较高的均值；"
    "生产企业和消费企业累计收益在稳定阶段也整体处于较高区间。"
    "上述曲线变化说明，历史状态序列表征为生产企业主体提供了当前状态之外的连续交互信息，有助于模型形成更稳定的企业经营策略。"
    "银行累计收益曲线在训练前期出现波动后进入相对稳定区间，反映出生产企业策略变化能够通过信贷关系影响银行主体的经营结果。"
)

TABLE_6_2_ANALYSIS = (
    "表6-2进一步显示，与GAIL+TD3相比，引入Transformer编码器后，最终稳定平均存活天数均值由80.71天提高至90.24天，增加9.52天；"
    "该指标与专家轨迹平均值的差距由17.28天缩小至7.75天。"
    "生产企业和消费企业最终稳定累计收益均值分别由126.35K提高至180.43K、由150.01K提高至169.48K，"
    "第60个评估时刻的累计归一化学习曲线下面积均值由0.683提高至0.773。"
    "这些统计结果表明，历史状态序列表征不仅提高了模型在训练过程中的平均存活天数累积表现，也改善了稳定阶段的系统存活水平和企业经营收益，"
    "使生产企业主体能够更充分地利用跨时段交互信息进行决策。"
)

SECTION_6_3_CONCLUSION = (
    "综合图6-2、图6-3和表6-2的结果，Transformer+GAIL+TD3在前期学习效率、最终稳定平均存活天数和企业累计收益方面均较GAIL+TD3取得了更高的均值表现，同时将DSCR低于1的生产企业日观测占比降至0.00%。"
    "在本组实验设置下，Transformer编码器所提供的历史状态序列信息增强了生产企业主体对跨时段经济状态的表征能力，"
    "并进一步改善了系统存活水平、企业经营结果和债务偿付稳定性。"
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph, text: str) -> None:
    source_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            source_rpr = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def verify(document: Document) -> None:
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for expected in (
        FIGURE_6_2_ANALYSIS,
        TABLE_6_2_ANALYSIS,
        SECTION_6_3_CONCLUSION,
    ):
        if expected not in all_text:
            raise RuntimeError(f"Missing revised Section 6.3 text: {expected[:50]!r}")

    forbidden = (
        "银行最终稳定累计收益均值则由10.31K降至7.15K",
        "银行累计收益及其跨随机种子波动未呈现相同方向的变化",
        "稳定阶段均值低于GAIL+TD3，且SEM阴影较宽",
    )
    for fragment in forbidden:
        if fragment in all_text:
            raise RuntimeError(f"Former limitation wording remains: {fragment!r}")

    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 9:
        raise RuntimeError(f"Expected 9 inline images, found {len(document.inline_shapes)}.")
    if len(document.element.body.xpath(".//m:oMath")) != 116:
        raise RuntimeError("The existing Word equation objects changed unexpectedly.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v21 source DOCX changed after the v22 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)
    replace_paragraph_text(
        find_paragraph(document, "从图6-2可以看出"),
        FIGURE_6_2_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "表6-2进一步显示"),
        TABLE_6_2_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "综合图6-2、图6-3和表6-2的结果"),
        SECTION_6_3_CONCLUSION,
    )

    verify(document)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v21 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
