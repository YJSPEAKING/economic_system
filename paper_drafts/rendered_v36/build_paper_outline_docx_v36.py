from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import importlib.util
from pathlib import Path
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image

BUNDLED_SITE_PACKAGES = Path(
    r"C:\Users\legion\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\python\Lib\site-packages"
)
if str(BUNDLED_SITE_PACKAGES) not in sys.path:
    sys.path.append(str(BUNDLED_SITE_PACKAGES))

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "reference"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_v1.docx"
)
OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT = OUTPUT_DIR / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v36.docx"
EXPECTED_SOURCE_SHA256 = (
    "753659DBA9394BA37EC057DF1C61A53450F779258635D8E2E30C234253B92421"
)
MML2OMML_XSL = Path(
    r"C:\Program Files\Microsoft Office\Root\Office16\MML2OMML.XSL"
)
FIGURE_OUTPUT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "paper_gail_component_validation"
)
NOISY_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "expert_vs_noisy_expert_discriminator_eval_noise_std_0p30"
)
ACTOR_RESULT_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "actor_snapshot_evolution"
)
DISCRIMINATOR_FIGURE = FIGURE_OUTPUT_DIR / "expert_vs_noisy_expert_seed184_paper.png"
ACTOR_FIGURE = FIGURE_OUTPUT_DIR / "actor_expert_js_evolution_mean_sem_paper.png"
ENGLISH_FIGURE_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "paper_v29_english_figures"
)
TD3_GAIL_FIGURE = ENGLISH_FIGURE_DIR / "td3_vs_gail_td3_three_metrics_sem.png"
GAIL_TRANSFORMER_FIGURE = (
    ENGLISH_FIGURE_DIR / "gail_td3_vs_transformer_three_metrics_sem.png"
)
DSCR_RISK_FIGURE = ENGLISH_FIGURE_DIR / "three_algorithm_dscr_risk.png"
ACTOR_STRUCTURE_FIGURE = (
    ROOT
    / "paper_drafts"
    / "figures"
    / "actor_network_structure"
    / "production_actor_network_structure.png"
)
CRITIC_STRUCTURE_FIGURE = (
    ROOT
    / "paper_drafts"
    / "figures"
    / "critic_network_structure"
    / "production_critic_network_structure.png"
)
DISCRIMINATOR_STRUCTURE_FIGURE = (
    ROOT
    / "paper_drafts"
    / "figures"
    / "discriminator_network_structure"
    / "production_discriminator_network_structure.png"
)
GAIL_TD3_TRAINING_FIGURE = (
    ROOT
    / "paper_drafts"
    / "figures"
    / "gail_td3_training_scheme"
    / "gail_td3_training_scheme_en.png"
)
TRANSFORMER_THREE_BRANCH_FIGURE = (
    ROOT
    / "paper_drafts"
    / "figures"
    / "transformer_three_branch_structure"
    / "transformer_three_branch_structure.png"
)
DAY_SIMULATION_FIGURE = (
    ROOT
    / "paper_drafts"
    / "reference"
    / "初步设计的仿真运行规则示意图_day.png"
)
TRAINING_EPISODE_FIGURE = (
    ROOT
    / "paper_drafts"
    / "reference"
    / "初步设计的仿真运行规则示意图.png"
)
LEARNING_CURVE_SCRIPT = (
    ROOT / "real_System_remake" / "plot_five_metrics_two_comparisons_aulc_sem.py"
)
LEARNING_CURVE_RAW_DIR = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "five_metrics_two_comparisons_aulc_sem"
    / "raw_metrics"
)
LATEST_EXPERIMENT_STATS: dict[str, dict[str, float | int | None]] = {}
DSCR_RISK_CSV = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "post_training_dscr"
    / "three_algorithm_recent_100_survival_gt_90"
    / "three_algorithm_dscr_below_1_share.csv"
)


NOISY_ACTION_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mover><mi>a</mi><mo>~</mo></mover><mi>i</mi><mi>E</mi></msubsup>
  <mo>=</mo><mi mathvariant="normal">clip</mi><mfenced>
  <msubsup><mover><mi>a</mi><mo>^</mo></mover><mi>i</mi><mi>E</mi></msubsup>
  <mo>+</mo><msub><mi>ε</mi><mi>i</mi></msub><mo>,</mo><mo>−</mo><mi>b</mi>
  <mo>,</mo><mi>b</mi></mfenced><mo>,</mo>
  <msub><mi>ε</mi><mi>i</mi></msub><mo>∼</mo>
  <mi mathvariant="normal">N</mi><mfenced><mn>0</mn><mo>,</mo>
  <msubsup><mi>σ</mi><mtext>noise</mtext><mn>2</mn></msubsup></mfenced>
</math>
"""

DISCRIMINATOR_SCORE_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mi>d</mi><mi>i</mi><mi>E</mi></msubsup><mo>=</mo>
  <msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mover><mi>a</mi><mo>^</mo></mover><mi>i</mi><mi>E</mi></msubsup></mfenced>
  <mo>,</mo><mspace width="1em"/>
  <msubsup><mi>d</mi><mi>i</mi><mi>N</mi></msubsup><mo>=</mo>
  <msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mover><mi>a</mi><mo>~</mo></mover><mi>i</mi><mi>E</mi></msubsup></mfenced>
</math>
"""

JS_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi>M</mi><mo>=</mo><mfrac><mrow><mi>P</mi><mo>+</mo><mi>Q</mi></mrow><mn>2</mn></mfrac><mo>,</mo><mspace width="1em"/>
  <mi mathvariant="normal">JS</mi><mfenced open="(" close=")">
  <mi>P</mi><mo>∥</mo><mi>Q</mi></mfenced>
  <mo>=</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">KL</mi><mfenced><mi>P</mi><mo>∥</mo><mi>M</mi></mfenced>
  <mo>+</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">KL</mi><mfenced><mi>Q</mi><mo>∥</mo><mi>M</mi></mfenced>
</math>
"""

OVERLAP_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi mathvariant="normal">OVL</mi><mfenced><mi>P</mi><mo>,</mo><mi>Q</mi></mfenced><mo>=</mo>
  <munderover><mo>∑</mo><mrow><mi>h</mi><mo>=</mo><mn>1</mn></mrow><mi>H</mi></munderover>
  <mi mathvariant="normal">min</mi><mfenced>
  <msub><mi>P</mi><mi>h</mi></msub><mo>,</mo>
  <msub><mi>Q</mi><mi>h</mi></msub></mfenced>
</math>
"""

AUC_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <mi mathvariant="normal">AUC</mi><mo>=</mo>
  <mi mathvariant="normal">Pr</mi><mfenced>
  <msubsup><mi>x</mi><mi>i</mi><mo>+</mo></msubsup><mo>&gt;</mo>
  <msubsup><mi>x</mi><mi>j</mi><mo>−</mo></msubsup></mfenced>
  <mo>+</mo><mfrac><mn>1</mn><mn>2</mn></mfrac>
  <mi mathvariant="normal">Pr</mi><mfenced>
  <msubsup><mi>x</mi><mi>i</mi><mo>+</mo></msubsup><mo>=</mo>
  <msubsup><mi>x</mi><mi>j</mi><mo>−</mo></msubsup></mfenced>
</math>
"""

ACTOR_SCORE_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mi>a</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup>
  <mo>=</mo><msub><mi>μ</mi><msub><mi>θ</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub></mfenced><mo>,</mo><mspace width="1em"/>
  <msubsup><mi>d</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup>
  <mo>=</mo><msub><mi>D</mi><mrow><msubsup><mi>φ</mi><mi>r</mi><mo>*</mo></msubsup></mrow></msub>
  <mfenced><msub><mover><mi>s</mi><mo>^</mo></mover><mi>i</mi></msub><mo>,</mo>
  <msubsup><mi>a</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi><mo>,</mo><mi>i</mi></mrow><mi>G</mi></msubsup></mfenced>
</math>
"""

ACTOR_JS_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mi>J</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub><mo>=</mo>
  <mi mathvariant="normal">JS</mi><mfenced>
  <msubsup><mi>P</mi><mi>r</mi><mi>E</mi></msubsup><mo>∥</mo>
  <msubsup><mi>P</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow><mi>G</mi></msubsup>
  </mfenced>
</math>
"""

ACTOR_MEAN_SEM_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msub><mover><mi>J</mi><mo>¯</mo></mover><mi>k</mi></msub><mo>=</mo>
  <mfrac><mn>1</mn><mi>R</mi></mfrac>
  <munderover><mo>∑</mo><mrow><mi>r</mi><mo>=</mo><mn>1</mn></mrow><mi>R</mi></munderover>
  <msub><mi>J</mi><mrow><mi>r</mi><mo>,</mo><mi>k</mi></mrow></msub><mo>,</mo><mspace width="1em"/>
  <msub><mi>SEM</mi><mi>k</mi></msub><mo>=</mo><mfrac><msub><mi>s</mi><mi>k</mi></msub><msqrt><mi>R</mi></msqrt></mfrac>
</math>
"""

ACTOR_NETWORK_MATHML = r"""
<math xmlns="http://www.w3.org/1998/Math/MathML">
  <msubsup><mi>h</mi><mi>t</mi><mn>1</mn></msubsup><mo>=</mo>
  <mi mathvariant="normal">tanh</mi><mfenced>
  <msub><mi>W</mi><mn>1</mn></msub>
  <msub><mover><mi>s</mi><mo>^</mo></mover><mi>t</mi></msub>
  <mo>+</mo><msub><mi>b</mi><mn>1</mn></msub></mfenced><mo>,</mo><mspace width="1em"/>
  <msubsup><mi>h</mi><mi>t</mi><mn>2</mn></msubsup><mo>=</mo>
  <mi mathvariant="normal">LeakyReLU</mi><mfenced>
  <msub><mi>W</mi><mn>2</mn></msub>
  <msubsup><mi>h</mi><mi>t</mi><mn>1</mn></msubsup>
  <mo>+</mo><msub><mi>b</mi><mn>2</mn></msub></mfenced><mo>,</mo><mspace width="1em"/>
  <msub><mi>a</mi><mi>t</mi></msub><mo>=</mo>
  <msub><mi>a</mi><mtext>max</mtext></msub>
  <mi mathvariant="normal">tanh</mi><mfenced>
  <msub><mi>W</mi><mn>3</mn></msub>
  <msubsup><mi>h</mi><mi>t</mi><mn>2</mn></msubsup>
  <mo>+</mo><msub><mi>b</mi><mn>3</mn></msub></mfenced>
</math>
"""

EQUATIONS = (
    ACTOR_NETWORK_MATHML,
    NOISY_ACTION_MATHML,
    DISCRIMINATOR_SCORE_MATHML,
    JS_MATHML,
    OVERLAP_MATHML,
    AUC_MATHML,
    ACTOR_SCORE_MATHML,
    ACTOR_JS_MATHML,
    ACTOR_MEAN_SEM_MATHML,
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def first_run_properties(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    source_rpr = first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def insert_paragraph_after(
    reference: Paragraph,
    text: str = "",
    *,
    format_source: Paragraph,
) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    if format_source._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(format_source._p.pPr))
    if text:
        run = paragraph.add_run(text)
        source_rpr = first_run_properties(format_source)
        if source_rpr is not None:
            run._r.insert(0, source_rpr)
    return paragraph


def insert_paragraph_before(
    reference: Paragraph,
    text: str = "",
    *,
    format_source: Paragraph,
) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addprevious(element)
    paragraph = Paragraph(element, reference._parent)
    if format_source._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(format_source._p.pPr))
    if text:
        run = paragraph.add_run(text)
        source_rpr = first_run_properties(format_source)
        if source_rpr is not None:
            run._r.insert(0, source_rpr)
    return paragraph


def mathml_to_omml(mathml: str):
    transform = etree.XSLT(etree.parse(str(MML2OMML_XSL)))
    result = transform(etree.fromstring(mathml.encode("utf-8")))
    return etree.fromstring(etree.tostring(result.getroot()))


def insert_equation_after(
    reference: Paragraph,
    mathml: str,
    *,
    equation_template: Paragraph,
) -> Paragraph:
    paragraph = insert_paragraph_after(
        reference, format_source=equation_template
    )
    paragraph._p.append(mathml_to_omml(mathml))
    return paragraph


def insert_picture_after(
    reference: Paragraph,
    image_path: Path,
    *,
    body_template: Paragraph,
    width_inches: float,
) -> Paragraph:
    paragraph = insert_paragraph_after(reference, format_source=body_template)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    return paragraph


def insert_results_table_after(
    document: Document,
    reference: Paragraph,
    rows: tuple[tuple[str, str], ...],
) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.style = document.tables[2].style
    table.autofit = False
    widths = (4.25, 2.25)
    header_rpr = first_run_properties(document.tables[2].cell(0, 0).paragraphs[0])
    body_rpr = first_run_properties(document.tables[2].cell(1, 0).paragraphs[0])

    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            source_rpr = header_rpr if row_index == 0 else body_rpr
            if source_rpr is not None:
                run._r.insert(0, deepcopy(source_rpr))

    reference._p.addnext(table._tbl)


def set_table_cell_borders(cell, *, top=None, bottom=None) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)

    for edge_name in ("left", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")

    for edge_name, edge_specification in (("top", top), ("bottom", bottom)):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        if edge_specification is None:
            edge.set(qn("w:val"), "nil")
        else:
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), str(edge_specification))
            edge.set(qn("w:space"), "0")
            edge.set(qn("w:color"), "000000")


def insert_transformer_parameter_table_after(
    document: Document,
    reference: Paragraph,
) -> None:
    rows = (
        (
            "分支",
            "历史输入",
            "序列长度（l）",
            "编码层数",
            "注意力头数",
            "表征维度",
            "前馈层维度",
            "Dropout",
        ),
        ("Actor", "状态序列", "5", "1", "4", "32", "64", "0.0"),
        ("Critic", "状态—动作序列", "5", "1", "4", "32", "64", "0.1"),
        (
            "判别器",
            "专家或生成状态—动作序列",
            "5",
            "1",
            "4",
            "100",
            "200",
            "0.1",
        ),
    )
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = document.tables[2].style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (0.55, 1.35, 0.70, 0.60, 0.70, 0.65, 0.75, 0.60)
    total_width_twips = int(sum(widths) * 1440)
    table_width = table._tbl.tblPr.first_child_found_in("w:tblW")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total_width_twips))
    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(int(width * 1440)))
    header_rpr = first_run_properties(document.tables[2].cell(0, 0).paragraphs[0])
    body_rpr = first_run_properties(document.tables[2].cell(1, 0).paragraphs[0])

    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            source_rpr = header_rpr if row_index == 0 else body_rpr
            if source_rpr is not None:
                run._r.insert(0, deepcopy(source_rpr))
            set_table_cell_borders(
                cell,
                top=10 if row_index == 0 else None,
                bottom=(6 if row_index == 0 else 10 if row_index == len(rows) - 1 else None),
            )

    reference._p.addnext(table._tbl)


def configure_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "font.serif": ["Times New Roman", "SimSun"],
            "axes.unicode_minus": False,
            "font.size": 9.5,
            "axes.labelsize": 10,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 7.7,
            "axes.linewidth": 1.0,
            "lines.linewidth": 1.7,
            "savefig.dpi": 300,
        }
    )


def style_axis(axis) -> None:
    axis.grid(True, linestyle="-.", linewidth=0.65, color="#a8a8a8", alpha=0.65)
    for spine in axis.spines.values():
        spine.set_color("black")
        spine.set_linewidth(1.0)


def load_learning_curve_module():
    spec = importlib.util.spec_from_file_location(
        "paper_v28_learning_curves", LEARNING_CURVE_SCRIPT
    )
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load plotting module: {LEARNING_CURVE_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def build_learning_data_from_cache(plotting):
    runs_by_group = {}
    for group_name in plotting.GROUP_DIRS:
        cache_dir = plotting.RAW_DIR / group_name.replace("+", "_plus_")
        csv_paths = sorted(cache_dir.glob("run-*_metrics.csv"))
        expected_count = plotting.EXPECTED_SELECTED_RUNS[group_name]
        if len(csv_paths) < expected_count:
            raise RuntimeError(
                f"{group_name} requires {expected_count} cached runs, "
                f"but only {len(csv_paths)} were found in {cache_dir}."
            )
        selected_paths = csv_paths[-expected_count:]
        runs_by_group[group_name] = [plotting.load_run(path) for path in selected_paths]

    all_survival = [
        run["values"]["survival"]
        for runs in runs_by_group.values()
        for run in runs
    ]
    common_steps = sorted(
        set.intersection(*(set(values.keys()) for values in all_survival))
    )
    for runs in runs_by_group.values():
        for run in runs:
            plotting.add_cumulative_normalized_aulc(run, common_steps)
    data = {
        group_name: {
            metric_id: plotting.aggregate_sem(runs, metric_id)
            for metric_id in plotting.METRIC_ORDER
        }
        for group_name, runs in runs_by_group.items()
    }
    return data, runs_by_group


def create_english_learning_curve_figures() -> None:
    global LATEST_EXPERIMENT_STATS

    plotting = load_learning_curve_module()
    plotting.METRIC_TITLES.update(
        {
            "survival": "(a) Survival Days",
            "production": "(b) Production Agent Return",
            "naulc": "(c) Cumulative nAULC",
        }
    )
    plotting.Y_LABELS["production"] = "Return"
    plotting.CHINESE_TITLE_FONT = plotting.FontProperties(
        family="Times New Roman", size=10.5
    )
    plotting.configure_matplotlib()
    if all(path.exists() for path in plotting.GROUP_DIRS.values()):
        data, runs_by_group = plotting.build_data()
    else:
        data, runs_by_group = build_learning_data_from_cache(plotting)

    def threshold_step(group_name: str, target: float) -> int | None:
        aggregate = data[group_name]["survival"]
        for index in range(len(aggregate["steps"]) - 2):
            if all(
                value >= target
                for value in aggregate["means"][index : index + 3]
            ):
                return int(aggregate["steps"][index])
        return None

    statistics_by_group = {}
    for group_name, runs in runs_by_group.items():
        group_statistics: dict[str, float | int | None] = {
            "n_runs": len(runs),
            "step_80": threshold_step(group_name, 80.0),
            "step_90": threshold_step(group_name, 90.0),
        }
        for metric_id in ("survival", "production"):
            common_metric_steps = sorted(
                set.intersection(
                    *(set(run["values"][metric_id]) for run in runs)
                )
            )
            stable_steps = common_metric_steps[-20:]
            run_values = [
                float(
                    np.mean(
                        [run["values"][metric_id][step] for step in stable_steps]
                    )
                )
                for run in runs
            ]
            mean_value = float(np.mean(run_values))
            sem_value = float(
                np.std(run_values, ddof=1) / np.sqrt(len(run_values))
            )
            group_statistics[f"{metric_id}_mean"] = mean_value
            group_statistics[f"{metric_id}_sem"] = sem_value
            group_statistics[f"{metric_id}_relative_sem"] = (
                sem_value / abs(mean_value) * 100.0 if mean_value else 0.0
            )

        final_step = min(max(run["values"]["naulc"]) for run in runs)
        naulc_values = [
            run["values"]["naulc"][final_step] for run in runs
        ]
        naulc_mean = float(np.mean(naulc_values))
        naulc_sem = float(
            np.std(naulc_values, ddof=1) / np.sqrt(len(naulc_values))
        )
        group_statistics.update(
            {
                "naulc_step": int(final_step),
                "naulc_mean": naulc_mean,
                "naulc_sem": naulc_sem,
                "naulc_relative_sem": naulc_sem / abs(naulc_mean) * 100.0,
            }
        )
        statistics_by_group[group_name] = group_statistics
    LATEST_EXPERIMENT_STATS = statistics_by_group

    comparisons = (
        (("TD3", "GAIL+TD3"), TD3_GAIL_FIGURE),
        (("GAIL+TD3", "Transformer+GAIL+TD3"), GAIL_TRANSFORMER_FIGURE),
    )
    metric_order = ("survival", "production", "naulc")
    for groups, output_path in comparisons:
        figure, axes = plt.subplots(1, 3, figsize=(12.75, 4.0))
        for axis, metric_id in zip(axes, metric_order):
            plotting.plot_metric(axis, metric_id, groups, data)
        figure.subplots_adjust(
            left=0.055,
            right=0.992,
            bottom=0.19,
            top=0.90,
            wspace=0.34,
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        figure.savefig(output_path, facecolor="white", dpi=300)
        plt.close(figure)


def create_english_dscr_risk_figure() -> None:
    from matplotlib.ticker import MultipleLocator

    rows = pd.read_csv(DSCR_RISK_CSV, encoding="utf-8-sig")
    values_by_algorithm = {
        str(row["algorithm"]): float(row["dscr_below_1_share_percent_all"])
        for _, row in rows.iterrows()
    }
    order = ("TD3", "GAIL+TD3", "Transformer+GAIL+TD3")
    missing = [name for name in order if name not in values_by_algorithm]
    if missing:
        raise RuntimeError(f"Missing DSCR risk rows: {missing}")

    colors = ("#4C86E8", "#F08A5D", "#63B995")
    values = [values_by_algorithm[name] for name in order]
    figure, axis = plt.subplots(figsize=(6.6, 5.6))
    bars = axis.bar(
        range(len(order)),
        values,
        width=0.56,
        color=colors,
        edgecolor="#555555",
        linewidth=0.8,
        zorder=3,
    )
    upper = max(1.0, float(np.ceil(max(values) * 1.22)))
    axis.set_ylim(0, upper)
    axis.yaxis.set_major_locator(MultipleLocator(1))
    axis.set_xticks(range(len(order)))
    axis.set_xticklabels(order)
    axis.set_ylabel("Share (%)", labelpad=10)
    axis.grid(
        axis="y",
        linestyle="-.",
        linewidth=0.65,
        color="#A8A8A8",
        alpha=0.65,
        zorder=0,
    )
    for spine in axis.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.0)
    axis.tick_params(direction="out", width=1.0, length=4, colors="black")
    label_offset = upper * 0.025
    for bar, value in zip(bars, values):
        axis.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + label_offset,
            f"{value:.2f}%",
            ha="center",
            va="bottom",
            fontsize=9,
            color="#222222",
        )
    figure.subplots_adjust(left=0.13, right=0.98, bottom=0.14, top=0.98)
    DSCR_RISK_FIGURE.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(DSCR_RISK_FIGURE, facecolor="white", dpi=300)
    plt.close(figure)


def replace_picture_before_caption(
    document: Document,
    caption_prefix: str,
    image_path: Path,
    *,
    width_inches: float | None = None,
) -> None:
    caption = find_paragraph(document, caption_prefix)
    element = caption._p.getprevious()
    while element is not None:
        blips = element.xpath(".//a:blip")
        if blips:
            relationship_id = blips[0].get(qn("r:embed"))
            image_part = document.part.related_parts[relationship_id]
            image_part._blob = image_path.read_bytes()
            if width_inches is not None:
                pixel_width, pixel_height = Image.open(image_path).size
                target_width = Inches(width_inches)
                target_height = int(target_width * pixel_height / pixel_width)
                for inline_shape in document.inline_shapes:
                    shape_blips = inline_shape._inline.xpath(".//a:blip")
                    if (
                        shape_blips
                        and shape_blips[0].get(qn("r:embed")) == relationship_id
                    ):
                        inline_shape.width = target_width
                        inline_shape.height = target_height
                        break
            return
        element = element.getprevious()
    raise RuntimeError(f"No picture found before caption: {caption_prefix}")


def create_paper_figures() -> None:
    FIGURE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    configure_matplotlib()

    noisy_samples = pd.read_csv(
        NOISY_RESULT_DIR / "seed_184" / "evaluation_samples_and_scores.csv"
    )
    noisy_summary = pd.read_csv(
        NOISY_RESULT_DIR / "three_seed_expert_noise_summary.csv"
    )
    noisy_row = noisy_summary.loc[noisy_summary["seed"].astype(int) == 184].iloc[0]
    figure, axis = plt.subplots(figsize=(6.4, 4.05))
    axis.hist(
        noisy_samples["expert_score"],
        bins=50,
        range=(0.0, 1.0),
        alpha=0.52,
        color="#2ca02c",
        label="Expert",
    )
    axis.hist(
        noisy_samples["noisy_expert_score"],
        bins=50,
        range=(0.0, 1.0),
        alpha=0.42,
        color="#9467bd",
        label="Noisy expert",
    )
    axis.set_xlabel("Discriminator confidence score")
    axis.set_ylabel("Sample count")
    axis.set_xlim(0.0, 1.0)
    style_axis(axis)
    axis.legend(loc="upper left", frameon=True, framealpha=0.95)
    figure.tight_layout()
    figure.savefig(DISCRIMINATOR_FIGURE, bbox_inches="tight")
    plt.close(figure)

    actor_summary = pd.read_csv(ACTOR_RESULT_DIR / "actor_expert_js_mean_sem.csv")
    x = actor_summary["evaluation_step"].to_numpy(dtype=float)
    mean = actor_summary["mean"].to_numpy(dtype=float)
    lower = actor_summary["mean_minus_sem"].to_numpy(dtype=float)
    upper = actor_summary["mean_plus_sem"].to_numpy(dtype=float)
    figure, axis = plt.subplots(figsize=(6.4, 4.05))
    axis.fill_between(x, lower, upper, color="#4c78a8", alpha=0.22, linewidth=0)
    axis.plot(x, mean, color="#1f4e79", linewidth=1.8)
    axis.set_xlabel("Evaluation step (100 episodes)")
    axis.set_ylabel("Actor-expert JS divergence")
    axis.set_xlim(1, 60)
    axis.set_ylim(0.0, max(0.7, float(upper.max()) * 1.03))
    style_axis(axis)
    figure.tight_layout()
    figure.savefig(ACTOR_FIGURE, bbox_inches="tight")
    plt.close(figure)


def renumber_existing_section(document: Document) -> None:
    replacements = (
        ("基于上一节的结果，本文进一步比较", "在完成上述生成对抗模仿学习有效性检验后，本文进一步比较"),
        ("图6-2给出了两种模型", "图6-4给出了两种模型"),
        ("图6-2(a)中的水平虚线", "图6-4(a)中的水平虚线"),
        ("图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比", "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比"),
        ("从图6-2(a)可以看出", "从图6-4(a)可以看出"),
        ("图6-2(b)显示", "图6-4(b)显示"),
        ("图6-2(c)显示", "图6-4(c)显示"),
        ("图6-3给出了三种方法", "图6-5给出了三种方法"),
        ("图6-3 三种主体训练方案下", "图6-5 三种主体训练方案下"),
        ("图6-3显示", "图6-5显示"),
        ("综合图6-2、图6-3和表6-2", "综合图6-4、图6-5和表6-2"),
    )
    for paragraph in document.paragraphs:
        revised = paragraph.text
        for old, new in replacements:
            revised = revised.replace(old, new)
        if revised != paragraph.text:
            replace_paragraph_text(paragraph, revised)


def insert_metric_definition(document: Document) -> None:
    body_template = find_paragraph(document, "（5）生产企业偿债能力")
    dscr_explanation = find_paragraph(document, "式中，分子表示生产企业第t天")
    equation_template = Paragraph(
        dscr_explanation._p.getprevious(), dscr_explanation._parent
    )
    anchor = dscr_explanation

    paragraphs_and_equations = (
        (
            "text",
            "（6）生成对抗模仿学习有效性。本文采用Jensen–Shannon散度、重叠系数和受试者工作特征曲线下面积评价生成对抗模仿学习的行为分布接近程度及判别器区分能力。其中，Jensen–Shannon散度（JS divergence）用于度量两组判别器得分概率分布之间的差异。设P和Q表示两组归一化判别器得分分布，其定义如下：",
        ),
        ("equation", JS_MATHML),
        (
            "text",
            "式中，M表示P和Q的等权混合分布，KL表示Kullback–Leibler散度。JS散度越小表示两组判别器得分分布越接近。重叠系数（OVL: Overlap Coefficient）用于度量两组归一化直方图的共同区域，其定义如下：",
        ),
        ("equation", OVERLAP_MATHML),
        (
            "text",
            "式中，h表示得分区间编号，H表示区间总数，P_h和Q_h分别表示两组样本落入第h个区间的概率。OVL取值范围为[0,1]，数值越大表示两组得分分布的重叠程度越高。受试者工作特征曲线下面积（AUC: Area Under the Receiver Operating Characteristic Curve）用于评价判别器对两类样本的排序能力，其定义如下：",
        ),
        ("equation", AUC_MATHML),
        (
            "text",
            "式中，Pr表示事件发生的概率，x_i^+和x_j^-分别表示从正类与负类样本中抽取的判别器得分。AUC=0.5对应随机排序，AUC大于0.5表示判别器更倾向于给予正类样本较高评分。",
        ),
    )
    for kind, content in paragraphs_and_equations:
        if kind == "text":
            anchor = insert_paragraph_after(
                anchor, content, format_source=body_template
            )
        else:
            anchor = insert_equation_after(
                anchor, content, equation_template=equation_template
            )


def insert_component_validation_section(document: Document) -> None:
    old_section_heading = find_paragraph(
        document, "6.3 GAIL+TD3与Transformer+GAIL+TD3的结果对比"
    )
    body_template = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    caption_template = find_paragraph(document, "图6-1 TD3与GAIL+TD3训练结果对比")
    table_caption_template = find_paragraph(document, "表6-2 三种主体训练方案的最终性能统计")
    dscr_explanation = find_paragraph(document, "式中，分子表示生产企业第t天")
    equation_template = Paragraph(
        dscr_explanation._p.getprevious(), dscr_explanation._parent
    )

    comparison_heading = find_paragraph(document, "6.2 TD3与GAIL+TD3的结果对比")
    first_comparison_paragraph = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    first_subheading = insert_paragraph_before(
        first_comparison_paragraph,
        "（1）生成对抗模仿学习的主体决策运行效果",
        format_source=body_template,
    )
    first_subheading.runs[0].bold = True

    anchor = Paragraph(old_section_heading._p.getprevious(), old_section_heading._parent)
    heading = insert_paragraph_after(
        anchor,
        "（2）生成对抗模仿学习有效性",
        format_source=body_template,
    )
    heading.runs[0].bold = True
    anchor = insert_paragraph_after(
        heading,
        "前述运行结果表明，引入生成对抗模仿学习后，生产企业主体能够以较少的训练回合达到较高的平均存活水平。为进一步判断这一变化是否对应着对专家行为的有效模仿，本文从行为分布层面对模仿效果进行独立评价。专家数据来源于训练前采集的历史仿真轨迹，而在线训练过程中各评估时刻的Actor依据当前仿真状态生成动作，两类数据之间不存在逐时刻配对的决策记录，因而难以直接计算相同决策条件下的动作差异。本文据此采用参数固定的判别器将不同来源的状态—动作样本映射至统一得分空间，并以得分分布差异衡量Actor生成行为与专家行为的接近程度。由于该评价以判别器输出为依据，本文首先检验判别器对专家行为与非专家行为的区分能力，再以具备相应区分能力的判别器作为固定评价尺度，分析历史Actor在训练过程中的模仿效果。",
        format_source=body_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "第一步检验判别器对专家行为与非专家行为的区分能力。本文选取一组独立运行在训练结束时保存的判别器，并在评价过程中固定其网络参数，后文将该判别器称为固定评价判别器。本文在留出测试集的16952条专家状态—动作样本上叠加标准差为0.30的独立高斯噪声，并将动作裁剪至[-0.5,0.5]，由此构造加噪专家样本，作为非专家行为的受控参照。留出测试集包含173个完整专家回合。加噪专家动作定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, NOISY_ACTION_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，i表示留出测试集中的样本序号；hat{a}_i^E表示第i个专家动作；tilde{a}_i^E表示加噪并裁剪后的第i个专家动作；ε_i表示第i个样本对应的零均值独立高斯噪声；σ_noise表示原始动作单位下的噪声标准差；b表示动作边界，本文取b=0.5。固定评价判别器对原始专家样本和加噪专家样本的评分分别定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, DISCRIMINATOR_SCORE_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，D表示判别器；hat{s}_i表示第i个专家状态；带星号的φ_r表示第r次独立运行在训练结束时保存并固定用于评价的判别器参数；d_i^E和d_i^N分别表示第i个专家样本与加噪专家样本的判别器得分。图6-2给出了两组样本的判别器得分分布。",
        format_source=body_template,
    )
    anchor = insert_picture_after(
        anchor,
        DISCRIMINATOR_FIGURE,
        body_template=body_template,
        width_inches=6.35,
    )
    figure_caption = insert_paragraph_after(
        anchor,
        "图6-2 专家样本与加噪专家样本的判别器得分分布",
        format_source=caption_template,
    )
    analysis = insert_paragraph_after(
        figure_caption,
        "图6-2与表6-3显示，固定评价判别器对专家样本与加噪专家样本的平均评分分别为0.8988和0.6111，平均评分差为0.2877。两组得分分布的JS散度为0.4100；由于JS散度为0时两组分布完全一致，因此该结果反映出专家样本与加噪专家样本的判别器得分分布存在差异。OVL为18.85%，表示两组归一化得分直方图的共同区域占18.85%，两组分布的重叠程度较低。以专家样本为正类的AUC为0.8186，表示随机抽取一个专家样本和一个加噪专家样本时，判别器给予专家样本更高评分的概率约为81.86%，高于随机排序对应的0.5。上述结果表明，固定评价判别器能够区分原始专家样本与作为非专家行为参照的加噪专家样本，并倾向于给予原始专家样本更高评分。",
        format_source=body_template,
    )
    table_caption = insert_paragraph_after(
        figure_caption,
        "表6-3 判别器对专家样本与加噪专家样本的评价结果",
        format_source=table_caption_template,
    )
    insert_results_table_after(
        document,
        table_caption,
        (
            ("统计量", "估计值"),
            ("专家样本平均评分", "0.8988"),
            ("加噪专家样本平均评分", "0.6111"),
            ("平均评分差", "0.2877"),
            ("JS散度", "0.4100"),
            ("OVL", "18.85%"),
            ("AUC", "0.8186"),
        ),
    )
    anchor = analysis
    anchor = insert_paragraph_after(
        anchor,
        "第二步评价历史Actor生成行为在训练过程中的模仿效果。本文选取3次独立运行，分别以各次运行在训练结束时保存的判别器作为固定评价尺度，并加载每100个训练回合保存的历史Actor快照。所有历史Actor均在同一批留出专家状态上以确定性方式生成动作，不加入探索噪声。对于第r次独立运行的第k个评估时刻，历史Actor生成的动作及其判别器得分定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, ACTOR_SCORE_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，μ表示生产企业Actor；θ_(r,k)表示第r次独立运行在第k个评估时刻保存的Actor参数；上标G表示Actor生成动作或相应判别器得分。根据专家得分分布P_r^E和历史Actor得分分布P_(r,k)^G，Actor—专家JS散度定义如下：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, ACTOR_JS_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，J_(r,k)表示第r次独立运行在第k个评估时刻的Actor—专家JS散度。较小的J_(r,k)表示历史Actor生成行为在固定判别器得分表示下更接近专家样本。本文在每个评估时刻计算3次独立运行的JS散度均值及其均值标准误：",
        format_source=body_template,
    )
    anchor = insert_equation_after(
        anchor, ACTOR_MEAN_SEM_MATHML, equation_template=equation_template
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，带横线的J表示第k个评估时刻3次独立运行的JS散度均值；s_k表示相应JS散度的样本标准差；R表示独立运行次数，本文取R=3。图6-3给出了60个评估时刻的Actor—专家JS散度均值，阴影表示均值上下1个SEM。",
        format_source=body_template,
    )
    anchor = insert_picture_after(
        anchor,
        ACTOR_FIGURE,
        body_template=body_template,
        width_inches=6.35,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-3 Actor与专家样本判别器得分分布的JS散度变化",
        format_source=caption_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图6-3显示，Actor—专家JS散度均值在第1个评估时刻为0.580±0.041，训练过程中总体下降，并在第19个评估时刻达到0.172±0.033。第60个评估时刻的均值为0.258±0.017，较第1个评估时刻降低约55.5%，且训练后期整体维持在低于训练初期的水平。固定评价判别器下的得分分布结果表明，历史Actor生成行为在训练过程中总体向专家样本接近。",
        format_source=body_template,
    )
    insert_paragraph_after(
        anchor,
        "两步评价分别从判别器对专家与非专家行为的区分能力，以及历史Actor生成行为与专家行为的分布接近程度两个方面分析了生成对抗模仿学习的实际效果。固定评价判别器能够区分原始专家样本与加噪专家样本，历史Actor相对于训练初期则形成了更接近专家得分分布的生成行为。上述结果从行为分布层面支持了生产企业主体在训练过程中实现有效模仿的结论，并为前述运行效果提供了相应证据。",
        format_source=body_template,
    )


def update_transitions_and_metric_count(document: Document) -> None:
    replace_paragraph_text(
        find_paragraph(document, "在上述训练过程中，本文跟踪"),
        "在上述训练过程中，本文跟踪平均存活天数、生产企业累计收益、累计归一化学习曲线下面积、学习速度、生产企业偿债能力和生成对抗模仿学习有效性六项评价指标。各项指标的定义如下。",
    )
    replace_paragraph_text(
        find_paragraph(document, "不过，GAIL+TD3的最终稳定平均存活天数为80.71天"),
        "不过，GAIL+TD3的最终稳定平均存活天数为80.71天，与专家轨迹平均值仍相差17.28天。该结果说明，模仿学习信号缩短了模型达到80天所需的训练过程，但稳定阶段的存活水平仍有进一步接近专家轨迹的空间。为进一步判断上述运行效果是否对应着对专家行为的有效模仿，下面从判别器区分能力和历史Actor行为分布两个方面进行分析。",
    )


def update_chapter_two(document: Document) -> None:
    replace_paragraph_text(
        find_paragraph(document, "2  仿真环境、目标与主体决策建模问题"),
        "2  仿真环境、目标与主体决策问题",
    )
    replace_paragraph_text(
        find_paragraph(document, "2.2  主体决策建模"),
        "2.2  主体决策问题",
    )
    replace_paragraph_text(
        find_paragraph(document, "图2-1  初步设计的仿真运行规则示意图"),
        "图2-1  单次仿真运行规则示意图",
    )
    replace_paragraph_text(
        find_paragraph(document, "如图2-1所示，在仿真过程中"),
        "如图2-1所示，一次仿真由连续推进的若干天组成。仿真开始后，系统进入当天的P1-P6六阶段流程；完成P6结算后，系统判断是否满足终止条件。若未满足终止条件，系统进入下一天并重复六阶段流程；若满足终止条件，当前仿真结束。终止条件包括以下三种情况：某个企业破产，市场上已经无法购买到支撑后续生产的某种商品，或者仿真达到预设的天数上限T。",
    )
    replace_paragraph_text(
        find_paragraph(document, "在每个回合中，系统按照含有 6 个阶段"),
        "在仿真实验中，系统按照含有6个阶段（P1-P6）的日循环运行。本文将一次P1-P6循环称为“一天”。一天中的6个阶段依次为P1企业决策阶段、P2银行决策阶段、P3贷款发放阶段、P4商品交易阶段、P5生产阶段和P6结算阶段。由于本文重点是主体决策模型的训练方案，具体生产公式、交易规则和银行约束沿用既有研究[1]，本文不展开重复推导。",
    )
    replace_paragraph_text(
        find_paragraph(document, "决策模型是智能主体建立“从观察到行动”映射的计算结构"),
        "在上述仿真环境中，企业主体和银行主体需要依据各自在当前决策阶段能够观测到的信息，从连续动作空间中确定经营决策。企业决策直接作用于贷款申请、商品采购和产品定价，银行决策直接作用于贷款配置；各主体动作经过贷款、交易、生产和结算环节共同改变系统状态。因此，主体决策问题可以表述为：在多主体持续交互和局部信息约束下，根据当前可观测状态确定连续动作，并使动作与后续环境演化相衔接。",
    )
    replace_paragraph_text(
        find_paragraph(document, "在进行上述决策时，主体能够利用的信息主要来自自身经营状态"),
        "上述主体决策问题具有连续动作、多主体耦合和时序依赖等特征。各主体在同一天作出的决策通过信贷和商品交易相互作用，并共同影响后续现金、库存、债务和市场价格状态；同一主体的当前可观测状态又由此前多个阶段的交互结果累积形成。决策机制需要在既定状态空间和动作空间下处理当前信息，并对持续交互过程中的历史关联进行表达。第3章和第4章分别给出单步状态输入下的生产企业训练方案及其历史状态表征扩展。",
    )
    for table in document.tables:
        if not table.rows or len(table.rows[0].cells) != 3:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header == ["主体", "Actor 输入", "Actor 输出"]:
            replace_paragraph_text(table.cell(0, 1).paragraphs[0], "状态输入")
            replace_paragraph_text(table.cell(0, 2).paragraphs[0], "动作输出")
            break
    else:
        raise RuntimeError("Unable to locate Table 2-1 state-action headers.")


def restructure_chapter_three(document: Document) -> None:
    training_heading = find_paragraph(document, "3.1 模型训练方案")
    sample_heading = find_paragraph(document, "3.2 状态、动作与经验样本")
    discriminator_heading = find_paragraph(
        document, "3.3 判别器目标与模仿奖励构造"
    )

    first_sample_content_element = sample_heading._p.getnext()
    if first_sample_content_element is None:
        raise RuntimeError("The state-action sample section has no body content.")
    first_sample_content = Paragraph(
        first_sample_content_element, sample_heading._parent
    )
    first_equation_element = first_sample_content_element.getnext()
    if first_equation_element is None:
        raise RuntimeError("The experience-batch equation is missing.")
    equation_template = Paragraph(first_equation_element, sample_heading._parent)

    copied_sample_elements = []
    element = first_equation_element
    while element is not None and element is not discriminator_heading._p:
        copied_sample_elements.append(deepcopy(element))
        element = element.getnext()
    if element is None:
        raise RuntimeError("Unable to locate the end of the sample section.")

    new_heading = insert_paragraph_before(
        training_heading,
        "3.1 主体决策模型",
        format_source=training_heading,
    )
    anchor = insert_paragraph_after(
        new_heading,
        "在生成对抗模仿学习与TD3的联合训练框架中，生产企业主体需要通过可训练的策略模型形成连续经营决策，并为后续的判别器训练与价值估计提供状态—动作样本。为此，本文首先对生产企业主体的Actor神经网络、状态输入、动作输出和经验样本进行统一定义。Actor采用多层感知机（MLP：Multilayer Perceptron）结构，将生产企业在当前决策时刻的可观测状态映射为连续经营动作。网络输入为经标准化处理的33维状态向量，网络输出为4维动作向量，分别对应贷款意愿（WNDF）、生产资料购买意愿（K）、消费品购买意愿（L）和次日定价（P）。生产企业Actor的网络结构如图3-1所示。",
        format_source=first_sample_content,
    )
    training_caption = find_paragraph(
        document, "图3-1  基于生成对抗模仿学习与TD3的主体训练方案"
    )
    actor_picture = insert_picture_after(
        anchor,
        ACTOR_STRUCTURE_FIGURE,
        body_template=first_sample_content,
        width_inches=6.0,
    )
    anchor = insert_paragraph_after(
        actor_picture,
        "图3-1  生产企业Actor神经网络结构",
        format_source=training_caption,
    )
    anchor = insert_paragraph_after(
        anchor,
        "Actor由三层可训练的全连接映射组成。第一层将33维状态向量映射为128维隐藏特征，并采用双曲正切函数进行非线性变换；第二层将128维隐藏特征映射为32维隐藏特征，并采用带泄漏线性整流函数（Leaky ReLU：Leaky Rectified Linear Unit）进行非线性变换；第三层将32维隐藏特征映射为4维动作向量。Actor在第t个决策时刻的前向计算过程定义如下：",
        format_source=first_sample_content,
    )
    anchor = insert_equation_after(
        anchor,
        ACTOR_NETWORK_MATHML,
        equation_template=equation_template,
    )
    anchor = insert_paragraph_after(
        anchor,
        "式中，hat{s}_t表示第t个决策时刻经标准化处理的生产企业状态；h_t^(1)和h_t^(2)分别表示第一隐藏层和第二隐藏层输出的特征向量；W_1、W_2和W_3表示三层全连接映射的权重矩阵；b_1、b_2和b_3表示对应的偏置向量；a_t表示Actor生成的动作；a_max表示动作边界，本文取a_max=0.5。输出层采用双曲正切函数，并按动作边界进行缩放，使Actor直接生成的各维动作均位于[-0.5,0.5]区间内。",
        format_source=first_sample_content,
    )
    anchor = insert_paragraph_after(
        anchor,
        "图3-1所示Actor以当前状态向量作为唯一输入。第一隐藏层形成128维中间特征，第二隐藏层进一步形成32维特征，其输出直接进入动作输出层并生成连续经营动作。该网络构成GAIL+TD3模型中的单步Actor。第4.3节所述扩展模型将在该结构基础上融合Transformer历史特征。",
        format_source=first_sample_content,
    )
    anchor = insert_paragraph_after(
        anchor,
        "除Actor外，生产企业主体的训练还涉及Critic与判别器。Critic用于估计生产企业状态—动作对的长期价值，并为Actor参数更新提供价值依据，其网络结构如图3-2所示。",
        format_source=first_sample_content,
    )
    critic_picture = insert_picture_after(
        anchor,
        CRITIC_STRUCTURE_FIGURE,
        body_template=first_sample_content,
        width_inches=6.0,
    )
    anchor.paragraph_format.keep_with_next = True
    critic_picture.paragraph_format.keep_with_next = True
    critic_picture.paragraph_format.first_line_indent = Inches(0)
    critic_picture.paragraph_format.left_indent = Inches(0)
    critic_picture.paragraph_format.right_indent = Inches(0)
    anchor = insert_paragraph_after(
        critic_picture,
        "图3-2  生产企业Critic双Q网络结构",
        format_source=training_caption,
    )
    anchor = insert_paragraph_after(
        anchor,
        "在Critic参数更新阶段，网络输入来自经验池采样批量中的生产企业状态—动作对。经标准化处理的33维状态向量与对应的4维动作向量首先拼接为37维输入，随后同时进入两个参数相互独立的Q值估计分支。每个分支均采用37-128-32-1的全连接结构，两个隐藏层均使用带泄漏线性整流函数，输出层分别得到Q_1与Q_2的标量估计。两个分支采用相同的输入维度和网络层级，但不共享参数。",
        format_source=first_sample_content,
    )
    anchor = insert_paragraph_after(
        anchor,
        "两个Q值分支在Critic更新时共同参与当前状态—动作价值估计；在Actor参数更新时，第一Q值分支进一步评价Actor依据同一批状态生成的策略动作，并将相应价值梯度传递至Actor。目标Critic采用与图3-2相同的双分支结构，并通过独立参数计算下一状态—动作对的目标Q值。Critic目标值与参数更新过程见第3.4节。",
        format_source=first_sample_content,
    )
    anchor = insert_paragraph_after(
        anchor,
        "判别器用于区分专家状态—动作样本与生产企业生成样本，并根据分类结果为生成样本构造模仿奖励，其网络结构如图3-3所示。",
        format_source=first_sample_content,
    )
    discriminator_picture = insert_picture_after(
        anchor,
        DISCRIMINATOR_STRUCTURE_FIGURE,
        body_template=first_sample_content,
        width_inches=6.0,
    )
    anchor.paragraph_format.keep_with_next = True
    discriminator_picture.paragraph_format.keep_with_next = True
    discriminator_picture.paragraph_format.first_line_indent = Inches(0)
    discriminator_picture.paragraph_format.left_indent = Inches(0)
    discriminator_picture.paragraph_format.right_indent = Inches(0)
    anchor = insert_paragraph_after(
        discriminator_picture,
        "图3-3  生产企业GAIL判别器网络结构",
        format_source=training_caption,
    )
    anchor = insert_paragraph_after(
        anchor,
        "专家样本与生成样本均由33维状态向量和4维动作向量组成，并通过同一个判别器完成评价。状态与动作依据专家数据统计量完成标准化后拼接为37维输入，随后依次经过37-100、100-100和100-1三层全连接映射。前两个隐藏层均采用双曲正切函数，两个100维隐藏层逐层形成状态—动作样本的中间特征表示；最后一层输出单个logits值，Sigmoid函数再将其转换为[0,1]区间内的专家样本判别概率。判别器的参数在两类样本之间共享，其训练目标与模仿奖励构造见第3.3节。",
        format_source=first_sample_content,
    )
    anchor = insert_paragraph_after(
        anchor,
        "为保持专家行为与在线生成行为的数据接口一致，每条专家记录由33维状态向量和4维动作向量组成，其中前33个字段对应专家状态，后4个字段对应专家动作。在线交互过程中，生产企业Actor根据当前状态生成动作，仿真环境执行该动作后返回环境奖励、下一状态和终止标记，并由此形成经验样本。令一次网络参数更新阶段从经验池采样得到的批量为：",
        format_source=first_sample_content,
    )
    for copied_element in copied_sample_elements:
        anchor._p.addnext(copied_element)
        anchor = Paragraph(copied_element, anchor._parent)

    replace_paragraph_text(training_heading, "3.2 模型训练方案")

    training_intro = find_paragraph(
        document, "本文将生成对抗模仿学习引入生产企业主体的 TD3 训练过程"
    )
    replace_paragraph_text(
        training_intro,
        "模型训练通过连续执行多次仿真推进。每次仿真均从环境初始化开始，经过若干天连续运行，直至触发第2.1节所述终止条件而结束。本文将这一完整过程定义为一个回合（episode）。由多个回合构成的整体训练流程如图3-4所示。",
    )
    episode_picture = insert_picture_after(
        training_intro,
        TRAINING_EPISODE_FIGURE,
        body_template=first_sample_content,
        width_inches=6.0,
    )
    episode_caption = insert_paragraph_after(
        episode_picture,
        "图3-4  模型训练中的多回合仿真流程",
        format_source=training_caption,
    )
    cross_episode_description = insert_paragraph_after(
        episode_caption,
        "当一个回合结束时，系统重置仿真环境和各主体属性，并以新的初始状态开始下一回合。生产企业主体的Actor、Critic和判别器参数以及经验池不随环境重置而清空，而是在跨回合交互中持续更新。交互过程中产生的状态、动作、奖励、下一状态和终止标记被写入经验池，从而将单个回合内的环境演化与跨回合的模型训练衔接起来。",
        format_source=first_sample_content,
    )
    insert_paragraph_after(
        cross_episode_description,
        "在上述多回合交互框架下，本文基于第3.1节定义的Actor、专家样本和在线经验样本，将生成对抗模仿学习引入生产企业主体的TD3训练过程。生产企业Actor同时承担TD3策略网络和GAIL生成器的作用；经验池中的生产企业交互样本为判别器训练、模仿奖励构造和Critic价值估计提供统一的数据接口。",
        format_source=first_sample_content,
    )

    element = sample_heading._p
    while element is not discriminator_heading._p:
        next_element = element.getnext()
        element.getparent().remove(element)
        element = next_element

    batch_alignment = find_paragraph(
        document, "本文实现中的一个重要细节是批量对齐"
    )
    replace_paragraph_text(
        batch_alignment,
        "本文采用批量对齐机制。每一次网络参数更新阶段只从经验池采样一次生产企业经验批量。该批量的状态-动作部分用于构造判别器中的生成样本，该批量的完整经验样本又用于TD3的Critic目标值计算。因此，判别器给出的模仿奖励与Critic使用的环境奖励和下一状态在样本索引上保持一致。",
    )
    training_description = find_paragraph(document, "图3-1给出了")
    replace_paragraph_text(
        training_description,
        "图3-5进一步展示了生成对抗模仿学习与TD3的联合训练流程。专家行为样本为判别器提供专家分布参照，生产企业在线交互形成的状态-动作对构成生成样本。判别器根据两类样本之间的差异产生模仿奖励；该奖励与环境奖励融合后进入Critic目标值计算，并通过Critic对策略动作的价值估计影响Actor更新。模仿学习信号在训练过程中的传递路径可以概括为“判别器—融合奖励—Critic—Actor”。",
    )
    replace_paragraph_text(
        training_caption,
        "图3-5  基于生成对抗模仿学习与TD3的主体训练方案",
    )


def update_chapter_four(document: Document) -> None:
    motivation_conclusion = find_paragraph(document, "后续实验将进一步检验单步状态表征")
    replace_paragraph_text(
        motivation_conclusion,
        "基于上述时序依赖特征，本文在第3章GAIL+TD3训练方案基础上引入Transformer编码器。该编码器首先将固定长度的历史状态或状态—动作序列转换为定长历史特征，随后将历史特征接入生产企业主体的Actor、Critic和判别器，使三类网络能够在保留当前输入的基础上利用局部历史交互信息。",
    )

    transformer_heading = find_paragraph(document, "4.2 Transformer 编码器结构")
    replace_paragraph_text(transformer_heading, "4.2 Transformer历史序列表征模块")

    sequence_intro = find_paragraph(
        document, "本文方法将历史窗口长度记为"
    )
    for text_element in sequence_intro._p.xpath(".//w:t"):
        if text_element.text and "本文方法将历史窗口长度记为" in text_element.text:
            text_element.text = text_element.text.replace(
                "本文方法将历史窗口长度记为",
                "本文将历史序列所包含的时间步数定义为序列长度，并记为",
            )

    sequence_description = find_paragraph(
        document, "其中， 表示截至第  天的历史状态窗口"
    )
    for text_element in sequence_description._p.xpath(".//w:t"):
        if text_element.text and "可用历史长度不足" in text_element.text:
            text_element.text = text_element.text.replace(
                "可用历史长度不足",
                "可用历史序列的时间步数不足",
            )
    body_template = sequence_description
    table_caption_template = find_paragraph(
        document, "表6-2 三种主体训练方案的最终性能统计"
    )
    parameter_description = insert_paragraph_after(
        sequence_description,
        "本文将线性序列映射、可学习位置嵌入、Transformer编码和末时刻特征读取组成的结构定义为Transformer历史序列表征模块。该模块首先把序列中每个时间步的输入投影至统一隐藏空间，并与对应的可学习位置嵌入相加；单层Transformer编码器随后利用多头自注意力和前馈子层提取序列内部的时序关联；编码器最后一个时间位置的输出经层归一化后形成定长历史特征。本文将该历史特征的维数定义为表征维度。三个分支均采用长度为5的历史序列、4个注意力头和高斯误差线性单元（GELU：Gaussian Error Linear Unit）激活函数。Actor与Critic分支的表征维度为32，前馈层维度为64；判别器分支的表征维度为100，前馈层维度为200。各分支分别维护独立的线性映射、位置嵌入、Transformer编码器和层归一化参数，彼此不共享参数。主要参数设置见表4-1。",
        format_source=body_template,
    )
    table_caption = insert_paragraph_after(
        parameter_description,
        "表4-1 Transformer表征模块参数设置",
        format_source=table_caption_template,
    )
    insert_transformer_parameter_table_after(document, table_caption)

    transformer_figure_caption = find_paragraph(
        document, "图4-1  生产企业主体三分支Transformer表征结构"
    )
    transformer_figure_element = transformer_figure_caption._p.getprevious()
    if transformer_figure_element is None or not transformer_figure_element.xpath(
        ".//a:blip"
    ):
        raise RuntimeError("Unable to locate the original Figure 4-1 picture.")
    transformer_figure_element.getparent().remove(transformer_figure_element)
    remove_paragraph(transformer_figure_caption)

    actor_description = find_paragraph(
        document, "Actor 侧的 Transformer 编码器以历史状态序列为输入"
    )
    section_heading = find_paragraph(
        document, "4.3 表征增强模块的训练接入与方法边界"
    )
    element = actor_description._p
    while element is not section_heading._p:
        next_element = element.getnext()
        element.getparent().remove(element)
        element = next_element

    remove_paragraph(find_paragraph(document, "[此处插入图4-2流程图]"))
    remove_paragraph(
        find_paragraph(
            document,
            "图4-2  基于Transformer表征增强的GAIL+TD3生产企业主体训练方案",
        )
    )
    training_connection = find_paragraph(document, "图4-2展示了")
    replace_paragraph_text(
        training_connection,
        "图4-1展示了Transformer历史序列表征模块在Actor、Critic和判别器三个分支中的实例化方式及其与原有网络输入端的连接关系。三个分支采用相同的表征流程，但根据各自任务使用不同的历史输入、表征维度和输出结构。",
    )

    replace_paragraph_text(section_heading, "4.3 历史表征与三类网络的融合")
    transformer_picture = insert_picture_after(
        training_connection,
        TRANSFORMER_THREE_BRANCH_FIGURE,
        body_template=body_template,
        width_inches=6.0,
    )
    insert_paragraph_after(
        transformer_picture,
        "图4-1  生产企业主体三分支Transformer表征结构",
        format_source=transformer_figure_caption,
    )

    online_training = find_paragraph(document, "在线训练中，专家数据和生成数据按回合位置构造")
    replace_paragraph_text(
        online_training,
        "Actor分支以长度为5的历史状态序列作为表征模块输入，并得到32维历史状态特征。该特征与33维当前状态向量拼接，形成65维联合输入；随后，一组独立的全连接动作头按照65-128-32-4的结构生成4维确定性动作。该连接方式保留当前时刻状态，同时利用历史交互信息补充动作生成所需的时序特征。",
    )

    method_boundary = find_paragraph(document, "该扩展作用于生产企业主体")
    replace_paragraph_text(
        method_boundary,
        "Critic分支把历史序列中每个时间步的33维状态与4维动作拼接为37维状态—动作向量，并通过表征模块得到32维历史状态—动作特征。该特征与当前37维状态—动作对拼接，形成69维联合输入；两个参数独立的Q值分支分别采用69-128-32-1的全连接结构输出Q_1与Q_2。历史特征由此直接参与当前状态—动作价值估计。",
    )
    insert_paragraph_after(
        method_boundary,
        "判别器分支以专家或生成轨迹的历史状态—动作序列作为输入。每个时间步的37维状态—动作向量先被映射至100维隐藏空间，表征模块输出100维历史状态—动作特征；层归一化与线性输出头随后产生单个logits值，并经Sigmoid函数转换为专家样本判别概率。除上述输入表征与网络输入维度调整外，判别器损失、模仿奖励构造、融合奖励以及Actor-Critic参数更新均采用第3章定义的训练方式。",
        format_source=body_template,
    )

    transformer_experiment = find_paragraph(
        document, "在完成上述生成对抗模仿学习有效性检验后"
    )
    replace_paragraph_text(
        transformer_experiment,
        "在完成上述生成对抗模仿学习有效性检验后，本文进一步比较GAIL+TD3与Transformer+GAIL+TD3，以分析历史状态序列表征对生产企业主体训练过程的影响。图6-4给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标上的均值曲线及SEM阴影，表6-2列出了对应的训练表现统计。图6-4(a)中的水平虚线同样表示专家轨迹的合计平均存活天数，用于比较两种模型与专家轨迹存活水平之间的差距。",
    )

    experiment_setup = find_paragraph(
        document, "本节实验从模仿学习信号与历史状态表征两个层面"
    )
    replace_paragraph_text(
        experiment_setup,
        "本节实验从模仿学习信号与历史状态表征两个层面，通过比较不同主体决策模型在复杂经济系统仿真中的训练表现，分析两类机制对模型训练过程及系统运行结果的影响。本文设置两组对比实验：第一组比较TD3与GAIL+TD3，用于考察模仿学习信号对生产企业主体训练过程的影响；第二组比较GAIL+TD3与Transformer+GAIL+TD3，用于考察历史状态序列表征对生产企业主体训练过程的影响。三种模型中，TD3作为基线模型；GAIL+TD3在生产企业主体的TD3训练过程中引入生成对抗模仿学习信号；Transformer+GAIL+TD3在GAIL+TD3基础上加入Transformer编码器。本次实验中的模仿学习与历史表征改进均应用于生产企业主体，消费企业主体和银行主体在三种模型中保持既有TD3决策流程，以控制非目标主体决策机制变化对比较结果的影响。GAIL+TD3和Transformer+GAIL+TD3使用第五章整理得到的同一组专家数据，三种模型采用相同的仿真环境、训练进度尺度和统计方法。",
    )


def update_latest_experiment_statistics(document: Document) -> None:
    if not LATEST_EXPERIMENT_STATS:
        raise RuntimeError("Learning-curve statistics were not calculated.")

    td3 = LATEST_EXPERIMENT_STATS["TD3"]
    gail = LATEST_EXPERIMENT_STATS["GAIL+TD3"]
    transformer = LATEST_EXPERIMENT_STATS["Transformer+GAIL+TD3"]

    replace_paragraph_text(
        find_paragraph(document, "神经网络参数初始化、动作探索"),
        "神经网络参数初始化、动作探索、训练样本抽取和经验池采样等过程均包含随机因素。本文将用于控制上述随机因素的联合配置统称为“随机种子”。TD3与GAIL+TD3分别在8组随机种子下独立运行，Transformer+GAIL+TD3的统计分析纳入7组独立运行。每种模型据此形成相应数量的指标序列。本文在每个评估时刻计算同一模型各次独立运行结果的算术平均值，并以该均值作为曲线值。均值标准误（SEM: Standard Error of the Mean）用于描述各模型独立运行均值的不确定性，其定义为：",
    )
    replace_paragraph_text(
        find_paragraph(document, "式中，s表示8次运行结果的样本标准差"),
        "式中，s表示同一模型各次独立运行结果的样本标准差，S表示该模型纳入统计的独立运行次数。训练曲线中的实线表示同一模型各次独立运行的均值，阴影表示均值上下1个SEM。",
    )
    replace_paragraph_text(
        find_paragraph(document, "式中，S表示独立运行次数，本文取S=8"),
        "式中，S表示对应模型纳入统计的独立运行次数；r表示独立运行编号；μ表示第k个评估时刻的平均存活天数均值；q表示目标存活天数，本文分别取80天和90天；T表示均值曲线达到目标值所需的评估时刻。较小的T表示模型以较少的训练回合达到目标存活水平。若均值曲线截至第60个评估时刻仍未满足连续3个评估时刻均不低于目标值的条件，则记为“未达到”。",
    )

    comparison_intro = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    replace_paragraph_text(
        comparison_intro,
        "本文首先比较TD3与GAIL+TD3的训练过程，以分析模仿学习信号对生产企业主体决策训练的影响。图6-1给出了两种模型在平均存活天数、生产企业累计收益和累计归一化学习曲线下面积三项指标上的均值曲线及SEM阴影。其中，图6-1(a)中的水平虚线表示表5-1所列专家轨迹的合计平均存活天数，其数值为97.99天。为定量比较三种模型的训练表现，本文对每次独立运行在最后20个评估时刻的平均存活天数和生产企业累计收益分别取均值，再在各模型纳入统计的独立运行之间计算总体均值和SEM；累计归一化学习曲线下面积采用第60个评估时刻的数值；学习速度采用平均存活天数均值曲线达到80天和90天所需的评估时刻。表6-2汇总了三种模型的相应统计结果。",
    )
    replace_paragraph_text(
        find_paragraph(document, "注：平均存活天数和生产企业累计收益"),
        "注：平均存活天数和生产企业累计收益为每次独立运行最后20个评估时刻的均值，再对各模型纳入统计的独立运行进行统计；nAULC采用第60个评估时刻的数值。括号内为相对SEM，K表示10³。学习速度依据各模型独立运行的平均存活天数均值曲线计算；“未达到”表示均值曲线截至第60个评估时刻仍未连续3个评估时刻达到相应目标值。",
    )

    for table in document.tables:
        if not table.rows:
            continue
        first_column = [row.cells[0].text.strip() for row in table.rows]
        if "随机种子组数" in first_column:
            for row in table.rows:
                item = row.cells[0].text.strip()
                if item == "随机种子组数":
                    replace_paragraph_text(
                        row.cells[1].paragraphs[0],
                        "TD3、GAIL+TD3各8组；Transformer+GAIL+TD3为7组",
                    )
                elif item == "统计方法":
                    replace_paragraph_text(
                        row.cells[1].paragraphs[0],
                        "各模型独立运行的均值及均值上下1个SEM",
                    )

        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header == ["评价指标", "TD3", "GAIL+TD3", "Transformer+GAIL+TD3"]:
            stats_by_column = (td3, gail, transformer)
            for row in table.rows[1:]:
                metric_name = row.cells[0].text.strip()
                for column_index, statistics in enumerate(stats_by_column, start=1):
                    if metric_name == "平均存活天数/天":
                        value = (
                            f"{statistics['survival_mean']:.2f}"
                            f"（±{statistics['survival_relative_sem']:.2f}%）"
                        )
                    elif metric_name == "生产企业累计收益":
                        value = (
                            f"{statistics['production_mean'] / 1000.0:.2f}K"
                            f"（±{statistics['production_relative_sem']:.2f}%）"
                        )
                    elif metric_name == "第60个评估时刻的nAULC":
                        value = (
                            f"{statistics['naulc_mean']:.3f}"
                            f"（±{statistics['naulc_relative_sem']:.2f}%）"
                        )
                    elif metric_name == "达到80天所需评估时刻":
                        value = (
                            "未达到"
                            if statistics["step_80"] is None
                            else str(statistics["step_80"])
                        )
                    elif metric_name == "达到90天所需评估时刻":
                        value = (
                            "未达到"
                            if statistics["step_90"] is None
                            else str(statistics["step_90"])
                        )
                    else:
                        continue
                    replace_paragraph_text(row.cells[column_index].paragraphs[0], value)

    expert_mean = 168_548 / 1_720
    gail_gap = abs(expert_mean - float(gail["survival_mean"]))
    transformer_gap = abs(expert_mean - float(transformer["survival_mean"]))
    gap_reduction = (gail_gap - transformer_gap) / gail_gap * 100.0
    naulc_increase = (
        (float(transformer["naulc_mean"]) - float(gail["naulc_mean"]))
        / float(gail["naulc_mean"])
        * 100.0
    )

    replace_paragraph_text(
        find_paragraph(document, "表6-2中的具体数值显示，Transformer+GAIL+TD3达到80天"),
        f"表6-2中的具体数值显示，Transformer+GAIL+TD3达到80天所需的评估时刻由GAIL+TD3的第{gail['step_80']}个缩短至第{transformer['step_80']}个，并在第{transformer['step_90']}个评估时刻达到连续3个评估时刻不低于90天的条件，而GAIL+TD3在前60个评估时刻内未达到该条件。引入Transformer编码器后，最终稳定平均存活天数由{gail['survival_mean']:.2f}天提高至{transformer['survival_mean']:.2f}天，与专家轨迹平均值的差距由{gail_gap:.2f}天缩小至{transformer_gap:.2f}天，缩小幅度约为{gap_reduction:.1f}%。第60个评估时刻的累计归一化学习曲线下面积由{gail['naulc_mean']:.3f}提高至{transformer['naulc_mean']:.3f}，增幅约为{naulc_increase:.1f}%；生产企业最终稳定累计收益由{gail['production_mean'] / 1000.0:.2f}K提高至{transformer['production_mean'] / 1000.0:.2f}K。这些统计结果表明，在本文实验设置下，历史状态表征增强了模型对连续交互信息的利用，并进一步缩小了模型运行结果与专家轨迹之间的差距。",
    )
    replace_paragraph_text(
        find_paragraph(document, "综合图6-4、图6-5和表6-2的结果"),
        f"综合图6-4、图6-5和表6-2的结果，Transformer+GAIL+TD3达到80天和90天所需的训练回合较少，最终稳定平均存活天数与专家轨迹平均值的差距由{gail_gap:.2f}天缩小至{transformer_gap:.2f}天，并将DSCR低于1的生产企业日观测占比降至0.00%。历史状态序列表征使生产企业主体能够利用跨时段交互信息，从而在学习速度、训练全过程的累积平均存活表现、稳定阶段存活水平和债务偿付状态等方面进一步利用专家行为信息。",
    )
    replace_paragraph_text(
        find_paragraph(document, "本文围绕复杂经济系统多主体仿真环境中的智能主体决策生成问题"),
        "本文围绕复杂经济系统多主体仿真环境中的智能主体决策生成问题，在既有TD3主体决策模型基础上引入生成对抗模仿学习，并采用Transformer编码器增强生产企业主体的历史状态表征。本文通过人工行为采集系统整理专家状态—动作数据，比较了TD3、GAIL+TD3和Transformer+GAIL+TD3三种主体训练方案的训练过程与系统运行结果。主要研究结论如下。",
    )
    replace_paragraph_text(
        find_paragraph(document, "首先，生成对抗模仿学习信号缩短了生产企业主体达到较高存活水平所需的训练过程"),
        "首先，生成对抗模仿学习信号提高了生产企业主体的前期学习效率。专家状态—动作样本经判别器形成的模仿奖励为策略更新提供了行为参照，使生产企业主体能够以较少的训练回合形成有效策略，并在训练全过程中保持较高的累积平均存活表现。",
    )
    replace_paragraph_text(
        find_paragraph(document, "其次，引入Transformer编码器后"),
        "其次，Transformer编码器增强了生产企业主体对历史交互信息的表征能力。历史状态序列的引入进一步改善了模型的学习速度、稳定阶段存活水平和生产企业经营表现，使系统平均存活水平更加接近专家轨迹。",
    )
    replace_paragraph_text(
        find_paragraph(document, "最后，生产企业偿债风险结果表明"),
        "最后，生产企业偿债风险分析表明，模仿学习信号与历史状态表征的引入对应着更低的当期偿债缺口观测比例。该结果说明，改进后的主体训练方案在促进策略学习的同时，有助于生产企业维持现金流与债务偿付之间的协调。",
    )


def verify(document: Document, source_equation_count: int) -> None:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    text = "\n".join(paragraphs)
    required = (
        "图2-1  单次仿真运行规则示意图",
        "一次仿真由连续推进的若干天组成",
        "在仿真实验中，系统按照含有6个阶段",
        "3.1 主体决策模型",
        "3.2 模型训练方案",
        "图3-1  生产企业Actor神经网络结构",
        "图3-2  生产企业Critic双Q网络结构",
        "图3-3  生产企业GAIL判别器网络结构",
        "图3-4  模型训练中的多回合仿真流程",
        "图3-5  基于生成对抗模仿学习与TD3的主体训练方案",
        "本文将这一完整过程定义为一个回合（episode）",
        "在跨回合交互中持续更新",
        "在上述多回合交互框架下",
        "表4-1 Transformer表征模块参数设置",
        "图4-1  生产企业主体三分支Transformer表征结构",
        "4.2 Transformer历史序列表征模块",
        "4.3 历史表征与三类网络的融合",
        "本文将线性序列映射、可学习位置嵌入、Transformer编码和末时刻特征读取组成的结构定义为Transformer历史序列表征模块",
        "本文将历史序列所包含的时间步数定义为序列长度，并记为",
        "可用历史序列的时间步数不足",
        "本次实验中的模仿学习与历史表征改进均应用于生产企业主体",
        "三个分支均采用长度为5的历史序列、4个注意力头",
        "第4.3节所述扩展模型将在该结构基础上融合Transformer历史特征",
        "Actor由三层可训练的全连接映射组成",
        "该网络构成GAIL+TD3模型中的单步Actor",
        "两个参数相互独立的Q值估计分支",
        "两个100维隐藏层逐层形成状态—动作样本的中间特征表示",
        "Actor分支以长度为5的历史状态序列作为表征模块输入",
        "Critic分支把历史序列中每个时间步的33维状态与4维动作拼接为37维状态—动作向量",
        "判别器分支以专家或生成轨迹的历史状态—动作序列作为输入",
        "（6）生成对抗模仿学习有效性",
        "（1）生成对抗模仿学习的主体决策运行效果",
        "（2）生成对抗模仿学习有效性",
        "图6-2 专家样本与加噪专家样本的判别器得分分布",
        "图6-3 Actor与专家样本判别器得分分布的JS散度变化",
        "6.3 GAIL+TD3与Transformer+GAIL+TD3的结果对比",
        "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        "图6-5 三种主体训练方案下DSCR 低于 1 的生产企业日观测占比",
        "表6-3 判别器对专家样本与加噪专家样本的评价结果",
        "第60个评估时刻的均值为0.258±0.017",
        "OVL为18.85%",
    )
    for value in required:
        if value not in text:
            raise RuntimeError(f"Required content is missing: {value}")
    forbidden = (
        "图2-1  初步设计的仿真运行规则示意图",
        "在每个回合中，系统按照含有 6 个阶段",
        "图3-2  基于生成对抗模仿学习与TD3的主体训练方案",
        "图3-3  基于生成对抗模仿学习与TD3的主体训练方案",
        "图3-4  基于生成对抗模仿学习与TD3的主体训练方案",
        "图3-2  生产企业Critic与判别器神经网络结构",
        "4.3 表征增强模块的训练接入与方法边界",
        "4.3 表征增强模块的训练接入",
        "4.3 历史表征模块与联合训练",
        "该扩展作用于生产企业主体",
        "注意力头数均为1",
        "本文方法将历史窗口长度记为",
        "可用历史长度不足",
        "本组实验仅在生产企业主体的Actor、Critic和判别器中接入Transformer历史序列表征模块",
        "2  仿真环境、目标与主体决策建模问题",
        "2.2  主体决策建模",
        "决策模型是智能主体建立“从观察到行动”映射的计算结构",
        "如果模型需要进一步利用历史信息",
        "3.1 模型训练方案",
        "3.2 状态、动作与经验样本",
        "图6-2 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        "图6-3 三种主体训练方案下DSCR 低于 1",
        "并非单调下降",
        "随机种子184",
        "随机种子652",
        "随机种子187",
        "p=0.0002",
        "95%区间[0.4012,0.4191]",
        "95%区间[0.8073,0.8293]",
        "[此处插入图4-2流程图]",
        "图4-2  基于Transformer表征增强的GAIL+TD3生产企业主体训练方案",
        "图4-2展示了",
        "每种模型8组",
        "本文取S=8",
        "生成对抗模仿学习组件有效性",
        "6.3 生成对抗模仿学习组件有效性分析",
        "6.4 GAIL+TD3与Transformer+GAIL+TD3的结果对比",
        "最终判别器",
        "对动作扰动的识别能力",
        "第6.4节",
    )
    for value in forbidden:
        if value in text:
            raise RuntimeError(f"Former numbering remains: {value}")
    if not any(
        cell.text.strip() == "序列长度（l）"
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ):
        raise RuntimeError("Table 4-1 sequence-length header was not updated.")
    if len(document.tables) != 6:
        raise RuntimeError(f"Expected 6 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 15:
        raise RuntimeError(f"Expected 15 inline figures, found {len(document.inline_shapes)}.")
    equation_count = len(document.element.body.xpath(".//m:oMath"))
    expected_equation_count = source_equation_count + len(EQUATIONS) - 30
    if equation_count != expected_equation_count:
        raise RuntimeError(
            f"Expected {expected_equation_count} equations, "
            f"found {equation_count}."
        )


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The source DOCX changed after this builder was prepared.")
    if not MML2OMML_XSL.exists():
        raise FileNotFoundError(MML2OMML_XSL)
    create_paper_figures()
    create_english_learning_curve_figures()
    create_english_dscr_risk_figure()
    if not DISCRIMINATOR_FIGURE.exists() or not ACTOR_FIGURE.exists():
        raise RuntimeError("Publication figures were not generated.")
    for figure_path in (
        DAY_SIMULATION_FIGURE,
        TRAINING_EPISODE_FIGURE,
        ACTOR_STRUCTURE_FIGURE,
        CRITIC_STRUCTURE_FIGURE,
        DISCRIMINATOR_STRUCTURE_FIGURE,
        GAIL_TD3_TRAINING_FIGURE,
        TRANSFORMER_THREE_BRANCH_FIGURE,
        TD3_GAIL_FIGURE,
        GAIL_TRANSFORMER_FIGURE,
        DSCR_RISK_FIGURE,
    ):
        if not figure_path.exists():
            raise RuntimeError(f"English publication figure was not generated: {figure_path}")

    document = Document(SOURCE)
    source_equation_count = len(document.element.body.xpath(".//m:oMath"))
    renumber_existing_section(document)
    update_transitions_and_metric_count(document)
    update_chapter_two(document)
    insert_metric_definition(document)
    insert_component_validation_section(document)
    restructure_chapter_three(document)
    update_chapter_four(document)
    update_latest_experiment_statistics(document)
    replace_picture_before_caption(
        document,
        "图2-1  单次仿真运行规则示意图",
        DAY_SIMULATION_FIGURE,
        width_inches=5.8,
    )
    replace_picture_before_caption(
        document,
        "图3-5  基于生成对抗模仿学习与TD3的主体训练方案",
        GAIL_TD3_TRAINING_FIGURE,
        width_inches=6.0,
    )
    replace_picture_before_caption(
        document,
        "图4-1  生产企业主体三分支Transformer表征结构",
        TRANSFORMER_THREE_BRANCH_FIGURE,
        width_inches=6.0,
    )
    replace_picture_before_caption(
        document,
        "图6-1 TD3与GAIL+TD3训练结果对比",
        TD3_GAIL_FIGURE,
    )
    replace_picture_before_caption(
        document,
        "图6-4 GAIL+TD3与Transformer+GAIL+TD3训练结果对比",
        GAIL_TRANSFORMER_FIGURE,
    )
    replace_picture_before_caption(
        document,
        "图6-5 三种主体训练方案下DSCR 低于 1 的生产企业日观测占比",
        DSCR_RISK_FIGURE,
    )
    verify(document, source_equation_count)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The source DOCX was modified unexpectedly.")
    verify(Document(OUTPUT), source_equation_count)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
