from copy import deepcopy
from pathlib import Path
import re

from docx import Document


DOCX = next(Path("paper_drafts/rendered_v6").glob("*.docx"))


def normalized(text: str) -> str:
    return " ".join((text or "").split())


def replace_child(parent, tag_name, new_child):
    old = getattr(parent, tag_name)
    if old is not None:
        old.getparent().remove(old)
    if new_child is not None:
        parent.insert(0, deepcopy(new_child))


def apply_heading_format(paragraph, ref_paragraph):
    # Copy paragraph-level spacing/alignment/style settings.
    replace_child(paragraph._p, "pPr", ref_paragraph._p.pPr)

    ref_run = ref_paragraph.runs[0] if ref_paragraph.runs else None
    ref_rpr = ref_run._r.rPr if ref_run is not None else None
    for run in paragraph.runs:
        replace_child(run._r, "rPr", ref_rpr)


doc = Document(str(DOCX))

ref_h1 = None
ref_h2 = None
for paragraph in doc.paragraphs:
    text = normalized(paragraph.text)
    if text == "2 仿真环境与主体决策建模问题":
        ref_h1 = paragraph
    elif text == "2.1 仿真环境":
        ref_h2 = paragraph

if ref_h1 is None or ref_h2 is None:
    raise RuntimeError("未找到参考标题“2”或参考小标题“2.1”。")

for paragraph in doc.paragraphs:
    text = normalized(paragraph.text)
    if not text:
        continue
    if paragraph.style.name == "Heading 1" or re.match(r"^\d+\s+", text) or text == "参考文献":
        apply_heading_format(paragraph, ref_h1)
    elif paragraph.style.name == "Heading 2" or re.match(r"^\d+\.\d+\s+", text):
        apply_heading_format(paragraph, ref_h2)

doc.save(str(DOCX))
print(DOCX)
