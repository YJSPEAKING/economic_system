from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v24"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v24.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v25.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "CF628BF5234D62C8CC93AD598B097BF7760295A8C21B55C3BD3BE950DA9E5062"
)

FIGURE_TREND_ANALYSIS = (
    "从图6-1(a)可以看出，引入模仿学习信号后，GAIL+TD3的平均存活天数曲线"
    "在训练前期较快上升，并较早转入相对稳定阶段；TD3的平均存活天数曲线在"
    "训练前期上升较缓，随后在训练中后期持续提高并进入相近的稳定区间。"
    "图6-1(b)显示，GAIL+TD3的生产企业累计收益曲线较早进入快速提升阶段并转入"
    "相对稳定状态，TD3的生产企业累计收益曲线则在训练中后期快速提高。"
    "图6-1(c)显示，GAIL+TD3的累计归一化学习曲线下面积整体高于TD3，且两条曲线"
    "的差距主要在训练前期形成并延续至训练末期。"
)

TABLE_QUANTITATIVE_ANALYSIS = (
    "表6-2中的具体数值显示，GAIL+TD3达到80天所需的评估时刻由TD3的44缩短至17，"
    "提前27个评估时刻，即2700个训练回合；两种模型的平均存活天数均值曲线在前60个"
    "评估时刻内均未达到连续3个评估时刻不低于90天的条件。GAIL+TD3在第60个评估时刻"
    "的累计归一化学习曲线下面积由TD3的0.435提高至0.683，增幅约为56.8%，表明该模型"
    "在相同训练区间内取得了更高的累积平均存活表现。结合图6-1(b)所示的曲线变化，"
    "模仿奖励能够较早引导生产企业主体形成有效策略。"
)

SECTION_TRANSITION = (
    "不过，GAIL+TD3的最终稳定平均存活天数为80.71天，与专家轨迹平均值仍相差"
    "17.28天。该结果表明，模仿学习信号缩短了模型达到80天所需的训练过程，但稳定"
    "阶段的存活水平仍有进一步接近专家轨迹的空间。因此，下一节引入历史状态序列表征，"
    "以考察模型利用连续交互信息后能否缩小这一差距。"
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    run_properties = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run_properties = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def verify(document: Document) -> None:
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for expected in (
        FIGURE_TREND_ANALYSIS,
        TABLE_QUANTITATIVE_ANALYSIS,
        SECTION_TRANSITION,
    ):
        if expected not in text:
            raise RuntimeError(f"Revised paragraph is missing: {expected[:60]}")

    old_wording = (
        "并在第17个评估时刻达到80天；TD3在第44个评估时刻达到相同水平",
        "两种模型的均值曲线在前60个评估时刻内均未达到连续3个评估时刻不低于90天的条件。图6-1(b)",
        "且其均值曲线在前60个评估时刻内未达到90天",
    )
    for wording in old_wording:
        if wording in text:
            raise RuntimeError(f"Former mixed qualitative/quantitative wording remains: {wording}")

    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 9:
        raise RuntimeError(f"Expected 9 inline figures, found {len(document.inline_shapes)}.")
    if len(document.element.body.xpath(".//m:oMath")) != 117:
        raise RuntimeError("The Word equation-object count changed unexpectedly.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v24 source changed after this builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)
    replace_paragraph_text(
        find_paragraph(document, "从图6-1(a)可以看出"),
        FIGURE_TREND_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "表6-2中的具体数值显示"),
        TABLE_QUANTITATIVE_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "不过，GAIL+TD3的最终稳定"),
        SECTION_TRANSITION,
    )
    verify(document)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v24 source was modified unexpectedly.")
    verify(Document(OUTPUT))
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
