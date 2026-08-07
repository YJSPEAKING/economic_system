from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT / "paper_drafts" / "rendered_v8"
OUT_DIR = ROOT / "paper_drafts" / "rendered_v9"


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_paragraph_text_with_format(paragraph, text, ppr=None, rpr=None):
    clear_runs(paragraph)
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if ppr is not None:
        paragraph._p.insert(0, deepcopy(ppr))
    run = paragraph.add_run(text)
    if rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(rpr))
    return run


def replace_style_format(style, ppr=None, rpr=None):
    style_el = style._element
    for tag in ("w:pPr", "w:rPr"):
        child = style_el.find(qn(tag))
        if child is not None:
            style_el.remove(child)
    if ppr is not None:
        style_el.append(deepcopy(ppr))
    if rpr is not None:
        style_el.append(deepcopy(rpr))


def first_non_empty_run_rpr(paragraph):
    for run in paragraph.runs:
        if run.text.strip():
            return deepcopy(run._r.rPr)
    return deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None


def set_cell_single_run(cell, text, rpr):
    paragraph = cell.paragraphs[0]
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    if rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(rpr))
    for extra in cell.paragraphs[1:]:
        clear_runs(extra)


def main():
    src_files = list(SRC_DIR.glob("*.docx"))
    if not src_files:
        raise FileNotFoundError(f"No docx found in {SRC_DIR}")
    src = src_files[0]

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / src.name.replace("v8", "v9")
    shutil.copy2(src, out)

    doc = Document(out)

    h1_ref = next(p for p in doc.paragraphs if "".join(r.text for r in p.runs).strip().startswith("1  "))
    h2_ref = next(p for p in doc.paragraphs if "".join(r.text for r in p.runs).strip().startswith("2.1  "))
    h1_ppr = deepcopy(h1_ref._p.pPr)
    h1_rpr = first_non_empty_run_rpr(h1_ref)
    h2_ppr = deepcopy(h2_ref._p.pPr)
    h2_rpr = first_non_empty_run_rpr(h2_ref)

    replace_style_format(doc.styles["Heading 1"], h1_ppr, h1_rpr)
    replace_style_format(doc.styles["Heading 2"], h2_ppr, h2_rpr)

    new_p21 = (
        "仅采用完全理性目标时，仿真主体通常围绕奖励最大化调整策略，并以收益水平、存活时间和系统稳定性作为主要评价依据。"
        "该目标有助于检验仿真环境是否可运行，也有助于检验主体能否利用奖励信号形成高收益策略。"
        "但是，经济系统仿真中的主体决策并不总是符合完全理性假设。真实市场主体往往受到信息不完全、风险态度、经验依赖、偏好差异和路径依赖等因素影响，其行为很难被简化为单一收益最大化过程。"
    )
    new_p22 = (
        "基于上述局限，本文进一步引入有限理性仿真目标。该目标不要求生产企业主体在每个状态下输出全局最优动作，"
        "而是关注生成行为是否接近专家样本中体现的决策模式，并考察动作分布、状态覆盖和动态适应性。"
        "模仿学习以专家行为为参照，能够为有限理性行为生成提供训练信号；强化学习仍保留环境交互和价值估计作用，使生成动作与仿真反馈保持关联。"
    )

    for paragraph in doc.paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip()
        if paragraph.style and paragraph.style.name == "Heading 1":
            replacement = "2  仿真环境、目标与主体决策建模问题" if text.startswith("2  ") else text
            set_paragraph_text_with_format(paragraph, replacement, h1_ppr, h1_rpr)
        elif paragraph.style and paragraph.style.name == "Heading 2":
            replacement = "2.3 仿真目标差异与有限理性专家数据表示" if text.startswith("2.3 ") else text
            set_paragraph_text_with_format(paragraph, replacement, h2_ppr, h2_rpr)
        elif text.startswith("复杂经济系统仿真一方面服务于理论经济学研究中的数据生成需求") or text.startswith("仅采用完全理性目标时"):
            clear_runs(paragraph)
            paragraph.add_run(new_p21)
        elif text.startswith("有限理性目标并不要求主体在每个状态下输出全局最优动作") or text.startswith("基于上述局限，本文进一步引入有限理性仿真目标"):
            clear_runs(paragraph)
            paragraph.add_run(new_p22)

    table = doc.tables[0]
    output_rpr = None
    for run in table.rows[1].cells[2].paragraphs[0].runs:
        if run.text.strip():
            output_rpr = deepcopy(run._r.rPr)
            break
    set_cell_single_run(
        table.rows[1].cells[1],
        "33 维状态向量：包含经营状态、需求信息、交易记录、价格信息和时间状态变量",
        output_rpr,
    )
    set_cell_single_run(
        table.rows[2].cells[1],
        "51 维状态向量：包含银行经营状态与企业主体信息",
        output_rpr,
    )

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
