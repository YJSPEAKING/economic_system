from __future__ import annotations

import csv
import math
from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v17"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v17.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v18.docx"
)
STATS_CSV = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "five_metrics_two_comparisons_aulc_sem"
    / "final_performance_last20_summary.csv"
)
EXPECTED_SOURCE_SHA256 = (
    "6F6629D79DE750F600AF15B9E95DBE5F34DD37A3FBABB53E107770F0F94A193C"
)

GROUPS = ("TD3", "GAIL+TD3", "Transformer+GAIL+TD3")
METRICS = ("survival", "production", "consumption", "bank", "naulc")
EXPERT_MEAN_SURVIVAL = 97.99


SECTION_6_2_INTRO = (
    "本文首先比较TD3与GAIL+TD3的训练过程，以分析模仿学习信号对生产企业主体决策训练的影响。"
    "图6-1给出了两种模型在平均存活天数、生产企业累计收益、消费企业累计收益、银行累计收益和累计归一化学习曲线下面积五项指标上的均值曲线及SEM阴影，表6-2汇总了三种模型的最终稳定性能。"
    "图6-1(a)中的水平虚线表示表5-1所列专家轨迹的合计平均存活天数，其数值为97.99天。"
)

FIGURE_6_1_ANALYSIS = (
    "从图6-1可以看出，GAIL+TD3的平均存活天数和三项主体收益曲线均早于TD3进入快速提升阶段，并较早转入相对稳定状态。"
    "其中，GAIL+TD3的平均存活天数在第20个评估时刻之前已接近后期水平，而TD3直至训练后期才达到相近区间；"
    "生产企业、消费企业和银行累计收益也呈现相似的前期差异。"
    "训练后期，TD3的生产企业和消费企业累计收益均值高于GAIL+TD3，且对应的SEM阴影更宽；"
    "银行累计收益则表现出相反关系，GAIL+TD3在完成前期提升后保持了高于TD3的均值。"
    "图6-1(e)所示的累计归一化学习曲线下面积综合记录了平均存活天数在整个训练区间内的提升速度和持续水平，其曲线差异进一步反映出GAIL+TD3在训练前期的学习效率优势。"
)

TABLE_6_2_ANALYSIS = (
    "表6-2对上述曲线差异进行了定量汇总。与TD3相比，GAIL+TD3在第60个评估时刻的累计归一化学习曲线下面积均值由0.435提高至0.683，表明模仿学习信号使模型在相同训练区间内更早形成了较高的平均存活水平。"
    "GAIL+TD3的银行最终稳定累计收益均值由TD3的5.28K提高至10.31K；"
    "生产企业和消费企业最终稳定累计收益均值则分别由205.43K降至126.35K、由185.81K降至150.01K。"
    "这些结果表明，模仿学习信号主要改变了生产企业策略的形成速度，并通过商品交易和信贷关系对其他主体产生了不同方向的影响，因此学习速度与各主体的后期收益水平需要分别评价。"
)

SECTION_6_2_CONCLUSION = (
    "但是，GAIL+TD3的最终稳定平均存活天数均值为80.71天，低于TD3的81.50天，二者相差0.79天；"
    "GAIL+TD3与专家轨迹平均存活天数之间仍相差17.28天。"
    "这一结果说明，模仿学习信号带来的前期学习效率优势没有同步转化为更高的最终稳定存活水平。"
    "因此，本文在下一节进一步引入历史状态序列表征，以考察生产企业主体利用连续交互信息后能否改善稳定阶段的系统存活水平。"
)

SECTION_6_3_INTRO = (
    "基于上一节的结果，本文进一步比较GAIL+TD3与Transformer+GAIL+TD3，以分析历史状态序列表征对生产企业主体训练过程的影响。"
    "图6-2给出了两种模型在五项指标上的均值曲线及SEM阴影，表6-2列出了对应的最终稳定性能。"
    "图6-2(a)中的水平虚线同样表示专家轨迹的合计平均存活天数，用于比较两种模型与专家轨迹存活水平之间的差距。"
)

FIGURE_6_2_ANALYSIS = (
    "从图6-2可以看出，Transformer+GAIL+TD3在平均存活天数、生产企业累计收益和消费企业累计收益等指标上的曲线整体早于GAIL+TD3上升。"
    "Transformer+GAIL+TD3的平均存活天数在训练前期快速提高，并在后期保持了高于GAIL+TD3的均值；"
    "两项企业累计收益在稳定阶段也整体处于较高区间。"
    "这一变化说明，历史状态序列表征为生产企业主体提供了当前状态之外的连续交互信息，并改善了企业经营策略的形成过程。"
    "银行累计收益的变化方向与企业收益并不完全一致，Transformer+GAIL+TD3在前期呈现较大波动，稳定阶段均值低于GAIL+TD3，且SEM阴影较宽，反映出生产企业策略变化对信贷关系及跨随机种子结果产生了间接影响。"
)

TABLE_6_2_TRANSFORMER_ANALYSIS = (
    "表6-2进一步显示，引入Transformer编码器后，最终稳定平均存活天数均值由80.71天提高至90.24天，增加9.52天；"
    "该指标与专家轨迹平均值的差距由17.28天缩小至7.75天。"
    "生产企业和消费企业最终稳定累计收益均值分别由126.35K提高至180.43K、由150.01K提高至169.48K，"
    "第60个评估时刻的累计归一化学习曲线下面积均值由0.683提高至0.773。"
    "银行最终稳定累计收益均值则由10.31K降至7.15K，其相对SEM由21.80%提高至50.02%，说明历史状态表征对不同主体绩效和跨随机种子波动产生的影响并不一致。"
)

FIGURE_6_3_ANALYSIS = (
    "图6-3显示，TD3、GAIL+TD3和Transformer+GAIL+TD3中DSCR小于1的生产企业日观测占比分别为6.52%、0.23%和0.00%。"
    "DSCR小于1表示生产企业当日可支配现金不足以覆盖当期应还本金与利息，因此，较低的风险占比对应更少的当期偿债缺口。"
    "与GAIL+TD3相比，Transformer+GAIL+TD3中DSCR小于1的日观测占比由0.23%降至0.00%，即本组稳定运行观测中未出现生产企业当期现金不足以覆盖应还本息的记录。"
    "该结果与图6-2(a)和表6-2所反映的稳定阶段存活水平一致，说明历史状态序列表征有助于生产企业主体在经营扩张、现金持有和债务偿付之间形成更稳定的决策关系。"
)

SECTION_6_3_CONCLUSION = (
    "综合图6-2、图6-3和表6-2的结果，Transformer+GAIL+TD3在前期学习效率、最终稳定平均存活天数、企业累计收益和生产企业偿债风险占比等方面均较GAIL+TD3取得了更好的均值结果。"
    "银行累计收益及其跨随机种子波动未呈现相同方向的变化，说明历史状态表征改善生产企业决策后，不同经济主体的收益目标仍存在差异。"
    "在本组实验设置下，Transformer编码器所提供的历史状态序列信息进一步弥补了单步状态表征在稳定阶段存活水平上的不足。"
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def copy_first_run_properties(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def replace_paragraph_text(paragraph, text: str) -> None:
    source_rpr = copy_first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def insert_body_after(document: Document, reference, text: str):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    source_rpr = copy_first_run_properties(reference)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)
    reference._p.addnext(paragraph._p)
    return paragraph


def set_run_font(run, *, size: float = 9.0, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")


def set_cell_text(cell, text: str, *, bold: bool = False, centered: bool = True) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=8.5, bold=bold)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths_cm: list[float]) -> None:
    widths_dxa = [round(width * 567) for width in widths_cm]
    total_width = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_width))

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            cell.width = Cm(widths_cm[index])


def load_statistics() -> dict[tuple[str, str], dict[str, float | str]]:
    if not STATS_CSV.exists():
        raise FileNotFoundError(STATS_CSV)
    result: dict[tuple[str, str], dict[str, float | str]] = {}
    with STATS_CSV.open("r", encoding="utf-8-sig", newline="") as file:
        for row in csv.DictReader(file):
            key = (row["group"], row["metric_id"])
            result[key] = {
                "basis": row["basis"],
                "start_step": float(row["start_step"]),
                "end_step": float(row["end_step"]),
                "n_runs": float(row["n_runs"]),
                "mean": float(row["mean"]),
                "sem": float(row["sem"]),
                "relative_sem_percent": float(row["relative_sem_percent"]),
            }

    missing = [
        (group, metric)
        for group in GROUPS
        for metric in METRICS
        if (group, metric) not in result
    ]
    if missing:
        raise RuntimeError(f"Missing statistics rows: {missing}")
    return result


def format_stat(metric_id: str, row: dict[str, float | str]) -> str:
    mean = float(row["mean"])
    relative_sem = float(row["relative_sem_percent"])
    if metric_id == "survival":
        value = f"{mean:.2f}"
    elif metric_id == "naulc":
        value = f"{mean:.3f}"
    else:
        value = f"{mean / 1000.0:.2f}K"
    return f"{value}（±{relative_sem:.2f}%）"


def insert_final_performance_table(document: Document, reference, statistics_rows) -> None:
    definition = document.add_paragraph(style="Normal")
    definition.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    definition.paragraph_format.first_line_indent = Cm(0.74)
    definition.paragraph_format.line_spacing = 1.5
    definition.paragraph_format.space_after = Pt(3)
    definition.add_run(
        "为定量比较三种模型在训练后期的表现，本文先对每次独立运行在第41至第60个评估时刻的平均存活天数和三项累计收益分别取均值，再在8次运行之间计算总体均值和SEM；累计归一化学习曲线下面积采用第60个评估时刻的数值。本文将上述统计结果称为最终稳定性能，具体结果见表6-2。"
    )
    reference._p.addprevious(definition._p)

    caption = document.add_paragraph(style="Normal")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = None
    caption.paragraph_format.line_spacing = 1.2
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(4)
    caption_run = caption.add_run("表6-2 三种主体训练方案的最终性能统计")
    set_run_font(caption_run, size=9)
    reference._p.addprevious(caption._p)

    headers = ("评价指标", *GROUPS)
    row_specs = (
        ("平均存活天数/天", "survival"),
        ("生产企业累计收益", "production"),
        ("消费企业累计收益", "consumption"),
        ("银行累计收益", "bank"),
        ("累计归一化学习曲线下面积", "naulc"),
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade_cell(table.rows[0].cells[index], "F2F2F2")
        set_cell_text(table.rows[0].cells[index], header, bold=True)
    for label, metric_id in row_specs:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, centered=False)
        for group_index, group_name in enumerate(GROUPS, start=1):
            set_cell_text(
                cells[group_index],
                format_stat(metric_id, statistics_rows[(group_name, metric_id)]),
            )
    set_table_geometry(table, [4.1, 3.3, 3.8, 4.4])
    reference._p.addprevious(table._tbl)

    note = document.add_paragraph(style="Normal")
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note.paragraph_format.first_line_indent = None
    note.paragraph_format.line_spacing = 1.2
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(4)
    note_run = note.add_run(
        "注：平均存活天数及三项累计收益为每次独立运行最后20个评估时刻的均值，再对8次运行进行统计；累计归一化学习曲线下面积采用第60个评估时刻的数值。括号内为相对SEM，K表示10³。"
    )
    set_run_font(note_run, size=8.5)
    reference._p.addprevious(note._p)


def verify(document: Document, statistics_rows) -> None:
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    expected_fragments = (
        "表6-2 三种主体训练方案的最终性能统计",
        "累计归一化学习曲线下面积均值由0.435提高至0.683",
        "最终稳定平均存活天数均值为80.71天，低于TD3的81.50天",
        "最终稳定平均存活天数均值由80.71天提高至90.24天",
        "DSCR小于1的日观测占比由0.23%降至0.00%",
        SECTION_6_3_CONCLUSION,
    )
    for fragment in expected_fragments:
        if fragment not in all_text:
            raise RuntimeError(f"Missing expected text: {fragment[:50]!r}")

    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    table = document.tables[-1]
    if len(table.rows) != 6 or len(table.columns) != 4:
        raise RuntimeError("Table 6-2 geometry is incorrect.")
    if table.rows[0].cells[0].text != "评价指标":
        raise RuntimeError("Table 6-2 header is incorrect.")

    expected_cells = {
        (1, 1): format_stat("survival", statistics_rows[("TD3", "survival")]),
        (1, 2): format_stat("survival", statistics_rows[("GAIL+TD3", "survival")]),
        (1, 3): format_stat(
            "survival", statistics_rows[("Transformer+GAIL+TD3", "survival")]
        ),
        (5, 1): format_stat("naulc", statistics_rows[("TD3", "naulc")]),
        (5, 2): format_stat("naulc", statistics_rows[("GAIL+TD3", "naulc")]),
        (5, 3): format_stat(
            "naulc", statistics_rows[("Transformer+GAIL+TD3", "naulc")]
        ),
    }
    for (row_index, col_index), expected in expected_cells.items():
        actual = table.rows[row_index].cells[col_index].text
        if actual != expected:
            raise RuntimeError(
                f"Unexpected Table 6-2 value at {(row_index, col_index)}: {actual!r}"
            )

    figure_6_1_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("从图6-1可以看出")
    ]
    figure_6_2_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("从图6-2可以看出")
    ]
    if len(figure_6_1_paragraphs) != 1 or len(figure_6_2_paragraphs) != 1:
        raise RuntimeError("Each learning-curve figure must retain one analysis paragraph.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v17 source DOCX changed after the v18 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    statistics_rows = load_statistics()
    document = Document(SOURCE)

    heading_6_2 = find_paragraph(document, "6.2 TD3与GAIL+TD3")
    insert_final_performance_table(document, heading_6_2, statistics_rows)

    replace_paragraph_text(
        find_paragraph(document, "本文首先比较TD3与GAIL+TD3"),
        SECTION_6_2_INTRO,
    )
    figure_6_1 = find_paragraph(document, "从图6-1可以看出")
    replace_paragraph_text(figure_6_1, FIGURE_6_1_ANALYSIS)
    table_6_2_analysis = insert_body_after(
        document, figure_6_1, TABLE_6_2_ANALYSIS
    )
    insert_body_after(document, table_6_2_analysis, SECTION_6_2_CONCLUSION)

    replace_paragraph_text(
        find_paragraph(document, "本文进一步比较GAIL+TD3与Transformer+GAIL+TD3"),
        SECTION_6_3_INTRO,
    )
    figure_6_2 = find_paragraph(document, "从图6-2可以看出")
    replace_paragraph_text(figure_6_2, FIGURE_6_2_ANALYSIS)
    insert_body_after(document, figure_6_2, TABLE_6_2_TRANSFORMER_ANALYSIS)

    figure_6_3 = find_paragraph(document, "图6-3显示")
    replace_paragraph_text(figure_6_3, FIGURE_6_3_ANALYSIS)
    insert_body_after(document, figure_6_3, SECTION_6_3_CONCLUSION)

    verify(document, statistics_rows)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v17 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check, statistics_rows)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
