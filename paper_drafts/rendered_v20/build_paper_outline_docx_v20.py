from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v19"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v19.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v20.docx"
)
FIGURE_6_3 = (
    ROOT
    / "real_System_remake"
    / "analysis_plots"
    / "post_training_dscr"
    / "three_algorithm_recent_100_survival_gt_90"
    / "three_algorithm_dscr_boxplot_and_risk.png"
)
EXPECTED_SOURCE_SHA256 = (
    "0982E883D4AA7AECDD5022152EA4F40B1C80C01E41F2E597C128877B58034EF3"
)

MERGED_SECTION_6_2_OPENING = (
    "本文首先比较TD3与GAIL+TD3的训练过程，以分析模仿学习信号对生产企业主体决策训练的影响。"
    "图6-1给出了两种模型在平均存活天数、生产企业累计收益、消费企业累计收益、银行累计收益和累计归一化学习曲线下面积五项指标上的均值曲线及SEM阴影。"
    "其中，图6-1(a)中的水平虚线表示表5-1所列专家轨迹的合计平均存活天数，其数值为97.99天。"
    "为定量比较三种模型的训练后期表现，本文先对每次独立运行在第41至第60个评估时刻的平均存活天数和三项累计收益分别取均值，再在8次独立运行之间计算总体均值和SEM；"
    "累计归一化学习曲线下面积采用第60个评估时刻的数值。"
    "表6-2汇总了三种模型的相应统计结果。"
)

TABLE_6_1_LEAD = (
    "本文以回合作为训练进度单位，并将每个评估时刻所包含的连续训练回合数记为E。"
    "模型每完成E个连续训练回合，本文汇总一次评价指标，并将该汇总位置记为一个评估时刻（evaluation step）。"
    "第k个评估时刻对应第k组连续训练回合的统计结果。"
    "表6-1列出了总训练天数、随机种子组数、每个评估时刻对应回合数、单回合最大天数和统计方法。"
)

FIGURE_6_3_CAPTION = (
    "图6-3 三种主体训练方案下DSCR 低于 1 的生产企业日观测占比"
)

TABLE_6_1_ROWS = (
    ("项目", "设定"),
    ("总训练天数", "180000天"),
    ("随机种子组数", "每种模型8组"),
    ("每个评估时刻对应回合数", "E=100个连续训练回合"),
    ("单回合最大天数", "100天"),
    ("统计方法", "8次独立运行的均值及均值上下1个SEM"),
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def first_run_properties(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def replace_paragraph_text(paragraph, text: str) -> None:
    source_rpr = first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def replace_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    source_rpr = first_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def find_table_6_1(document: Document):
    matches = []
    for table in document.tables:
        if len(table.rows) != len(TABLE_6_1_ROWS) or len(table.columns) != 2:
            continue
        if table.rows[0].cells[0].text == "项目" and any(
            row.cells[0].text == "对比模型" for row in table.rows
        ):
            matches.append(table)
    if len(matches) != 1:
        raise RuntimeError(f"Expected one Table 6-1, found {len(matches)}.")
    return matches[0]


def replace_figure_image(document: Document, caption, image_path: Path) -> str:
    image_paragraph_element = caption._p.getprevious()
    if image_paragraph_element is None:
        raise RuntimeError("The Figure 6-3 image paragraph is missing.")
    blips = image_paragraph_element.xpath(".//a:blip")
    if len(blips) != 1:
        raise RuntimeError(f"Expected one Figure 6-3 image relationship, found {len(blips)}.")
    relationship_id = blips[0].get(qn("r:embed"))
    image_part = document.part.related_parts[relationship_id]
    new_blob = image_path.read_bytes()
    image_part._blob = new_blob
    return sha256(new_blob).hexdigest()


def verify(document: Document, expected_image_hash: str) -> None:
    opening = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    if opening.text != MERGED_SECTION_6_2_OPENING:
        raise RuntimeError("The merged Section 6.2 opening is incorrect.")
    if any(
        paragraph.text.startswith("为定量比较三种模型在训练后期的表现")
        for paragraph in document.paragraphs
    ):
        raise RuntimeError("The former separate Section 6.2 paragraph remains.")

    lead = find_paragraph(document, "本文以回合作为训练进度单位")
    if lead.text != TABLE_6_1_LEAD:
        raise RuntimeError("The Table 6-1 lead paragraph is incorrect.")

    table = document.tables[2]
    actual_rows = tuple(
        (row.cells[0].text, row.cells[1].text) for row in table.rows
    )
    if actual_rows != TABLE_6_1_ROWS:
        raise RuntimeError(f"Table 6-1 rows are incorrect: {actual_rows!r}")

    caption = find_paragraph(document, "图6-3 ")
    if caption.text != FIGURE_6_3_CAPTION:
        raise RuntimeError(f"Unexpected Figure 6-3 caption: {caption.text!r}")
    image_paragraph_element = caption._p.getprevious()
    blip = image_paragraph_element.xpath(".//a:blip")[0]
    relationship_id = blip.get(qn("r:embed"))
    image_part = document.part.related_parts[relationship_id]
    embedded_hash = sha256(image_part.blob).hexdigest()
    if embedded_hash != expected_image_hash:
        raise RuntimeError("The embedded Figure 6-3 image was not replaced correctly.")

    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 9:
        raise RuntimeError(f"Expected 9 inline images, found {len(document.inline_shapes)}.")
    if len(document.element.body.xpath(".//m:oMath")) != 116:
        raise RuntimeError("The existing Word equation objects changed unexpectedly.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not FIGURE_6_3.exists():
        raise FileNotFoundError(FIGURE_6_3)
    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v19 source DOCX changed after the v20 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)

    former_definition = find_paragraph(
        document, "为定量比较三种模型在训练后期的表现"
    )
    former_comparison = find_paragraph(document, "本文首先比较TD3与GAIL+TD3")
    replace_paragraph_text(former_definition, MERGED_SECTION_6_2_OPENING)
    remove_paragraph(former_comparison)

    replace_paragraph_text(
        find_paragraph(document, "本文以回合作为训练进度单位"),
        TABLE_6_1_LEAD,
    )
    table_6_1 = find_table_6_1(document)
    for row, (item, setting) in zip(table_6_1.rows, TABLE_6_1_ROWS):
        replace_cell_text(row.cells[0], item)
        replace_cell_text(row.cells[1], setting)

    figure_caption = find_paragraph(
        document, "图6-3 三种主体训练方案下生产企业偿债风险对比"
    )
    replace_paragraph_text(figure_caption, FIGURE_6_3_CAPTION)
    expected_image_hash = replace_figure_image(document, figure_caption, FIGURE_6_3)

    verify(document, expected_image_hash)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v19 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check, expected_image_hash)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
