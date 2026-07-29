from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "paper_drafts"
    / "rendered_v20"
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v20.docx"
)
OUTPUT = (
    Path(__file__).resolve().parent
    / "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v21.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "05967C448C014BD2259A29F59CACBB33963533F84E86ACACEEC4A503A3FE28E7"
)

FIGURE_ANALYSIS = (
    "从图6-1可以看出，GAIL+TD3的平均存活天数和三项主体收益曲线均早于TD3进入快速提升阶段，并较早转入相对稳定状态。"
    "其中，GAIL+TD3的平均存活天数在第20个评估时刻之前已接近后期水平，而TD3直至训练后期才达到相近区间；"
    "生产企业、消费企业和银行累计收益也呈现相似的前期差异。"
    "图6-1(e)显示，GAIL+TD3的累计归一化学习曲线下面积整体高于TD3，且两条曲线的差距主要在训练前期形成，并延续至训练末期。"
)

TABLE_ANALYSIS = (
    "表6-2中的具体数值显示，GAIL+TD3在第60个评估时刻的累计归一化学习曲线下面积均值由TD3的0.435提高至0.683，增幅约为56.8%；"
    "银行最终稳定累计收益均值由5.28K提高至10.31K，增幅约为95.4%。"
    "与此同时，生产企业和消费企业最终稳定累计收益均值分别由205.43K降至126.35K、由185.81K降至150.01K。"
    "这些结果表明，GAIL+TD3虽然并未在所有主体收益指标上取得更高的最终稳定均值，但其提高了平均存活天数的累积表现和银行收益，"
    "在训练速度与不同主体的后期收益目标之间呈现出不同于TD3的权衡关系。"
)

LIMITATION_AND_TRANSITION = (
    "不过，表6-2和图6-1也显示，GAIL+TD3的最终稳定平均存活天数均值为80.71天，低于TD3的81.50天，二者相差0.79天；"
    "GAIL+TD3与专家轨迹平均存活天数之间仍相差17.28天。"
    "此外，GAIL+TD3最终稳定平均存活天数的相对SEM为5.38%，高于TD3的3.05%，说明其稳定阶段的平均存活天数在不同随机种子之间具有更大的相对波动。"
    "上述差异表明，模仿学习信号虽然加快了训练前期的策略形成过程，但在最终存活水平和跨随机种子稳定性方面仍存在改进空间。"
    "因此，本文在下一节进一步引入历史状态序列表征，以考察生产企业主体利用连续交互信息后能否改善稳定阶段的系统存活水平。"
)


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {prefix!r}, found {len(matches)}."
        )
    return matches[0]


def replace_paragraph_text(paragraph, text: str) -> None:
    source_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            source_rpr = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if source_rpr is not None:
        run._r.insert(0, source_rpr)


def verify(document: Document) -> None:
    expected = (
        FIGURE_ANALYSIS,
        TABLE_ANALYSIS,
        LIMITATION_AND_TRANSITION,
    )
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for paragraph_text in expected:
        if paragraph_text not in all_text:
            raise RuntimeError(f"Missing revised analysis: {paragraph_text[:50]!r}")

    forbidden = (
        "图6-1(e)所示的累计归一化学习曲线下面积综合记录了",
        "表6-2对上述曲线差异进行了定量汇总",
        "但是，GAIL+TD3的最终稳定平均存活天数",
    )
    for fragment in forbidden:
        if fragment in all_text:
            raise RuntimeError(f"Former wording remains: {fragment!r}")

    if len(document.tables) != 4:
        raise RuntimeError(f"Expected 4 tables, found {len(document.tables)}.")
    if len(document.inline_shapes) != 9:
        raise RuntimeError(f"Expected 9 inline images, found {len(document.inline_shapes)}.")
    if len(document.element.body.xpath(".//m:oMath")) != 116:
        raise RuntimeError("The existing Word equation objects changed unexpectedly.")


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = file_sha256(SOURCE)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The v20 source DOCX changed after the v21 builder was prepared. "
            f"Expected {EXPECTED_SOURCE_SHA256}, got {source_hash}."
        )

    document = Document(SOURCE)
    replace_paragraph_text(
        find_paragraph(document, "从图6-1可以看出"), FIGURE_ANALYSIS
    )
    replace_paragraph_text(
        find_paragraph(document, "表6-2对上述曲线差异进行了定量汇总"),
        TABLE_ANALYSIS,
    )
    replace_paragraph_text(
        find_paragraph(document, "但是，GAIL+TD3的最终稳定平均存活天数"),
        LIMITATION_AND_TRANSITION,
    )

    verify(document)
    document.save(OUTPUT)

    if file_sha256(SOURCE) != source_hash:
        raise RuntimeError("The v20 source DOCX was modified unexpectedly.")

    check = Document(OUTPUT)
    verify(check)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
