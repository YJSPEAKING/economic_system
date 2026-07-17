from __future__ import annotations

import copy
import hashlib
import os
import re
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = os.path.join(
    "paper_drafts",
    "rendered_v12",
    "基于模仿学习的复杂经济系统主体决策模型小论文大纲稿_修改版v12.docx",
)
TEMP_PATH = DOCX_PATH + ".tmp.docx"


P23 = (
    "为了让模型学习有限理性行为，本文先将人工和辅助采集到的决策过程整理为专家样本。本文具体使用两类专家数据：第一类来自网页"
    "实验系统中的人工参与式采集，参与者扮演生产企业的经理，并依据每日可见状态填写贷款、生产资料采购、消费品采购和生产资料"
    "次日定价动作；第二类来自同一网页实验系统中的Codex交互式辅助采集。Codex 5.5通过Control Chrome读取每日可见状态和动作"
    "边界，依据角色与目标、环境规则、实时观察、动作约束和反馈修正提示生成四项业务动作，再由插件写入页面并提交。两类轨迹均"
    "经过环境执行、动作边界检查和存活天数筛选，并以生产企业状态—动作对形式保存，用于后续GAIL与TD3训练。"
)

P96 = (
    "本文使用两类数据来源。第一类为网页人工采集数据。参与者按照生产企业角色阅读当天可见状态，填写贷款、生产资料采购、"
    "消费品采购和生产资料次日定价。第二类为Codex交互式辅助采集数据。Codex 5.5通过Control Chrome读取同一网页中的可见信息"
    "和动作边界，依据第5.3节所述分层提示逐日生成并提交四项业务动作。为控制环境随机性和非目标主体策略变化对采集结果的影响，"
    "两类采集均采用统一的实验随机种子，并固定消费企业和银行的预训练Actor参数。采集阶段不对Codex、消费企业或银行进行在线"
    "训练，仅执行模型推理和环境交互，从而避免实时训练带来的额外计算和等待开销。所有轨迹均按照相同的环境规则执行，并依据存活"
    "天数进行筛选。"
)

P131 = (
    "表5-1根据筛选后的人工轨迹和Codex交互式辅助轨迹统计。网页人工采集部分目前包含8个存活超过90天的回合和782条人工提交"
    "记录；Codex辅助采集尚未完成，成功回合数和状态—动作样本数将在采集结束后补充。最终数据集由两类筛选后的轨迹合并形成，"
    "每条记录由33维状态和4维动作组成，可供第3章和第4章所述模型读取。"
)

P133 = (
    "专家轨迹可能包含局部失误、经验性捷径或不稳定动作。本文仅保留存活天数大于90天的回合，该条件能够排除在第90天及以前"
    "终止的轨迹，但不能证明保留轨迹中的每个动作均为最优或无误。网页端的动作边界、环境执行检查和回合级筛选共同构成基础质量"
    "控制。对于Codex交互式辅助轨迹，存活天数筛选同样只能排除明显失败的轨迹，不能单独证明生成行为与人类行为一致。后续实验"
    "仍需从动作分布、状态条件下的响应方向、动作波动和时间序列特征等方面比较人工轨迹与Codex交互式辅助轨迹。"
)


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def base_run_properties(paragraph: Paragraph):
    for run in paragraph.runs:
        if run.text and run._r.rPr is not None:
            return copy.deepcopy(run._r.rPr)
    return None


def replace_text(paragraph: Paragraph, text: str) -> None:
    run_props = base_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_props is not None:
        run._r.insert(0, run_props)


def format_inline_math(paragraph: Paragraph, segments) -> None:
    run_props = base_run_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    for segment in segments:
        if isinstance(segment, str):
            run = paragraph.add_run(segment)
            if run_props is not None:
                run._r.insert(0, copy.deepcopy(run_props))
            continue
        base, subscript, plain_subscript = segment
        base_run = paragraph.add_run(base)
        if run_props is not None:
            base_run._r.insert(0, copy.deepcopy(run_props))
        base_run.italic = True
        if subscript is not None:
            sub_run = paragraph.add_run(subscript)
            if run_props is not None:
                sub_run._r.insert(0, copy.deepcopy(run_props))
            sub_run.font.subscript = True
            sub_run.italic = False if plain_subscript else True


def math_run(text: str, plain: bool = False):
    run = OxmlElement("m:r")
    if plain:
        run_props = OxmlElement("m:rPr")
        style = OxmlElement("m:sty")
        style.set(qn("m:val"), "p")
        run_props.append(style)
        run.append(run_props)
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def math_subscript(base: str, subscript: str, plain_subscript: bool = False):
    element = OxmlElement("m:sSub")
    base_node = OxmlElement("m:e")
    base_node.append(math_run(base))
    sub_node = OxmlElement("m:sub")
    sub_node.append(math_run(subscript, plain=plain_subscript))
    element.append(base_node)
    element.append(sub_node)
    return element


def replace_formula(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    math_para = OxmlElement("m:oMathPara")
    math_para_props = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "center")
    math_para_props.append(justification)
    math_para.append(math_para_props)

    equation = OxmlElement("m:oMath")
    equation.append(math_subscript("a", "t"))
    equation.append(math_run(" = "))
    equation.append(math_subscript("f", "Codex", plain_subscript=True))
    equation.append(math_run("("))
    equation.append(math_subscript("P", "role", plain_subscript=True))
    equation.append(math_run(", "))
    equation.append(math_subscript("P", "rule", plain_subscript=True))
    equation.append(math_run(", "))
    equation.append(math_subscript("o", "t"))
    equation.append(math_run(", "))
    equation.append(math_subscript("h", "t"))
    equation.append(math_run(", "))
    equation.append(math_subscript("b", "t"))
    equation.append(math_run(")"))
    math_para.append(equation)
    paragraph._p.append(math_para)


def replace_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    paragraph = cell.paragraphs[0]
    replace_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        element = extra._element
        element.getparent().remove(element)


def image_hashes(document: Document):
    return sorted(
        hashlib.sha256(rel.target_part.blob).hexdigest()
        for rel in document.part.rels.values()
        if "image" in rel.reltype
    )


def xml_signature(element):
    if element is None:
        return None
    return (
        element.tag,
        tuple(sorted(element.attrib.items())),
        element.text,
        tuple(xml_signature(child) for child in element),
    )


def is_caption(paragraph: Paragraph) -> bool:
    return bool(re.match(r"^(图|表)\s*\d+[-－—–]\d+\s+", paragraph.text.strip()))


def main() -> None:
    document = Document(DOCX_PATH)
    paragraphs = document.paragraphs

    p23 = find_paragraph(document, "为了让模型学习有限理性行为")
    p95 = find_paragraph(document, "参与者或辅助采集程序在第")
    p96 = find_paragraph(document, "本文使用两类数据来源")
    p131 = find_paragraph(document, "表5-1根据筛选后的人工轨迹")
    p133 = find_paragraph(document, "专家轨迹可能包含局部失误")
    section_53_index = next(i for i, p in enumerate(paragraphs) if p.text.startswith("5.3 "))
    section_54_index = next(i for i, p in enumerate(paragraphs) if p.text.startswith("5.4 "))
    formula = next(
        p for i, p in enumerate(paragraphs)
        if section_53_index < i < section_54_index and p._p.xpath(".//m:oMathPara")
    )
    explanation = find_paragraph(document, "式中，a")

    changed_elements = {p23._p, p95._p, p96._p, formula._p, explanation._p, p131._p, p133._p}
    changed_indices = {i for i, p in enumerate(paragraphs) if p._p in changed_elements}
    if len(changed_indices) != 7:
        raise RuntimeError(f"Expected seven changed paragraph locations, found {sorted(changed_indices)}")
    unchanged_paragraph_signatures = {
        i: xml_signature(p._p) for i, p in enumerate(paragraphs) if i not in changed_indices
    }
    original_paragraph_count = len(paragraphs)
    original_images = image_hashes(document)
    original_inline_shapes = len(document.inline_shapes)
    original_captions = [p.text for p in paragraphs if is_caption(p)]
    original_heading_texts = [p.text for p in paragraphs if p.style.name.startswith("Heading")]
    original_table_count = len(document.tables)
    table = document.tables[1]
    original_table_dimensions = [(len(t.rows), len(t.columns)) for t in document.tables]
    original_other_cells = {
        (ri, ci): xml_signature(cell._tc)
        for ri, row in enumerate(table.rows)
        for ci, cell in enumerate(row.cells)
        if (ri, ci) not in {(1, 3), (2, 0), (2, 1), (2, 2), (2, 3), (3, 1), (3, 2), (3, 3)}
    }

    replace_text(p23, P23)
    format_inline_math(
        p95,
        [
            "人工参与者或Codex在第", ("t", None, False), "日根据状态", ("s", "t", False),
            "提交生产企业动作", ("a", "t", False),
            "。对于Codex采集，Control Chrome负责读取页面状态并执行动作填写和提交。系统补全消费企业和银行的动作，并将",
            ("s", "t", False), "与", ("a", "t", False),
            "保存为专家样本。随后，环境依次完成贷款发放、商品交易、商品生产和清算等阶段，生成下一状态",
            ("s", "t+1", False),
            "和回合终止标志。若回合继续，网页将", ("s", "t+1", False),
            "作为下一日经营信息呈现；若回合终止，系统显示并汇总存活天数等回合信息。",
        ],
    )
    replace_text(p96, P96)
    replace_formula(formula)
    format_inline_math(
        explanation,
        [
            "式中，", ("a", "t", False), "表示第", ("t", None, False),
            "日的生产企业动作；", ("f", "Codex", True),
            "表示Codex 5.5在提示和网页状态约束下形成的隐式决策映射；",
            ("P", "role", True), "表示角色与经营目标提示；",
            ("P", "rule", True), "表示环境规则提示；", ("o", "t", False),
            "表示第", ("t", None, False), "日网页可见状态；", ("h", "t", False),
            "表示截至第", ("t", None, False), "日的近期状态与动作记录；",
            ("b", "t", False), "表示网页在第", ("t", None, False),
            "日给出的动态动作边界。该式不是具有预设系数的显式决策函数，也不涉及公式参数初始化或在线更新。本文预先设置角色提示、环境规则和输出约束，Codex在各决策日根据变化后的输入重新推理并生成动作。",
        ],
    )
    replace_text(p131, P131)
    replace_text(p133, P133)

    replace_cell_text(
        table.cell(1, 3),
        "由参与者通过网页逐日提交动作；保留8个存活天数大于90天的回合，共782条状态—动作样本。",
    )
    replace_cell_text(table.cell(2, 0), "Codex交互式辅助采集数据")
    replace_cell_text(table.cell(2, 1), "待补充")
    replace_cell_text(table.cell(2, 2), "待补充")
    replace_cell_text(
        table.cell(2, 3),
        "Codex 5.5通过Control Chrome逐日读取网页状态、生成并提交动作；按照相同的生存条件筛选。",
    )
    replace_cell_text(table.cell(3, 1), "待补充")
    replace_cell_text(table.cell(3, 2), "待补充")
    replace_cell_text(
        table.cell(3, 3),
        "由网页人工轨迹与Codex辅助轨迹合并形成；每条样本包含33维状态和4维动作。",
    )

    document.save(TEMP_PATH)
    with zipfile.ZipFile(TEMP_PATH, "r") as archive:
        if archive.testzip() is not None:
            raise RuntimeError("The generated DOCX archive failed its CRC check.")

    checked = Document(TEMP_PATH)
    if len(checked.paragraphs) != original_paragraph_count:
        raise RuntimeError("The paragraph count changed unexpectedly.")
    for i, signature in unchanged_paragraph_signatures.items():
        if xml_signature(checked.paragraphs[i]._p) != signature:
            raise RuntimeError(f"Unrelated paragraph {i} changed unexpectedly.")
    if image_hashes(checked) != original_images or len(checked.inline_shapes) != original_inline_shapes:
        raise RuntimeError("Images changed unexpectedly.")
    if len(checked.tables) != original_table_count:
        raise RuntimeError("The table count changed unexpectedly.")
    if [(len(t.rows), len(t.columns)) for t in checked.tables] != original_table_dimensions:
        raise RuntimeError("Table dimensions changed unexpectedly.")
    checked_table = checked.tables[1]
    for key, signature in original_other_cells.items():
        if xml_signature(checked_table.cell(*key)._tc) != signature:
            raise RuntimeError(f"Unrelated table cell {key} changed unexpectedly.")
    if [p.text for p in checked.paragraphs if is_caption(p)] != original_captions:
        raise RuntimeError("Figure or table captions changed unexpectedly.")
    if [p.text for p in checked.paragraphs if p.style.name.startswith("Heading")] != original_heading_texts:
        raise RuntimeError("Headings changed unexpectedly.")

    checked_53_index = next(i for i, p in enumerate(checked.paragraphs) if p.text.startswith("5.3 "))
    checked_54_index = next(i for i, p in enumerate(checked.paragraphs) if p.text.startswith("5.4 "))
    formula_checked = next(
        p for i, p in enumerate(checked.paragraphs)
        if checked_53_index < i < checked_54_index and p._p.xpath(".//m:oMathPara")
    )
    formula_text = "".join(node.text or "" for node in formula_checked._p.xpath(".//m:t"))
    if formula_text != "at = fCodex(Prole, Prule, ot, ht, bt)":
        raise RuntimeError(f"Unexpected formula content: {formula_text!r}")

    body_prefixes = [
        "为了让模型学习有限理性行为",
        "人工参与者或Codex在第",
        "本文使用两类数据来源",
        "式中，a",
        "表5-1根据筛选后的人工轨迹",
        "专家轨迹可能包含局部失误",
    ]
    for prefix in body_prefixes:
        paragraph = find_paragraph(checked, prefix)
        indent = paragraph._p.pPr.ind if paragraph._p.pPr is not None else None
        if indent is None or indent.get(qn("w:firstLineChars")) != "200":
            raise RuntimeError(f"Modified body paragraph lacks a two-character first-line indent: {prefix}")

    global_text = "\n".join(p.text for p in checked.paragraphs)
    global_text += "\n" + "\n".join(cell.text for row in checked.tables[1].rows for cell in row.cells)
    forbidden = [
        "大模型辅助批量采集",
        "参与者或辅助采集程序",
        "大模型辅助设计的类人行为采集数据",
        "该程序参考采集目录中的成功轨迹",
        "大模型辅助设计轨迹",
        "大模型辅助设计的类人批量采集数据",
        "程序提交记录",
        "survival_days",
        "expert_data_production1_collected.csv",
        "无表头",
    ]
    if any(item in global_text for item in forbidden):
        raise RuntimeError("An obsolete collection-method statement remains in the document.")
    required = [
        "Codex交互式辅助采集",
        "Codex 5.5通过Control Chrome",
        "不涉及公式参数初始化或在线更新",
        "Codex交互式辅助采集数据",
    ]
    if any(item not in global_text for item in required):
        raise RuntimeError("A required aligned collection-method statement is missing.")

    os.replace(TEMP_PATH, DOCX_PATH)
    print(f"Updated: {DOCX_PATH}")
    print("Aligned sections: 2.3, 5.1, 5.3, 5.4, and Table 5-1")
    print("Verified: formula, paragraph indents, images, headings, captions, and unrelated content")


if __name__ == "__main__":
    try:
        main()
    finally:
        if os.path.exists(TEMP_PATH):
            os.remove(TEMP_PATH)
